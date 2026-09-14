"""Full synthetic regression for v3.1.1.
Covers NEW/EXISTING weekly rendering, 2D~6D text, multi-image collages,
page-1 summary/title, occurrence stage/date, and grouped/ungrouped marker layouts.
"""
from pathlib import Path
import io, sys
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'app'))
import main_v311 as app
import main_v310 as v310

OUT = ROOT / 'test_output_v311'
OUT.mkdir(exist_ok=True)


def png(label):
    im = Image.new('RGB', (500, 260), 'white')
    d = ImageDraw.Draw(im); d.rectangle((5,5,495,255), outline='black', width=4); d.text((30,110), label, fill='black')
    b = io.BytesIO(); im.save(b, 'PNG'); return b.getvalue()


def textbox(sl, text, x, y, w=2, h=.45):
    sh = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); sh.text = text; return sh


def make_template(path, grouped):
    prs = Presentation(); blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    tb = s1.shapes.add_table(2,5,Inches(.3),Inches(.3),Inches(9.5),Inches(1)).table
    for i,h in enumerate(['이슈','과제명','문제/현상','진행 현황','Signal']): tb.cell(0,i).text = h
    textbox(s1, '000팀 주요 논의 사항', .5, 1.5, 5)
    s2 = prs.slides.add_slide(blank)
    pos=[('2D',1,1,'문제/현상'),('3D',3,2,'임시조치'),('4D',1,5,'원인분석'),('5D',7,2,'개선대책'),('6D',8,5,'효과검증')]
    original={}
    for lab,x,y,title in pos:
        a=s2.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.5),Inches(.5)); a.text=lab
        b=textbox(s2,title,x+.55,y,1.5,.5)
        if grouped:
            g=s2.shapes.add_group_shape([a,b]); g.name='TEMPLATE_GROUP_'+lab; original[lab]=(g.left,g.top)
        else: original[lab]=(a.left,a.top)
    textbox(s2,'발생단계 (발생일자)',.5,.5,2.5,.4); textbox(s2,'',3.1,.5,2,.4)
    logo=s2.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(9.5),Inches(.2),Inches(1),Inches(.4)); logo.text='KEEP_LOGO'; logo.name='LOGO'
    prs.save(path); return original


def all_text(prs):
    vals=[]
    for sl in prs.slides:
        for sh in v310.walk(sl):
            vals.append(getattr(sh,'text',''))
            if getattr(sh,'has_table',False):
                for row in sh.table.rows:
                    vals.extend(c.text for c in row.cells)
    return '\n'.join(vals)


def run_case(mode, grouped):
    src=OUT/f'template_{mode}_{grouped}.pptx'; out=OUT/f'output_{mode}_{grouped}.pptx'; original=make_template(src,grouped)
    d={'task_name':'TestTask','issue_name':'Ford_TestIssue','problem':'2D 문제 현상 원문','temporary_action':'3D 임시조치 원문','customer_response':'고객 대응 완료','cause_4d':'4D 발생원인 원문','leak_cause':'4D 유출원인 원문','system_cause':'4D 시스템원인 원문','action_5d':'5D 개선대책 원문','verification_6d':'6D 검증 완료 원문','status':'개선 완료','occurrence_date':'2026-09-10','_section_images':{}}
    for k in ['2D','3D','4D','5D','6D']:
        d['_section_images'][k]=[(png(k+'-1'),(0,0,1,1),1),(png(k+'-2'),(0,0,1,1),1)]
    g={'team':'Pack개발품질1','owner':'홍길동','stage':'양산'}
    _, saved=app.weekly(str(src),str(out),d,g,mode)
    prs=Presentation(saved); assert len(prs.slides)==2
    text=all_text(prs)
    for expected in ['Pack개발품질1팀 주요 논의 사항','Ford_TestIssue','TestTask','2D 문제 현상 원문','3D 임시조치 원문','4D 발생원인 원문','4D 유출원인 원문','4D 시스템원인 원문','5D 개선대책 원문','6D 검증 완료 원문','양산 (2026-09-10)','KEEP_LOGO']:
        assert expected in text, expected
    assert len([s for s in prs.slides[1].shapes if s.name.startswith('AUTO_8D_IMG_')]) == 5
    for lab in ['2D','3D','4D','5D','6D']:
        marker,parent=v310.find_marker(prs.slides[1],lab); obj=parent or marker
        if mode=='신규':
            x,y=v310.NEW_MARKERS[lab]; assert abs(obj.left-Inches(x))<2000 and abs(obj.top-Inches(y))<2000
        else:
            assert (obj.left,obj.top)==original[lab]
    return saved


if __name__=='__main__':
    for grouped in (True,False):
        for mode in ('신규','기존'):
            print('PASS', mode, 'grouped=', grouped, run_case(mode,grouped))
    print('FULL SYNTHETIC PASS')
