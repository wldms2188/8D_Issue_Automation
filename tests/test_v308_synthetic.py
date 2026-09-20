from pathlib import Path
import io, sys, traceback
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'app'))
import main_v310 as app

OUT = ROOT / 'test_output'
OUT.mkdir(exist_ok=True)

def make_png(label):
    im = Image.new('RGB', (500, 260), 'white'); d = ImageDraw.Draw(im)
    d.rectangle((5,5,495,255), outline='black', width=4); d.text((30,110), label, fill='black')
    b=io.BytesIO(); im.save(b,'PNG'); return b.getvalue()

def add_text(slide,text,x,y,w=3,h=.5,size=12):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); sh.text=text
    for p in sh.text_frame.paragraphs:
        for r in p.runs:r.font.size=Pt(size)
    return sh

def make_weekly(path):
    prs=Presentation(); prs.slide_width=Inches(11); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    s1=prs.slides.add_slide(blank); tb=s1.shapes.add_table(2,6,Inches(.4),Inches(.5),Inches(10),Inches(1.2)).table
    for i,h in enumerate(['이슈','과제명','문제/현상','진행 현황','Signal','기타']):tb.cell(0,i).text=h
    add_text(s1,'000팀 주요 논의 사항',.5,2,5,.5,16)
    s2=prs.slides.add_slide(blank)
    for label,x,y in [('2D',.2,2.2),('3D',3,3),('4D',.3,5),('5D',7,2),('6D',8,5)]:add_text(s2,label,x,y,.55,.45,12)
    add_text(s2,'발생단계 (발생일자)',.5,.5,3,.4,10); add_text(s2,'',3.6,.5,2,.4,10)
    prs.save(path)

def make_source(path):
    prs=Presentation(); prs.slide_width=Inches(11); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    first=prs.slides.add_slide(blank)
    add_text(first,'고객사: Ford',.5,.3,3,.4); add_text(first,'Ford_TestTask',.5,.8,4,.4)
    sections=[('2D','Problem: battery leakage observed'),('3D','Temporary action: quarantine lot and inspect'),('4D','Cause: sealing process variation'),('5D','Action: update process control and fixture'),('6D','Verification: reliability test passed')]
    for label,text in sections:
        sl=prs.slides.add_slide(blank); add_text(sl,label,.5,.4,1,.4,16); add_text(sl,text,.5,1,6,.7,12)
        sl.shapes.add_picture(io.BytesIO(make_png(label+' IMAGE 1')),Inches(7),Inches(1),width=Inches(3),height=Inches(1.8))
        sl.shapes.add_picture(io.BytesIO(make_png(label+' IMAGE 2')),Inches(7),Inches(3.1),width=Inches(3),height=Inches(1.8))
    prs.save(path)

def run_case(mode):
    src=OUT/'source_8d.pptx'; weekly=OUT/f'weekly_{mode}.pptx'; out=OUT/f'output_{mode}.pptx'
    make_source(src); make_weekly(weekly)
    d={'customer':'Ford','issue_name':'Ford_TestIssue','task_name':'TestTask','problem':'Problem: battery leakage observed','temporary_action':'Temporary action: quarantine lot and inspect','customer_response':'Customer response: containment completed','cause_4d':'Cause: sealing process variation','leak_cause':'Leak cause: sealing defect','system_cause':'System cause: process control gap','action_5d':'Action: update process control and fixture','verification_6d':'Verification: reliability test passed','team':'Pack개발품질1','occurrence_stage':'양산','occurrence_date':'2026-09-10','status':'개선 완료','_section_images':{}}
    for label in ['2D','3D','4D','5D','6D']:
        d['_section_images'][label]=[(make_png(label+' TEST 1'),(0,0,1,1),1),(make_png(label+' TEST 2'),(0,0,1,1),1)]
    g={'team':'Pack개발품질1','owner':'홍길동','stage':'양산'}
    result=app.base.weekly(str(weekly),str(out),d,g,mode); assert Path(result[1]).exists()
    prs=Presentation(result[1]); assert len(prs.slides)==2
    p1='\n'.join(sh.text for sh in prs.slides[0].shapes if hasattr(sh,'text'))
    assert 'TestTask' in p1 and 'Ford_TestIssue' in p1 and 'Pack개발품질1팀 주요 논의 사항' in p1
    p2='\n'.join(sh.text for sh in prs.slides[1].shapes if hasattr(sh,'text'))
    for x in ['Problem: battery leakage observed','Temporary action: quarantine lot and inspect','Cause: sealing process variation','Action: update process control and fixture','Verification: reliability test passed']:
        assert x in p2, x
    assert any('발생단계' in sh.text for sh in prs.slides[1].shapes if hasattr(sh,'text'))
    return result[1]

if __name__=='__main__':
    try:
        print('TEST v3.1 synthetic NEW'); print(run_case('신규'))
        print('TEST v3.1 synthetic EXISTING'); print(run_case('기존'))
        print('RESULT: PASS')
    except Exception:
        traceback.print_exc(); print('RESULT: FAIL'); raise
