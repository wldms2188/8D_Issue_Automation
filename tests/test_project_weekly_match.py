import sys,unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
import project_weekly_match_final as w
import main_recovery_step14_fix2 as core
import main_v319 as v319

class ProjectWeeklyMatchTests(unittest.TestCase):
 def test_customer_prefix_is_ignored(self):
  self.assertEqual(w._project_key('MBAG_EB-L(EU)'),w._project_key('Other_EB-L(EU)'))
 def test_eu_us_remain_distinct(self):
  self.assertNotEqual(w._project_key('MBAG_EB-L(EU)'),w._project_key('MBAG_EB-L(US)'))
 def test_no_underscore_project_is_whole(self):
  self.assertEqual(w._project_key('Model Care 25'),'modelcare25')
 def test_underscore_inside_parentheses_is_not_customer_split(self):
  self.assertEqual(w._project_key('JF2S Delta (Set_Biz)'),'jf2sdeltasetbiz')
 def test_customer_project_is_not_prefixed_twice(self):
  d={'customer':'MBAG','task_name':'MBAG_EB-L(EU)'}
  self.assertEqual(w.s13._customer_task(d),'MBAG_EB-L(EU)')
  self.assertEqual(w.s13.v319._weekly_task(d),'MBAG_EB-L(EU)')

 def test_active_base_task_uses_canonical_customer_project(self):
  d={'customer':'MBAG','task_name':'MBAG_EB-L(EU)'}
  self.assertEqual(w.s13.base.task(d),'MBAG_EB-L(EU)')
  d2={'customer':'OLD','task_name':'Ford_V710'}
  self.assertEqual(w.s13._customer_task(d2),'Ford_V710')

 def test_repeated_customer_prefix_is_collapsed(self):
  d={'customer':'MBAG','task_name':'MBAG_MBAG_EB-L(EU)'}
  self.assertEqual(w.s13._customer_task(d),'MBAG_EB-L(EU)')
  self.assertEqual(w.s13.v319._weekly_task(d),'MBAG_EB-L(EU)')

 def test_detail_title_customer_project_is_not_prefixed_twice(self):
  d={'customer':'MBAG','task_name':'MBAG_EB565M','issue_name':'MBAG_EB565M_Scratch'}
  self.assertEqual(w.s13.v319._weekly_task(d),'MBAG_EB565M')
  self.assertEqual(w.s13.step4.step1._page1_issue(d),'Scratch')

 def test_detail_title_repeated_customer_prefix_collapses(self):
  d={'customer':'MBAG','task_name':'MBAG_MBAG_EB565M','issue_name':'MBAG_EB565M_Scratch'}
  self.assertEqual(w.s13.v319._weekly_task(d),'MBAG_EB565M')

 def test_issue_display_removes_8d_and_report_words(self):
  self.assertEqual(v319._clean_issue_label('8D_Report_Crack 발생'),'Crack 발생')

 def test_cloned_detail_shell_removes_old_content_shapes_but_keeps_d_labels(self):
  prs=Presentation(); sl=prs.slides.add_slide(prs.slide_layouts[6])
  label=sl.shapes.add_textbox(Inches(.42),Inches(2.18),Inches(.7),Inches(.3))
  label.text='2D'
  old=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.0),Inches(2.60),Inches(2.0),Inches(.65))
  old.text='OLD ISSUE CONTENT'
  callout=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(3.1),Inches(2.62),Inches(1.5),Inches(.55))
  callout.text=''
  core._clear_cloned_issue_content(sl)
  texts=[str(getattr(sh,'text','') or '') for sh in sl.shapes]
  self.assertIn('2D',texts)
  self.assertNotIn('OLD ISSUE CONTENT',texts)
  self.assertNotIn(old._element,[sh._element for sh in sl.shapes])
  self.assertNotIn(callout._element,[sh._element for sh in sl.shapes])

 def test_create_native_section_for_new_project(self):
  prs=Presentation(); prs.slides.add_slide(prs.slide_layouts[6])
  sec=core._create_native_section(prs,'NEW_PROJECT',0)
  self.assertIsNotNone(sec)
  sections=core._native_sections(prs)
  self.assertTrue(any(x['name']=='NEW_PROJECT' and 0 in x['indices'] for x in sections))

if __name__=='__main__':unittest.main()
