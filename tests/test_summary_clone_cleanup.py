import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
import tempfile
import main_recovery_step14 as s14

def test_new_summary_clone_removes_picture_and_keeps_table():
    prs=Presentation()
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    tb=sl.shapes.add_table(2,5,Inches(.5),Inches(1.5),Inches(8),Inches(1.2)).table
    for i,v in enumerate(['과제명','이슈','현상','진행사항','Signal']): tb.cell(0,i).text=v
    tb.cell(1,0).text='OLD_TASK'
    path=Path(tempfile.gettempdir())/'summary_old_issue.png'
    Image.new('RGB',(60,60),'white').save(path)
    sl.shapes.add_picture(str(path),Inches(1),Inches(3),Inches(5),Inches(2))
    pages=s14._summary_pages(prs)
    idx,tb2,hr,row=s14._prepare_new_summary_page(prs,pages,{},template_index=0)
    cloned=prs.slides[idx]
    assert any(getattr(sh,'has_table',False) for sh in cloned.shapes)
    assert not any(getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.PICTURE for sh in cloned.shapes)
    assert tb2.cell(row,0).text==''
