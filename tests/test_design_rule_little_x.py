"""Regression tests for automatic Issue DB design-rule Little x classification."""
import unittest

from openpyxl import Workbook

import design_rule_little_x as dr


class DesignRuleLittleXTests(unittest.TestCase):
    def test_full_photo_reference_contains_all_36_yyx_and_examples(self):
        text = dr.full_design_rule_reference_text()
        for rule in dr.RULES:
            with self.subTest(code=rule["code"]):
                self.assertIn(f"Y.y.x {rule['code']}", text)

        # Big Y / Little y / photo example text must be visible, not just rule names.
        self.assertIn("팩/모듈 제품 불량 유형 및 규칙", text)
        self.assertIn("Little y 하위 요소/셀 불량 유형", text)
        self.assertIn("Little x 팩 기인 원인 유형", text)
        self.assertIn("사용 중, 셀 실링 손상(누액) 없을 것", text)
        self.assertIn("고온저장 중 구조물 가스켓 간섭에 의한 크랙/벤트", text)
        self.assertIn("접촉면 Air Gap", text)
        self.assertIn("패드 누락, 냉각 유로 막힘", text)
        self.assertIn("TP시험 중 셀 외부 단락 발생", text)
        self.assertIn("접지 연속성 불만족", text)
        self.assertIn("부품/셀 사양 사전 정합화 미흡", text)
        self.assertIn("부품 열화로 인한 고정 후크/클립 파손", text)

    def test_all_36_rule_titles_map_to_their_own_code(self):
        self.assertEqual(len(dr.RULES), 36)
        for rule in dr.RULES:
            with self.subTest(code=rule["code"], name=rule["name"]):
                self.assertIn(rule["code"], dr.classify_little_x(rule["name"]))

    def test_all_36_rules_have_english_alias_and_match(self):
        self.assertEqual(set(dr.ENGLISH_RULE_ALIASES), {r["code"] for r in dr.RULES})
        for rule in dr.RULES:
            aliases = dr.ENGLISH_RULE_ALIASES[rule["code"]]
            examples = list(aliases.get("phrases") or ())
            self.assertTrue(examples, msg=f'No English phrase for {rule["code"]}')
            with self.subTest(code=rule["code"], phrase=examples[0]):
                self.assertIn(rule["code"], dr.classify_little_x(examples[0]))

    def test_2d_yyx_cause_wording_reverse_maps_to_little_y(self):
        cases = {
            "볼트 풀림 발생": "3.3",
            "접촉면 Air Gap 발생": "3.3",
            "접지 누락 확인": "5.3",
            "냉각수 O-ring 누수 발생": "3.1",
            "Bolt loosening was detected.": "3.3",
            "Grounding missing was found.": "5.3",
            "Coolant leakage due to O-ring sealing failure.": "3.1",
        }
        for problem, ycode in cases.items():
            with self.subTest(problem=problem):
                ys = [x["code"] for x in dr.classify_little_y(problem)]
                self.assertIn(ycode, ys)

    def test_2d_yyx_wording_then_4d_cause_recommendation(self):
        rec = dr.recommendation(
            "볼트 풀림 발생",
            "체결력 부족으로 버스바 연결부 풀림",
        )
        self.assertIn("3.3", rec["little_y"])
        self.assertEqual(rec["recommended"], "3.3.1")

    def test_two_stage_korean_2d_to_little_y_then_cause_to_yyx(self):
        rec = dr.recommendation(
            "팩 부품 이탈 및 연결불량 발생",
            "볼트 풀림으로 체결력 부족 발생",
        )
        self.assertEqual(rec["little_y"], ["3.3"])
        self.assertEqual(rec["recommended"], "3.3.1")
        self.assertEqual(rec["candidates"], ["3.3.1"])

    def test_two_stage_contact_failure_selects_332(self):
        rec = dr.recommendation(
            "버스바 연결 불량 및 접촉 불량 발생",
            "접촉면 Air Gap으로 접촉저항 증가",
        )
        self.assertIn("3.3", rec["little_y"])
        self.assertEqual(rec["recommended"], "3.3.2")
        self.assertEqual(rec["candidates"], ["3.3.2"])

    def test_two_stage_english_symptom_and_cause(self):
        rec = dr.recommendation(
            "Component connection failure was detected at the busbar.",
            "Bolt loosening caused insufficient fastening force.",
        )
        self.assertEqual(rec["little_y"], ["3.3"])
        self.assertEqual(rec["recommended"], "3.3.1")

    def test_two_stage_filters_out_cause_from_unrelated_little_y(self):
        rec = dr.recommendation(
            "팩 부품 이탈 발생",
            "Ground continuity failure due to missing grounding connection.",
        )
        self.assertEqual(rec["little_y"], ["3.3"])
        self.assertNotIn("5.3.1", rec["candidates"])
        self.assertEqual(set(rec["candidates"]), {"3.3.1", "3.3.2"})
        self.assertTrue(rec["ambiguous"])

    def test_vague_2d_falls_back_to_cause_across_all_rules(self):
        rec = dr.recommendation(
            "이상 발생",
            "Ground continuity failure due to missing grounding connection.",
        )
        self.assertEqual(rec["little_y"], [])
        self.assertEqual(rec["recommended"], "5.3.1")
        self.assertEqual(rec["candidates"], ["5.3.1"])

    def test_representative_english_8d_causes(self):
        cases = {
            "Bolt loosening caused insufficient fastening force at the busbar joint.": "3.3.1",
            "High contact resistance was caused by an air gap on the contact surface.": "3.3.2",
            "Coolant leakage occurred due to O-ring sealing failure.": "3.1.1",
            "Ground continuity failure due to missing grounding connection.": "5.3.1",
            "Thermal runaway propagation occurred because adjacent cell insulation was insufficient.": "4.1.1",
            "EOL voltage measurement failure was repeatedly detected.": "6.1.3",
            "The warning label was missing because the component specification was not reflected.": "3.6.2",
            "Moisture ingress caused corrosion inside the pack.": "3.4.2",
        }
        for cause, code in cases.items():
            with self.subTest(cause=cause):
                self.assertIn(code, dr.classify_little_x(cause))

    def test_korean_english_mixed_cause_is_supported(self):
        cause = "Busbar 체결부 bolt loosening으로 fastening force 부족 발생"
        self.assertIn("3.3.1", dr.classify_little_x(cause))

    def test_representative_pack_cause_examples(self):
        cases = {
            "가스켓 간섭에 의한 캔 크랙 발생": "1.1.1",
            "레이저 오조사로 실링부 손상": "1.1.3",
            "전극 돌출 확인": "1.1.2",
            "볼트 풀림으로 버스바 체결력 불량": "3.3.1",
            "접촉면 Air Gap 발생으로 접촉 불량": "3.3.2",
            "접지 누락으로 접지 연속성 불만족": "5.3.1",
            "LQC 공정에서 오조립 불량 다발": "6.1.2",
            "EOL 전압 측정 불량 다발": "6.1.3",
        }
        for cause, code in cases.items():
            with self.subTest(cause=cause):
                self.assertIn(code, dr.classify_little_x(cause))

    def test_multiple_matching_rules_are_preserved(self):
        codes = dr.classify_little_x(
            "볼트 풀림으로 체결력 불량 발생, 접촉면 Air Gap으로 접촉 불량 동시 확인"
        )
        self.assertIn("3.3.1", codes)
        self.assertIn("3.3.2", codes)

    def test_unrelated_cause_returns_no_rule(self):
        self.assertEqual(
            dr.classify_little_x("원인 분석 진행 중이며 현재 명확한 원인 미확인"),
            [],
        )

    def test_ambiguous_match_is_recommendation_not_final_multi_write(self):
        rec = dr.recommendation(
            "볼트 풀림으로 체결력 불량 발생, 접촉면 Air Gap으로 접촉 불량 동시 확인"
        )
        self.assertTrue(rec["ambiguous"])
        self.assertIn("3.3.1", rec["candidates"])
        self.assertIn("3.3.2", rec["candidates"])

        wb = Workbook()
        ws = wb.active
        ws.cell(6, 19).value = "규칙 위반"
        ws.cell(6, 20).value = "Y.y.x"
        result = dr.write_confirmed_rule_classification(
            ws, 8, "3.3.2", confirmed=True
        )
        self.assertEqual(result["violation"], "Y")
        self.assertEqual(result["codes"], ["3.3.2"])
        self.assertEqual(ws.cell(8, 19).value, "Y")
        self.assertEqual(ws.cell(8, 20).value, "3.3.2")
        self.assertNotIn("3.3.1", ws.cell(8, 20).value)

    def test_single_match_is_still_only_a_recommendation_until_user_confirms(self):
        rec = dr.recommendation("볼트 풀림으로 체결력 불량 발생")
        self.assertFalse(rec["ambiguous"])
        self.assertEqual(rec["recommended"], "3.3.1")

        wb = Workbook()
        ws = wb.active
        ws.cell(6, 19).value = "규칙 위반"
        ws.cell(6, 20).value = "Y.y.x"
        ws.cell(8, 19).value = "기존값"
        ws.cell(8, 20).value = "기존규칙"
        result = dr.write_confirmed_rule_classification(
            ws, 8, rec["recommended"], confirmed=False
        )
        self.assertIsNone(result["violation"])
        self.assertEqual(ws.cell(8, 19).value, "기존값")
        self.assertEqual(ws.cell(8, 20).value, "기존규칙")

    def test_user_can_confirm_no_violation(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(6, 19).value = "규칙 위반"
        ws.cell(6, 20).value = "Y.y.x"
        result = dr.write_confirmed_rule_classification(
            ws, 8, dr.NONE_CODE, confirmed=True
        )
        self.assertEqual(result["violation"], "N")
        self.assertEqual(result["codes"], [])
        self.assertEqual(ws.cell(8, 19).value, "N")
        self.assertEqual(ws.cell(8, 20).value, "")

    def test_exact_issue_db_headers_are_detected(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(4, 21).value = "규칙 위반"
        ws.cell(4, 22).value = "Y.y.x"
        self.assertEqual(dr.find_rule_columns(ws), (21, 22))

    def test_header_detection_and_y_code_write(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(6, 19).value = "규칙 위반"
        ws.cell(6, 20).value = "Y.y.x"
        result = dr.write_rule_classification(
            ws, 8, "볼트 풀림으로 체결력 불량 발생"
        )
        self.assertEqual((result["violation_col"], result["little_x_col"]), (19, 20))
        self.assertEqual(ws.cell(8, 19).value, "Y")
        self.assertIn("3.3.1", ws.cell(8, 20).value)

    def test_no_match_writes_n_and_blank_little_x(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(5, 22).value = "규칙 위반"
        ws.cell(5, 23).value = "Y.y.x"
        dr.write_rule_classification(ws, 8, "원인 미확인")
        self.assertEqual(ws.cell(8, 22).value, "N")
        self.assertEqual(ws.cell(8, 23).value, "")


if __name__ == "__main__":
    unittest.main()
