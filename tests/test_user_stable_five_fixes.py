import io
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import main_recovery_step12 as s12
import main_recovery_step13 as s13
import main_recovery_step14 as s14
import user_stable_weekly_fixes as fix


def add_summary_slide(prs, task, issue="old", top=0.7, rows=3):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    tb = sl.shapes.add_table(
        rows, 5, Inches(0.5), Inches(top), Inches(9.5), Inches(1.4)
    ).table
    headers = ["과제명", "이슈명", "현상", "진행사항", "Signal"]
    for c, h in enumerate(headers):
        tb.cell(0, c).text = h
    tb.cell(1, 0).text = task
    tb.cell(1, 1).text = issue
    return sl, tb


class UserStableFiveFixesTest(unittest.TestCase):
    def test_detail_title_exact_user_format(self):
        d = {"customer": "GM", "task_name": "GM_MBAG"}
        g = {
            "task_name": "GM_MBAG",
            "sample": "B2",
            "stage": "DV",
            "occurrence_site": "제품 생산",
        }
        self.assertEqual(
            fix._detail_title_user_format(d, g),
            "GM_MBAG_B2 샘플_DV_제품 생산 이슈 발생",
        )

    def test_summary_uses_last_exact_matching_page(self):
        prs = Presentation()
        add_summary_slide(prs, "GM_MBAG", "old-1")
        prs.slides.add_slide(prs.slide_layouts[6])
        add_summary_slide(prs, "GM_MBAG", "old-2")

        d = {
            "customer": "GM",
            "task_name": "GM_MBAG",
            "issue_name": "new",
        }
        g = {"task_name": "GM_MBAG"}

        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm["task"]).text = s13._customer_task(_d)
                tb.cell(row, hm["issue"]).text = "NEW"
            s14._write_summary_row = writer
            si, row, action = s14._update_summary_by_task(
                prs, d, g, "new"
            )
        finally:
            s14._write_summary_row = old_writer

        self.assertEqual(si, 2)
        self.assertIn("마지막", action)
        self.assertEqual(
            prs.slides[2].shapes[0].table.cell(row, 0).text,
            "GM_MBAG",
        )

    def test_confirmed_similar_section_reuses_real_summary_label(self):
        prs = Presentation()
        _, tb = add_summary_slide(prs, "GM_MBAG E~", "old")

        d = {
            "customer": "GM",
            "task_name": "GM_MBAG",
            "issue_name": "new",
        }
        g = {
            "task_name": "GM_MBAG",
            "_weekly_section_override_name": "GM_MBAG E~",
        }

        old_writer = s14._write_summary_row
        try:
            def writer(tb0, row, hr, _d, _g):
                hm = s13._summary_map(tb0, hr)
                tb0.cell(row, hm["task"]).text = s13._customer_task(_d)
                tb0.cell(row, hm["issue"]).text = "NEW"
            s14._write_summary_row = writer
            si, row, _ = s14._update_summary_by_task(
                prs, d, g, "new"
            )
        finally:
            s14._write_summary_row = old_writer

        self.assertEqual(si, 0)
        self.assertEqual(tb.cell(row, 0).text, "GM_MBAG E~")

    def test_similarity_without_confirmed_section_does_not_force_match(self):
        prs = Presentation()
        add_summary_slide(prs, "GM_MBAG E~", "old")

        d = {
            "customer": "GM",
            "task_name": "GM_MBAG",
            "issue_name": "new",
        }
        g = {"task_name": "GM_MBAG"}

        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm["task"]).text = s13._customer_task(_d)
            s14._write_summary_row = writer
            _, _, action = s14._update_summary_by_task(
                prs, d, g, "new"
            )
        finally:
            s14._write_summary_row = old_writer

        self.assertIn("신규 요약 페이지", action)

    def test_summary_overflow_clones_immediately_after_last_match(self):
        prs = Presentation()
        add_summary_slide(prs, "OTHER", "old")
        prs.slides.add_slide(prs.slide_layouts[6])
        add_summary_slide(prs, "GM_MBAG", "old")

        d = {
            "customer": "GM",
            "task_name": "GM_MBAG",
            "issue_name": "new",
        }
        g = {"task_name": "GM_MBAG"}

        old_fit = s14._summary_row_insert_fits
        old_writer = s14._write_summary_row
        try:
            s14._summary_row_insert_fits = lambda *args, **kwargs: False

            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm["task"]).text = s13._customer_task(_d)
                tb.cell(row, hm["issue"]).text = "NEW"

            s14._write_summary_row = writer
            si, _, action = s14._update_summary_by_task(
                prs, d, g, "new"
            )
        finally:
            s14._summary_row_insert_fits = old_fit
            s14._write_summary_row = old_writer

        self.assertEqual(si, 3)
        self.assertIn("바로 다음 페이지", action)

    def test_section_match_levels_are_strict(self):
        d = {"customer": "GM", "task_name": "GM_MBAG"}
        self.assertEqual(fix._section_match_level("GM_MBAG", d), 3)
        self.assertEqual(fix._section_match_level("MBAG", d), 2)
        self.assertEqual(fix._section_match_level("GM_MBAG E~", d), 1)

    def test_clone_visual_cleanup_preserves_table_and_frame(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.shapes.add_table(
            2, 2, Inches(0.5), Inches(0.5), Inches(3), Inches(1)
        )
        sl.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.3),
            Inches(0.3),
            Inches(9),
            Inches(5),
        )

        image = Image.new("RGB", (20, 20), "white")
        bio = io.BytesIO()
        image.save(bio, format="PNG")
        bio.seek(0)
        sl.shapes.add_picture(
            bio, Inches(1), Inches(2), Inches(1), Inches(1)
        )

        before = len(sl.shapes)
        fix._remove_copied_visuals(sl)
        after = len(sl.shapes)

        self.assertEqual(after, before - 1)
        self.assertTrue(
            any(getattr(sh, "has_table", False) for sh in sl.shapes)
        )

    def test_7d_moves_below_6d_absolute_target(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        marker = sl.shapes.add_textbox(
            Inches(0.5), Inches(2.0), Inches(0.5), Inches(0.3)
        )
        marker.text = "7D"
        title = sl.shapes.add_textbox(
            Inches(1.1), Inches(2.0), Inches(1.2), Inches(0.3)
        )
        title.text = "수평전개"

        old_layout = s12.step11._adaptive_cascade_layout
        try:
            s12.step11._adaptive_cascade_layout = lambda d, imgs: (
                {"6D": {"x": 5.5, "y": 4.0, "w": 5.0, "h": 1.0}},
                {},
                {},
            )
            fix._adjust_7d_after_6d(sl, {})
        finally:
            s12.step11._adaptive_cascade_layout = old_layout

        self.assertGreater(float(marker.top) / fix.v310.EMU, 5.0)
        self.assertAlmostEqual(
            float(marker.top) / fix.v310.EMU,
            float(title.top) / fix.v310.EMU,
            places=2,
        )


if __name__ == "__main__":
    unittest.main()
