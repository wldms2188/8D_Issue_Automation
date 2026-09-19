import sys,unittest,tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
import english_8d_final as e

class English8DTests(unittest.TestCase):
 def test_common_fields_translate_without_changing_identifiers(self):
  d={'problem':'Crack occurred at MBAG_EB-L(EU) Sample DUT3, 2.5 mm.',
     'cause_4d':'Insufficient design margin and tolerance interference confirmed.',
     'action_5d':'Corrective action: drawing change for MBAG_EB-L(EU).'}
  x=e.enhance_dict(d)
  self.assertIn('MBAG_EB-L(EU)',x['problem']); self.assertIn('DUT3',x['problem']); self.assertIn('2.5 mm',x['problem'])
  self.assertIn('크랙',x['problem']); self.assertIn('설계마진',x['cause_4d']); self.assertIn('공차',x['cause_4d'])
 def test_korean_is_not_rewritten(self):
  s='체결토크 산포로 조립불량 발생'; self.assertEqual(e.enhance_dict({'cause_4d':s})['cause_4d'],s)
 def test_mixed_preserves_model(self):
  s='Root cause: torque 산포 at BDI_BOLT'; x=e.enhance_dict({'cause_4d':s}); self.assertIn('BDI_BOLT',x['cause_4d'])
 def test_incomplete_translation_is_flagged(self):
  x=e.enhance_dict({'problem':'Unexpected electrical behavior remained after repeated vehicle evaluation.'})
  self.assertTrue(x['_english_translation_incomplete'])
  self.assertEqual(e.apply_english_choice(x,True)['problem'],'Unexpected electrical behavior remained after repeated vehicle evaluation.')
 def test_technical_identifiers_do_not_trigger_failure(self):
  x=e.enhance_dict({'problem':'Crack occurred at MBAG_EB-L(EU) DUT3.'})
  self.assertFalse(x['_english_translation_incomplete'])
 def test_numbered_sections_keep_all_content_and_stop_at_next_d(self):
  blocks=[
   '2D Problem Description\nCrack found at weld\nLeak observed',
   '3D',
   'Containment\nStop shipment\n100% inspection',
   '4D Root Cause\nDesign margin insufficient\nTolerance stack-up',
   '5D Corrective Action\nDrawing changed\nTolerance revised',
   '6D Verification\nDV test passed\nNo recurrence',
   '7D Prevent Recurrence\nUpdate lesson learned',
   '8D Closure\nClosed']
  x=e.extract_sections_from_blocks(blocks)
  self.assertIn('Crack found at weld',x['problem'])
  self.assertIn('Leak observed',x['problem'])
  self.assertIn('Stop shipment',x['temporary_action'])
  self.assertIn('100% inspection',x['temporary_action'])
  self.assertIn('Design margin insufficient',x['cause_4d'])
  self.assertIn('Tolerance stack-up',x['cause_4d'])
  self.assertIn('Drawing changed',x['action_5d'])
  self.assertIn('DV test passed',x['verification_6d'])
  self.assertNotIn('7D',x['verification_6d'])
  self.assertNotIn('Closure',x['verification_6d'])

 def test_unsplit_4d_defaults_to_occurrence_cause(self):
  x=e.extract_sections_from_blocks(['4D','Root cause item A','Root cause item B','5D','Action'])
  self.assertEqual(x['cause_4d'],'Root cause item A\nRoot cause item B')
  self.assertEqual(x['leak_cause'],'')

 def test_explicit_escape_cause_splits_4d(self):
  x=e.extract_sections_from_blocks(['4D Root Cause','Occurrence A','Escape Cause','Detection control missing','5D Corrective Action','Fix'])
  self.assertIn('Occurrence A',x['cause_4d'])
  self.assertIn('Detection control missing',x['leak_cause'])

 def test_english_pending_status_is_open(self):
  self.assertEqual(e.judge_status_bilingual({'verification_6d':'Validation is in progress.'})[0],'open')
  self.assertEqual(e.judge_status_bilingual({'verification_6d':'Verification completed and passed.'})[0],'close')

 def test_supplied_english_layout_boundaries_and_photo_captions(self):
  blocks=[
   '1D','Team build','Model EB565M (E122A)','2D','Problem description','Defect Phenomenon','Scratch on metal strap',
   '[Close-up Photo]','3D','Containment','Stop shipment','4D','Root cause','Guide interference confirmed','Non-defect : N/A',
   '5D','Corrective action','Guide position improved','[Unloading Process]','[After Improvement]',
   '6D','Validation','DV validation passed','7D','Preventive action','Standard update','8D','Follow up','Improvement : N/A']
  x=e.extract_sections_from_blocks(blocks)
  self.assertIn('Scratch on metal strap',x['problem'])
  self.assertNotIn('Team build',x['problem'])
  self.assertNotIn('[Close-up Photo]',x['problem'])
  self.assertIn('Stop shipment',x['temporary_action'])
  self.assertIn('Guide interference confirmed',x['cause_4d'])
  self.assertIn('Guide position improved',x['action_5d'])
  self.assertNotIn('[Unloading Process]',x['action_5d'])
  self.assertNotIn('[After Improvement]',x['action_5d'])
  self.assertIn('DV validation passed',x['verification_6d'])
  self.assertNotIn('Preventive action',x['verification_6d'])
  self.assertNotIn('Follow up',x['verification_6d'])

 def test_varied_english_section_names(self):
  variants=[
   (['Problem Definition','Scratch found','Immediate Action','Stop shipment','Cause Investigation','Guide interference','Countermeasure','Guide revised','Effectiveness Check','Passed DV'],
    {'problem':'Scratch found','temporary_action':'Stop shipment','cause_4d':'Guide interference','action_5d':'Guide revised','verification_6d':'Passed DV'}),
   (['Failure Phenomenon','Dent found','Customer Protection','100% sorting','5 Why','Fixture gap','Permanent Action','Fixture changed','Validation Results','No recurrence'],
    {'problem':'Dent found','temporary_action':'100% sorting','cause_4d':'Fixture gap','action_5d':'Fixture changed','verification_6d':'No recurrence'}),
   (['Issue Description','Leak found','Short-Term Action','Quarantine lot','Failure Cause','Seal damage','Corrective Measure','Seal changed','Effect Confirmation','Test passed'],
    {'problem':'Leak found','temporary_action':'Quarantine lot','cause_4d':'Seal damage','action_5d':'Seal changed','verification_6d':'Test passed'}),
  ]
  for blocks,expected in variants:
   x=e.extract_sections_from_blocks(blocks)
   for key,val in expected.items():self.assertIn(val,x[key],msg=(blocks,key,x))

 def test_semantic_7d_8d_boundaries_stop_6d_without_numbers(self):
  x=e.extract_sections_from_blocks([
   '6D','Validation','DV passed',
   'Customer Response','Horizontal deployment done',
   'Request Items','Customer request A'])
  self.assertIn('DV passed',x['verification_6d'])
  self.assertNotIn('Horizontal deployment done',x['verification_6d'])
  self.assertNotIn('Customer request A',x['verification_6d'])

 def test_authoritative_spatial_fields_clear_legacy_bleed(self):
  d={'problem':'Member Alice\nQuality Bob','verification_6d':'Customer Response\nRequest Items'}
  x=e.enhance_dict(d,section_blocks=['2D','Scratch in strap','6D','DV passed'],
                   authoritative_keys={'problem','verification_6d'})
  self.assertEqual(x['problem'],'Scratch in strap')
  self.assertEqual(x['verification_6d'],'DV passed')

 def test_spatial_cell_mapping_prevents_1d_and_7d_8d_bleed(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'layout.pptx'
   prs=Presentation(); sl=prs.slides.add_slide(prs.slide_layouts[6])
   marker_y=[0.5,2.0,3.5,5.0]
   for idx,y in enumerate(marker_y,1):
    sh=sl.shapes.add_textbox(Inches(0.1),Inches(y),Inches(0.4),Inches(0.3)); sh.text=f'{idx}D'
   for idx,y in enumerate(marker_y,5):
    sh=sl.shapes.add_textbox(Inches(5.1),Inches(y),Inches(0.4),Inches(0.3)); sh.text=f'{idx}D'
   lt=sl.shapes.add_table(4,2,Inches(0.7),Inches(0.5),Inches(4.0),Inches(6.0)).table
   rt=sl.shapes.add_table(4,2,Inches(5.7),Inches(0.5),Inches(4.0),Inches(6.0)).table
   left=[('Member','Alice'),('Defect Phenomenon','Scratch in strap'),('Containment','Stop shipment'),('Root cause','Guide interference')]
   right=[('Corrective action','Guide revised'),('Validation','DV passed'),('Customer Response','Horizontal deployment done'),('Request Items','Customer request A')]
   for r,(a,b) in enumerate(left):
    lt.rows[r].height=Inches(1.5); lt.cell(r,0).text=a; lt.cell(r,1).text=b
   for r,(a,b) in enumerate(right):
    rt.rows[r].height=Inches(1.5); rt.cell(r,0).text=a; rt.cell(r,1).text=b
   prs.save(p)
   blocks=e._spatial_section_blocks(p)
   x=e.extract_sections_from_blocks(blocks)
   self.assertIn('Scratch in strap',x['problem'])
   self.assertNotIn('Alice',x['problem'])
   self.assertNotIn('Member',x['problem'])
   self.assertIn('DV passed',x['verification_6d'])
   self.assertNotIn('Horizontal deployment done',x['verification_6d'])
   self.assertNotIn('Customer request A',x['verification_6d'])

if __name__=='__main__':unittest.main()
