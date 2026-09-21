import sys
import unittest
from pathlib import Path

APP = Path(__file__).resolve().parents[1] / 'app'
if str(APP) not in sys.path:
    sys.path.insert(0, str(APP))

from pptx import Presentation
from pptx.util import Inches

import main_recovery_step14 as s14
import main_recovery_step13 as s13
import continuation_fix_final as fix
import main_recovery_step1 as step1
import main_recovery_step4 as step4
import main_recovery_step14_fix2 as step14fix
import main_recovery_step12 as step12
import main_v310 as v310


HEADERS = ['과제명', '이슈', '현상', '진행현황', 'Signal']


def add_summary_slide(prs, task, issue='old', top=1.0, rows=3):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    shape = sl.shapes.add_table(rows, len(HEADERS), Inches(.4), Inches(top), Inches(12.2), Inches(.75))
    tb = shape.table
    for c, h in enumerate(HEADERS):
        tb.cell(0, c).text = h
    if rows > 1:
        tb.cell(1, 0).text = task
        tb.cell(1, 1).text = issue
    return sl, shape, tb


class ContinuationFixTests(unittest.TestCase):
    def test_detail_title_uses_gui_values_and_space_before_issue(self):
        d = {'customer': 'OLD', 'task_name': 'OLD_TASK'}
        g = {
            'task_name': 'GM_A1',
            'sample': 'B2',
            'stage': 'DV',
            'occurrence_site': '제품 생산',
        }
        self.assertEqual(
            fix._detail_title_from_user_inputs(d, g),
            'GM_A1_B2_DV_제품 생산 이슈 발생',
        )
        self.assertNotIn('_이슈 발생', fix._detail_title_from_user_inputs(d, g))

    def test_detail_title_falls_back_per_field_not_whole_title(self):
        d = {
            'customer': 'GM',
            'task_name': 'A1',
            'sample': 'AUTO_B2',
            'development_stage': 'PD',
            'occurrence_site': '부품 생산',
        }
        g = {
            'task_name': 'GM_A1',
            'sample': '',
            'stage': '',
            'occurrence_site': '',
        }
        self.assertEqual(
            fix._detail_title_from_user_inputs(d, g),
            'GM_A1_AUTO_B2_PD_부품 생산 이슈 발생',
        )

    def test_same_task_is_appended_on_last_matching_summary_page(self):
        prs = Presentation()
        blank = prs.slide_layouts[6]
        # remove default first slide if any is not present; Presentation() starts empty.
        add_summary_slide(prs, 'GM_A1', 'issue-1', top=1.0, rows=3)
        add_summary_slide(prs, 'OTHER_X', 'other', top=1.0, rows=3)
        add_summary_slide(prs, 'GM_A1', 'issue-2', top=1.0, rows=3)

        d = {'customer': 'GM', 'task_name': 'A1', 'issue_name': 'issue-new', 'problem': 'P'}
        g = {}

        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm['task']).text = 'GM_A1'
                tb.cell(row, hm['issue']).text = 'NEW'
            s14._write_summary_row = writer
            si, row, action = s14._update_summary_by_task(prs, d, g, 'new')
        finally:
            s14._write_summary_row = old_writer

        self.assertEqual(si, 2)
        self.assertEqual(prs.slides[2].shapes[0].table.cell(row, 1).text, 'NEW')
        self.assertNotEqual(prs.slides[0].shapes[0].table.cell(2, 1).text, 'NEW')

    def test_full_last_task_page_inserts_continuation_immediately_after_it(self):
        prs = Presentation()
        add_summary_slide(prs, 'GM_A1', 'issue-1', top=1.0, rows=2)
        add_summary_slide(prs, 'OTHER_X', 'other', top=1.0, rows=2)
        # Put the last matching table near the bottom so adding one row cannot fit.
        add_summary_slide(prs, 'GM_A1', 'issue-2', top=7.0, rows=2)
        # A detail/non-summary page after it; continuation must be inserted before this page.
        detail = prs.slides.add_slide(prs.slide_layouts[6])
        detail.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = 'DETAIL_SENTINEL'

        d = {'customer': 'GM', 'task_name': 'A1', 'issue_name': 'issue-new', 'problem': 'P'}
        g = {}

        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm['task']).text = 'GM_A1'
                tb.cell(row, hm['issue']).text = 'NEW'
            s14._write_summary_row = writer
            si, row, action = s14._update_summary_by_task(prs, d, g, 'new')
        finally:
            s14._write_summary_row = old_writer

        self.assertEqual(si, 3)
        self.assertEqual(len(prs.slides), 5)
        # New continuation summary is now index 3; old detail shifted to index 4.
        summary_text = '\n'.join(
            getattr(sh, 'text', '') for sh in prs.slides[3].shapes
        )
        detail_text = '\n'.join(
            getattr(sh, 'text', '') for sh in prs.slides[4].shapes
        )
        self.assertNotIn('DETAIL_SENTINEL', summary_text)
        self.assertIn('DETAIL_SENTINEL', detail_text)

        # Confirm the inserted continuation page contains NEW.
        found_new = False
        for sh in prs.slides[3].shapes:
            if getattr(sh, 'has_table', False):
                tb = sh.table
                for r in range(len(tb.rows)):
                    for c in range(len(tb.columns)):
                        if tb.cell(r, c).text == 'NEW':
                            found_new = True
        self.assertTrue(found_new)

    def test_weekly_prefers_exact_customer_project_then_project_only(self):
        prs = Presentation()
        add_summary_slide(prs, 'OTHER_A1', 'wrong-customer', top=1.0, rows=3)
        add_summary_slide(prs, 'GM_A1', 'exact-customer-project', top=1.0, rows=3)
        add_summary_slide(prs, 'OTHER_A1', 'later-project-only', top=1.0, rows=3)

        d = {'customer': 'OLD', 'task_name': 'OLD', 'issue_name': 'new', 'problem': 'P'}
        g = {'task_name': 'GM_A1'}

        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm['task']).text = 'GM_A1'
                tb.cell(row, hm['issue']).text = 'NEW'
            s14._write_summary_row = writer
            si, row, action = s14._update_summary_by_task(prs, d, g, 'new')
        finally:
            s14._write_summary_row = old_writer

        # Exact GM_A1 must win globally even though OTHER_A1 appears on a later page.
        self.assertEqual(si, 1)
        self.assertEqual(prs.slides[1].shapes[0].table.cell(row, 1).text, 'NEW')

    def test_weekly_falls_back_to_project_only_when_customer_project_absent(self):
        prs = Presentation()
        add_summary_slide(prs, 'OLD_CUSTOMER_A1', 'project-match', top=1.0, rows=3)
        d = {'customer': 'OLD', 'task_name': 'OLD', 'issue_name': 'new', 'problem': 'P'}
        g = {'task_name': 'GM_A1'}

        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm['task']).text = 'GM_A1'
                tb.cell(row, hm['issue']).text = 'NEW'
            s14._write_summary_row = writer
            si, row, action = s14._update_summary_by_task(prs, d, g, 'new')
        finally:
            s14._write_summary_row = old_writer

        self.assertEqual(si, 0)

    def test_weekly_similar_project_uses_existing_summary_label_verbatim(self):
        prs = Presentation()
        add_summary_slide(prs, 'MBAG E~', 'old', top=1.0, rows=3)
        d = {'customer': '', 'task_name': 'MBAG', 'issue_name': 'new', 'problem': 'P'}
        g = {'task_name': 'MBAG', '_weekly_similar_section_confirmed': '1', '_weekly_confirmed_section_name': 'MBAG E~'}

        captured = {}
        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                # The patched customer-task accessor must expose the exact existing label.
                value = s13._customer_task(_d)
                captured['task'] = value
                tb.cell(row, hm['task']).text = value
                tb.cell(row, hm['issue']).text = 'NEW'
            s14._write_summary_row = writer
            si, row, action = s14._update_summary_by_task(prs, d, g, 'new')
        finally:
            s14._write_summary_row = old_writer

        self.assertEqual(si, 0)
        self.assertEqual(captured.get('task'), 'MBAG E~')
        self.assertEqual(prs.slides[0].shapes[0].table.cell(row, 0).text, 'MBAG E~')


    def test_weekly_similarity_is_not_used_without_section_confirmation(self):
        prs = Presentation()
        add_summary_slide(prs, 'MBAG E~', 'old', top=1.0, rows=3)
        d = {'customer': '', 'task_name': 'MBAG', 'issue_name': 'new', 'problem': 'P'}
        g = {'task_name': 'MBAG'}

        # Without explicit "해당 구역 업데이트", MBAG E~ must not be treated
        # as the target project. The updater should create/use the normal new path.
        old_writer = s14._write_summary_row
        try:
            def writer(tb, row, hr, _d, _g):
                hm = s13._summary_map(tb, hr)
                tb.cell(row, hm['task']).text = s13._customer_task(_d)
                tb.cell(row, hm['issue']).text = 'NEW'
            s14._write_summary_row = writer
            si, row, action = s14._update_summary_by_task(prs, d, g, 'new')
        finally:
            s14._write_summary_row = old_writer

        self.assertIn('신규 요약 페이지', action)


    def test_summary_issue_removes_routing_customer_and_project(self):
        d = {
            'customer': 'GM',
            'task_name': 'MBAG',
            'issue_name': 'IF_AA_Module_GM_MBAG_Cell swelling',
        }
        self.assertEqual(step1._page1_issue(d), 'Cell swelling')

    def test_detail_issue_keeps_customer_project_but_removes_routing(self):
        d = {
            'customer': 'GM',
            'task_name': 'MBAG',
            'issue_name': 'IF_ES_GM_MBAG_Cell swelling',
        }
        self.assertEqual(
            step4._strip_selected_project_prefix(d['issue_name'], d),
            'GM_MBAG_Cell swelling',
        )

    def test_blink_jf2_routing_prefix_is_removed_from_both_surfaces(self):
        d = {
            'customer': 'B Link',
            'task_name': 'B Link_JF2',
            'issue_name': 'IF_ES_B Link_JF2_Cell swelling',
        }
        self.assertEqual(step1._page1_issue(d), 'Cell swelling')
        self.assertEqual(
            step4._strip_selected_project_prefix(d['issue_name'], d),
            'JF2_Cell swelling',
        )

    def test_jf2_project_anchor_does_not_depend_on_customer_field(self):
        d = {
            'customer': 'wrong extracted customer',
            'task_name': 'B Link_JF2',
            'issue_name': 'IF_ES_B Link_JF2_voltage drop',
        }
        self.assertEqual(step1._page1_issue(d), 'voltage drop')
        self.assertEqual(
            step4._strip_selected_project_prefix(d['issue_name'], d),
            'JF2_voltage drop',
        )

    def test_section_customer_and_project_match_ignores_separator_style(self):
        d = {'customer': 'GM', 'task_name': 'MBAG'}
        self.assertEqual(step14fix._section_match_level('GM_MBAG', d), 3)
        self.assertEqual(step14fix._section_match_level('GM MBAG', d), 3)
        self.assertEqual(step14fix._section_match_level('GM-MBAG', d), 3)

    def test_section_priority_project_exact_before_similarity(self):
        d = {'customer': 'GM', 'task_name': 'MBAG'}
        self.assertEqual(step14fix._section_match_level('MBAG', d), 2)
        self.assertEqual(step14fix._section_match_level('MBAG E~', d), 1)

    def test_new_section_name_prefers_confirmed_customer_project(self):
        d = {'customer': 'GM', 'task_name': 'MBAG'}
        g = {'task_name': 'GM_MBAG'}
        self.assertEqual(step14fix._new_section_name(d, g), 'GM_MBAG')


    def test_7d_marker_unit_moves_below_6d_without_limiting_6d(self):
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        # Synthetic native 7D marker and its nearby title start too high.
        marker = sl.shapes.add_textbox(Inches(5.59), Inches(6.05), Inches(.35), Inches(.25))
        marker.text = '7D'
        title = sl.shapes.add_textbox(Inches(5.98), Inches(6.05), Inches(1.3), Inches(.25))
        title.text = '수평전개'

        old_layout = step12.step11._adaptive_cascade_layout
        old_render = step12.v319._render
        old_meta = step12.v319._page2_meta
        old_ensure = step12.v315._ensure_4d_units
        try:
            zones = {
                k: dict(v) for k, v in v310.ZONES.items()
            }
            zones['6D']['y'] = 5.90
            zones['6D']['h'] = .95
            texts = {k: '' for k in zones}
            fonts = {k: 8 for k in zones}
            step12.step11._adaptive_cascade_layout = lambda d, imgs: (zones, texts, fonts)
            step12.v319._render = lambda *args, **kwargs: None
            step12.v319._page2_meta = lambda *args, **kwargs: None
            step12.v315._ensure_4d_units = lambda *args, **kwargs: None

            step12._update_page2_step12(sl, {}, {}, 'new')
        finally:
            step12.step11._adaptive_cascade_layout = old_layout
            step12.v319._render = old_render
            step12.v319._page2_meta = old_meta
            step12.v315._ensure_4d_units = old_ensure

        moved, parent = v310.find_marker(sl, '7D')
        self.assertIsNotNone(moved)
        self.assertGreaterEqual(float(moved.top) / v310.EMU, 5.90 + .95 + .22 - .01)


if __name__ == '__main__':
    unittest.main()
