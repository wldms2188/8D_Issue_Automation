import sys,tempfile,unittest,base64
from pathlib import Path
from pptx import Presentation
from openpyxl import Workbook,load_workbook
from openpyxl.drawing.image import Image as XLImage
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
import final_output_polish as f
import output_variant_final as o

class OutputVersionTests(unittest.TestCase):
 def test_version_path_increments_existing_versions(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)
   (p/'weekly_v0.1.pptx').touch(); (p/'weekly_v0.2.pptx').touch()
   self.assertEqual(f._version_path(p/'weekly_업데이트.pptx').name,'weekly_v0.3.pptx')
 def test_latest_version_finds_highest_number(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td); out=p/'자동화_결과'; out.mkdir()
   src=p/'weekly.pptx'; src.touch()
   (out/'weekly_v0.2.pptx').touch(); (out/'weekly_v0.10.pptx').touch()
   self.assertEqual(o._latest_version(src).name,'weekly_v0.10.pptx')
 def test_excel_update_only_keeps_only_inserted_or_changed_row(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td); src=p/'issue.xlsx'; dst=p/'issue_v0.1.xlsx'
   a=Workbook(); ws=a.active; ws.title='Sheet1'
   for r in range(1,7):
    ws.cell(r,1).value=f'HEADER{r}'
   ws.cell(7,10).value='A_PROJECT1'; ws.cell(7,14).value='old1'
   ws.cell(8,10).value='A_PROJECT2'; ws.cell(8,14).value='old2'
   ws.cell(9,10).value='A_PROJECT3'; ws.cell(9,14).value='old3'
   a.save(src)

   b=load_workbook(src); w=b['Sheet1']
   w.insert_rows(8,1)
   w.cell(8,10).value='A_NEWPROJECT'; w.cell(8,14).value='new issue'
   b.save(dst)

   reduced,n=o._excel_update_only(src,dst)
   self.assertEqual(n,1)
   rw=load_workbook(reduced)['Sheet1']
   self.assertEqual(rw.max_row,7)
   self.assertEqual(rw.cell(7,10).value,'A_NEWPROJECT')
   self.assertEqual(rw.cell(7,14).value,'new issue')

 def test_excel_reduced_removes_other_row_images_and_keeps_newest_updated_image(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td); src=p/'issue_img.xlsx'; dst=p/'issue_img_v0.1.xlsx'
   png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2nX0AAAAASUVORK5CYII=')
   img1=p/'i1.png'; img2=p/'i2.png'; img3=p/'i3.png'; imgnew=p/'inew.png'
   for q in (img1,img2,img3,imgnew): q.write_bytes(png)

   wb=Workbook(); ws=wb.active; ws.title='Sheet1'
   for r in range(1,7): ws.cell(r,1).value=f'H{r}'
   for r,name in [(7,'OLD1'),(8,'OLD2'),(9,'OLD3')]:
    ws.cell(r,10).value=name
   ws.add_image(XLImage(str(img1)),'P7')
   ws.add_image(XLImage(str(img2)),'P8')
   ws.add_image(XLImage(str(img3)),'P9')
   wb.save(src)

   full=load_workbook(src); fw=full['Sheet1']
   fw.cell(8,10).value='UPDATED'
   # Simulate full-output behavior that leaves old image and appends the new one.
   fw.add_image(XLImage(str(imgnew)),'P8')
   full.save(dst)

   reduced,n=o._excel_update_only(src,dst,preferred_row=8)
   self.assertEqual(n,1)
   rw=load_workbook(reduced)['Sheet1']
   self.assertEqual(rw.max_row,7)
   self.assertEqual(rw.cell(7,10).value,'UPDATED')
   data_images=[img for img in rw._images if o._anchor_row(img) and o._anchor_row(img)>=7]
   self.assertEqual(len(data_images),1)
   self.assertEqual(o._anchor_row(data_images[0]),7)

 def test_update_only_ignores_unchanged_slides_that_shift_index(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td); src=p/'weekly.pptx'; dst=p/'weekly_v0.1.pptx'
   a=Presentation(); a.slides.add_slide(a.slide_layouts[6]).shapes.add_textbox(0,0,100,100).text='Summary'
   a.slides.add_slide(a.slide_layouts[6]).shapes.add_textbox(0,0,100,100).text='Detail A'
   a.slides.add_slide(a.slide_layouts[6]).shapes.add_textbox(0,0,100,100).text='Detail B'
   a.save(src)
   b=Presentation(src)
   b.slides[0].shapes[0].text='Summary updated'
   # Insert a new changed detail before old details, causing their indices to shift.
   new=b.slides.add_slide(b.slide_layouts[6]); new.shapes.add_textbox(0,0,100,100).text='New detail'
   sid=b.slides._sldIdLst[-1]; b.slides._sldIdLst.remove(sid); b.slides._sldIdLst.insert(1,sid)
   b.save(dst)
   reduced,n=o._ppt_update_only(src,dst)
   self.assertEqual(n,2)
   r=Presentation(reduced)
   texts=[' '.join(sh.text for sh in sl.shapes if hasattr(sh,'text')) for sl in r.slides]
   self.assertEqual(len(texts),2)
   self.assertTrue(any('Summary updated' in x for x in texts))
   self.assertTrue(any('New detail' in x for x in texts))

if __name__=='__main__':unittest.main()
