import sys, unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))

import main_enterprise as ent
import main_recovery_step12 as step12
import main_enterprise_v2 as v2
import main_enterprise_v3 as v3
import main_recovery_step11 as step11
import main_recovery_step4 as step4
import main_v315 as v315
import main_v310 as v310

class UIStateTests(unittest.TestCase):
    def test_excel_safe_payload_removes_controls_recursively(self):
        payload={
            'problem':'Phenomenon\x0bB2',
            'cause_4d':'Root cause\x0cBolt',
            'nested':['A\x00B',('C\x0bD',)],
        }
        clean=ent.excel_safe_payload(payload)
        self.assertEqual(clean['problem'],'PhenomenonB2')
        self.assertEqual(clean['cause_4d'],'Root causeBolt')
        self.assertEqual(clean['nested'][0],'AB')
        self.assertEqual(clean['nested'][1][0],'CD')

    def test_selected_8d_replaces_initial_prompt(self):
        s=ent.target_ready_status(r'C:\\work\\sample.pptx','READY  ·  8D 원본을 선택해 주세요.')
        self.assertIn('8D 원본 선택 완료',s)
        self.assertNotIn('선택해 주세요',s)

    def test_selected_8d_ready_message_is_explicit(self):
        s=ent.target_ready_status('sample.pptx','READY  ·  8D 원본을 선택해 주세요.')
        self.assertEqual(s,'READY  ·  8D 원본 선택 완료 · 미리보기 가능')

    def test_preview_complete_is_not_final_update_complete(self):
        pct,label=v2.progress_state('READY · 8D 추출 완료',15)
        self.assertEqual(pct,25)
        self.assertEqual(label,'8D 추출 완료')

    def test_final_update_complete_is_100(self):
        pct,label=v2.progress_state('COMPLETE · 선택한 자료의 업데이트가 완료되었습니다.',85)
        self.assertEqual((pct,label),(100,'업데이트 완료'))

    def test_english_weekly_native_zones_never_overlap(self):
        long_2d=('Scratch was observed during visual OQC inspection after unloading process. ' * 10).strip()
        long_3d=('Containment action included sorting, reinspection, customer protection and shipment hold. ' * 8).strip()
        d={
            '_english_mode':True,
            'problem':long_2d,
            'temporary_action':long_3d,
            'customer_response':'Customer notified and protected.',
            'cause_4d':'Bolt loosening due to fastening torque variation.',
            'leak_cause':'Inspection control missed the condition.',
            'system_cause':'Control plan linkage gap.',
            'action_5d':'Corrective action implemented.',
            'verification_6d':'No additional abnormalities were observed.',
        }
        zones,texts,fonts=step11._english_template_anchored_layout(d,{})
        self.assertEqual(zones['2D']['y'],v310.ZONES['2D']['y'])
        self.assertEqual(zones['3D']['y'],v310.ZONES['3D']['y'])
        self.assertLessEqual(zones['2D']['y']+zones['2D']['h']+.12,zones['3D']['y']+.001)
        self.assertLessEqual(zones['3D']['y']+zones['3D']['h']+.12,zones['4D_CAUSE']['y']+.001)
        self.assertLessEqual(zones['4D_LEAK']['y']+zones['4D_LEAK']['h']+.12,zones['5D']['y']+.001)
        self.assertLessEqual(zones['5D']['y']+zones['5D']['h']+.12,zones['6D']['y']+.001)
        self.assertLessEqual(zones['6D']['y']+zones['6D']['h'],7.47)
        self.assertGreaterEqual(fonts['2D'],6.0)
        self.assertGreaterEqual(fonts['3D'],6.0)

    def test_initial_window_keeps_status_message_visible(self):
        w,h=v3.EnterpriseAppV3._initial_window_size(1920,1080)
        self.assertEqual((w,h),(1360,950))
        self.assertLess(h,1000)

    def test_weekly_status_follows_agreed_6d_priority(self):
        self.assertEqual(ent.weekly_recommended_status({'verification_6d':'검증 완료'}),'개선 완료')
        self.assertEqual(ent.weekly_recommended_status({'verification_6d':'완료 예정'}),'개선 검증중')
        self.assertEqual(ent.weekly_recommended_status({'verification_6d':'Validation in progress.'}),'개선 검증중')
        self.assertEqual(ent.weekly_recommended_status({'cause_4d':'원인 확인','action_5d':''}),'개선 검증중')
        self.assertEqual(ent.weekly_recommended_status({'cause_4d':'','action_5d':'','verification_6d':''}),'원인/개선 미확인')
        self.assertEqual(ent.weekly_status_from_choice({'verification_6d':'DV abnormal'},'close'),'개선 완료')

    def test_explicit_result_after_progress_label_has_priority(self):
        completed=[
            '진행 : 이상없음',
            '진행 중 : 이상 없음',
            '검증 진행 : 정상',
            '평가 진행 : 문제 없음',
            'Progress: No abnormalities',
            'Validation in progress: No additional abnormalities',
            'Verification ongoing: Result normal',
            'Test progress: All tests passed',
        ]
        for text in completed:
            state,_=ent.weekly_verification_state(text)
            self.assertEqual(state,'complete',msg=text)
            status,_=ent.issue_db_recommended_status({
                'action_5d':'Corrective action completed.',
                'verification_6d':text,
            })
            self.assertEqual(status,'close',msg=text)

    def test_explicit_result_abnormal_or_pending_still_wins(self):
        cases=[
            ('진행 중 : 이상 발생','abnormal'),
            ('검증 진행 : NG','abnormal'),
            ('Progress: Final result failed','abnormal'),
            ('진행 : 추가 검증 예정','pending'),
            ('Progress: Additional validation required','pending'),
        ]
        for text,expected in cases:
            state,_=ent.weekly_verification_state(text)
            self.assertEqual(state,expected,msg=text)

    def test_so_far_no_abnormality_remains_provisional(self):
        provisional=[
            '진행 중 : 현재까지 이상 없음',
            'Validation in progress: No abnormalities so far',
            'Verification ongoing: No abnormality to date',
        ]
        for text in provisional:
            state,_=ent.weekly_verification_state(text)
            self.assertEqual(state,'pending',msg=text)

    def test_no_additional_abnomalities_is_complete(self):
        text='No additional abnomalities were observed after the verification.'
        state,_=ent.weekly_verification_state(text)
        self.assertEqual(state,'complete')
        self.assertEqual(ent.weekly_recommended_status({'verification_6d':text}),'개선 완료')

    def test_project_confirmation_value_uses_filename_only_when_extraction_missing(self):
        self.assertEqual(
            ent.project_confirmation_value('', 'MBAG_EB565M'),
            'MBAG_EB565M'
        )
        self.assertEqual(
            ent.project_confirmation_value('MBAG_EB-L(EU)', 'MBAG_EB565M'),
            'MBAG_EB-L(EU)'
        )
        self.assertEqual(
            ent.project_confirmation_value('', ''),
            '(과제명 추출 못함)'
        )

    def test_filename_catalog_confirmation_exact_match_needs_no_warning(self):
        mismatch,candidates=ent.project_confirmation_mismatch(
            {'customer':'','task_name':''},
            'MBAG_EB565M',
            r'C:\\tmp\\8D_Report_MBAG_EB565M_Scratch.pptx',
            '파우치형Pack개발품질1팀'
        )
        self.assertFalse(mismatch)
        self.assertEqual(candidates,['MBAG_EB565M'])

    def test_filename_catalog_confirmation_shared_word_shows_candidates(self):
        mismatch,candidates=ent.project_confirmation_mismatch(
            {'customer':'','task_name':''},
            'MBAG_EB-L(EU)',
            r'C:\\tmp\\8D_Report_EB-L_issue.pptx',
            '원통형Pack개발품질팀'
        )
        self.assertTrue(mismatch)
        self.assertIn('MBAG_EB-L(EU)',candidates)
        self.assertIn('MBAG_EB-L(US)',candidates)

    def test_project_match_in_8d_filename_returns_literal_match(self):
        self.assertEqual(
            ent.project_match_in_8d_filename(
                r'C:\\tmp\\8D_Report_MBAG_EB565M_Scratch.pptx',
                'MBAG_EB565M'
            ),
            'MBAG_EB565M'
        )
        self.assertEqual(
            ent.project_match_in_8d_filename(
                r'C:\\tmp\\8D_Report_CustomerA_Model Care 25_Issue.pptx',
                'Model Care 25'
            ),
            'Model Care 25'
        )
        self.assertEqual(
            ent.project_match_in_8d_filename(
                r'C:\\tmp\\8D_Report_OtherProject.pptx',
                'MBAG_EB565M'
            ),
            ''
        )

    def test_project_only_selection_compares_only_project_part(self):
        mismatch,extracted,selected=ent.customer_project_mismatch(
            {'customer':'CustomerA','task_name':'Model Care 25'},'Model Care 25'
        )
        self.assertFalse(mismatch)
        self.assertEqual(extracted,'CustomerA_Model Care 25')

    def test_selected_project_warns_when_8d_project_is_missing_or_slightly_different(self):
        mismatch,extracted,selected=ent.customer_project_mismatch(
            {'customer':'','task_name':''},'MBAG_EB565M'
        )
        self.assertTrue(mismatch)
        self.assertEqual(extracted,'')
        self.assertEqual(selected,'MBAG_EB565M')

        mismatch,extracted,selected=ent.customer_project_mismatch(
            {'customer':'MBAG','task_name':'EB-565M'},'MBAG_EB565M'
        )
        self.assertTrue(mismatch)
        self.assertEqual(extracted,'MBAG_EB-565M')

    def test_detail_title_drops_selected_project_prefix_when_owner_overlap_risk(self):
        prs=Presentation(); sl=prs.slides.add_slide(prs.slide_layouts[6])
        title=sl.shapes.add_textbox(Inches(.2),Inches(.10),Inches(2.0),Inches(.42))
        title.text='과제명_이슈 제목'
        owner=sl.shapes.add_textbox(Inches(2.45),Inches(.10),Inches(2.6),Inches(.42))
        owner.text='00팀 담당자 : 이름'
        d={
            'customer':'MBAG',
            'task_name':'MBAG_VERY_LONG_SELECTED_PROJECT_NAME',
            'issue_name':'MBAG_EB565M_Scratch',
            'problem':'Scratch',
        }
        g={'team':'Pack개발품질1팀','owner':'홍길동'}
        step4._force_page2_header(sl,d,g)
        self.assertNotIn('VERY_LONG_SELECTED_PROJECT_NAME',title.text)
        self.assertIn('MBAG_EB565M_Scratch',title.text)

    def test_4d_marker_and_title_are_grouped_and_follow_zone_positions(self):
        prs=Presentation(); sl=prs.slides.add_slide(prs.slide_layouts[6])
        for x,y in ((1.0,1.0),(7.0,1.0)):
            marker=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.42),Inches(.42))
            marker.text='4D'
            title=sl.shapes.add_textbox(Inches(x+.48),Inches(y),Inches(1.2),Inches(.42))
            title.text='원인 분석'
        d={
            'problem':'현상',
            'temporary_action':'임시조치',
            'cause_4d':'발생원인',
            'leak_cause':'유출원인',
            'system_cause':'',
            'action_5d':'개선대책',
            'verification_6d':'검증 완료',
            '_section_images':{},
        }
        step12._update_page2_step12(sl,d,{},'existing')
        units=v315._top_4d_units(sl)
        self.assertEqual(len(units),2)
        self.assertTrue(all(u.shape_type==MSO_SHAPE_TYPE.GROUP for u in units))
        zones,_,_=step11._layout_with_4d_placeholders(d,{})
        expected=[
            (max(.10,zones['4D_CAUSE']['x']-.18),max(.10,zones['4D_CAUSE']['y']-.35)),
            (max(.10,zones['4D_LEAK']['x']-.18),max(.10,zones['4D_LEAK']['y']-.35)),
        ]
        actual=sorted((round(float(u.left)/v310.EMU,2),round(float(u.top)/v310.EMU,2)) for u in units)
        wanted=sorted((round(x,2),round(y,2)) for x,y in expected)
        self.assertEqual(actual,wanted)

    def test_customer_project_mismatch_uses_canonical_8d_value(self):
        d={'customer':'MBAG','task_name':'EB565M'}
        mismatch,extracted,selected=ent.customer_project_mismatch(d,'MBAG_EB565M')
        self.assertFalse(mismatch)
        self.assertEqual(extracted,'MBAG_EB565M')
        mismatch,_,_=ent.customer_project_mismatch(d,'GM_A1')
        self.assertTrue(mismatch)

    def test_weekly_english_completion_variants(self):
        completed=[
            'Verification completed.',
            'Validation is complete.',
            'All tests passed.',
            'No abnormality observed after verification.',
            'No abnomality detected.',
            'No abnormalities found.',
            'No recurrence observed.',
            'Result is within specification.',
            'Effectiveness confirmed.',
        ]
        for text in completed:
            state,_=ent.weekly_verification_state(text)
            self.assertEqual(state,'complete',msg=text)
            self.assertEqual(ent.weekly_recommended_status({'verification_6d':text}),'개선 완료',msg=text)

    def test_issue_db_comprehensive_completed_5d_6d_matrix(self):
        action_done=[
            'Corrective action completed.',
            'Countermeasure implemented.',
            'Change successfully implemented.',
            'Design revision completed.',
            'Process parameter updated.',
            'New control applied.',
            'Improvement action released.',
            '개선 완료 및 적용 완료',
        ]
        verification_done=[
            'No additional abnormalities were observed.',
            'No additional abnomalities were observed.',
            'No further abnormalities detected.',
            'No abnormality occurred after implementation.',
            'No abnormalities found during validation.',
            'No recurrence observed.',
            'No repeat issue was found.',
            'Verification completed and passed.',
            'Validation completed successfully.',
            'Effectiveness confirmed.',
            'All acceptance criteria met.',
            'Result is within specification.',
            'Evaluation completed with normal result.',
            'All tests passed.',
        ]
        total=0
        for a in action_done:
            for v in verification_done:
                status,_=ent.issue_db_recommended_status({'action_5d':a,'verification_6d':v})
                self.assertEqual(status,'close',msg=(a,v))
                total+=1
        self.assertEqual(total,112)

    def test_issue_db_comprehensive_pending_matrix_stays_open(self):
        action_done=[
            'Corrective action completed.',
            'Countermeasure implemented.',
            'Design revision completed.',
            'Process parameter updated.',
        ]
        verification_pending=[
            'Verification in progress.',
            'Validation pending.',
            'Verification is scheduled next week.',
            'To be verified after DV.',
            'Not yet validated.',
            'Additional verification required.',
            'Remaining validation is in progress.',
            'Awaiting verification result.',
            'Verification to follow.',
            'Under validation.',
        ]
        total=0
        for a in action_done:
            for v in verification_pending:
                status,_=ent.issue_db_recommended_status({'action_5d':a,'verification_6d':v})
                self.assertEqual(status,'open',msg=(a,v))
                total+=1
        self.assertEqual(total,40)

    def test_issue_db_comprehensive_abnormal_matrix_stays_open(self):
        abnormal=[
            'Final validation failed.',
            'Abnormality observed during DV.',
            'Defect detected after implementation.',
            'Issue recurred after the action.',
            'Result was NG.',
            'Result is out of spec.',
            'Acceptance criteria not met.',
            'Requirement not met.',
        ]
        for v in abnormal:
            status,_=ent.issue_db_recommended_status({
                'action_5d':'Corrective action completed.',
                'verification_6d':v
            })
            self.assertEqual(status,'open',msg=v)

    def test_no_additional_abnormalities_is_not_misread_as_abnormal(self):
        variants=[
            'No additional abnormalities.',
            'No additional abnormalities were observed.',
            'No additional abnomalities were observed.',
            'No further abnormalities were detected.',
            'No abnormality occurred.',
            'No abnormalities found.',
        ]
        for v in variants:
            state,_=ent.weekly_verification_state(v)
            self.assertEqual(state,'complete',msg=v)
            status,_=ent.issue_db_recommended_status({
                'action_5d':'Corrective action completed.',
                'verification_6d':v
            })
            self.assertEqual(status,'close',msg=v)

    def test_weekly_english_pending_variants(self):
        pending=[
            'Verification in progress.',
            'Validation pending.',
            'Verification is scheduled.',
            'To be completed after DV.',
            'Not completed yet.',
            'Under validation.',
            'Monitoring ongoing.',
            'Awaiting verification result.',
            'Expected to be completed next week.',
        ]
        for text in pending:
            state,_=ent.weekly_verification_state(text)
            self.assertEqual(state,'pending',msg=text)
            self.assertEqual(ent.weekly_recommended_status({'verification_6d':text}),'개선 검증중',msg=text)

    def test_weekly_failed_result_beats_completed_word(self):
        text='Verification completed, but final result failed / abnormal.'
        state,_=ent.weekly_verification_state(text)
        self.assertEqual(state,'abnormal')
        self.assertEqual(ent.weekly_recommended_status({'verification_6d':text}),'개선 검증중')

    def test_issue_db_status_popup_when_5d_or_6d_exists(self):
        self.assertTrue(ent.issue_db_status_confirmation_required({'action_5d':'대책 있음','verification_6d':''}))
        self.assertTrue(ent.issue_db_status_confirmation_required({'action_5d':'','verification_6d':'검증 내용'}))
        self.assertTrue(ent.issue_db_status_confirmation_required({'action_5d':'대책','verification_6d':'검증'}))
        self.assertFalse(ent.issue_db_status_confirmation_required({'action_5d':'','verification_6d':''}))

    def test_weekly_popup_only_for_final_open_issue_db(self):
        self.assertTrue(ent.weekly_status_confirmation_required('open'))
        self.assertFalse(ent.weekly_status_confirmation_required('close'))
        self.assertFalse(ent.weekly_status_confirmation_required(None))

    def test_weekly_selected_status_is_independent_from_issue_db(self):
        d={'action_5d':'fix','verification_6d':'DV abnormal'}
        g={'_issue_status_selected':'open','_weekly_status_selected':'개선 완료'}
        self.assertEqual(step12._signal_status_from_g(d,g),'개선 완료')

    def test_status_reasons_are_separate(self):
        d={'action_5d':'Guide revised','verification_6d':'Validation is in progress.'}
        db_reason=ent.issue_db_status_reason(d,'open')
        weekly_reason=ent.weekly_status_reason(d,'개선 검증중')
        self.assertIn('진행 중',db_reason)
        self.assertIn('진행 중',weekly_reason)

if __name__=='__main__':
    unittest.main()
