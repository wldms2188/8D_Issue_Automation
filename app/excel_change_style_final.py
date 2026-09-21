"""Issue DB change-aware text styling."""
import datetime
from difflib import SequenceMatcher
from openpyxl import load_workbook
import os
import subprocess
from pathlib import Path
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font
import weekly_style_final as weekly

base = weekly.core.base
BLUE = "0033FF"
BLACK = "000000"

# Every field written by write_row must participate in change highlighting.
# Column 2 (최종 수정일) is stored as plain YYYY-MM-DD display text.  This avoids
# Excel applying a time-bearing date format and showing 00:00:00 on first open.
MANAGED_COLS = (2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 16, 17, 18)

def _plain(value): return "" if value is None else str(value)
def _norm(value): return "\n".join(line.rstrip() for line in _plain(value).replace("\r\n","\n").replace("\r","\n").splitlines()).strip()
def _inline_font(cell,color):
    f=cell.font
    return InlineFont(rFont=f.name or "맑은 고딕",sz=f.sz,b=f.b,i=f.i,strike=f.strike,color=color)
def _set_font_color(cell,color):
    old=cell.font
    cell.font=Font(name=old.name,sz=old.sz,b=old.b,i=old.i,vertAlign=old.vertAlign,underline=old.underline,strike=old.strike,color=color,family=old.family,scheme=old.scheme,charset=old.charset,outline=old.outline,shadow=old.shadow,condense=old.condense,extend=old.extend)
def _set_plain_black(cell,text):
    cell.value=text; _set_font_color(cell,BLACK)
def _set_change_rich_text(cell,old_text,new_text):
    old_text=_plain(old_text); new_text=_plain(new_text)
    if not new_text: cell.value=""; return
    if _norm(old_text)==_norm(new_text): _set_plain_black(cell,new_text); return
    black=_inline_font(cell,BLACK); blue=_inline_font(cell,BLUE); rich=CellRichText(); matcher=SequenceMatcher(None,old_text,new_text,autojunk=False)
    for tag,_i1,_i2,j1,j2 in matcher.get_opcodes():
        if j1==j2: continue
        rich.append(TextBlock(black if tag=="equal" else blue,new_text[j1:j2]))
    cell.value=rich if len(rich) else new_text

def _date_text(value):
    """Return only YYYY-MM-DD, regardless of source date/datetime/Excel text."""
    if isinstance(value, datetime.datetime): return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date): return value.strftime("%Y-%m-%d")
    text=_plain(value).strip()
    if not text:return ""
    for fmt in ("%Y-%m-%d %H:%M:%S","%Y-%m-%d","%Y/%m/%d %H:%M:%S","%Y/%m/%d"):
        try:return datetime.datetime.strptime(text,fmt).strftime("%Y-%m-%d")
        except ValueError:pass
    # If Excel/openpyxl supplied an ISO-like value, discard a trailing time only
    # when the leading part is a valid date.
    m=__import__('re').match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:\s+.*)?$",text)
    if m:
        try:return datetime.date(int(m.group(1)),int(m.group(2)),int(m.group(3))).strftime("%Y-%m-%d")
        except ValueError:pass
    return text

def _style_modified_date(cell,old_value,new_value,is_new=False):
    """Store final modified date as plain text so Excel can never render 00:00:00."""
    new_text=_date_text(new_value)
    old_text=_date_text(old_value)
    cell.value=new_text
    cell.number_format="@"
    changed=is_new or old_text!=new_text
    _set_font_color(cell,BLUE if changed else BLACK)

def _find_target_row(src,d,g):
    wb=load_workbook(src,rich_text=True); ws=wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active; row,_=base.find(ws,d,g); return row

_original_update_excel=base.update_excel

def update_excel_change_aware(src,out,d,g,new=False):
    # Preserve the long-stable change-highlighting path.  Do not introduce a new
    # Excel COM save here; today's WMF symptom is upstream in openpyxl's drawing
    # reader, not evidence that the mature blue-text pass caused it.
    target_row=None; before={}
    if not new:
        target_row=_find_target_row(src,d,g)
        if target_row:
            wb0=load_workbook(src,rich_text=True); ws0=wb0["Sheet1"] if "Sheet1" in wb0.sheetnames else wb0.active
            before={c:ws0.cell(target_row,c).value for c in MANAGED_COLS}
    msg,saved=_original_update_excel(src,out,d,g,new=new)
    wb=load_workbook(saved,rich_text=True); ws=wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
    if new:
        row,_=base.find(ws,d,g)
        if row:
            for c in MANAGED_COLS:
                cell=ws.cell(row,c)
                if c==2:
                    _style_modified_date(cell,None,cell.value,is_new=True)
                    continue
                text=_plain(cell.value)
                if text: cell.value=CellRichText([TextBlock(_inline_font(cell,BLUE),text)])
    elif target_row:
        for c in MANAGED_COLS:
            cell=ws.cell(target_row,c)
            if c==2:
                _style_modified_date(cell,before.get(c),cell.value)
            else:
                _set_change_rich_text(cell,before.get(c,""),_plain(cell.value))
    wb.save(saved)
    return msg+(" / 신규 입력 내용 파란색 표시" if new else " / 변경 내용 파란색 표시"),saved

base.update_excel=update_excel_change_aware
