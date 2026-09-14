import os
import copy
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook
from pptx.util import Pt

import main_recovery_step7 as step7
import main_v310 as v310

base=step7.base
N=v310.N

ORIGINS=('부품','설계','공정','기타','논의 중')


def _cause_text(d):
    parts=[]
    for label,key in [('발생원인','cause_4d'),('유출원인','leak_cause'),('시스템원인','system_cause')]:
        s=N(d.get(key))
        if s:
            parts.append(f'{label} : {s}')
    return '\n'.join(parts)


def _recommend_origin(d):
    text=_cause_text(d)
    if not text:
        return 'TBD','4D 원인 내용이 아직 없어 TBD로 표시합니다.'

    q=text.lower().replace(' ','')
    keys={
        '부품':('부품','자재','소재','원재료','supplier','협력사','입고','component','cell','셀불량'),
        '설계':('설계','design','도면','공차','사양','spec','구조','치수','강성','간섭','설계마진'),
        '공정':('공정','작업','조립','체결','토크','용접','검사','검출','설비','조건관리','작업자','생산','가공','도포','압착','공정조건'),
        '기타':('운송','보관','취급','고객','외부','환경','사용조건'),
    }
    scores={cat:sum(1 for w in words if w.lower().replace(' ','') in q) for cat,words in keys.items()}
    best=max(scores,key=scores.get)
    if scores[best]==0:
        return '논의 중','원인 내용은 있으나 부품/설계/공정/기타로 명확히 분류할 근거가 부족합니다.'

    tied=[k for k,v in scores.items() if v==scores[best]]
    if len(tied)>1:
        return '논의 중','원인 내용에 여러 이슈기인 범주의 단서가 함께 있어 추가 논의가 필요합니다.'

    evidence=[w for w in keys[best] if w.lower().replace(' ','') in q][:3]
    return best, f"8D 원인 내용에서 {', '.join(evidence)} 관련 표현이 확인되어 {best}을(를) 추천합니다."


class OriginDialog(tk.Toplevel):
    def __init__(self,parent,d):
        super().__init__(parent)
        self.result=None
        self.title('이슈기인 확인')
        self.geometry('760x590')
        self.transient(parent)
        self.grab_set()

        cause=_cause_text(d)
        rec,reason=_recommend_origin(d)

        ttk.Label(self,text='8D 원인 내용을 기준으로 이슈기인을 추천했습니다.',font=('',10,'bold')).pack(anchor='w',padx=20,pady=(18,8))
        ttk.Label(self,text='8D 원인 내용').pack(anchor='w',padx=20)
        box=tk.Text(self,height=12,wrap='word')
        box.pack(fill='both',expand=True,padx=20,pady=(4,10))
        box.insert('1.0',cause or '(4D 원인 내용 없음)')
        box.configure(state='disabled')

        ttk.Label(self,text=f'추천 이슈기인 : {rec}',font=('',11,'bold')).pack(anchor='w',padx=20,pady=(4,2))
        ttk.Label(self,text='추천 이유 : '+reason,wraplength=700,justify='left').pack(anchor='w',padx=20,pady=(0,10))

        self.var=tk.StringVar(value=rec)
        if rec=='TBD':
            ttk.Label(self,text='원인이 없으므로 이슈기인은 TBD로 반영됩니다.').pack(anchor='w',padx=20,pady=5)
        else:
            row=ttk.Frame(self)
            row.pack(anchor='w',padx=20,pady=5)
            for value in ORIGINS:
                label=value + ('  ← 추천' if value==rec else '')
                ttk.Radiobutton(row,text=label,variable=self.var,value=value).pack(side='left',padx=(0,12))

        buttons=ttk.Frame(self)
        buttons.pack(pady=16)
        ttk.Button(buttons,text='확인',width=16,command=self.ok).pack(side='left',padx=6)
        ttk.Button(buttons,text='취소',width=16,command=self.cancel).pack(side='left',padx=6)
        self.protocol('WM_DELETE_WINDOW',self.cancel)
        self.wait_visibility()
        self.focus_set()
        self.wait_window(self)

    def ok(self):
        self.result=self.var.get()
        self.destroy()

    def cancel(self):
        self.result=None
        self.destroy()


_old_weekly=base.weekly


def _first_run(cell):
    try:
        for p in cell.text_frame.paragraphs:
            if p.runs:
                return p.runs[0]
    except Exception:
        pass
    return None


def _reference_value_cell(tb,origin_row,origin_col):
    # Use the value cell from '발생단계' first because it is the closest normal
    # metadata value style. Avoid Signal because its bullet/color is special.
    for r in range(len(tb.rows)):
        for c in range(len(tb.columns)):
            if base.compact(tb.cell(r,c).text)=='발생단계' and c+1<len(tb.columns):
                ref=tb.cell(r,c+1)
                if _first_run(ref) is not None:
                    return ref
    # Fallback: another non-empty value cell in the same metadata table.
    for r in range(len(tb.rows)):
        for c in range(1,len(tb.columns)):
            if r==origin_row and c==origin_col:
                continue
            ref=tb.cell(r,c)
            if N(ref.text) and _first_run(ref) is not None and base.compact(ref.text)!='●':
                return ref
    return None


def _set_cell_like_reference(cell,text,reference=None,size=9):
    # Keep the target cell itself (fill/border/margins). Only replace text and
    # copy the surrounding metadata text style, then force the requested 9 pt.
    tf=cell.text_frame
    old_run=_first_run(cell)
    old_rpr=copy.deepcopy(old_run._r.get_or_add_rPr()) if old_run is not None else None

    ref_run=_first_run(reference) if reference is not None else None
    ref_rpr=copy.deepcopy(ref_run._r.get_or_add_rPr()) if ref_run is not None else None
    ref_p=reference.text_frame.paragraphs[0] if reference is not None and reference.text_frame.paragraphs else None

    cell.text=N(text)
    p=tf.paragraphs[0]
    run=p.runs[0] if p.runs else p.add_run()
    if not run.text:
        run.text=N(text)

    # Prefer the normal surrounding value-cell style; otherwise retain target style.
    source_rpr=ref_rpr if ref_rpr is not None else old_rpr
    if source_rpr is not None:
        try:
            current=run._r.rPr
            if current is not None:
                run._r.remove(current)
            run._r.insert(0,copy.deepcopy(source_rpr))
        except Exception:
            pass

    if ref_p is not None:
        try:
            p.alignment=ref_p.alignment
        except Exception:
            pass
        try:
            tf.vertical_anchor=reference.text_frame.vertical_anchor
        except Exception:
            pass

    # User-requested fixed size.
    for para in tf.paragraphs:
        for r in para.runs:
            r.font.size=Pt(size)


def _set_origin_in_ppt(path,origin):
    from pptx import Presentation
    prs=Presentation(path)
    if len(prs.slides)<2:
        return False

    sl=prs.slides[1]
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False):
            continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                if base.compact(tb.cell(r,c).text)=='이슈기인' and c+1<len(tb.columns):
                    target=tb.cell(r,c+1)
                    reference=_reference_value_cell(tb,r,c+1)
                    _set_cell_like_reference(target,origin,reference,9)
                    prs.save(path)
                    return True
    return False


def weekly_step8(src,out,d,g,mode):
    msg,saved=_old_weekly(src,out,d,g,mode)
    origin=N(g.get('_issue_origin_selected'))
    if origin:
        _set_origin_in_ppt(saved,origin)
    return msg,saved


base.weekly=weekly_step8


class RecoveryStep8App(step7.RecoveryStep7App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP8 FIX2')

    def run(self):
        g=self.gui()
        if any(not g[k] or not os.path.exists(g[k]) for k in ['ppt8d','pptweekly','xlsx']):
            return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')

        try:
            d=base.extract(g['ppt8d'])
            mode=self.mode.get()

            od=OriginDialog(self,d)
            if od.result is None:
                self.log.delete('1.0','end')
                self.log.insert('end','업데이트가 취소되었습니다.')
                return
            g['_issue_origin_selected']=od.result

            summary=step7._db_summary(d)
            pd=step7.ProblemChoiceDialog(self,summary)
            if pd.result is None:
                self.log.delete('1.0','end')
                self.log.insert('end','업데이트가 취소되었습니다.')
                return
            g['_db_problem_selected']=summary if pd.result=='summary' else N(d.get('problem'))

            out=Path(g['xlsx']).parent/'자동화_결과'
            out.mkdir(exist_ok=True)
            xo=out/(Path(g['xlsx']).stem+'_업데이트.xlsx')
            po=out/(Path(g['pptweekly']).stem+'_업데이트.pptx')

            wb=load_workbook(g['xlsx'])
            ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active
            r,_=base.find(ws,d,g)
            if mode=='existing' and not r:
                return messagebox.showwarning('업데이트 중단','기존 이슈를 Excel에서 정확하게 찾지 못했습니다. PMS/PLM 이슈번호를 입력하거나 이슈 정보를 확인하세요.')
            if mode=='new' and r and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?'):
                return

            a,_=base.update_excel(g['xlsx'],xo,d,g,new=(mode=='new'))
            b,saved_ppt=base.weekly(g['pptweekly'],po,d,g,mode)

            phenomenon_label='현상 요약' if pd.result=='summary' else '전체 내용'
            self.log.delete('1.0','end')
            self.log.insert('end',
                f'=== 완료 ===\n'
                f'이슈 구분: {mode}\n'
                f'Issue DB 현상: {phenomenon_label}\n'
                f'이슈기인: {od.result}\n'
                f'상태: {base.status(d)}\n\n'
                f'{a}\n{b}\n\n'
                f'결과: {out}'
            )
            messagebox.showinfo('완료',f'자동 업데이트가 완료되었습니다.\n\n이슈기인: {od.result}\n\n{out}')

        except Exception as e:
            messagebox.showerror('실행 오류',repr(e))


if __name__=='__main__':
    RecoveryStep8App().mainloop()
