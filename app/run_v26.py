"""v2.6 mapping wrapper.
The numbered areas in the supplied example templates are development annotations only.
Real files do not need those numbers. This wrapper patches the v2.5 app with semantic
2D~7D extraction, Excel header mapping, and reference-region mapping for the weekly PPT.
"""
import re, copy, datetime
from pathlib import Path
import main
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Pt

MAP={
 '1':('2D 현상',['현상','문제 현황','불량 현상']),
 '2':('3D 임시대책',['임시대책','임시 대책','임시 조치']),
 '3-1':('4D 발생원인',['원인분석','원인 분석','발생 원인']),
 '3-2':('4D 유출원인',['유출 원인','검출 원인','escape']),
 '4':('5D 개선대책',['개선대책','개선 대책','영구 개선']),
 '5':('6D 유효성검증',['유효성점검','유효성 검증','효과 검증','효과성 검증']),
 '6':('7D 수평전개',['수평전개','수평 전개','재발방지','재발 방지'])}
WEEKLY_GEOM={'1':(0.53,2.53,4.76,0.98),'2':(0.53,3.91,4.76,1.14),'3-1':(0.53,5.38,4.76,1.79),'3-2':(5.77,2.47,4.76,0.87),'4':(5.77,3.74,4.76,1.77),'5':(5.77,5.91,4.78,0.94)}

def norm(x): return re.sub(r'\n+','\n',re.sub(r'[ \t]+',' ',str(x or '').replace('\r','\n'))).strip()
def one(x,n=9999):
 s=re.sub(r'\s+',' ',norm(x)); return s if len(s)<=n else s[:n-1]+'…'
def ck(x): return re.sub(r'[^0-9a-z가-힣]','',norm(x).lower())
def blank(x): return not norm(x) or norm(x) in {'000','00','’00.00.00','담당자명','수동 기입','수동기입'}

def split_sections(s):
 s=norm(s); hits=[]
 patterns=[('2d',r'(?<!\d)2\s*D\s*(?:현상|문제\s*현황)?\s*[:：]?'),('3d',r'(?<!\d)3\s*D\s*(?:임시\s*대책|임시\s*조치)?\s*[:：]?'),('4d',r'(?<!\d)4\s*D\s*(?:원인\s*분석|원인\s*파악)?\s*[:：]?'),('5d',r'(?<!\d)5\s*D\s*(?:개선\s*대책)?\s*[:：]?'),('6d',r'(?<!\d)6\s*D\s*(?:유효성\s*(?:점검|검증)|효과\s*검증)?\s*[:：]?'),('7d',r'(?<!\d)7\s*D\s*(?:수평\s*전개|재발\s*방지)?\s*[:：]?')]
 for k,p in patterns:
  for m in re.finditer(p,s,re.I): hits.append((m.start(),m.end(),k))
 hits.sort(); out={}
 for i,(a,b,k) in enumerate(hits):
  e=hits[i+1][0] if i+1<len(hits) else len(s); v=s[b:e].strip(' :-\n')
  if v: out[k]=v
 return out

def table_value(sl,labels):
 labs=[ck(x) for x in labels]
 for sh in sl.shapes:
  if not getattr(sh,'has_table',False): continue
  tb=sh.table
  for r,row in enumerate(tb.rows):
   for c,cell in enumerate(row.cells):
    raw=norm(cell.text); k=ck(raw)
    if not k or not any(l and (l in k or k in l) for l in labs): continue
    vals=[]
    if ':' in raw or '：' in raw: vals.append(re.split(r'[:：]',raw,1)[1].strip())
    vals += [norm(tb.cell(r,cc).text) for cc in range(c+1,len(tb.columns))]
    vals += [norm(tb.cell(rr,c).text) for rr in range(r+1,min(r+3,len(tb.rows)))]
    for v in vals:
     if v and ck(v) not in labs and not blank(v): return v
 return ''

def extract(path):
 prs=Presentation(path); d={k:'' for k in ['issue_name','task_name','customer','occurrence_site','occurrence_date','receipt_date','problem','temporary_action','customer_response','cause_4d','leak_cause','action_5d','verification_6d','spread_7d']}; all=[]; combined={}
 for sl in prs.slides:
  for sh in sl.shapes:
   if sh.shape_type==MSO_SHAPE_TYPE.GROUP: shapes=list(sh.shapes)
   else: shapes=[sh]
   for x in shapes:
    t=norm(getattr(x,'text',''))
    if t: all.append(t); combined.update(split_sections(t))
   pairs=[(['과제명'],'task_name'),(['Customer','고객사'],'customer'),(['발생 Site'],'occurrence_site'),(['발생 일자'],'occurrence_date'),(['접수 일자'],'receipt_date'),(['불량 현상','문제 현황'],'problem'),(['임시 조치','임시대책'],'temporary_action'),(['고객 대응'],'customer_response'),(['발생 원인'],'cause_4d'),(['유출 원인'],'leak_cause'),(['개선 대책','대책'],'action_5d'),(['유효성 검증','유효성점검','효과 검증','효과검증'],'verification_6d'),(['수평전개','수평 전개','재발방지'],'spread_7d')]
   for labels,k in pairs:
    if not d[k]:
     v=table_value(sl,labels)
     if v:d[k]=v
 full='\n'.join(all)
 m=re.search(r'이슈명\s*[:：]\s*(.+)',full)
 if m:d['issue_name']=m.group(1).strip()
 for k,src in [('problem','2d'),('temporary_action','3d'),('cause_4d','4d'),('action_5d','5d'),('verification_6d','6d'),('spread_7d','7d')]:
  if not d[k] and src in combined:d[k]=combined[src]
 if d['cause_4d'] and '유출 원인' in d['cause_4d'] and not d['leak_cause']:
  m=re.search(r'유출\s*원인\s*[:：]?',d['cause_4d'])
  if m:d['leak_cause']=d['cause_4d'][m.end():].strip(' :-');d['cause_4d']=d['cause_4d'][:m.start()].strip(' :-')
 return d

def status(d):
 v=one(d.get('verification_6d'))
 if re.search(r'완료\s*(예정|추정)',v):return '개선 검증중'
 if re.search(r'검증\s*완료|(?<![가-힣])완료(?![가-힣])',v):return '개선 완료'
 if re.search(r'진행\s*중|진행중|검증\s*중|검증중|예정|추정',v):return '개선 검증중'
 if blank(d.get('cause_4d')) and blank(d.get('action_5d')):return '원인/개선 미확인'
 return '개선 검증중'

def header_map(ws):
 for r in range(1,min(15,ws.max_row)+1):
  mp={ck(ws.cell(r,c).value):c for c in range(1,ws.max_column+1) if ck(ws.cell(r,c).value)}
  if any('최종수정일' in k for k in mp) and any('이슈상태' in k for k in mp):return r,mp
 return 6,{ck(ws.cell(6,c).value):c for c in range(1,ws.max_column+1)}
def col(mp,*names):
 for n in names:
  q=ck(n)
  for k,c in mp.items():
   if q==k or q in k or k in q:return c
 return None

def update_excel(src,out,d,g,new=False):
 wb=main.load_workbook(src);ws=wb.active;hr,mp=header_map(ws)
 c={k:col(mp,*v) for k,v in {'date':['최종 수정일'],'plm':['PLM 이슈 번호'],'form':['폼 팩터'],'ptype':['제품 타입'],'year':['발생 연도'],'month':['발생 월'],'team':['담당팀'],'owner':['담당자'],'task':['고객사 과제명'],'sample':['발생 샘플'],'stage':['개발 단계'],'site':['발생처'],'problem':['이슈 현상 요약'],'cause':['팩 기인 원인','4M 근본 원인'],'action':['적용된 영구 개선 대책'],'status':['이슈 상태']}.items()}
 def val(r,k):return ck(ws.cell(r,c[k]).value) if c[k] else ''
 match=None;plm=ck(g.get('plm_no'))
 if plm and c['plm']:
  for r in range(hr+1,ws.max_row+1):
   if val(r,'plm')==plm:match=r;break
 if match is None and not new:
  target=ck(' / '.join(x for x in [d.get('customer'),d.get('task_name')] if x))
  for r in range(hr+1,ws.max_row+1):
   if target and target==val(r,'task'):match=r;break
 if match is None:
  match=ws.max_row+1
  if match>hr+1:
   ws.row_dimensions[match].height=ws.row_dimensions[match-1].height
   for cc in range(1,ws.max_column+1):ws.cell(match,cc)._style=copy.copy(ws.cell(match-1,cc)._style)
  msg='신규 이슈 추가'
 else:msg=f'기존 이슈 업데이트 (row {match})'
 y,m=main.parse_year_month(d.get('occurrence_date'));task=' / '.join(x for x in [one(d.get('customer')),one(d.get('task_name'))] if x)
 vals={'date':datetime.date.today(),'plm':g.get('plm_no',''),'form':g.get('form_factor',''),'ptype':g.get('product_type',''),'year':y,'month':m,'team':g.get('team',''),'owner':g.get('owner',''),'task':task,'sample':g.get('sample',''),'stage':g.get('stage',''),'site':d.get('occurrence_site',''),'problem':one(d.get('problem'),180),'cause':one(' / '.join(x for x in [d.get('cause_4d'),d.get('leak_cause')] if x),220),'action':one(d.get('action_5d'),220),'status':status(d)}
 for k,v in vals.items():
  if c[k]:ws.cell(match,c[k]).value=v
 Path(out).parent.mkdir(parents=True,exist_ok=True)
 try:wb.save(out);saved=out
 except PermissionError:
  p=Path(out);saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix));wb.save(saved)
 return msg,saved

def recursive(sl):
 out=[]
 def rec(sh):
  if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
   for x in sh.shapes:rec(x)
  else:out.append(sh)
 for sh in sl.shapes:rec(sh)
 return out

def settext(sh,text):
 if not hasattr(sh,'text_frame'):return
 sh.text_frame.clear();p=sh.text_frame.paragraphs[0];p.text=text
 for r in p.runs:r.font.name='맑은 고딕';r.font.size=Pt(8)

def update_weekly(src,out,d,g):
 prs=Presentation(src);st=status(d);title=' / '.join(x for x in [one(d.get('customer')),one(d.get('task_name'))] if x);issue=one(d.get('issue_name') or title,100);done=[]
 contents={'1':one(d.get('problem'),900) or '검토 중','2':one(d.get('temporary_action'),700) or '검토 중','3-1':one(d.get('cause_4d'),900) or '검토 중','3-2':one(d.get('leak_cause'),700) or '검토 중','4':one(d.get('action_5d'),900) or '검토 중','5':one(d.get('verification_6d'),700) or '검토 중','6':one(d.get('spread_7d'),900) or '검토 중'}
 for sl in prs.slides:
  for sh in recursive(sl):
   if not hasattr(sh,'text_frame'):continue
   old=norm(sh.text);new=old
   if '과제명_이슈 제목' in old:new=f'{title}_{issue}'
   elif old.startswith('이슈명') and ':' in old:new='이슈명 : '+issue
   elif '00팀 담당자' in old:new=f"{g.get('team','')} 담당자 : {g.get('owner','')}"
   if new!=old:settext(sh,new)
  ann={norm(sh.text):sh for sh in recursive(sl) if norm(getattr(sh,'text','')) in MAP}
  if ann:
   for aid,sh in ann.items():
    x,y,w,h=sh.left,sh.top,sh.width,sh.height;sl.shapes._spTree.remove(sh._element);tb=sl.shapes.add_textbox(x,y,w,h);settext(tb,contents[aid])
   done.append('번호 지정 영역')
  elif len(prs.slides)>1 and sl==prs.slides[1]:
   used=[]
   for aid,(x,y,w,h) in WEEKLY_GEOM.items():
    rx,ry,rw,rh=x*914400,y*914400,w*914400,h*914400;best=None
    for sh in recursive(sl):
     if sh in used or sh.shape_type not in (1,17) or norm(getattr(sh,'text','')):continue
     dist=abs(sh.left-rx)+abs(sh.top-ry)+abs(sh.width-rw)+abs(sh.height-rh)
     if abs(sh.width-rw)/max(rw,1)<.5 and abs(sh.height-rh)/max(rh,1)<.65 and dist<1.5*914400:
      if best is None or dist<best[0]:best=(dist,sh)
    if best:settext(best[1],contents[aid]);used.append(best[1]);done.append(aid)
  for sh in recursive(sl):
   if not getattr(sh,'has_table',False):continue
   tb=sh.table
   for r in range(len(tb.rows)):
    heads=[ck(tb.cell(r,c).text) for c in range(len(tb.columns))]
    if any('과제명' in h for h in heads) and any('현상' in h for h in heads) and r+1<len(tb.rows):
     for c,h in enumerate(heads):
      if '과제명' in h:tb.cell(r+1,c).text=one(title,80)
      elif h=='이슈':tb.cell(r+1,c).text=issue
      elif '현상' in h:tb.cell(r+1,c).text=one(d.get('problem'),250)
      elif '진행사항' in h:tb.cell(r+1,c).text='- 원인 : '+one(d.get('cause_4d') or d.get('leak_cause'),160)+'\n- 진행 현황 : '+one(' / '.join(x for x in [d.get('action_5d'),d.get('verification_6d')] if x),220)
      elif h=='signal':
       cl=tb.cell(r+1,c);cl.text='●';cl.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(*main.SIGNAL[st])
 Path(out).parent.mkdir(parents=True,exist_ok=True)
 try:prs.save(out);saved=out
 except PermissionError:
  p=Path(out);saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix));prs.save(saved)
 return done,saved

main.extract=extract
main.status=status
main.update_excel=update_excel
main.update_weekly=update_weekly
if __name__=='__main__':main.App().mainloop()
