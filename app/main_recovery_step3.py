# Recovery STEP3: text-only fixes on top of STEP2.
# Preserve v3.2.0 image/layout/font/Excel behavior.
import re, datetime
from pathlib import Path
from pptx import Presentation

import main_recovery_step2 as step2
import main_recovery_step1 as step1
import main_v320 as v320
import main_v319 as v319
import main_v310 as v310

base=step2.base
N=v310.N; C=v310.C


def _event_name(d):
    """Test token is the complete underscore sequence after customer/task through the token ending in 시험.
    Example customer_task_00_00시험 중_x -> 00_00시험.
    """
    issue=N(d.get('issue_name'))
    customer=N(d.get('customer')); task=N(d.get('task_name'))
    toks=[x.strip() for x in issue.split('_') if x.strip()]
    start=0

    # Current GUI task_name may already be canonical CUSTOMER_PROJECT. Strip that
    # full prefix first; the older V1 logic assumed customer and task were separate.
    canonical=v319._weekly_task(d)
    ptoks=[x.strip() for x in canonical.split('_') if x.strip()]
    if ptoks and len(toks)>=len(ptoks) and all(C(toks[i])==C(ptoks[i]) for i in range(len(ptoks))):
        start=len(ptoks)
    elif customer and task:
        for i in range(len(toks)-1):
            if C(toks[i])==C(customer) and C(toks[i+1])==C(task):
                start=i+2
                break
    tail=toks[start:]
    parts=[]
    for tok in tail:
        # Normalize '00시험 중' / '00시험중' to token ending exactly in 시험.
        m=re.match(r'^(.*?시험)\s*(?:중|진행\s*중|완료)?(?:\s+.*)?$',tok,re.I)
        if m:
            final=m.group(1).strip(' _-')
            if final:
                parts.append(final)
                return '시험','_'.join(parts)
        # Stop carrying obvious post-event failure detail only after no test marker yet.
        parts.append(tok.strip())
    # fallback search across full issue/problem, preserving underscores immediately before 시험
    for src in (issue,N(d.get('problem'))):
        m=re.search(r'((?:[^_\s]+_)*[^_\s]*시험)\s*(?:중|진행\s*중|완료)?',src,re.I)
        if m:
            name=m.group(1).strip('_ ')
            # remove customer_task prefix if regex captured it
            pref=f'{customer}_{task}_' if customer and task else ''
            if pref and C(name).startswith(C(pref)):
                name=name[len(pref):]
            return '시험',name
    for src in (issue,N(d.get('problem'))):
        m=re.search(r'(?<![A-Za-z0-9])([A-D]\d+)(?![A-Za-z0-9])',src,re.I)
        if m: return '빌드',m.group(1).upper()
    return '',''

# Strict symptom mode: do not retain generic descriptive lines.
# Only lines/clauses explicitly expressing a symptom are accepted.
_EXTRA_EXCLUDE=(
    '발생경위','확인사항','시험조건','시험 조건','시험방법','시험 방법','시험절차','시험 절차',
    '시험목적','시험 목적','시험환경','시험 환경','시험일자','시험 일자','시험장소','시험 장소',
    '시험장비','시험 장비','평가조건','평가 조건','평가방법','평가 방법','평가경위','확인경위',
    '배경','목적','조건','절차','방법','경위'
)


def _actual_symptoms(d):
    problem=N(d.get('problem')); kind,event=_event_name(d)
    kept=[]
    for raw in problem.split('\n'):
        line=step2._clean_piece(raw)
        if not line or step2._is_heading(line): continue
        pieces=[step2._clean_piece(x) for x in re.split(r'\s*[;/|]\s*',line) if step2._clean_piece(x)]
        for p in pieces:
            q=C(p)
            if not q: continue
            if any(C(x) in q for x in _EXTRA_EXCLUDE): continue
            if step2._is_condition_or_meta(p): continue
            if event and q in (C(event),C(event+'중')): continue
            # STRICT: actual failure/symptom wording is required.
            if any(C(w) in q for w in step1._SYMPTOM_WORDS):
                kept.append(p)
    out=[]; seen=set()
    for x in kept:
        k=C(x)
        if k not in seen: seen.add(k); out.append(x)
    return out


def _short_problem(d):
    symptoms=_actual_symptoms(d); kind,event=_event_name(d)
    lines=['1) 현상']
    lines.extend(symptoms or ['(현상 내용 미확인)'])
    if kind=='시험' and event: lines += ['2) 시험명',event]
    elif kind=='빌드' and event: lines += ['2) 빌드명',event]
    return '\n'.join(lines)


def _title_suffix(d):
    kind,event=_event_name(d)
    if kind=='시험' and event: return f'{event} 이슈 발생'
    if kind=='빌드' and event: return f'{event} 빌드 이슈 발생'
    return v319._trim_before_customer(d.get('issue_name'),d.get('customer'))


def _page2_title_task(d):
    customer=N(d.get('customer')); task=N(d.get('task_name'))
    return f'{customer}_{task}' if customer and task else (customer or task)


def _force_page2_header(sl,d,g):
    task=_page2_title_task(d); suffix=_title_suffix(d)
    title_sh,issue_sh=v319._find_page2_header_shapes(sl)
    if title_sh is not None:
        # Do not rely on intermediate weekly_display values: use original extracted customer/task directly.
        v319._replace_title_runs(title_sh,task,suffix)
    if issue_sh is not None:
        issue=v319._trim_before_customer(d.get('issue_name'),d.get('customer'))
        v319._set_issue_line(issue_sh,issue)
    step1._update_team_owner(sl,g)


def weekly(src,out,d,g,mode):
    dd=v319._weekly_display(d)
    p1=dict(dd); p1['issue_name']=step1._page1_issue(d); p1['problem']=_short_problem(d)
    prs=Presentation(src)
    v319._update_page1(prs,p1,g)
    if len(prs.slides)>1:
        v320._update_page2(prs.slides[1],dd,g,mode)
        _force_page2_header(prs.slides[1],d,g)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try: prs.save(out); saved=out
    except PermissionError:
        p=Path(out); saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix)); prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.weekly=weekly

class RecoveryStep3App(base.App):
    def __init__(self):
        super().__init__(); self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP3')

if __name__=='__main__': RecoveryStep3App().mainloop()
