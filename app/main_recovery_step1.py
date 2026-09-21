# Recovery STEP1 based on v3.2.0 rendering.
# ONLY fixes weekly meeting fields requested after recovery:
# 1) page1 issue column removes customer_task prefix
# 2) page1 problem is shortened to actual symptom (plus test/build name when detected)
# 3) page2 title uses specific test/build name
# 4) page2 top-right team/owner placeholder is updated
# Image extraction, D-section layout, fonts, Excel writing, and v3.2.0 adaptive layout are untouched.

import re
import datetime
from pathlib import Path
from pptx import Presentation

import main_v320 as v320
import main_v319 as v319
import main_v310 as v310

base = v320.base
N = v310.N
C = v310.C
EMU = v310.EMU


# ---------- text helpers ----------
def _event_name(d):
    """Return ('시험'|'빌드'|'', name). Build format is A-D + digits."""
    customer=C(d.get('customer')); task=C(d.get('task_name'))
    sources=[N(d.get('issue_name')), N(d.get('problem'))]
    for src in sources:
        toks=[x.strip() for x in re.split(r'[_/|]+',src) if x.strip()]
        for t in toks:
            if C(t) in (customer,task):
                continue
            m=re.search(r'([^\s,;:()]+시험)',t)
            if m:
                return '시험',m.group(1).strip()
    for src in sources:
        m=re.search(r'(?<![A-Za-z0-9])([A-D]\d+)(?![A-Za-z0-9])',src,re.I)
        if m:
            return '빌드',m.group(1).upper()
    return '',''


def _page1_issue(d):
    """Weekly-summary issue: remove routing labels AND customer/project prefix."""
    issue=N(d.get('issue_name'))
    if not issue:
        return ''
    customer=N(d.get('customer'))
    task=N(d.get('task_name'))
    project=task
    if customer and project:
        project=re.sub(r'^\s*'+re.escape(customer)+r'\s*[_\-/／|:： ]+\s*','',project,flags=re.I)
    project=N(project).strip(' _-/／|:：')

    parts=[x for x in re.split(r'[_/|:：\\-]+',issue) if N(x)]
    pk=re.sub(r'[^0-9A-Za-z가-힣]+','',project).lower()
    for idx,part in enumerate(parts):
        q=re.sub(r'[^0-9A-Za-z가-힣]+','',N(part)).lower()
        if pk and q and (pk==q or (min(len(pk),len(q))>=4 and (pk in q or q in pk))):
            # Summary table already has a separate customer/project column,
            # so its Issue cell starts AFTER the matched project token.
            rest='_'.join(parts[idx+1:]).strip(' _-/／|:：')
            if rest:
                return v319._clean_issue_label(rest)
            break

    return v319._clean_issue_label(v319._strip_markers_before_customer(issue,customer))


_SYMPTOM_WORDS=(
    '발생','불량','이상','파손','파단','누액','누수','누설','변형','부풀','스웰','swelling',
    '전압저하','전압 저하','voltage drop','drop','단선','단락','쇼트','short','통신불가','통신 불가',
    '미동작','오동작','미점등','소손','탄화','크랙','crack','이탈','탈락','변색','과열','발열',
    '편차','오차','미검출','검출불가','정지','fail','failure','ng'
)
_CONDITION_WORDS=(
    '시험조건','시험 조건','조건','soc','dod','soh','온도','습도','전류','전압 조건','전압조건',
    '충전조건','방전조건','충방전조건','시간','cycle','사이클','샘플수','sample 수','시험방법','시험 방법',
    '시험절차','시험 절차','프로파일','profile','c-rate','crate'
)


def _clean_piece(s):
    s=N(s)
    s=re.sub(r'^[-•·▪◦]\s*','',s)
    s=re.sub(r'^\(?\d+\)?[.)]\s*','',s)
    return s.strip()


def _short_problem(d):
    """Keep actual symptom. Add test/build name only when applicable."""
    problem=N(d.get('problem'))
    kind,event=_event_name(d)
    kept=[]
    for raw in problem.split('\n'):
        line=_clean_piece(raw)
        if not line:
            continue
        # If a line contains several clauses, retain symptom clauses and discard condition-only clauses.
        pieces=[_clean_piece(x) for x in re.split(r'\s*[;/|]\s*',line) if _clean_piece(x)]
        symptom=[p for p in pieces if any(C(w) in C(p) for w in _SYMPTOM_WORDS)]
        if symptom:
            kept.extend(symptom)
            continue
        # Remove obvious condition-only lines.
        if any(C(w) in C(line) for w in _CONDITION_WORDS):
            continue
        # Preserve descriptive non-condition lines, but not a standalone event name.
        if not event or C(line)!=C(event):
            kept.append(line)

    # de-duplicate, source order
    out=[]; seen=set()
    for x in kept:
        k=C(x)
        if k and k not in seen:
            seen.add(k); out.append(x)

    # Safety fallback: never erase the problem.
    if not out and problem:
        out=[_clean_piece(problem.split('\n')[0])]

    if kind=='시험' and event:
        out.append(f'시험명: {event}')
    elif kind=='빌드' and event:
        out.append(f'빌드명: {event}')
    return '\n'.join(x for x in out if x).strip()


def _title_issue(d, issue_from_customer):
    kind,event=_event_name(d)
    if kind=='시험' and event:
        return f'{event} 이슈 발생'
    if kind=='빌드' and event:
        return f'{event} 이슈 발생'
    return v319._title_issue(issue_from_customer)


# ---------- page2 metadata patch only ----------
def _set_preserve_runs(sh,text):
    if not hasattr(sh,'text_frame'):
        return False
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    if runs:
        runs[0].text=N(text)
        for r in runs[1:]:
            r.text=''
    else:
        sh.text_frame.text=N(text)
    return True


def _update_team_owner(sl,g):
    team=N(g.get('team')); owner=N(g.get('owner'))
    if not (team or owner):
        return False
    target=f'{team} 담당자 : {owner}'.strip()
    candidates=[]
    for sh in v310.walk(sl):
        if not hasattr(sh,'text_frame'):
            continue
        t=N(getattr(sh,'text',''))
        if '담당자' not in t:
            continue
        try:
            top=float(sh.top)/EMU; left=float(sh.left)/EMU
        except Exception:
            top=99; left=0
        # Prefer the page-2 header placeholder in the upper-right.
        score=(0 if ('00팀' in t or '000' in t) else 1, top, -left)
        candidates.append((score,sh))
    if not candidates:
        return False
    candidates.sort(key=lambda x:x[0])
    return _set_preserve_runs(candidates[0][1],target)


_original_page2_meta=v319._page2_meta

def _page2_meta_step1(sl,d,g):
    _original_page2_meta(sl,d,g)
    _update_team_owner(sl,g)


# ---------- weekly wrapper; renderer/layout/image logic remains v3.2.0 ----------
def weekly(src,out,d,g,mode):
    dd=v319._weekly_display(d)
    dd['_title_issue']=_title_issue(d,dd.get('issue_name'))

    # Page 1 gets its own display-only issue/problem values.
    p1=dict(dd)
    p1['issue_name']=_page1_issue(d)
    p1['problem']=_short_problem(d)

    prs=Presentation(src)
    v319._update_page1(prs,p1,g)

    if len(prs.slides)>1:
        old=v319._page2_meta
        try:
            v319._page2_meta=_page2_meta_step1
            v320._update_page2(prs.slides[1],dd,g,mode)
        finally:
            v319._page2_meta=old

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved


# Preserve recovery baseline extraction/Excel; replace only weekly output.
base.weekly=weekly


class RecoveryStep1App(base.App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP1')


if __name__=='__main__':
    RecoveryStep1App().mainloop()
