"""Final output type selector: full updated files or update-only reduced files."""
import copy
import re
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

import main_enterprise as ent
import ui_enterprise as ui

_original_info = ui.info


def _latest_version(source_path):
    src=Path(source_path)
    out=src.parent/'자동화_결과'
    root=src.stem
    pat=re.compile(rf'^{re.escape(root)}_v0[.](\d+){re.escape(src.suffix)}$',re.I)
    found=[]
    if out.exists():
        for p in out.iterdir():
            m=pat.match(p.name)
            if m:
                found.append((int(m.group(1)),p))
    return max(found,key=lambda x:x[0])[1] if found else None


def _slide_signature(slide):
    """Compare slide content while ignoring package-level metadata."""
    try:return slide._element.xml
    except Exception:return ''


def _remove_slide(prs,index):
    slide_id=prs.slides._sldIdLst[index]
    rel_id=slide_id.rId
    prs.part.drop_rel(rel_id)
    del prs.slides._sldIdLst[index]


def _ppt_update_only(source,saved):
    src=Presentation(source); dst=Presentation(saved)
    keep=[]
    for i,slide in enumerate(dst.slides):
        if i>=len(src.slides) or _slide_signature(slide)!=_slide_signature(src.slides[i]):
            keep.append(i)
    if not keep:
        return None,0
    for i in range(len(dst.slides)-1,-1,-1):
        if i not in keep:_remove_slide(dst,i)
    target=Path(saved).with_name(Path(saved).stem+'_업데이트사항만'+Path(saved).suffix)
    dst.save(target)
    return target,len(keep)


def _excel_update_only(source,saved):
    src=load_workbook(source,rich_text=True)
    dst=load_workbook(saved,rich_text=True)
    ws0=src['Sheet1'] if 'Sheet1' in src.sheetnames else src.active
    ws=dst['Sheet1'] if 'Sheet1' in dst.sheetnames else dst.active
    changed=[]
    max_row=max(ws0.max_row,ws.max_row); max_col=max(ws0.max_column,ws.max_column)
    for r in range(1,max_row+1):
        different=False
        for c in range(1,max_col+1):
            a=ws0.cell(r,c).value if r<=ws0.max_row and c<=ws0.max_column else None
            b=ws.cell(r,c).value if r<=ws.max_row and c<=ws.max_column else None
            if str(a or '')!=str(b or ''):
                different=True; break
        if different:changed.append(r)
    if not changed:return None,0
    # Issue DB contract: rows 1-6 are fixed header/template rows.
    # In update-only output, always preserve them exactly and keep only changed data rows from row 7 onward.
    changed=[r for r in changed if r>=7]
    if not changed:return None,0
    keep=set(range(1,7))|set(changed)
    for r in range(ws.max_row,0,-1):
        if r not in keep:ws.delete_rows(r,1)
    target=Path(saved).with_name(Path(saved).stem+'_업데이트사항만'+Path(saved).suffix)
    dst.save(target)
    return target,len(changed)


class OutputTypeDialog(ent.tk.Toplevel):
    def __init__(self,parent):
        super().__init__(parent); self.withdraw(); self.result=None
        self.title('업데이트 파일 유형 선택'); self.configure(bg='white'); self.resizable(False,False); self.transient(parent)
        head=ent.tk.Frame(self,bg=ui.NAVY,height=58); head.pack(fill='x'); head.pack_propagate(False)
        ent.tk.Label(head,text='업데이트 파일 유형 선택',bg=ui.NAVY,fg='white',font=('Malgun Gothic',12,'bold')).pack(side='left',padx=22)
        body=ent.tk.Frame(self,bg='white'); body.pack(fill='both',expand=True,padx=24,pady=20)
        ent.tk.Label(body,text='저장할 Output 파일 유형을 선택해 주세요.',bg='white',fg=ui.TEXT,font=('Malgun Gothic',10,'bold')).pack(anchor='w',pady=(0,14))
        self.var=ent.tk.StringVar(value='full')
        ent.ttk.Radiobutton(body,text='1. 전체 업데이트 본',variable=self.var,value='full',style='Mode.TRadiobutton').pack(anchor='w',pady=8)
        ent.tk.Label(body,text='기존 주간회의 전체 페이지 / Issue DB 전체 행을 유지합니다.',bg='white',fg=ui.MUTED,font=('Malgun Gothic',8)).pack(anchor='w',padx=24)
        ent.ttk.Radiobutton(body,text='2. 업데이트 사항 이외 삭제본',variable=self.var,value='reduced',style='Mode.TRadiobutton').pack(anchor='w',pady=(16,8))
        ent.tk.Label(body,text='주간회의는 업데이트된 페이지만, Issue DB는 헤더와 업데이트된 행만 남깁니다.',bg='white',fg=ui.MUTED,font=('Malgun Gothic',8)).pack(anchor='w',padx=24)
        foot=ent.tk.Frame(self,bg='#F6F8FA',height=64); foot.pack(fill='x'); foot.pack_propagate(False)
        ent.tk.Button(foot,text='선택 완료',command=self.ok,bg=ui.BLUE,fg='white',bd=0,font=('Malgun Gothic',9,'bold'),width=14,pady=8).pack(side='right',padx=20,pady=13)
        self.protocol('WM_DELETE_WINDOW',self.cancel); ui.center_window(self,parent,620,330); self.deiconify(); self.grab_set(); self.focus_force(); parent.wait_window(self)
    def ok(self):self.result=self.var.get(); self.destroy()
    def cancel(self):self.result='full'; self.destroy()


def _select_and_finish(parent,title,message):
    if title!='업데이트 완료':
        return _original_info(parent,title,message)
    dlg=OutputTypeDialog(parent)
    if dlg.result=='reduced':
        made=[]; errors=[]
        weekly=parent.vars.get('pptweekly').get().strip() if parent.vars.get('pptweekly') else ''
        excel=parent.vars.get('xlsx').get().strip() if parent.vars.get('xlsx') else ''
        try:
            if weekly:
                saved=_latest_version(weekly)
                if saved:
                    reduced,n=_ppt_update_only(weekly,saved)
                    if reduced:
                        made.append(f'주간회의: 업데이트 페이지 {n}개')
                        try:Path(saved).unlink()
                        except Exception:pass
        except Exception as e:errors.append('주간회의 축약본 생성 실패: '+str(e))
        try:
            if excel:
                saved=_latest_version(excel)
                if saved:
                    reduced,n=_excel_update_only(excel,saved)
                    if reduced:
                        made.append(f'Issue DB: 업데이트 행 {n}개')
                        try:Path(saved).unlink()
                        except Exception:pass
        except Exception as e:errors.append('Issue DB 축약본 생성 실패: '+str(e))
        if errors:
            extra='\n'.join(made+errors)
            return _original_info(parent,'업데이트 완료','일부 축약본 생성에 실패했습니다. 전체 업데이트 본은 보존했습니다.\n\n'+extra)
        if not made:
            return _original_info(parent,'업데이트 완료','변경된 행/페이지를 찾지 못해 축약본을 만들지 않았습니다. 전체 업데이트 본은 보존했습니다.')
        extra='\n'.join(made)
        return _original_info(parent,'업데이트 완료','업데이트 사항 이외 삭제본으로 저장했습니다.\n\n'+extra)
    return _original_info(parent,'업데이트 완료',message+'\n\n선택 유형: 전체 업데이트 본')


ui.info=_select_and_finish
