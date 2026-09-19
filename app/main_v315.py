# 8D Issue Automation v3.1.5
# Strict 4D marker policy:
# - occurrence cause -> left-bottom 4D
# - leak/system cause -> right-top 4D
# - missing side -> remove that 4D marker/title unit
import copy, datetime
from pathlib import Path

import main_v314 as v314
import main_v310 as v310
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

base = v314.base
N = v310.N
box = v310.box
walk = v310.walk


def _top_4d_units(sl):
    """Return unique top-level units that contain the visible 4D marker."""
    units=[]
    for sh in sl.shapes:
        if N(getattr(sh,'text','')) == '4D':
            units.append(sh)
            continue
        if getattr(sh,'shape_type',None) == MSO_SHAPE_TYPE.GROUP:
            for ch in walk(sh):
                if ch is sh:
                    continue
                if N(getattr(ch,'text','')) == '4D':
                    units.append(sh)
                    break
    # unique by element identity
    out=[]; seen=set()
    for u in units:
        k=id(u._element)
        if k not in seen:
            seen.add(k); out.append(u)
    return out


def _nearby_title(sl, marker):
    """Prefer the native 4D item-name text next to the 4D circle."""
    try:
        mx,my,mw,mh=box(marker)
        ranked=[]
        for sh in sl.shapes:
            if sh is marker or not hasattr(sh,'text_frame'):
                continue
            t=N(getattr(sh,'text',''))
            if not t or t=='4D' or '담당자' in t or 'signal' in t.lower():
                continue
            sx,sy,sw,shh=box(sh)
            if abs(sy-my)>.85*v310.EMU or sx<mx-.35*v310.EMU or sx>mx+3.0*v310.EMU:
                continue
            q=t.lower().replace(' ','')
            semantic=0 if any(k in q for k in ('원인','분석','cause','analysis','root')) else 1
            score=semantic*10*v310.EMU+abs(sy-my)+abs(sx-(mx+mw))
            ranked.append((score,sh))
        if ranked:
            ranked.sort(key=lambda x:x[0])
            return ranked[0][1]
        return v310.nearby_title(sl, marker, '4D')
    except Exception:
        return None

def _group_4d_marker_titles(sl):
    """Group each ungrouped 4D circle with its nearby item-name before moving/cloning."""
    units=list(_top_4d_units(sl))
    for unit in units:
        if getattr(unit,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:
            continue
        title=_nearby_title(sl,unit)
        if title is None:
            continue
        try:
            group=sl.shapes.add_group_shape([unit,title])
            group.name='AUTO_8D_4D_UNIT'
        except Exception:
            pass


def _delete_shape(sh):
    try:
        el=sh._element
        parent=el.getparent()
        if parent is not None:
            parent.remove(el)
            return True
    except Exception:
        pass
    return False


def _delete_4d_unit(sl, unit):
    """Delete marker+title as one unit. If ungrouped, also delete the nearby title."""
    if getattr(unit,'shape_type',None) == MSO_SHAPE_TYPE.GROUP:
        _delete_shape(unit)
        return
    title=_nearby_title(sl,unit)
    if title is not None:
        _delete_shape(title)
    _delete_shape(unit)


def _move_unit(unit,x,y):
    bx,by,_,_=box(unit)
    unit.left += Inches(x)-int(bx)
    unit.top += Inches(y)-int(by)


def _clone_unit(sl,unit,x,y):
    try:
        newel=copy.deepcopy(unit._element)
        sl.shapes._spTree.insert_element_before(newel,'p:extLst')
        # locate wrapper for inserted element
        dup=None
        for sh in sl.shapes:
            if sh._element is newel:
                dup=sh; break
        if dup is None:
            dup=list(sl.shapes)[-1]
        dup.name='AUTO_8D_4D_LEFT'
        _move_unit(dup,x,y)
        return dup
    except Exception:
        return None


def _ensure_4d_units(sl, has_cause, has_leak, left_xy, right_xy):
    """Make visible 4D marker+title groups match the two 4D content zones."""
    _group_4d_marker_titles(sl)
    units=_top_4d_units(sl)

    # No 4D content: remove every visible 4D marker/title unit.
    if not has_cause and not has_leak:
        for u in list(units):
            _delete_4d_unit(sl,u)
        return

    # Need a source unit. If template has none, create a marker + item-name group.
    if not units:
        x,y=(right_xy if has_leak else left_xy)
        sh=sl.shapes.add_shape(1,Inches(x),Inches(y),Inches(.42),Inches(.42))
        sh.name='AUTO_8D_4D_MARKER'; v310.set_text(sh,'4D',8)
        title=sl.shapes.add_textbox(Inches(x+.48),Inches(y),Inches(1.15),Inches(.42))
        title.name='AUTO_8D_4D_TITLE'; v310.set_text(title,'원인 분석',8)
        try:
            group=sl.shapes.add_group_shape([sh,title])
            group.name='AUTO_8D_4D_UNIT'
            units=[group]
        except Exception:
            units=[sh]

    if has_cause and has_leak:
        # Keep one unit on the right and one on the left; remove any extras.
        source=units[0]
        _move_unit(source,right_xy[0],right_xy[1])
        if len(units)>=2:
            left=units[1]
            _move_unit(left,left_xy[0],left_xy[1])
            for extra in units[2:]:
                _delete_4d_unit(sl,extra)
        else:
            dup=_clone_unit(sl,source,left_xy[0],left_xy[1])
            if dup is None:
                # fallback only if cloning the existing marker/title unit fails
                sh=sl.shapes.add_shape(1,Inches(left_xy[0]),Inches(left_xy[1]),Inches(.42),Inches(.42))
                sh.name='AUTO_8D_4D_LEFT'; v310.set_text(sh,'4D',8)
        return

    # Exactly one side exists: keep exactly one 4D unit and delete the others.
    keep=units[0]
    target=left_xy if has_cause else right_xy
    _move_unit(keep,target[0],target[1])
    for extra in units[1:]:
        _delete_4d_unit(sl,extra)


def update_page2(sl,d,g,mode):
    # Let v3.1.4 perform image/text rendering, metadata, and dynamic layout first.
    v314.update_page2(sl,d,g,mode)

    imgs=d.get('_section_images',{}) or {}
    zones,texts,font=v314._layout(d,imgs)
    has_cause=bool(N(d.get('cause_4d')))
    has_leak=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))

    zl=zones['4D_CAUSE']
    zr=zones['4D_LEAK']
    left_xy=(max(.10,zl['x']-.18), max(.10,zl['y']-.35))
    right_xy=(max(.10,zr['x']-.18), max(.10,zr['y']-.35))

    # Enforce the user's exact 4D placement/content rule for NEW and EXISTING outputs.
    _ensure_4d_units(sl,has_cause,has_leak,left_xy,right_xy)


def weekly(src,out,d,g,mode):
    prs=Presentation(src)
    v314.update_page1(prs,d,g)
    if len(prs.slides)>1:
        update_page2(prs.slides[1],d,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.weekly=weekly
base.extract=v314.extract

if __name__=='__main__':
    base.App().mainloop()
