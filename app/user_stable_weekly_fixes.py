"""Five narrowly-scoped weekly-meeting fixes on stable baseline 667623cc.

This module is imported LAST by main_enterprise_final.py.
It deliberately leaves the stable weekly writer, attachment flow, Excel logic,
and native-section creation implementation untouched. Only helper functions
for title text, summary placement/matching, 7D positioning, and clone cleanup
are patched.
"""

import re

from pptx.enum.shapes import MSO_SHAPE_TYPE

import main_recovery_step4 as s4
import main_recovery_step12 as s12
import main_recovery_step13 as s13
import main_recovery_step14 as s14
import main_recovery_step14_fix2 as core
import main_v310 as v310
import project_autocomplete_final as catalog

N = v310.N


# ---------------------------------------------------------------------------
# 1) Detail title:
#    고객사_과제명_발생샘플 샘플_개발단계_발생처 이슈 발생
# ---------------------------------------------------------------------------
_original_detail_title_text = s4._detail_title_text


def _title_customer_project(d, g):
    selected = N((g or {}).get("task_name")) or N(s13._customer_task(d or {}))
    customer, project = catalog.split_customer_task(selected)
    if not customer:
        customer = N((d or {}).get("customer"))
        project = N(project or selected)
    return N(customer), N(project)


def _sample_with_suffix(value):
    value = N(value)
    if not value:
        return ""
    if re.search(r"(?:샘플|sample)\s*$", value, re.I):
        return value
    return value + " 샘플"


def _detail_title_user_format(d, g):
    d = d or {}
    g = g or {}
    customer, project = _title_customer_project(d, g)
    sample = N(g.get("sample") or d.get("sample") or d.get("occurrence_sample"))
    stage = N(
        g.get("stage")
        or d.get("development_stage")
        or d.get("occurrence_stage")
        or d.get("stage")
    )
    site = N(
        g.get("occurrence_site")
        or d.get("_origin_occurrence_site")
        or d.get("occurrence_site")
    )

    if all((customer, project, sample, stage, site)):
        return f"{customer}_{project}_{_sample_with_suffix(sample)}_{stage}_{site} 이슈 발생"

    return _original_detail_title_text(d, g)


s4._detail_title_text = _detail_title_user_format


# ---------------------------------------------------------------------------
# 3) Keep the native 7D circle + 수평전개 unit below the dynamic 6D region.
#    Recompute an absolute target each run, so repeated updates do not drift.
# ---------------------------------------------------------------------------
_original_update_page2 = s12._update_page2_step12


def _adjust_7d_after_6d(sl, d):
    try:
        imgs = (d or {}).get("_section_images", {}) or {}
        zones, _, _ = s12.step11._adaptive_cascade_layout(d or {}, imgs)
        z6 = zones.get("6D")
        if not z6:
            return

        marker, parent = v310.find_marker(sl, "7D")
        if marker is None:
            return
        unit = parent if parent is not None else marker
        x7, _, _, h7 = v310.box(unit)

        try:
            slide_h = float(
                sl.part.package.presentation_part.presentation.slide_height
            ) / v310.EMU
        except Exception:
            slide_h = 7.5

        desired_y = float(z6["y"]) + float(z6["h"]) + 0.28
        max_y = max(0.10, slide_h - float(h7) / v310.EMU - 0.06)
        desired_y = min(max(0.10, desired_y), max_y)
        v310.move_marker_unit(sl, "7D", float(x7) / v310.EMU, desired_y)
    except Exception:
        pass


def _update_page2_with_7d(sl, d, g, mode):
    _original_update_page2(sl, d, g, mode)
    _adjust_7d_after_6d(sl, d)


s12._update_page2_step12 = _update_page2_with_7d


# ---------------------------------------------------------------------------
# 5) When a summary/detail page is CLONED, strip copied image/media objects.
#    Tables, text, rectangles/frames, lines and other template structure stay.
# ---------------------------------------------------------------------------
_VISUAL_TYPES = {
    x
    for x in (
        getattr(MSO_SHAPE_TYPE, "PICTURE", None),
        getattr(MSO_SHAPE_TYPE, "LINKED_PICTURE", None),
        getattr(MSO_SHAPE_TYPE, "CHART", None),
        getattr(MSO_SHAPE_TYPE, "MEDIA", None),
        getattr(MSO_SHAPE_TYPE, "OLE_OBJECT", None),
    )
    if x is not None
}


def _remove_shape(sh):
    try:
        el = sh._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
            return True
    except Exception:
        pass
    return False


def _clean_group_visuals(group):
    for child in list(getattr(group, "shapes", ())):
        st = getattr(child, "shape_type", None)
        if st in _VISUAL_TYPES:
            _remove_shape(child)
        elif st == MSO_SHAPE_TYPE.GROUP:
            _clean_group_visuals(child)


def _remove_copied_visuals(sl):
    for sh in list(sl.shapes):
        if getattr(sh, "has_table", False):
            continue
        st = getattr(sh, "shape_type", None)
        if st in _VISUAL_TYPES:
            _remove_shape(sh)
        elif st == MSO_SHAPE_TYPE.GROUP:
            _clean_group_visuals(sh)


_original_prepare_new_summary_page = s14._prepare_new_summary_page


def _prepare_new_summary_page_stable(prs, pages, g, template_index=None):
    """Clone overflow immediately after the LAST matched summary page.

    New-task behavior stays on the stable baseline path.
    """
    if template_index is None:
        result = _original_prepare_new_summary_page(
            prs, pages, g, template_index=None
        )
    else:
        original_picker = s14._first_summary_section_insert_index
        try:
            s14._first_summary_section_insert_index = (
                lambda _prs, _pages: min(
                    len(_prs.slides), int(template_index) + 1
                )
            )
            result = _original_prepare_new_summary_page(
                prs, pages, g, template_index=template_index
            )
        finally:
            s14._first_summary_section_insert_index = original_picker

    try:
        _remove_copied_visuals(prs.slides[result[0]])
    except Exception:
        pass
    return result


s14._prepare_new_summary_page = _prepare_new_summary_page_stable


_original_clone_detail_shell = core._clone_detail_shell


def _clone_detail_shell_clean(prs, d, insert_at, matched_section=None):
    result = _original_clone_detail_shell(prs, d, insert_at, matched_section)
    try:
        _remove_copied_visuals(result[0])
    except Exception:
        pass
    return result


core._clone_detail_shell = _clone_detail_shell_clean


# ---------------------------------------------------------------------------
# 4) Native-section resolution:
#    1. customer + project strict
#    2. project-only strict
#    3. similar section only as a USER-CONFIRMABLE suggestion
# ---------------------------------------------------------------------------
def _project_parts(d):
    d = d or {}
    full_raw = N(s13._customer_task(d))
    selected = N(d.get("task_name")) or full_raw
    selected_customer, project_raw = catalog.split_customer_task(selected)
    customer_raw = N(d.get("customer")) or N(selected_customer)
    project_raw = N(project_raw or selected)
    return s13._k(full_raw), s13._k(project_raw), s13._k(customer_raw)


def _section_match_level(name, d):
    q = s13._k(name)
    _, project, customer = _project_parts(d)
    if not q:
        return 0
    if customer and project and customer in q and project in q:
        return 3
    if project and q == project:
        return 2
    if project and q and min(len(project), len(q)) >= 3:
        if project in q or q in project:
            return 1
    return 0


def _matching_native_section_strict(prs, d):
    for level in (3, 2):
        matches = [
            sec
            for sec in core._native_sections(prs)
            if _section_match_level(sec.get("name"), d) == level
        ]
        if matches:
            matches.sort(
                key=lambda x: len(x.get("indices", [])), reverse=True
            )
            return matches[0]
    return None


def _section_resolution_strict(prs, d):
    try:
        s13._clear_slide_text_cache()
    except Exception:
        pass

    sections = core._native_sections(prs)

    for level, reason in (
        (3, "고객사와 과제명이 모두 일치하는 구역을 찾았습니다."),
        (2, "과제명이 정확히 일치하는 구역을 찾았습니다."),
    ):
        matches = [
            sec
            for sec in sections
            if _section_match_level(sec.get("name"), d) == level
        ]
        if matches:
            matches.sort(
                key=lambda x: len(x.get("indices", [])), reverse=True
            )
            sec = matches[0]
            return {
                "mode": "exact",
                "name": N(sec.get("name")),
                "reason": reason,
                "section": sec,
            }

    similar = [
        sec
        for sec in sections
        if _section_match_level(sec.get("name"), d) == 1
    ]
    if similar:
        similar.sort(
            key=lambda x: len(x.get("indices", [])), reverse=True
        )
        sec = similar[0]
        return {
            "mode": "suggest",
            "name": N(sec.get("name")),
            "reason": "과제명과 유사한 기존 구역이 확인되었습니다.",
            "section": sec,
        }

    return {
        "mode": "new",
        "name": N(s13._customer_task(d))
        or N((d or {}).get("task_name"))
        or N((d or {}).get("customer"))
        or "신규 과제",
        "reason": (
            "일치하거나 유사한 기존 구역을 찾지 못했습니다. "
            "신규 구역을 생성합니다."
        ),
        "section": None,
    }


core._project_parts = _project_parts
core._section_match_level = _section_match_level
core._matching_native_section = _matching_native_section_strict
core.section_resolution = _section_resolution_strict


# ---------------------------------------------------------------------------
# 2 + 4) Summary placement/matching across ALL summary pages.
#    1. exact customer+project
#    2. exact project
#    3. only if the chosen native section was SIMILAR and the user selected
#       "해당 구역 업데이트", match that section name against real summary labels.
#       If absent, use the normal new-summary path with customer+project.
# ---------------------------------------------------------------------------
def _summary_hits(prs, selected):
    pages = s14._summary_pages(prs)
    full_key = catalog._norm(selected)
    project_key = catalog.project_key(selected, False)
    full_hits = []
    project_hits = []

    for si, tb, hr in pages:
        hm = s13._summary_map(tb, hr)
        if "task" not in hm:
            continue
        fr = []
        pr = []
        for r in range(hr + 1, len(tb.rows)):
            raw = s13._row_text(tb, r, hm["task"])
            if full_key and catalog._norm(raw) == full_key:
                fr.append(r)
            if (
                project_key
                and catalog.project_key(raw, False) == project_key
            ):
                pr.append(r)
        if fr:
            full_hits.append((si, tb, hr, hm, fr))
        if pr:
            project_hits.append((si, tb, hr, hm, pr))

    return pages, (full_hits if full_hits else project_hits)


def _confirmed_section_summary_hits(pages, d, g):
    name = N((g or {}).get("_weekly_section_override_name"))
    if not name:
        return []

    # Level 1 means this section was only available through the confirmation
    # dialog; exact sections are levels 3/2 and do not unlock fuzzy summary use.
    if _section_match_level(name, d) != 1:
        return []

    confirmed_project = catalog.project_key(name, False)
    if not confirmed_project:
        return []

    hits = []
    for si, tb, hr in pages:
        hm = s13._summary_map(tb, hr)
        if "task" not in hm:
            continue
        rows = []
        for r in range(hr + 1, len(tb.rows)):
            raw = s13._row_text(tb, r, hm["task"])
            raw_project = catalog.project_key(raw, False)
            if not raw_project:
                continue
            if raw_project == confirmed_project:
                rows.append(r)
                continue
            if min(len(raw_project), len(confirmed_project)) >= 4:
                if (
                    raw_project in confirmed_project
                    or confirmed_project in raw_project
                ):
                    rows.append(r)
        if rows:
            hits.append((si, tb, hr, hm, rows))
    return hits


def _existing_summary_display(selected, task_hits):
    if not task_hits:
        return N(selected)
    si, tb, hr, hm, rows = sorted(
        task_hits, key=lambda x: x[0]
    )[-1]
    raw = N(s13._row_text(tb, max(rows), hm["task"]))
    return raw or N(selected)


def _update_summary_exact_then_confirmed(prs, d, g, mode):
    pages = s14._summary_pages(prs)
    if not pages:
        raise ValueError(
            "주간회의 PPT에서 과제명/Signal 요약 양식 페이지를 찾지 못했습니다."
        )

    selected = N((g or {}).get("task_name")) or N(s13._customer_task(d))
    issue = s13._issue_display(d)

    _, task_hits = _summary_hits(prs, selected)
    if not task_hits:
        task_hits = _confirmed_section_summary_hits(pages, d, g)

    display = _existing_summary_display(selected, task_hits)
    original_customer_task = s13._customer_task

    try:
        if display:
            s13._customer_task = lambda _d: display

        if mode == "existing" and task_hits:
            best = None
            for si, tb, hr, hm, rows in task_hits:
                for r in rows:
                    sc = s13._row_match_score(
                        tb, r, hm, display, issue
                    )
                    if best is None or sc > best[0]:
                        best = (sc, si, tb, hr, r)
            if best and best[0] >= 140:
                _, si, tb, hr, row = best
                s14._write_summary_row(tb, row, hr, d, g)
                return si, row, "기존 행 업데이트"

        if task_hits:
            # Continue from the LAST page where this task appears.
            si, tb, hr, hm, rows = sorted(
                task_hits, key=lambda x: x[0]
            )[-1]
            after = max(rows)
            if s14._summary_row_insert_fits(prs, si, tb, after):
                row = s14._insert_row_after(tb, after)
                s14._write_summary_row(tb, row, hr, d, g)
                return (
                    si,
                    row,
                    "확정 과제의 마지막 요약 행 바로 아래 삽입",
                )

            nsi, ntb, nhr, nrow = s14._prepare_new_summary_page(
                prs, pages, g, template_index=si
            )
            s14._write_summary_row(ntb, nrow, nhr, d, g)
            return (
                nsi,
                nrow,
                "확정 과제 요약 공간 초과로 바로 다음 페이지에 추가",
            )

        # No exact or user-confirmed similar summary label: do not force fuzzy.
        if selected:
            s13._customer_task = lambda _d: selected
        si, tb, hr, row = s14._prepare_new_summary_page(prs, pages, g)
        s14._write_summary_row(tb, row, hr, d, g)
        return (
            si,
            row,
            "고객사/과제명 매칭 없음: 신규 요약 페이지 생성",
        )
    finally:
        s13._customer_task = original_customer_task


s14._update_summary_by_task = _update_summary_exact_then_confirmed
