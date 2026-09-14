import os
from pathlib import Path
import tkinter as tk
from tkinter import messagebox
from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.utils.units import pixels_to_EMU

import main_recovery_step8 as step8
import main_recovery_step7 as step7
import main_v310 as v310

base=step8.base
N=v310.N

OCCURRENCE_SITES=('etc.','고객/상위 Claim','고객/상위 생산','고객/상위 시험','고객/상위 운송/보관','부품 생산','제품 생산','제품 시험','제품 운송/보관')

_old_write_row=base.write_row


def _judge_issue_status(d):
    """Issue DB status rule requested by user. Returns (status, 6D text)."""
    text=N(d.get('verification_6d'))
    if not text:
        return 'open',''
    q=text.replace(' ','').lower()
    # Ongoing/future wording has priority over any other wording.
    if '중' in q or '예정' in q:
        return 'open',text
    # Any meaningful completed 6D content is treated as close when it is not ongoing/future.
    return 'close',text


def _set_two_cell_anchor(ws,row):
    """Make representative photo behave as Excel 'Move and size with cells'."""
    for img in list(getattr(ws,'_images',[])):
        try:
            anchor=getattr(img,'anchor',None)
            from_marker=getattr(anchor,'_from',None)
            # Representative image is column O (zero-based 14) on this row.
            if from_marker is not None:
                if from_marker.row != row-1 or from_marker.col != 14:
                    continue
            else:
                # String anchor fallback from older writer.
                if str(anchor).upper() != f'O{row}'.upper():
                    continue

            w=max(1,int(getattr(img,'width',120) or 120))
            h=max(1,int(getattr(img,'height',80) or 80))
            fr=AnchorMarker(col=14,row=row-1,colOff=pixels_to_EMU(4),rowOff=pixels_to_EMU(4))
            to=AnchorMarker(col=14,row=row-1,colOff=pixels_to_EMU(4+w),rowOff=pixels_to_EMU(4+h))
            img.anchor=TwoCellAnchor(editAs='twoCell',_from=fr,to=to)
        except Exception:
            pass


def write_row_step9(ws,r,d,g):
    result=_old_write_row(ws,r,d,g)
    # 1) 발생처 is always the GUI-selected value, never the source 8D text.
    ws.cell(r,13).value=N(g.get('occurrence_site'))
    # 2) Issue DB status is strictly open/close from the confirmed STEP9 decision.
    status=N(g.get('_issue_status_selected')) or _judge_issue_status(d)[0]
    ws.cell(r,18).value=status
    # 3) Representative photo = Move and size with cells.
    _set_two_cell_anchor(ws,r)
    return status

base.write_row=write_row_step9


class RecoveryStep9App(step8.RecoveryStep8App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP9')
        # STEP6 already converted occurrence_site to readonly combobox.
        # Re-assert the exact requested option order in case an inherited widget changed it.
        var=getattr(self,'vars',{}).get('occurrence_site')
        if var is not None:
            def walk(w):
                for ch in w.winfo_children():
                    yield ch
                    yield from walk(ch)
            from tkinter import ttk
            for widget in walk(self):
                if isinstance(widget,ttk.Combobox):
                    try:
                        if str(widget.cget('textvariable'))==str(var):
                            widget.configure(values=OCCURRENCE_SITES,state='readonly')
                            if var.get() not in OCCURRENCE_SITES:
                                var.set(OCCURRENCE_SITES[0])
                            break
                    except Exception:
                        pass

    def run(self):
        g=self.gui()
        if any(not g[k] or not os.path.exists(g[k]) for k in ['ppt8d','pptweekly','xlsx']):
            return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')

        try:
            d=base.extract(g['ppt8d'])
            mode=self.mode.get()

            # STEP8 issue-origin recommendation/selection.
            od=step8.OriginDialog(self,d)
            if od.result is None:
                self.log.delete('1.0','end'); self.log.insert('end','업데이트가 취소되었습니다.'); return
            g['_issue_origin_selected']=od.result

            # STEP7 Issue DB phenomenon full/summary choice.
            summary=step7._db_summary(d)
            pd=step7.ProblemChoiceDialog(self,summary)
            if pd.result is None:
                self.log.delete('1.0','end'); self.log.insert('end','업데이트가 취소되었습니다.'); return
            g['_db_problem_selected']=summary if pd.result=='summary' else N(d.get('problem'))

            # STEP9 status decision uses 6D only.
            judged,reason=_judge_issue_status(d)
            final_status=judged
            if judged=='close':
                sixd=reason or '(6D 내용 없음)'
                prompt=(
                    '6D까지 작성되어 있으며, 아래 6D 내용에 진행 중/예정 표현이 없어 close로 판단됩니다.\n\n'
                    '판단 이유(6D 내용)\n'
                    '────────────────────\n'
                    f'{sixd}\n'
                    '────────────────────\n\n'
                    '이슈 상태를 close로 처리하시겠습니까?\n'
                    '아니오를 선택하면 open으로 처리합니다.'
                )
                if not messagebox.askyesno('이슈 상태 확인',prompt,parent=self):
                    final_status='open'
            g['_issue_status_selected']=final_status

            out=Path(g['xlsx']).parent/'자동화_결과'
            out.mkdir(exist_ok=True)
            xo=out/(Path(g['xlsx']).stem+'_업데이트.xlsx')
            po=out/(Path(g['pptweekly']).stem+'_업데이트.pptx')

            wb=load_workbook(g['xlsx'])
            ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active
            r,_=base.find(ws,d,g)
            if mode=='existing' and not r:
                return messagebox.showwarning('업데이트 중단','기존 이슈를 Excel에서 정확하게 찾지 못했습니다. PMS/PLM 이슈번호를 입력하거나 이슈 정보를 확인하세요.')
            if mode=='new' and r and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?'):
                return

            a,_=base.update_excel(g['xlsx'],xo,d,g,new=(mode=='new'))
            b,_=base.weekly(g['pptweekly'],po,d,g,mode)

            phenomenon_label='현상 요약' if pd.result=='summary' else '전체 내용'
            self.log.delete('1.0','end')
            self.log.insert('end',
                f'=== 완료 ===\n'
                f'이슈 구분: {mode}\n'
                f'발생처: {g.get("occurrence_site","")}\n'
                f'Issue DB 현상: {phenomenon_label}\n'
                f'이슈기인: {od.result}\n'
                f'Issue DB 상태: {final_status}\n\n'
                f'{a}\n{b}\n\n결과: {out}'
            )
            messagebox.showinfo('완료',
                f'자동 업데이트가 완료되었습니다.\n\n'
                f'발생처: {g.get("occurrence_site","")}\n'
                f'이슈기인: {od.result}\n'
                f'Issue DB 상태: {final_status}\n\n{out}'
            )
        except Exception as e:
            messagebox.showerror('실행 오류',repr(e))


if __name__=='__main__':
    RecoveryStep9App().mainloop()
