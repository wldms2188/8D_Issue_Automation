# Recovery STEP4: restore full page1 problem and force exact page2 title by geometry.
# Preserve v3.2.0 image/layout/font/Excel behavior.
import datetime
import re
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

import main_recovery_step3 as step3
import main_recovery_step1 as step1
import main_v320 as v320
import main_v319 as v319
import main_v310 as v310

base=step3.base
N=v310.N
EMU=v310.EMU


def _full_problem(d):
    return N(d.get('problem'))


def _title_shape_by_geometry(sl):
    """Find page2 title even after placeholder text was already overwritten.
    Prefer a text box in the upper-left header region and exclude the issue-name line.
    """
    candidates=[]
    for sh in v310.walk(sl):
        if not hasattr(sh,'text_frame'):
            continue
        t=N(getattr(sh,'text',''))
        try:
            x=float(sh.left)/EMU; y=float(sh.top)/EMU; w=float(sh.width)/EMU; h=float(sh.height)/EMU
        except Exception:
            continue
        if y>0.55 or x>4.0:
            continue
        q=t.replace(' ','').lower()
        if q.startswith('이슈명') or '담당자' in t or 'signal' in q:
            continue
        # actual template title is around x=.17 y=.11 w=2.11 h=.44
        score=abs(x-.17)+abs(y-.11)+.15*abs(w-2.11)+.15*abs(h-.44)
        candidates.append((score,sh))
    if not candidates:
        return None
    candidates.sort(key=lambda z:z[0])
    return candidates[0][1]


def _owner_shape(sl):
    candidates=[]
    for sh in v310.walk(sl):
        if not hasattr(sh,'text_frame'):
            continue
        t=N(getattr(sh,'text',''))
        if '담당자' not in t:
            continue
        try:
            x=float(sh.left)/EMU; y=float(sh.top)/EMU
        except Exception:
            continue
        if y<=0.80:
            candidates.append((-x,y,sh))
    if not candidates:
        return None
    candidates.sort(key=lambda z:(z[1],z[0]))
    return candidates[0][2]

def _estimated_text_width_in(sh,text):
    """Conservative one-line width estimate for the visible detail title."""
    pt=18.0
    try:
        runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
        sizes=[r.font.size.pt for r in runs if r.font.size]
        if sizes:
            # Older generated pages may already have a shrunken font. Do not let
            # that hide a collision that would occur with the normal template title.
            pt=max(14.0,max(sizes))
    except Exception:
        pass
    em=pt/72.0
    units=0.0
    for ch in N(text):
        if ord(ch)>=0x2E80:
            units+=1.0
        elif ch.isspace():
            units+=0.35
        elif ch in '._-()/[]':
            units+=0.45
        else:
            units+=0.58
    return units*em

def _title_owner_overlap_risk(sl,title_sh,title_text):
    owner=_owner_shape(sl)
    if title_sh is None:
        return False
    try:
        tx=float(title_sh.left)/EMU; ty=float(title_sh.top)/EMU
        tw=float(title_sh.width)/EMU; th=float(title_sh.height)/EMU
    except Exception:
        return False

    estimated=_estimated_text_width_in(title_sh,title_text)
    if owner is None:
        # Even without a detected owner box, do not let a one-line title visibly
        # run out of its own template title area.
        return estimated>max(.25,tw*.98)

    try:
        ox=float(owner.left)/EMU; oy=float(owner.top)/EMU
        ow=float(owner.width)/EMU; oh=float(owner.height)/EMU
    except Exception:
        return estimated>max(.25,tw*.98)

    # Header boxes are sometimes a few pixels vertically offset even though the
    # rendered text is on the same visual line. Treat near-aligned boxes as colliding.
    tc=ty+th/2; oc=oy+oh/2
    same_header_line=abs(tc-oc)<=max(.35,(th+oh)*.75)
    if not same_header_line:
        return estimated>max(.25,tw*.98)

    available=max(.25,ox-tx-.12)
    rendered_right=tx+min(max(estimated,.01),max(tw,estimated))
    return estimated>available or rendered_right>ox-.08

def _strip_selected_project_prefix(issue,d):
    """Remove routing markers and visible customer/project prefix from detail title."""
    s=v319._trim_before_customer(issue,d.get('customer'))
    s=v319._clean_issue_label(s)
    if not s:
        return ''

    prefix=N(v319._weekly_task(d))
    if prefix:
        m=re.match(r'^\s*'+re.escape(prefix)+r'\s*[_\-/／|:：]*\s*',s,re.I)
        if m:
            rest=s[m.end():].strip(' _-/／|:：')
            if rest:
                return rest

    customer=N(d.get('customer'))
    if customer:
        # Old generated labels can contain a project name different from the GUI
        # selection. Remove the visible CUSTOMER_<project> prefix rather than
        # leaving it in a collision fallback title.
        m=re.match(
            r'^\s*'+re.escape(customer)+r'\s*[_\-/／|:：]\s*'
            r'([^_／|:：]+)\s*[_／|:：]\s*(.+)$',
            s,re.I
        )
        if m and N(m.group(2)):
            return N(m.group(2)).strip(' _-/／|:：')

    return s


def _detail_title_suffix(d,g):
    """Unified weekly detail-title tail: 발생샘플_발생처_이슈 발생."""
    sample=N((g or {}).get('sample') or d.get('sample') or d.get('occurrence_sample'))
    site=N((g or {}).get('occurrence_site') or d.get('occurrence_site'))
    parts=[x for x in (sample,site) if x]
    parts.append('이슈 발생')
    return '_'.join(parts)


def _detail_title_text(d,g):
    """고객사_과제명_발생샘플_발생처_이슈 발생."""
    project=N(v319._weekly_task(d))
    tail=_detail_title_suffix(d,g)
    return '_'.join(x for x in (project,tail) if x)


def _title_base_font_pt(sh):
    vals=[]
    try:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    vals.append(float(r.font.size.pt))
    except Exception:
        pass
    return max(vals) if vals else 14.0


def _set_title_exact(sh,text):
    if sh is None or not hasattr(sh,'text_frame'):
        return False
    tf=sh.text_frame
    runs=[r for p in tf.paragraphs for r in p.runs]
    if runs:
        runs[0].text=N(text)
        for r in runs[1:]:
            r.text=''
    else:
        tf.text=N(text)
    try:
        tf.word_wrap=False
    except Exception:
        pass
    return True


def _fit_full_title_before_owner(sl,sh,text):
    """Keep the full unified title and shrink it instead of deleting fields."""
    if sh is None or not hasattr(sh,'text_frame'):
        return False

    base_pt=_title_base_font_pt(sh)
    _set_title_exact(sh,text)

    available=None
    owner=_owner_shape(sl)
    if owner is not None:
        try:
            tx=float(sh.left)/EMU
            ox=float(owner.left)/EMU
            available=max(.55,ox-tx-.12)
            sh.width=int(available*EMU)
        except Exception:
            available=None

    if available is None:
        try:
            available=max(.55,float(sh.width)/EMU)
        except Exception:
            available=3.0

    # Manual first-pass sizing makes the result deterministic even when PowerPoint
    # ignores TEXT_TO_FIT_SHAPE until the file is opened.
    natural=max(.01,_estimated_text_width_in(sh,text))
    target=base_pt
    if natural>available:
        target=max(6.0,min(base_pt,base_pt*available/natural))
    try:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size=Pt(target)
    except Exception:
        pass

    try:
        sh.text_frame.word_wrap=False
        sh.text_frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass
    return True


def _force_page2_header(sl,d,g):
    customer=v319._customer_core(d.get('customer'))
    title=_detail_title_text(d,g)

    title_sh,issue_sh=v319._find_page2_header_shapes(sl)
    if title_sh is None:
        title_sh=_title_shape_by_geometry(sl)

    # Keep every title field. If it approaches the 담당자 area, shrink the font
    # and usable title width rather than dropping 고객사/과제명/샘플/발생처.
    step1._update_team_owner(sl,g)
    _fit_full_title_before_owner(sl,title_sh,title)

    if issue_sh is not None:
        issue=v319._trim_before_customer(d.get('issue_name'),customer)
        v319._set_issue_line(issue_sh,issue)


def weekly(src,out,d,g,mode):
    dd=v319._weekly_display(d)
    p1=dict(dd)
    p1['issue_name']=step1._page1_issue(d)
    p1['problem']=_full_problem(d)

    prs=Presentation(src)
    v319._update_page1(prs,p1,g)

    if len(prs.slides)>1:
        v320._update_page2(prs.slides[1],dd,g,mode)
        _force_page2_header(prs.slides[1],d,g)

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.weekly=weekly

class RecoveryStep4App(base.App):
    def __init__(self):
        super().__init__(); self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP4')

if __name__=='__main__':
    RecoveryStep4App().mainloop()
