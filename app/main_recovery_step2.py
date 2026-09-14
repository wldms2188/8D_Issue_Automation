# Recovery STEP2 based on v3.2.0 rendering + STEP1 safe field fixes.
# ONLY weekly meeting text/meta refinements; image/layout/font/Excel logic remain untouched.
import re
import datetime
from pathlib import Path
from pptx import Presentation

import main_recovery_step1 as step1
import main_v320 as v320
import main_v319 as v319
import main_v310 as v310

base=step1.base
N=v310.N
C=v310.C
EMU=v310.EMU


# ---------- event parsing ----------
def _event_name(d):
    """Return ('시험'|'빌드'|'', clean name).
    Example: 'ABC시험중' -> ('시험','ABC'); 'A1' -> ('빌드','A1').
    """
    customer=C(d.get('customer')); task=C(d.get('task_name'))
    sources=[N(d.get('issue_name')),N(d.get('problem'))]

    # Test: capture the name immediately before 시험 / 시험중, excluding the word 시험 itself.
    for src in sources:
        toks=[x.strip() for x in re.split(r'[_/|]+',src) if x.strip()]
        for t in toks:
            if C(t) in (customer,task):
                continue
            # Remove status suffix after 시험 (시험중, 시험 중, 시험완료, etc.) and return only name before 시험.
            m=re.search(r'([^\s,;:()]+?)\s*시험(?:\s*중|중|\s*완료|완료|\s*진행\s*중|진행중)?',t,re.I)
            if m:
                name=m.group(1).strip(' _-')
                if name:
                    return '시험',name

    # Build: A-D followed by digits.
    for src in sources:
        m=re.search(r'(?<![A-Za-z0-9])([A-D]\d+)(?![A-Za-z0-9])',src,re.I)
        if m:
            return '빌드',m.group(1).upper()
    return '',''


# ---------- symptom-only formatting ----------
_SECTION_HEADINGS=(
    '발생경위및확인사항','발생경위','확인사항','시험조건','시험조건및방법','시험방법','시험절차',
    '현상내용','불량현상','문제현상','현상','상세내용','참고사항','비고'
)
_META_WORDS=(
    '시험일자','시험장소','시험장비','장비','설비','샘플수','sample수','lot','로트',
    '온도','습도','soc','dod','soh','전류','전압조건','충전조건','방전조건','충방전조건',
    '시간','cycle','사이클','프로파일','profile','c-rate','crate'
)
_SYMPTOM_WORDS=step1._SYMPTOM_WORDS


def _clean_piece(s):
    s=N(s)
    s=re.sub(r'^[-•·▪◦]\s*','',s)
    s=re.sub(r'^\(?\d+\)?[.)]\s*','',s)
    return s.strip()


def _is_heading(line):
    q=C(line)
    if not q:
        return True
    return q.rstrip(':：') in _SECTION_HEADINGS


def _is_condition_or_meta(line):
    q=C(line)
    if not q:
        return False
    if any(C(w) in q for w in _META_WORDS):
        # If it also clearly states the failure symptom, keep it.
        if any(C(w) in q for w in _SYMPTOM_WORDS):
            return False
        return True
    return False


def _actual_symptoms(d):
    problem=N(d.get('problem'))
    kind,event=_event_name(d)
    kept=[]
    for raw in problem.split('\n'):
        line=_clean_piece(raw)
        if not line or _is_heading(line):
            continue
        # Split obvious compound lines, but preserve original symptom wording.
        pieces=[_clean_piece(x) for x in re.split(r'\s*[;/|]\s*',line) if _clean_piece(x)]
        for p in pieces:
            if not p or _is_heading(p):
                continue
            if event and C(p) in (C(event),C(event+'시험'),C(event+'시험중')):
                continue
            if _is_condition_or_meta(p):
                continue
            # Prefer actual symptom statements. If not a known symptom word, keep only concise descriptive text.
            if any(C(w) in C(p) for w in _SYMPTOM_WORDS):
                kept.append(p)
            elif len(p)<=70 and not re.search(r'[:：]\s*[0-9]',p):
                kept.append(p)

    # de-duplicate while preserving order
    out=[]; seen=set()
    for x in kept:
        k=C(x)
        if k and k not in seen:
            seen.add(k); out.append(x)

    # Safe fallback: first non-heading, non-condition line only.
    if not out:
        for raw in problem.split('\n'):
            line=_clean_piece(raw)
            if line and not _is_heading(line) and not _is_condition_or_meta(line):
                if not event or C(line) not in (C(event),C(event+'시험'),C(event+'시험중')):
                    out=[line]; break
    return out


def _short_problem(d):
    symptoms=_actual_symptoms(d)
    kind,event=_event_name(d)
    lines=['1) 현상']
    lines.extend(symptoms or ['(현상 내용 미기재)'])
    if kind=='시험' and event:
        lines.extend(['2) 시험명',event])
    elif kind=='빌드' and event:
        lines.extend(['2) 빌드명',event])
    return '\n'.join(lines)


def _title_suffix(d):
    kind,event=_event_name(d)
    if kind=='시험' and event:
        return f'{event} 시험 이슈 발생'
    if kind=='빌드' and event:
        return f'{event} 빌드 이슈 발생'
    # Non-test/build issue keeps existing issue categorization.
    issue=v319._trim_before_customer(d.get('issue_name'),d.get('customer'))
    return issue


def _page2_title_task(d):
    customer=N(d.get('customer')); task=N(d.get('task_name'))
    if customer and task:
        return f'{customer}_{task}'
    return customer or task


def _force_page2_header(sl,d,g):
    """Explicitly set page2 title and owner after normal metadata update."""
    task=_page2_title_task(d)
    suffix=_title_suffix(d)
    title_sh,issue_sh=v319._find_page2_header_shapes(sl)
    if title_sh is not None:
        v319._replace_title_runs(title_sh,task,suffix)
    # Keep detailed issue line from customer onward as recovery baseline intended.
    if issue_sh is not None:
        issue=v319._trim_before_customer(d.get('issue_name'),d.get('customer'))
        v319._set_issue_line(issue_sh,issue)
    step1._update_team_owner(sl,g)


# ---------- weekly wrapper ----------
def weekly(src,out,d,g,mode):
    dd=v319._weekly_display(d)
    dd['_title_issue']=_title_suffix(d)

    # Page 1 only: remove customer_task prefix from issue and use structured short problem.
    p1=dict(dd)
    p1['issue_name']=step1._page1_issue(d)
    p1['problem']=_short_problem(d)

    prs=Presentation(src)
    v319._update_page1(prs,p1,g)

    if len(prs.slides)>1:
        # Keep v3.2.0 rendering/layout/images exactly as-is.
        v320._update_page2(prs.slides[1],dd,g,mode)
        _force_page2_header(prs.slides[1],d,g)

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved


base.weekly=weekly


class RecoveryStep2App(base.App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP2')


if __name__=='__main__':
    RecoveryStep2App().mainloop()
