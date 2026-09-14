# 8D Issue Automation v2.9.6 - stable integration build
# IMPORTANT: source 8D extraction uses the ORIGINAL PPT directly.
# Representative image selection is restricted to PAGE 1 and the 2D/4D regions.
# No temporary/copy PPT is created for image extraction.
import io, datetime, math, os
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image as PILImage

APP_DIR=Path(__file__).resolve().parent
if str(APP_DIR) not in __import__('sys').path: __import__('sys').path.insert(0,str(APP_DIR))
import main_v293 as _impl
base=_impl.base
v29=_impl.v29
EMU=914400

BOXES={'1':[.53,2.53,4.76,.98,'problem'],'2':[.53,3.91,4.76,1.14,'temporary_action'],'3-1':[.53,5.38,4.76,1.79,'cause_4d'],'3-2':[5.77,2.47,4.76,.87,'leak_cause'],'4':[5.77,3.74,4.76,1.77,'action_5d'],'5':[5.77,5.91,4.78,.94,'verification_6d']}
LEFT=['1','2','3-1']; RIGHT=['3-2','4','5']
NAMES={'1':'2D 현상','2':'3D 임시대책(필요시)','3-1':'4D 원인분석','3-2':'4D 원인분석','4':'5D 개선대책','5':'6D 유효성점검'}
DL={'1':'1D','2':'2D','3-1':'3D','3-2':'4D','4':'5D','5':'6D'}

def norm(x): return '\n'.join(str(x or '').replace('\r\n','\n').replace('\r','\n').splitlines()).strip()
def font(r,s=8):
    r.font.name='맑은 고딕';r.font.size=Pt(s)
    try:
        rp=r._r.get_or_add_rPr();rp.set('a:latin','맑은 고딕');rp.set('a:ea','맑은 고딕');rp.set('a:cs','맑은 고딕')
    except Exception:pass
def put(sh,text,size=8):
    tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=Inches(.05);tf.margin_right=Inches(.05);tf.margin_top=Inches(.03);tf.margin_bottom=Inches(.03)
    ls=norm(text).split('\n') if norm(text) else ['']
    for i,line in enumerate(ls):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        for r in p.runs:font(r,size)
def lines(text,w,size=8):
    c=max(8,int(w*10.5*(8/size)));return sum(max(1,math.ceil(len(x)/c)) for x in norm(text).split('\n'))
def need(text,w,size=8):return max(.45,lines(text,w,size)*(size/72*1.20)+.10)

def pics(sh):
    if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.PICTURE:return [sh]
    if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:
        z=[]
        for c in sh.shapes:z+=pics(c)
        return z
    return []
def overlap(a,b):
    ax,ay,aw,ah=a;bx,by,bw,bh=b
    return max(0,min(ax+aw,bx+bw)-max(ax,bx))*max(0,min(ay+ah,by+bh)-max(ay,by))
def box(sh):return float(sh.left),float(sh.top),float(sh.width),float(sh.height)

def composite(g):
    ps=pics(g)
    if len(ps)<2:return None
    x=min(p.left for p in ps);y=min(p.top for p in ps);xx=max(p.left+p.width for p in ps);yy=max(p.top+p.height for p in ps);sc=140/EMU
    imout=PILImage.new('RGB',(max(1,int((xx-x)*sc)),max(1,int((yy-y)*sc))),'white')
    for p in ps:
        try:
            with PILImage.open(io.BytesIO(p.image.blob)) as im:
                im=im.convert('RGB');im.thumbnail((max(1,int(p.width*sc)),max(1,int(p.height*sc))),PILImage.Resampling.LANCZOS);imout.paste(im,(int((p.left-x)*sc),int((p.top-y)*sc)))
        except Exception:pass
    b=io.BytesIO();imout.save(b,'PNG');return b.getvalue()

# ---------- representative image: PAGE 1 / 2D or 4D only ----------
def walk_shapes(shapes):
    for sh in shapes:
        yield sh
        if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(sh.shapes)

def region_labels(slide):
    """Find the actual 2D/4D labeled regions on page 1."""
    regs=[]
    for sh in walk_shapes(slide.shapes):
        q=base.compact(getattr(sh,'text',''))
        if not q: continue
        if ('2d' in q or '문제현상' in q or '문제현황' in q or '불량현상' in q):
            regs.append(('2D',box(sh)))
        elif ('4d' in q or '원인분석' in q or '발생원인' in q or '유출원인' in q):
            regs.append(('4D',box(sh)))
    return regs

def page1_visuals(slide):
    out=[]
    for sh in slide.shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
            ps=pics(sh)
            if len(ps)>=2:
                b=composite(sh)
                if b:
                    x,y,w,h=box(sh);out.append(('group',len(ps),(x,y,w,h),b))
            elif len(ps)==1:
                p=ps[0]
                try:
                    with PILImage.open(io.BytesIO(p.image.blob)) as im:
                        if im.width*im.height>=10000:out.append(('picture',1,box(sh),p.image.blob))
                except Exception:pass
        elif sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
            try:
                with PILImage.open(io.BytesIO(sh.image.blob)) as im:
                    if im.width*im.height>=10000:out.append(('picture',1,box(sh),sh.image.blob))
            except Exception:pass
    return out

def pick_page1_2d4d(prs):
    """Choose only a visual that overlaps the labeled 2D/4D area on page 1.
    Priority: 2D > 4D, grouped visual > single picture, larger overlap > area.
    """
    if not prs.slides:return None
    sl=prs.slides[0]
    regs=region_labels(sl)
    visuals=page1_visuals(sl)
    if not regs or not visuals:return None
    scored=[]
    for kind,n,vbox,blob in visuals:
        for label,rbox in regs:
            ov=overlap(vbox,rbox)
            if ov<=0:continue
            vx,vy,vw,vh=vbox;rx,ry,rw,rh=rbox
            # Overlap ratio is more reliable than raw area when the image is large.
            vr=max(vw*vh,1.0);rr=max(rw*rh,1.0)
            ratio=ov/min(vr,rr)
            priority=2 if label=='2D' else 1
            grouped=1 if kind=='group' else 0
            scored.append((priority,grouped,ratio,ov,vw*vh,blob,label))
    if not scored:return None
    scored.sort(key=lambda z:(z[0],z[1],z[2],z[3],z[4]),reverse=True)
    return scored[0][5]

def extract_v296(path):
    # DO NOT clone/copy the source PPT. Read the original package directly.
    d=_impl.extract_v293(path)
    try:
        b=pick_page1_2d4d(Presentation(path))
        d['_images']=[(1,1,1,b)] if b else []
    except Exception:
        # Text/Excel automation must not fail just because a representative image
        # cannot be found. The image is optional; the extracted 8D data remains valid.
        d['_images']=[]
    return d

base.extract=extract_v296

def walk(shapes):
    for sh in shapes:
        yield sh
        if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:yield from walk(sh.shapes)
def remove(sh):
    try:e=sh._element;e.getparent().remove(e)
    except Exception:pass

def make_detail(sl,d):
    for sh in list(walk(sl.shapes)):
        try:
            if str(sh.name).startswith('AUTO_8D296_'):remove(sh)
        except Exception:pass
    def one(k):return v29.v29_one(d.get(k)) or '검토 중'
    txt={'1':one('problem'),'2':one('temporary_action'),'3-1':'- 발생원인\n'+one('cause_4d'),'3-2':'- 유출원인\n'+one('leak_cause'),'4':one('action_5d'),'5':one('verification_6d')}
    for col,mx in ((LEFT,.25),(RIGHT,5.49)):
        pos={k:list(BOXES[k][:4]) for k in col}
        for i,k in enumerate(col):
            x,y,w,h=pos[k];pos[k][3]=max(h,need(txt[k],w,8))
            if i+1<len(col):pos[col[i+1]][1]=max(pos[col[i+1]][1],y+pos[k][3]+.14+.30+.10)
        for k in col:
            x,y,w,h=pos[k];dl=DL[k]
            m=sl.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL,Inches(mx),Inches(y-.38),Inches(.30),Inches(.30));m.name='AUTO_8D296_MARKER_'+dl;m.text=dl
            try:m.fill.background();m.line.width=Pt(.8)
            except Exception:pass
            for p in m.text_frame.paragraphs:
                p.alignment=1
                for r in p.runs:font(r,6.5)
            hd=sl.shapes.add_textbox(Inches(mx+.38),Inches(y-.39),Inches(2.5),Inches(.34));hd.name='AUTO_8D296_HEADING_'+dl;put(hd,NAMES[k],8)
            size=8 if need(txt[k],w,8)<=h+.02 else 7
            b=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));b.name='AUTO_8D296_DETAIL_'+k.replace('-','_');put(b,txt[k],size)
    for sh in walk(sl.shapes):
        if hasattr(sh,'text_frame'):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:font(r,8)

def weekly_v296(src,out,d,g,mode):
    prs=Presentation(src);st=base.status(d);title=base.task(d);issue=v29.v29_one(d.get('issue_name'));summary=None;idx=0
    for i,sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if getattr(sh,'has_table',False):
                heads=' '.join(base.norm(sh.table.cell(0,c).text) for c in range(len(sh.table.columns)))
                if '과제명' in heads and 'Signal' in heads:summary=sh.table;idx=i;break
        if summary:break
    if summary:
        row=1
        for r in range(1,len(summary.rows)):
            if not any(base.norm(summary.cell(r,c).text) for c in range(min(5,len(summary.columns))) if c!=4):row=r;break
        vals=[title,issue,norm(d.get('problem')),v29.v29_progress(d),'●']
        for c,v in enumerate(vals):
            if c<len(summary.columns):v29.v29_set_cell_text(summary.cell(row,c),v,8)
        if len(summary.columns)>4:base.signal(summary.cell(row,4),st)
    for sl in prs.slides[idx+1:]:make_detail(sl,d)
    Path(out).parent.mkdir(exist_ok=True)
    try:prs.save(out);saved=out
    except PermissionError:
        p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);prs.save(saved)
    return '주간회의 PPT 업데이트: '+st,saved

base.extract=extract_v296;base.weekly=weekly_v296
if __name__=='__main__':base.App().mainloop()
