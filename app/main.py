import os,re,copy,datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk,filedialog,messagebox
from pptx import Presentation
from pptx.dml.color import RGBColor
from openpyxl import load_workbook

STATUS={'complete':'개선 완료','verify':'개선 검증중','unknown':'원인/개선 미확인'}
SIGNAL={'개선 완료':(0,176,80),'개선 검증중':(255,192,0),'원인/개선 미확인':(255,0,0)}
STD=['D-FMEA','P-FMEA','Control plan','작업 표준','제품규격']

def norm(x):
    if x is None:return ''
    return re.sub(r'\n+','\n',re.sub(r'[ \t]+',' ',str(x).replace('\r','\n'))).strip()
def one(x,n=None):
    s=re.sub(r'\s+',' ',norm(x))
    return s if not n or len(s)<=n else s[:n-1]+'…'
def blank(x):
    x=norm(x); return not x or x in ['000','00','‘00.00.00','00호기','담당자/ 반영일','유첨 5why','대책에 대한 효과가 검증되어','10개 Lot 모니터링등','본 이슈로 인한 비용 및 산출 근거']
def key(x):return re.sub(r'[\s_/\\\-:.,·]+','',one(x).lower())
def cell(tb,r,c):
    return norm(tb.cell(r,c).text) if 0<=r<len(tb.rows) and 0<=c<len(tb.columns) else ''
def near(tb,labels):
    labels=[norm(x).lower() for x in labels]
    for r,row in enumerate(tb.rows):
        for c,cl in enumerate(row.cells):
            raw=norm(cl.text); low=raw.lower()
            if raw and any(x in low for x in labels):
                cand=[]
                if ':' in raw:cand.append(raw.split(':',1)[1].strip())
                cand += [cell(tb,r,cc) for cc in range(c+1,len(tb.columns))]
                cand += [cell(tb,rr,c) for rr in range(r+1,min(r+3,len(tb.rows)))]
                for v in cand:
                    if v and not blank(v):return v
    return ''

def extract(p):
    prs=Presentation(p); d={k:'' for k in ['issue_name','task_name','customer','occurrence_site','occurrence_date','receipt_date','problem','temporary_action','customer_response','cause_4d','leak_cause','system_cause','action_5d','verification_6d','spread_7d']}; d['standards']={x:'' for x in STD}
    texts=[]
    for sl in prs.slides:
        for sh in sl.shapes:
            if hasattr(sh,'text') and norm(sh.text):
                t=norm(sh.text);texts.append(t)
                if '이슈명' in t and ':' in t and not d['issue_name']:d['issue_name']=t.split(':',1)[1].strip()
            if getattr(sh,'has_table',False):
                tb=sh.table
                mp={'task_name':['과제명'],'customer':['Customer','고객사'],'occurrence_site':['발생 Site'],'occurrence_date':['발생 일자'],'receipt_date':['접수 일자'],'problem':['불량 현상','문제 현황'],'temporary_action':['임시 조치'],'customer_response':['고객 대응'],'cause_4d':['발생 원인'],'leak_cause':['유출 원인'],'system_cause':['IT 시스템'],'action_5d':['개선 대책','대책'],'verification_6d':['효과검증','효과 검증'],'spread_7d':['수평전개','재발방지']}
                for k,ls in mp.items():d[k]=d[k] or near(tb,ls)
                for s in STD:d['standards'][s]=d['standards'][s] or near(tb,[s])
    full='\n'.join(texts)
    for k,ls in {'issue_name':['이슈명'],'task_name':['과제명'],'customer':['Customer','고객사'],'occurrence_site':['발생 Site'],'problem':['불량 현상'],'cause_4d':['발생 원인'],'action_5d':['개선 대책'],'verification_6d':['효과검증'],'spread_7d':['수평전개']}.items():
        if not d[k]:
            for ln in full.splitlines():
                for lab in ls:
                    if ln.lower().startswith(lab.lower()) and ':' in ln:
                        d[k]=ln.split(':',1)[1].strip();break
                if d[k]:break
    return d

def status(d):
    v=one(d.get('verification_6d'))
    if re.search(r'완료\s*(예정|추정)',v):return STATUS['verify']
    if re.search(r'검증\s*완료|(?<![가-힣])완료(?![가-힣])',v):return STATUS['complete']
    if re.search(r'진행\s*중|진행중|검증\s*중|검증중|예정|추정',v):return STATUS['verify']
    if blank(d.get('cause_4d')) and blank(d.get('action_5d')):return STATUS['unknown']
    return STATUS['verify']
def progress(d):
    a=[];c=one(d.get('cause_4d') or d.get('leak_cause') or d.get('system_cause'),100);x=one(d.get('action_5d'),100);v=one(d.get('verification_6d'),100)
    if c:a.append('원인: '+c)
    if x or v:a.append('진행 현황: '+' / '.join(q for q in [x,v] if q))
    miss=[s for s in STD if blank(d.get('standards',{}).get(s,''))]
    if miss:a.append('요청 사항: '+', '.join(miss)+' 관련 표준 개정 요청')
    return '\n'.join(a) or '원인/개선 내용 미기재'
def task(d):
    return ('%s / %s'%(one(d.get('customer')),one(d.get('task_name')))).strip(' /')
def sig(ws,r):return {'plm':key(ws.cell(r,3).value),'task':key(ws.cell(r,10).value),'problem':key(ws.cell(r,14).value),'cause':key(ws.cell(r,16).value),'action':key(ws.cell(r,17).value)}
def find(ws,d,g):
    p=key(g.get('plm_no'))
    if p:
        for r in range(7,ws.max_row+1):
            if sig(ws,r)['plm']==p:return r,100
    tt,pp,cc,aa=key(task(d)),key(d.get('problem')),key(d.get('cause_4d')),key(d.get('action_5d'));best=(None,0)
    for r in range(7,ws.max_row+1):
        s=sig(ws,r);score=(5 if tt and s['task']==tt else 0)+(5 if pp and s['problem']==pp else 0)+(2 if cc and s['cause']==cc else 0)+(2 if aa and s['action']==aa else 0)
        if score>best[1]:best=(r,score)
    return best if best[1]>=10 else (None,0)
def copyfmt(ws,a,b):
    ws.row_dimensions[b].height=ws.row_dimensions[a].height
    for c in range(1,ws.max_column+1):ws.cell(b,c)._style=copy.copy(ws.cell(a,c)._style)
def write_row(ws,r,d,g):
    dt=one(d.get('occurrence_date')); m=re.match(r'(\d{4})[-./](\d{1,2})',dt);st=status(d)
    vals={2:datetime.date.today(),3:g.get('plm_no',''),4:g.get('form_factor',''),5:g.get('product_type',''),6:m.group(1) if m else '',7:m.group(2) if m else '',8:g.get('team',''),9:g.get('owner',''),10:task(d),11:g.get('sample',''),12:g.get('stage',''),13:d.get('occurrence_site',''),14:one(d.get('problem'),120),15:'',16:one(' / '.join(x for x in [d.get('cause_4d'),d.get('leak_cause'),d.get('system_cause')] if x),160),17:one(d.get('action_5d'),160),18:st}
    for c,v in vals.items():ws.cell(r,c).value=v
    return st
def update_excel(src,out,d,g,new=False):
    wb=load_workbook(src);ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active;r,score=find(ws,d,g)
    if r and not new:
        old=ws.cell(r,3).value;write_row(ws,r,d,g)
        if not g.get('plm_no') and old:ws.cell(r,3).value=old
        msg=f'기존 이슈 업데이트 (row {r}, match score {score})'
    else:
        r=ws.max_row+1
        if r>7:copyfmt(ws,r-1,r)
        write_row(ws,r,d,g);msg=f'신규 이슈 추가 (row {r})'
    Path(out).parent.mkdir(exist_ok=True);wb.save(out);return msg
def signal(c,s):
    c.text='●';rgb=SIGNAL[s]
    for p in c.text_frame.paragraphs:
        for run in p.runs:run.font.color.rgb=RGBColor(*rgb);run.font.bold=True
def clone(prs,sl):
    ns=prs.slides.add_slide(sl.slide_layout)
    for sh in sl.shapes:ns.shapes._spTree.insert_element_before(copy.deepcopy(sh.element),'p:extLst')
    return ns
def weekly(src,out,d,g,mode):
    prs=Presentation(src);st=status(d);t=task(d);i=one(d.get('issue_name'));summary=None
    for sl in prs.slides:
        for sh in sl.shapes:
            if getattr(sh,'has_table',False) and sh.table.rows and '과제명' in cell(sh.table,0,0) and 'Signal' in ' '.join(cell(sh.table,0,c) for c in range(len(sh.table.columns))):summary=sh.table;break
        if summary:break
    if summary:
        r=None
        for x in range(1,len(summary.rows)):
            if key(t) in key(cell(summary,x,0)) and key(i) in key(cell(summary,x,1)):r=x;break
        if r is None:
            for x in range(1,len(summary.rows)):
                if not any(cell(summary,x,c) for c in range(len(summary.columns)) if c!=4):r=x;break
        if r is None:summary._tbl.append(copy.deepcopy(summary.rows[-1]._tr));r=len(summary.rows)-1
        for c,v in enumerate([t,i,one(d.get('problem'),100),progress(d),'●']):summary.cell(r,c).text=v
        signal(summary.cell(r,4),st)
    target=None
    if mode=='existing':
        for sl in prs.slides:
            z=key(' '.join(norm(getattr(sh,'text','')) for sh in sl.shapes if hasattr(sh,'text')))
            if key(t) and key(i) and key(t) in z and key(i) in z:target=sl;break
    if target is None:target=clone(prs,prs.slides[1] if len(prs.slides)>1 else prs.slides[0])
    for sh in target.shapes:
        if hasattr(sh,'text'):
            if '과제명_이슈 제목' in sh.text:sh.text=f'{t}_{i}'
            elif '이슈명 :' in sh.text:sh.text='이슈명 : '+i
            elif '00팀 담당자 : 000' in sh.text:sh.text=f"{g.get('team','')} 담당자 : {g.get('owner','')}"
    tables=[sh.table for sh in target.shapes if getattr(sh,'has_table',False)]
    if tables:
        tb=tables[0]
        if len(tb.rows)>=2:tb.cell(1,1).text='\n'.join([f'2D 현상: {one(d.get("problem"),180)}',f'3D 임시대책: {one(d.get("temporary_action"),160)}',f'4D 원인분석: {one(d.get("cause_4d"),180)}',f'5D 개선대책: {one(d.get("action_5d"),180)}',f'6D 유효성점검: {one(d.get("verification_6d"),180)}',f'7D 수평전개: {one(d.get("spread_7d"),180)}'])
    if len(tables)>1 and len(tables[1].rows)>=3:
        tables[1].cell(0,1).text='●';signal(tables[1].cell(0,1),st);tables[1].cell(1,1).text=one(d.get('task_name') or d.get('issue_name'));tables[1].cell(2,1).text=one(d.get('occurrence_date'))
    Path(out).parent.mkdir(exist_ok=True);prs.save(out);return '주간회의 PPT 업데이트: '+st

class App(tk.Tk):
    def __init__(self):
        super().__init__();self.title('8D 이슈 자동화 v2.3');self.geometry('780x680');self.vars={};self.mode=tk.StringVar(value='existing')
        f=[('8D PPT','ppt8d',1),('주간회의 PPT','pptweekly',1),('이슈 DB Excel','xlsx',1),('담당팀','team',0),('담당자','owner',0),('폼팩터','form_factor',0),('제품 타입','product_type',0),('발생 샘플','sample',0),('개발 단계','stage',0),('PMS/PLM 이슈번호(기존값이 있으면 입력)','plm_no',0)]
        q=ttk.Frame(self);q.pack(fill='x',padx=18,pady=8);ttk.Label(q,text='이슈 구분',width=31).pack(side='left');ttk.Radiobutton(q,text='기존 이슈',variable=self.mode,value='existing').pack(side='left');ttk.Radiobutton(q,text='신규 이슈',variable=self.mode,value='new').pack(side='left')
        for lab,k,isfile in f:
            self.vars[k]=tk.StringVar();r=ttk.Frame(self);r.pack(fill='x',padx=18,pady=4);ttk.Label(r,text=lab,width=31).pack(side='left');ttk.Entry(r,textvariable=self.vars[k],width=66).pack(side='left')
            if isfile:ttk.Button(r,text='찾기',command=lambda x=k:self.pick(x)).pack(side='left',padx=5)
        self.log=tk.Text(self,height=14,width=95);self.log.pack(padx=18,pady=10);ttk.Button(self,text='8D 내용 미리보기',command=self.preview).pack(pady=4);ttk.Button(self,text='자동 업데이트 실행',command=self.run).pack(pady=4)
    def pick(self,k):
        ft=[('Excel','*.xlsx')] if k=='xlsx' else [('PowerPoint','*.pptx')];f=filedialog.askopenfilename(filetypes=ft)
        if f:self.vars[k].set(f)
    def gui(self):return {k:v.get().strip() for k,v in self.vars.items()}
    def preview(self):
        g=self.gui()
        if not g['ppt8d']:return messagebox.showwarning('확인','8D PPT를 선택하세요.')
        try:
            d=extract(g['ppt8d']);self.log.delete('1.0','end');self.log.insert('end','=== 8D 추출 결과 ===\n')
            for k in ['issue_name','task_name','customer','occurrence_site','occurrence_date','problem','cause_4d','action_5d','verification_6d','spread_7d']:self.log.insert('end',f'{k}: {d.get(k)}\n')
            self.log.insert('end',f'\n=== 상태 판정 ===\n{status(d)}\n\n=== 주간회의 진행사항 ===\n{progress(d)}')
        except Exception as e:messagebox.showerror('오류',repr(e))
    def run(self):
        g=self.gui()
        if any(not g[k] or not os.path.exists(g[k]) for k in ['ppt8d','pptweekly','xlsx']):return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')
        try:
            d=extract(g['ppt8d']);mode=self.mode.get();out=Path(g['xlsx']).parent/'자동화_결과';out.mkdir(exist_ok=True);xo=out/(Path(g['xlsx']).stem+'_업데이트.xlsx');po=out/(Path(g['pptweekly']).stem+'_업데이트.pptx');wb=load_workbook(g['xlsx']);ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active;r,_=find(ws,d,g)
            if mode=='existing' and not r:return messagebox.showwarning('업데이트 중단','기존 이슈를 Excel에서 정확하게 찾지 못했습니다. PMS/PLM 이슈번호를 입력하거나 이슈 정보를 확인하세요.')
            if mode=='new' and r and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?'):return
            a=update_excel(g['xlsx'],xo,d,g,new=(mode=='new'));b=weekly(g['pptweekly'],po,d,g,mode);self.log.delete('1.0','end');self.log.insert('end',f'=== 완료 ===\n이슈 구분: {mode}\n상태: {status(d)}\n\n{a}\n{b}\n\n결과: {out}');messagebox.showinfo('완료',f'자동 업데이트가 완료되었습니다.\n\n{out}')
        except Exception as e:messagebox.showerror('실행 오류',repr(e))
if __name__=='__main__':App().mainloop()
