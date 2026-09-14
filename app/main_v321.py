# 8D Issue Automation v3.2.1
# Guided user inputs + weekly page refinements + Issue DB formatting rules.
import os, re, io, copy, datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from pptx import Presentation
from pptx.util import Pt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage

import main_v320 as v320
import main_v319 as v319
import main_v318 as v318
import main_v314 as v314
import main_v310 as v310

base = v320.base
N = v310.N
C = v310.C

FORM_FACTORS = ('원통형','파우치형')
PRODUCT_TYPES = ('ESS BPU','ESS Link','EV Cell-Unit','EV Module','EV Pack','IT Pack','LEV Pack','PHEV Pack')
OCCURRENCE_SITES = ('etc.','고객/상위 Claim','고객/상위 생산','고객/상위 시험','고객/상위 운송/보관','부품 생산','제품 생산','제품 시험','제품 운송/보관')
STAGES = ('CV','DV','PD','MP')
ISSUE_ORIGINS = ('부품','설계','공정','기타','논의 중')


def _cause_known(d):
    txt=' '.join(N(d.get(k)) for k in ('cause_4d','leak_cause','system_cause'))
    q=C(txt)
    if not q:
        return False
    return not any(x in q for x in ('tbd','미확인','확인중','검토중','원인미상'))


def _event_name(d):
    """Return ('시험'|'빌드'|'', specific event name). Prefer issue filename/detail over problem text."""
    customer=N(d.get('customer')); task=N(d.get('task_name'))
    sources=[N(d.get('issue_name')),N(d.get('problem'))]
    for src in sources:
        if not src: continue
        toks=[x.strip() for x in re.split(r'[_/|]+',src) if x.strip()]
        # Ignore customer/task tokens and choose the first explicit test/build token.
        filtered=[t for t in toks if C(t) not in (C(customer),C(task))]
        for t in filtered:
            m=re.search(r'([^\s,;:()]+?시험)',t)
            if m: return '시험',m.group(1).strip()
        for t in filtered:
            m=re.search(r'([^\s,;:()]+?(?:빌드|build))',t,re.I)
            if m: return '빌드',m.group(1).strip()
    # Fallback: category only.
    full=' '.join(sources)
    if '시험' in full: return '시험','시험'
    if '빌드' in full or re.search(r'\bbuild\b',full,re.I): return '빌드','빌드'
    return '',''


def _title_issue(d):
    kind,name=_event_name(d)
    if kind and name:
        return f'{name} 이슈 발생'
    return v319._title_issue(v319._trim_before_customer(d.get('issue_name'),d.get('customer')))


def _short_problem(d):
    kind,name=_event_name(d)
    if kind=='시험': return f'시험명: {name}'
    if kind=='빌드': return f'빌드명: {name}'
    # If it is neither test nor build, keep only the first meaningful line.
    for line in N(d.get('problem')).split('\n'):
        if line.strip(): return line.strip()
    return N(d.get('problem'))


def _problem_is_long(d):
    p=N(d.get('problem'))
    return len(p)>=90 or len([x for x in p.split('\n') if x.strip()])>=5


def _weekly_display(d,short_problem=False):
    dd=v319._weekly_display(d)
    dd['_title_issue']=_title_issue(d)
    if short_problem:
        dd['problem']=_short_problem(d)
    return dd


def _set_text_preserve_first_run(sh,text):
    if not hasattr(sh,'text_frame'): return
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    if runs:
        runs[0].text=N(text)
        for r in runs[1:]: r.text=''
    else:
        sh.text_frame.text=N(text)


def _update_team_owner(sl,g):
    target=f"{N(g.get('team'))} 담당자 : {N(g.get('owner'))}".strip()
    for sh in v310.walk(sl):
        if not hasattr(sh,'text_frame'): continue
        t=N(getattr(sh,'text',''))
        if '담당자' in t and float(getattr(sh,'top',99999999))/v310.EMU < .75:
            _set_text_preserve_first_run(sh,target)
            return True
    return False


def _set_issue_origin(sl,value):
    value=N(value) or 'TBD'
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False): continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                if '이슈기인' in N(tb.cell(r,c).text) and c+1<len(tb.columns):
                    v314._set_cell(tb.cell(r,c+1),value,8)
                    return True
    return False


_original_page2_meta=v319._page2_meta

def _page2_meta(sl,d,g):
    _original_page2_meta(sl,d,g)
    _update_team_owner(sl,g)
    _set_issue_origin(sl,g.get('issue_origin') or 'TBD')


def weekly(src,out,d,g,mode):
    dd=_weekly_display(d,short_problem=bool(g.get('_short_problem')))
    prs=Presentation(src)
    v319._update_page1(prs,dd,g)
    if len(prs.slides)>1:
        old=v319._page2_meta
        try:
            v319._page2_meta=_page2_meta
            v320._update_page2(prs.slides[1],dd,g,mode)
        finally:
            v319._page2_meta=old
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out); saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix)); prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved


# ---------- Issue DB rules ----------
def _excel_status(d):
    txt=' '.join(N(d.get(k)) for k in ('action_5d','verification_6d'))
    return 'close' if re.search(r'(개선\s*완료|검증\s*완료)',txt,re.I) else 'open'


def _month_label(m):
    try: return f'{int(str(m).strip())}월'
    except Exception: return (str(m).strip()+'월') if str(m).strip() else ''


def _add_excel_photo(ws,row,blob):
    if not blob: return
    try:
        for old in list(getattr(ws,'_images',[])):
            try:
                a=old.anchor._from
                if a.row==row-1 and a.col==14: ws._images.remove(old)
            except Exception: pass
        with PILImage.open(io.BytesIO(blob)) as im:
            w,h=im.size
        col_width=ws.column_dimensions['O'].width or 20
        max_w=max(60,int(col_width*7)-8); max_h=95
        scale=min(max_w/max(w,1),max_h/max(h,1))
        new_w=max(1,int(w*scale)); new_h=max(1,int(h*scale))
        img=XLImage(io.BytesIO(blob)); img.width=new_w; img.height=new_h
        # TwoCellAnchor(editAs='twoCell') = Excel 'Move and size with cells'.
        fr=AnchorMarker(col=14,row=row-1,colOff=pixels_to_EMU(4),rowOff=pixels_to_EMU(4))
        to=AnchorMarker(col=14,row=row-1,colOff=pixels_to_EMU(4+new_w),rowOff=pixels_to_EMU(4+new_h))
        img.anchor=TwoCellAnchor(editAs='twoCell',_from=fr,to=to)
        ws.add_image(img)
        ws.row_dimensions[row].height=max(float(ws.row_dimensions[row].height or 15),min(105,new_h*.75+8))
    except Exception:
        # Keep output usable even when a workbook has unusual drawing XML.
        try:
            img=XLImage(io.BytesIO(blob)); ws.add_image(img,f'O{row}')
        except Exception: pass


def _write_row(ws,r,d,g):
    y,m=base.parse_year_month(d.get('occurrence_date'))
    cause='\n'.join(x for x in [d.get('cause_4d'),d.get('leak_cause'),d.get('system_cause')] if N(x))
    vals={
        2:datetime.date.today(),3:g.get('plm_no',''),4:g.get('form_factor',''),5:g.get('product_type',''),
        6:y,7:_month_label(m),8:g.get('team',''),9:g.get('owner',''),10:base.task(d),11:g.get('sample',''),
        12:g.get('stage',''),13:g.get('occurrence_site') or d.get('occurrence_site',''),14:N(d.get('problem')),15:'',
        16:cause,17:N(d.get('action_5d')),18:_excel_status(d)
    }
    for c,v in vals.items(): ws.cell(r,c).value=v
    _add_excel_photo(ws,r,base.representative_image(d))
    return vals[18]

# update_excel resolves these module globals through the shared base module object.
base.write_row=_write_row
base.add_excel_photo=_add_excel_photo
base.weekly=weekly
base.extract=v318.extract


class ChoiceDialog(tk.Toplevel):
    def __init__(self,parent,title,prompt,values):
        super().__init__(parent); self.title(title); self.resizable(False,False); self.result=None
        self.transient(parent); self.grab_set()
        ttk.Label(self,text=prompt,justify='left').pack(padx=20,pady=(18,8))
        self.var=tk.StringVar(value=values[0])
        cb=ttk.Combobox(self,textvariable=self.var,values=values,state='readonly',width=28); cb.pack(padx=20,pady=8); cb.current(0)
        row=ttk.Frame(self); row.pack(pady=(5,18))
        ttk.Button(row,text='확인',command=self.ok).pack(side='left',padx=5)
        ttk.Button(row,text='취소',command=self.cancel).pack(side='left',padx=5)
        self.protocol('WM_DELETE_WINDOW',self.cancel); self.wait_visibility(); self.focus_set(); self.wait_window(self)
    def ok(self): self.result=self.var.get(); self.destroy()
    def cancel(self): self.result=None; self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__(); self.title('8D 이슈 자동화 v3.2.1'); self.geometry('860x880'); self.vars={}; self.mode=tk.StringVar(value='existing')
        q=ttk.Frame(self); q.pack(fill='x',padx=18,pady=8)
        ttk.Label(q,text='이슈 구분',width=31).pack(side='left')
        ttk.Radiobutton(q,text='기존 이슈',variable=self.mode,value='existing').pack(side='left')
        ttk.Radiobutton(q,text='신규 이슈',variable=self.mode,value='new').pack(side='left')
        self._file_row('8D PPT','ppt8d','ppt')
        self._file_row('주간회의 PPT','pptweekly','ppt')
        self._file_row('이슈 DB Excel','xlsx','xlsx')
        self._entry_row('담당팀','team'); self._entry_row('담당자','owner')
        self._combo_row('폼팩터','form_factor',FORM_FACTORS)
        self._combo_row('제품 타입','product_type',PRODUCT_TYPES)
        self._entry_row('발생 샘플','sample')
        self._combo_row('발생 단계','stage',STAGES)
        self._combo_row('발생처','occurrence_site',OCCURRENCE_SITES)
        self._entry_row('PMS/PLM 이슈번호(기존값이 있으면 입력)','plm_no')
        self.log=tk.Text(self,height=15,width=104); self.log.pack(padx=18,pady=10)
        ttk.Button(self,text='8D 내용 미리보기',command=self.preview).pack(pady=4)
        ttk.Button(self,text='자동 업데이트 실행',command=self.run).pack(pady=4)

    def _row(self,lab):
        r=ttk.Frame(self); r.pack(fill='x',padx=18,pady=3); ttk.Label(r,text=lab,width=31).pack(side='left'); return r
    def _entry_row(self,lab,k):
        self.vars[k]=tk.StringVar(); r=self._row(lab); ttk.Entry(r,textvariable=self.vars[k],width=69).pack(side='left')
    def _combo_row(self,lab,k,values):
        self.vars[k]=tk.StringVar(); r=self._row(lab); cb=ttk.Combobox(r,textvariable=self.vars[k],values=values,state='readonly',width=66); cb.pack(side='left'); cb.current(0)
    def _file_row(self,lab,k,kind):
        self.vars[k]=tk.StringVar(); r=self._row(lab); ttk.Entry(r,textvariable=self.vars[k],width=60).pack(side='left'); ttk.Button(r,text='찾기',command=lambda:self.pick(k,kind)).pack(side='left',padx=5)
    def pick(self,k,kind):
        ft=[('Excel','*.xlsx')] if kind=='xlsx' else [('PowerPoint','*.pptx')]
        f=filedialog.askopenfilename(filetypes=ft)
        if f:self.vars[k].set(f)
    def gui(self): return {k:v.get().strip() for k,v in self.vars.items()}

    def preview(self):
        g=self.gui()
        if not g.get('ppt8d'): return messagebox.showwarning('확인','8D PPT를 선택하세요.')
        try:
            d=base.extract(g['ppt8d']); kind,event=_event_name(d)
            self.log.delete('1.0','end'); self.log.insert('end','=== 8D 추출 결과 ===\n')
            for k in ['issue_name','task_name','customer','occurrence_date','problem','cause_4d','leak_cause','system_cause','action_5d','verification_6d','spread_7d']:
                self.log.insert('end',f'{k}: {d.get(k)}\n')
            self.log.insert('end',f'\n시험/빌드 판정: {kind} / {event}\n이슈기인 기본값: {"선택 필요" if _cause_known(d) else "TBD"}\nDB 상태: {_excel_status(d)}')
        except Exception as e: messagebox.showerror('오류',repr(e))

    def _prepare_choices(self,d,g):
        if _cause_known(d):
            dlg=ChoiceDialog(self,'이슈기인 선택','원인이 확인되어 있습니다. 이슈기인을 선택하세요.',ISSUE_ORIGINS)
            if dlg.result is None: return None
            g['issue_origin']=dlg.result
        else:
            g['issue_origin']='TBD'
        if _problem_is_long(d):
            ans=messagebox.askyesno('현상 내용 확인','현상 내용이 깁니다.\n1페이지 주요 논의 사항의 현상 칸에는 시험명/빌드명 중심으로 짧게 기재할까요?\n\n예: 시험명: 고온충전시험')
            g['_short_problem']=bool(ans)
        else:
            g['_short_problem']=False
        return g

    def run(self):
        g=self.gui()
        if any(not g.get(k) or not os.path.exists(g[k]) for k in ('ppt8d','pptweekly','xlsx')):
            return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')
        try:
            d=base.extract(g['ppt8d'])
            g=self._prepare_choices(d,g)
            if g is None: return
            # 발생처는 사용자가 선택한 값을 최종값으로 사용한다.
            d['occurrence_site']=g.get('occurrence_site') or d.get('occurrence_site','')
            mode=self.mode.get(); out=Path(g['xlsx']).parent/'자동화_결과'; out.mkdir(exist_ok=True)
            xo=out/(Path(g['xlsx']).stem+'_업데이트.xlsx'); po=out/(Path(g['pptweekly']).stem+'_업데이트.pptx')
            wb=load_workbook(g['xlsx']); ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active; r,_=base.find(ws,d,g)
            if mode=='existing' and not r:
                return messagebox.showwarning('업데이트 중단','기존 이슈를 Excel에서 정확하게 찾지 못했습니다. PMS/PLM 이슈번호를 입력하거나 이슈 정보를 확인하세요.')
            if mode=='new' and r and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?'):
                return
            a,_=base.update_excel(g['xlsx'],xo,d,g,new=(mode=='new'))
            b,_=weekly(g['pptweekly'],po,d,g,mode)
            self.log.delete('1.0','end'); self.log.insert('end',f'=== 완료 ===\n이슈 구분: {mode}\nDB 상태: {_excel_status(d)}\n이슈기인: {g["issue_origin"]}\n\n{a}\n{b}\n\n결과: {out}')
            messagebox.showinfo('완료',f'자동 업데이트가 완료되었습니다.\n\n{out}')
        except Exception as e:
            messagebox.showerror('실행 오류',repr(e))


if __name__=='__main__':
    App().mainloop()
