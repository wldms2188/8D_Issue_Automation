import copy
import datetime
from pathlib import Path

from pptx import Presentation

import main_recovery_step13 as step13
import main_recovery_step12 as step12
import main_recovery_step10 as step10
import main_recovery_step4 as step4
import main_v319 as v319
import main_v310 as v310

base=step13.base
N=v310.N
C=v310.C


def _summary_pages(prs):
    """Return every summary-style page, not only slide 1.

    A summary page is identified from the actual table header (과제명 + Signal),
    so the customer/task summary can live on slide 1, 2, 3, ... .
    """
    out=[]
    for si,sl in enumerate(prs.slides):
        for sh in v310.walk(sl):
            if not getattr(sh,'has_table',False):
                continue
            tb=sh.table
            for hr in range(min(4,len(tb.rows))):
                qs=[C(tb.cell(hr,c).text) for c in range(len(tb.columns))]
                if any('과제명' in q for q in qs) and any(q=='signal' for q in qs):
                    out.append((si,tb,hr))
                    break
            else:
                continue
            break
    return out


def _task_rows(tb,hr,hm,task):
    if 'task' not in hm:
        return []
    tk=step13._k(task)
    rows=[]
    for r in range(hr+1,len(tb.rows)):
        rt=step13._k(step13._row_text(tb,r,hm['task']))
        if tk and rt==tk:
            rows.append(r)
    return rows


def _insert_row_after(tb,after_row):
    """Insert a styled row directly after after_row and return its row index."""
    source=after_row
    new_tr=copy.deepcopy(tb.rows[source]._tr)
    tb.rows[source]._tr.addnext(new_tr)
    r=after_row+1
    for c in range(len(tb.columns)):
        try:
            tb.cell(r,c).text=''
        except Exception:
            pass
    return r


def _clear_summary_data(tb,hr):
    """Keep header + one styled data row; remove old customer rows on a cloned page."""
    # Ensure there is at least one data row to preserve as the style template.
    if len(tb.rows)<=hr+1:
        tr=copy.deepcopy(tb.rows[hr]._tr)
        tb.rows[hr]._tr.addnext(tr)

    # python-pptx row collections do NOT support negative indexing.
    # Keep the first styled data row and remove all later rows from bottom to top.
    keep=hr+1
    while len(tb.rows)>keep+1:
        last_index=len(tb.rows)-1
        tr=tb.rows[last_index]._tr
        tr.getparent().remove(tr)

    for c in range(len(tb.columns)):
        try:
            tb.cell(keep,c).text=''
        except Exception:
            pass
    return keep


def _first_summary_section_insert_index(prs,pages):
    """Create a new summary page inside the first (summary) section.

    The first section is treated as the leading block of summary-style pages.
    The new page is inserted immediately after the last summary page in that
    leading block, i.e. before detail/second-section pages begin.
    """
    if not pages:
        raise ValueError('주간회의 PPT에서 과제명/Signal 요약 양식 페이지를 찾지 못했습니다.')

    indices=[x[0] for x in pages]
    start=min(indices)
    last=start
    idxset=set(indices)
    while last+1 in idxset:
        last+=1
    return last+1


def _prepare_new_summary_page(prs,pages,g,template_index=None):
    # Normally use the first summary page as the exact user template. When a
    # matching project's table is full, clone that project's last summary page
    # so the continuation page keeps the same local layout/style.
    if template_index is None:
        template_index=pages[0][0]
    insert_at=_first_summary_section_insert_index(prs,pages)
    step13._clone_slide_with_rels(prs,template_index)
    step13._move_last_slide_to(prs,insert_at)

    sl=prs.slides[insert_at]
    found=None
    for sh in v310.walk(sl):
        if getattr(sh,'has_table',False):
            tb=sh.table
            for hr in range(min(4,len(tb.rows))):
                qs=[C(tb.cell(hr,c).text) for c in range(len(tb.columns))]
                if any('과제명' in q for q in qs) and any(q=='signal' for q in qs):
                    found=(tb,hr)
                    break
        if found:
            break
    if not found:
        raise ValueError('복제한 주간회의 요약 페이지에서 과제명/Signal 표를 다시 찾지 못했습니다.')

    tb,hr=found
    row=_clear_summary_data(tb,hr)

    # Keep the shared page style and only replace the team token when present.
    team=N(g.get('team'))
    if team:
        for sh in v310.walk(sl):
            if not hasattr(sh,'text_frame'):
                continue
            old=N(getattr(sh,'text',''))
            if '000팀' in old:
                try:
                    sh.text=old.replace('000팀',team+'팀')
                except Exception:
                    pass
    return insert_at,tb,hr,row


def _write_summary_row(tb,row,hr,d,g):
    hm=step13._summary_map(tb,hr)
    vals={
        'task':step13._customer_task(d),
        'issue':step13._issue_display(d),
        'problem':step4._full_problem(d),
        'progress':v319._progress_text(d),
    }
    for k,v in vals.items():
        if k in hm:
            base.set_cell_text(tb.cell(row,hm[k]),N(v),8)
    if 'signal' in hm:
        base.signal(tb.cell(row,hm['signal']),step12._signal_status_from_g(d,g))


def _summary_table_shape(sl,tb):
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False):
            continue
        try:
            if sh.table._tbl is tb._tbl:
                return sh
        except Exception:
            try:
                if sh.table._tbl==tb._tbl:
                    return sh
            except Exception:
                pass
    return None


def _summary_row_insert_fits(prs,slide_index,tb,after_row):
    """True when cloning after_row still fits inside the current slide."""
    try:
        sl=prs.slides[slide_index]
        sh=_summary_table_shape(sl,tb)
        if sh is None:
            return True
        row_h=int(tb.rows[after_row].height or 0)
        total_h=sum(int(r.height or 0) for r in tb.rows)
        current_h=max(int(getattr(sh,'height',0) or 0),total_h)
        projected_bottom=int(sh.top)+current_h+row_h
        # Leave a small visual safety margin at the bottom.
        safe_bottom=int(prs.slide_height)-int(.12*v310.EMU)
        return projected_bottom<=safe_bottom
    except Exception:
        # If geometry cannot be read, prefer the existing page rather than
        # unnecessarily creating a continuation page.
        return True


def _update_summary_by_task(prs,d,g,mode):
    """Find customer_task across ALL summary pages, then insert/update there."""
    pages=_summary_pages(prs)
    if not pages:
        raise ValueError('주간회의 PPT에서 과제명/Signal 요약 양식 페이지를 찾지 못했습니다.')

    task=step13._customer_task(d)
    issue=step13._issue_display(d)

    # 1) Find every page containing this customer_task.
    task_hits=[]
    for si,tb,hr in pages:
        hm=step13._summary_map(tb,hr)
        rows=_task_rows(tb,hr,hm,task)
        if rows:
            task_hits.append((si,tb,hr,hm,rows))

    # 2) Existing issue: update the best matching issue row, but only inside
    #    pages/rows belonging to this customer_task.
    if mode=='existing' and task_hits:
        best=None
        for si,tb,hr,hm,rows in task_hits:
            for r in rows:
                sc=step13._row_match_score(tb,r,hm,task,issue)
                if best is None or sc>best[0]:
                    best=(sc,si,tb,hr,r)
        if best and best[0]>=140:
            _,si,tb,hr,row=best
            _write_summary_row(tb,row,hr,d,g)
            return si,row,'기존 행 업데이트'

    # 3) New issue (or existing issue missing from summary): if customer_task exists,
    #    add immediately after the LAST row of that customer_task on its last page.
    #    Only create a continuation page when one more styled row would exceed
    #    the usable slide height.
    if task_hits:
        si,tb,hr,hm,rows=sorted(task_hits,key=lambda x:x[0])[-1]
        after=max(rows)
        if _summary_row_insert_fits(prs,si,tb,after):
            row=_insert_row_after(tb,after)
            _write_summary_row(tb,row,hr,d,g)
            return si,row,'동일 고객사_과제명 마지막 행 바로 아래 삽입'

        nsi,ntb,nhr,nrow=_prepare_new_summary_page(prs,pages,g,template_index=si)
        _write_summary_row(ntb,nrow,nhr,d,g)
        return nsi,nrow,'동일 고객사_과제명 표 공간 초과로 다음 요약 페이지에 추가'

    # 4) customer_task does not exist at all: create a new summary page INSIDE
    #    the first summary section, and write this issue there.
    si,tb,hr,row=_prepare_new_summary_page(prs,pages,g)
    _write_summary_row(tb,row,hr,d,g)
    return si,row,'첫 구역에 신규 요약 페이지 생성'


def weekly_step14(src,out,d,g,mode):
    prs=Presentation(src)

    summary_index,_row,summary_action=_update_summary_by_task(prs,d,g,mode)

    target_index=None
    if mode=='existing':
        target_index=step13._find_existing_detail(prs,d,summary_index)

    detail_action='업데이트'
    if target_index is None:
        template_index=step13._find_detail_template(prs,summary_index)
        insert_at=step13._new_detail_position(prs,d,summary_index)
        step13._clone_slide_with_rels(prs,template_index)
        step13._move_last_slide_to(prs,insert_at)
        target_index=insert_at
        detail_action='추가'

    step13._update_detail_slide(prs.slides[target_index],d,g,mode)

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)

    return (
        '주간회의 PPT: 요약 '+summary_action+
        ' / 상세 '+detail_action+
        f' (고객사_과제명={step13._customer_task(d)})',
        saved
    )


base.weekly=weekly_step14


class RecoveryStep14App(step10.RecoveryStep10App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP14 FIX1')


if __name__=='__main__':
    RecoveryStep14App().mainloop()
