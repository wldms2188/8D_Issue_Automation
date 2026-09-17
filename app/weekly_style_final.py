"""Final weekly-PPT presentation polish with change-fragment highlighting."""
from difflib import SequenceMatcher
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
import re
import main_recovery_step14_fix2 as core
import main_recovery_step14 as s14
import main_recovery_step13 as s13
import main_v310 as v310

UPDATE_BLUE=RGBColor(0x00,0x33,0xFF); BLACK=RGBColor(0,0,0)
def _color_text_frame(tf,color):
    if tf is None:return
    for p in tf.paragraphs:
        for run in p.runs:run.font.color.rgb=color
def _color_cell(cell,color):
    try:_color_text_frame(cell.text_frame,color)
    except Exception:pass
def _color_shape(shape,color):
    try:
        if hasattr(shape,'text_frame'):_color_text_frame(shape.text_frame,color)
    except Exception:pass
def _norm(x):return '\n'.join(str(x or '').replace('\r\n','\n').replace('\r','\n').splitlines()).strip()
def _identity_norm(x):return re.sub(r'\s+','',str(x or '').replace('\u00a0',' ')).casefold()
def _auto_text_snapshot(sl):
    out={}
    for sh in v310.walk(sl):
        name=str(getattr(sh,'name','') or '')
        if name.startswith('AUTO_8D_TEXT_'):out[name[len('AUTO_8D_TEXT_'):].upper()]=_norm(getattr(sh,'text',''))
    return out

def _diff_color_cell(cell,old_text,new_text):
    """Keep equal text black and color only inserted/replaced fragments blue."""
    old_text=str(old_text or ''); new_text=str(new_text or '')
    if old_text==new_text:_color_cell(cell,BLACK); return
    tf=cell.text_frame
    # Capture the template-created run's basic font before rebuilding runs.
    sample=None
    try:
        for p in tf.paragraphs:
            if p.runs:sample=p.runs[0].font; break
    except Exception:pass
    tf.clear(); p=tf.paragraphs[0]
    for tag,_i1,_i2,j1,j2 in SequenceMatcher(None,old_text,new_text,autojunk=False).get_opcodes():
        if j1==j2:continue
        run=p.add_run(); run.text=new_text[j1:j2]
        try:
            if sample is not None:
                run.font.name=sample.name; run.font.size=sample.size; run.font.bold=sample.bold; run.font.italic=sample.italic
        except Exception:pass
        run.font.color.rgb=BLACK if tag=='equal' else UPDATE_BLUE

_original_signal=core.base.signal
def centered_signal(cell,status):
    _original_signal(cell,status)
    try:
        tf=cell.text_frame; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for p in tf.paragraphs:
            p.alignment=PP_ALIGN.CENTER; p.space_before=Pt(0); p.space_after=Pt(0)
            for run in p.runs:run.font.size=Pt(10); run.font.bold=True
    except Exception:pass
core.base.signal=centered_signal

_original_write_summary_row=s14._write_summary_row
def refined_write_summary_row(tb,row,hr,d,g):
    try:
        hm0=s13._summary_map(tb,hr); old={k:_norm(tb.cell(row,c).text) for k,c in hm0.items() if k in ('task','issue','problem','progress')}
    except Exception:old={}
    _original_write_summary_row(tb,row,hr,d,g)
    try:
        hm=s13._summary_map(tb,hr)
        for key in ('task','issue','problem','progress'):
            if key not in hm:continue
            cell=tb.cell(row,hm[key]); new=_norm(cell.text); previous=old.get(key,'')
            if key in ('task','issue'):
                changed=(_identity_norm(previous)!=_identity_norm(new))
                _color_cell(cell,UPDATE_BLUE if changed else BLACK)
            elif previous:
                _diff_color_cell(cell,previous,new)
            else:
                # A genuinely new summary value has no unchanged fragment.
                _color_cell(cell,UPDATE_BLUE if new else BLACK)
    except Exception:pass
s14._write_summary_row=refined_write_summary_row

_original_update_detail_slide=s13._update_detail_slide
def refined_update_detail_slide(sl,d,g,mode):
    before=_auto_text_snapshot(sl); _original_update_detail_slide(sl,d,g,mode); after=_auto_text_snapshot(sl); existing=(s13._k(mode) in ('existing','기존','기존이슈'))
    for sh in v310.walk(sl):
        name=str(getattr(sh,'name','') or '')
        if not name.startswith('AUTO_8D_TEXT_'):continue
        key=name[len('AUTO_8D_TEXT_'):].upper(); new_text=after.get(key,_norm(getattr(sh,'text',''))); old_text=before.get(key,'')
        if existing:
            changed=False if key not in before else old_text!=new_text; _color_shape(sh,UPDATE_BLUE if changed else BLACK)
        else:_color_shape(sh,UPDATE_BLUE if bool(new_text) else BLACK)
    label_keys={s13._k(x) for x in {'이슈기인','발생단계'}}
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False):continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                if s13._k(tb.cell(r,c).text) in label_keys and c+1<len(tb.columns):_color_cell(tb.cell(r,c+1),BLACK)
s13._update_detail_slide=refined_update_detail_slide
