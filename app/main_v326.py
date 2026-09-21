# 8D Issue Automation v3.2.6
# Fixes requested:
# 1) Weekly summary searches ALL summary pages and appends to the LAST page containing the same task.
#    If that page has room, add a row there; otherwise insert a new summary slide immediately after it.
# 2) Detail title uses GUI inputs when all five are provided:
#    고객사_과제명_발생샘플_개발단계_발생처 이슈 발생
#    (no underscore before "이슈 발생"). If inputs are incomplete, keep the existing title logic.
# 3) Customer/task GUI inputs are restored; "발생 단계" label is changed to "개발단계".

import os, re, copy, datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox

from pptx import Presentation
from pptx.util import Pt

import main_v325 as v325
import main_v321 as v321
import main_v320 as v320
import main_v319 as v319
import main_v314 as v314

base = v325.base
N = v321.N
C = v321.C
EMU = 914400

SUMMARY_BOTTOM_IN = 7.35
SUMMARY_MARGIN_IN = 0.12


def _headers(tb):
    best = None
    for hr in range(min(4, len(tb.rows))):
        hm = {}
        for c in range(len(tb.columns)):
            q = C(tb.cell(hr, c).text)
            if '과제명' in q:
                hm['task'] = c
            elif q == '이슈' or '이슈명' in q:
                hm['issue'] = c
            elif '현상' in q or '문제' in q:
                hm['problem'] = c
            elif '진행사항' in q or '진행현황' in q:
                hm['progress'] = c
            elif 'signal' in q:
                hm['signal'] = c
        if best is None or len(hm) > len(best[1]):
            best = (hr, hm)
    return best if best and 'task' in best[1] and 'issue' in best[1] else None


def _task_eq(a, b):
    return bool(C(a) and C(a) == C(b))


def _row_blank(tb, r):
    return not any(N(tb.cell(r, c).text) for c in range(len(tb.columns)))


def _write_summary_row(tb, row, hm, d):
    vals = {
        'task': N(d.get('task_name')),
        'issue': N(d.get('issue_name')),
        'problem': N(d.get('problem')),
        'progress': v319._progress_text(d),
    }
    for k, val in vals.items():
        if k in hm:
            v314._set_cell(tb.cell(row, hm[k]), val, 8)
    if 'signal' in hm:
        try:
            base.signal(tb.cell(row, hm['signal']), base.status(d))
        except Exception:
            pass


def _clear_summary_row(tb, r):
    for c in range(len(tb.columns)):
        v314._set_cell(tb.cell(r, c), '', 8)


def _can_grow_table(shape, tb):
    try:
        rh = float(tb.rows[-1].height or (0.36 * EMU))
        bottom = float(shape.top + shape.height + rh) / EMU
        return bottom <= SUMMARY_BOTTOM_IN
    except Exception:
        return False


def _append_table_row(shape, tb):
    # Clone the last data row so borders/fills/font styles remain identical.
    src = tb.rows[-1]._tr
    new_tr = copy.deepcopy(src)
    tb._tbl.append(new_tr)
    try:
        rh = int(tb.rows[-1].height or (0.36 * EMU))
        shape.height = int(shape.height) + rh
    except Exception:
        pass
    return len(tb.rows) - 1


def _find_last_task_location(prs, task):
    found = []
    for si, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if not getattr(sh, 'has_table', False):
                continue
            tb = sh.table
            info = _headers(tb)
            if not info:
                continue
            hr, hm = info
            rows = [r for r in range(hr + 1, len(tb.rows)) if _task_eq(tb.cell(r, hm['task']).text, task)]
            if rows:
                found.append((si, sl, sh, tb, hr, hm, max(rows)))
    return found[-1] if found else None


def _find_first_summary_table(sl):
    for sh in sl.shapes:
        if getattr(sh, 'has_table', False):
            info = _headers(sh.table)
            if info:
                return sh, sh.table, info[0], info[1]
    return None


def _duplicate_slide_after(prs, src_index):
    src = prs.slides[src_index]
    blank = prs.slide_layouts[6]
    new = prs.slides.add_slide(blank)

    # Copy shapes preserving formatting.
    for sh in src.shapes:
        el = copy.deepcopy(sh.element)
        new.shapes._spTree.insert_element_before(el, 'p:extLst')

    # Copy non-notes relationships required by copied shapes.
    for rel in src.part.rels.values():
        if 'notesSlide' in rel.reltype:
            continue
        try:
            new.part.rels.add_relationship(rel.reltype, rel._target, rel.rId)
        except Exception:
            pass

    # Move the newly added slide from the end to immediately after src_index.
    sldIdLst = prs.slides._sldIdLst
    moved = sldIdLst[-1]
    sldIdLst.remove(moved)
    sldIdLst.insert(src_index + 1, moved)
    return prs.slides[src_index + 1]


def _prepare_new_summary_slide(prs, src_index, d):
    new_sl = _duplicate_slide_after(prs, src_index)
    found = _find_first_summary_table(new_sl)
    if not found:
        return False
    sh, tb, hr, hm = found

    # Keep header/style, clear all existing data rows.
    for r in range(hr + 1, len(tb.rows)):
        _clear_summary_row(tb, r)

    row = hr + 1 if hr + 1 < len(tb.rows) else _append_table_row(sh, tb)
    _write_summary_row(tb, row, hm, d)
    return True


def _update_summary_all_pages(prs, d, g):
    task = N(d.get('task_name'))
    loc = _find_last_task_location(prs, task)

    if loc is None:
        # No existing task: keep legacy behavior, but write into the first summary table.
        for sl in prs.slides:
            found = _find_first_summary_table(sl)
            if not found:
                continue
            sh, tb, hr, hm = found
            row = None
            for r in range(hr + 1, len(tb.rows)):
                if _row_blank(tb, r):
                    row = r
                    break
            if row is None and _can_grow_table(sh, tb):
                row = _append_table_row(sh, tb)
            if row is not None:
                _write_summary_row(tb, row, hm, d)
                return
        return

    si, sl, sh, tb, hr, hm, last_row = loc

    # 1) Prefer an existing blank row BELOW the last row of the same task.
    for r in range(last_row + 1, len(tb.rows)):
        if _row_blank(tb, r):
            _write_summary_row(tb, r, hm, d)
            return

    # 2) If there is physical space, append a styled row to this same last task page.
    if _can_grow_table(sh, tb):
        r = _append_table_row(sh, tb)
        _clear_summary_row(tb, r)
        _write_summary_row(tb, r, hm, d)
        return

    # 3) No room: insert a new summary slide immediately after this last task page.
    _prepare_new_summary_slide(prs, si, d)


def _legacy_title_meta(sl, d, g):
    return _original_page2_meta(sl, d, g)


_original_page2_meta = v321._page2_meta


def _detail_title_from_inputs(d, g):
    customer = N(g.get('customer'))
    task = N(g.get('task_name'))
    sample = N(g.get('sample'))
    stage = N(g.get('stage'))
    site = N(g.get('occurrence_site'))

    if all((customer, task, sample, stage, site)):
        # Exactly requested: underscore through occurrence site, then SPACE + "이슈 발생".
        return f'{customer}_{task}_{sample}_{stage}_{site} 이슈 발생'
    return ''


def _page2_meta_v326(sl, d, g):
    # First preserve every existing metadata rule.
    _original_page2_meta(sl, d, g)

    title = _detail_title_from_inputs(d, g)
    if not title:
        return

    title_sh, _ = v319._find_page2_header_shapes(sl)
    if title_sh is None:
        return

    # Reuse the existing fitting logic, but pass the complete requested title as one title segment.
    v319._replace_title_runs(title_sh, '', title)


# Patch the functions actually called by the inherited v3.2.5 workflow.
v319._update_page1 = _update_summary_all_pages
v321._page2_meta = _page2_meta_v326


class App(v325.App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.6')

        # Rename existing stage label.
        sample_frame = None
        for child in self.winfo_children():
            if isinstance(child, ttk.Frame):
                labels = [x for x in child.winfo_children() if isinstance(x, ttk.Label)]
                for lab in labels:
                    txt = str(lab.cget('text'))
                    if txt == '발생 단계':
                        lab.configure(text='개발단계')
                    if txt == '발생 샘플':
                        sample_frame = child

        # Restore customer/task inputs and place them immediately before 발생 샘플 when possible.
        if 'customer' not in self.vars:
            self.vars['customer'] = tk.StringVar()
            r = ttk.Frame(self)
            ttk.Label(r, text='고객사', width=31).pack(side='left')
            ttk.Entry(r, textvariable=self.vars['customer'], width=69).pack(side='left')
            if sample_frame is not None:
                r.pack(fill='x', padx=18, pady=3, before=sample_frame)
            else:
                r.pack(fill='x', padx=18, pady=3, before=self.log)

        if 'task_name' not in self.vars:
            self.vars['task_name'] = tk.StringVar()
            r = ttk.Frame(self)
            ttk.Label(r, text='과제명', width=31).pack(side='left')
            ttk.Entry(r, textvariable=self.vars['task_name'], width=69).pack(side='left')
            if sample_frame is not None:
                r.pack(fill='x', padx=18, pady=3, before=sample_frame)
            else:
                r.pack(fill='x', padx=18, pady=3, before=self.log)

    def run(self):
        g = self.gui()
        if any(not g.get(k) or not os.path.exists(g[k]) for k in ('ppt8d','pptweekly','xlsx')):
            return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')
        try:
            d = base.extract(g['ppt8d'])

            # User-entered customer/task override extracted values only when entered.
            if N(g.get('customer')):
                d['customer'] = N(g.get('customer'))
            if N(g.get('task_name')):
                d['task_name'] = N(g.get('task_name'))

            g = self._prepare_choices(d, g)
            if g is None:
                return

            d['occurrence_site'] = g.get('occurrence_site') or d.get('occurrence_site','')
            mode = self.mode.get()
            out = Path(g['xlsx']).parent / '자동화_결과'
            out.mkdir(exist_ok=True)
            xo = out / (Path(g['xlsx']).stem + '_업데이트.xlsx')
            po = out / (Path(g['pptweekly']).stem + '_업데이트.pptx')

            from openpyxl import load_workbook
            wb = load_workbook(g['xlsx'])
            ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active
            r, _ = base.find(ws, d, g)

            if mode == 'existing' and not r:
                return messagebox.showwarning('업데이트 중단','기존 이슈를 Excel에서 정확하게 찾지 못했습니다. PMS/PLM 이슈번호를 입력하거나 이슈 정보를 확인하세요.')
            if mode == 'new' and r and not messagebox.askyesno('중복 가능성', f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?'):
                return

            a, _ = base.update_excel(g['xlsx'], xo, d, g, new=(mode == 'new'))
            b, _ = v321.weekly(g['pptweekly'], po, d, g, mode)

            self.log.delete('1.0','end')
            self.log.insert('end', f'=== 완료 ===\n이슈 구분: {mode}\nDB 상태: {v325._excel_status(d,g)}\n이슈기인: {g["issue_origin"]}\n\n{a}\n{b}\n\n결과: {out}')
            messagebox.showinfo('완료', f'자동 업데이트가 완료되었습니다.\n\n{out}')
        except Exception as e:
            messagebox.showerror('실행 오류', repr(e))


if __name__ == '__main__':
    App().mainloop()
