import sys,re
from pathlib import Path
APP_DIR=Path(__file__).resolve().parent
if str(APP_DIR) not in sys.path: sys.path.insert(0,str(APP_DIR))
import main_v306 as impl
base=impl.base

def _n(x): return str(x or '').strip()
def _c(x):
    try:return base.compact(x)
    except:return re.sub(r'[^0-9A-Za-z가-힣]','',_n(x)).lower()
_prev=base.extract
def extract(path):
    d=_prev(path)
    try:
        from pptx import Presentation
        prs=Presentation(path); vals=[]; customer=_n(d.get('customer')); issue=_n(d.get('issue_name'))
        for sl in prs.slides:
            for sh in sl.shapes:
                if getattr(sh,'has_table',False):
                    tb=sh.table
                    for r in range(len(tb.rows)):
                        row=[_n(tb.cell(r,c).text) for c in range(len(tb.columns))]
                        for i,v in enumerate(row):
                            if '고객사' in _c(v): vals += [x for x in row[i+1:i+5] if x]
                            if not issue and '이슈명' in _c(v):
                                for x in row[i+1:i+4]:
                                    if x and '과제명' not in _c(x): issue=x;break
                t=_n(getattr(sh,'text',''))
                if t and '고객사' in _c(t): vals += [x.strip() for x in t.splitlines() if x.strip()]
        if customer:
            p=re.compile(re.escape(customer)+r'\s*[_\-/／|:：]\s*([^\n,;|]+)',re.I)
            for x in vals+[issue]:
                m=p.search(x)
                if m and _n(m.group(1)): d['task_name']=m.group(1).strip(' _-/／|:：');break
        if issue: d['issue_name']=issue
    except Exception: pass
    return d
base.extract=extract

_old=base.weekly
def weekly(src,out,d,g,mode):
    result=_old(src,out,d,g,mode)
    try:
        from pptx import Presentation
        from pptx.util import Pt
        prs=Presentation(result[1]); team=_n(g.get('team')); stage=_n(g.get('stage')); date=_n(d.get('occurrence_date'))
        stage_value=(stage+' ('+date+')') if stage and date else stage or date
        def walk(sl):
            for sh in sl.shapes:
                yield sh
                if hasattr(sh,'shapes'): yield from walk(sh)
        for si,sl in enumerate(prs.slides):
            for sh in walk(sl):
                if hasattr(sh,'text_frame'):
                    t=_n(getattr(sh,'text',''))
                    if team and '팀 주요 논의 사항' in t: sh.text=re.sub(r'^[^\n]*?팀\s*주요\s*논의\s*사항',team+'팀 주요 논의 사항',t)
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:r.font.name='맑은 고딕';r.font.size=Pt(8)
                if getattr(sh,'has_table',False):
                    tb=sh.table
                    for r in range(len(tb.rows)):
                        for c in range(len(tb.columns)):
                            q=_c(tb.cell(r,c).text)
                            if si==0 and q in ('이슈','이슈명') and c+1<len(tb.columns): tb.cell(r,c+1).text=_n(d.get('issue_name'))
                            if si>0 and '발생단계' in q and c+1<len(tb.columns): tb.cell(r,c+1).text=stage_value
        prs.save(result[1])
    except Exception: pass
    return result
base.weekly=weekly
if __name__=='__main__':base.App().mainloop()
