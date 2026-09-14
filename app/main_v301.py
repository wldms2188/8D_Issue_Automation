# 8D Issue Automation v3.0.1
# Final launcher for the v3.0 weekly renderer.
# Adds a stricter metadata parser for the user's format:
#   이슈명 area -> 고객사 -> next value = 과제명

import sys, re
from pathlib import Path

APP_DIR = Path(__file__).resolve().parent
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

import main_v300 as impl
base = impl.base

from pptx import Presentation


def _n(x):
    return str(x or '').replace('\r\n', '\n').replace('\r', '\n').strip()


def _c(x):
    try:
        return base.compact(x)
    except Exception:
        return re.sub(r'[^0-9A-Za-z가-힣]', '', _n(x)).lower()


def _valid_task(x):
    q = _c(x)
    if not x or len(_n(x)) < 2:
        return False
    blocked = ('고객사','이슈명','과제명','model','packer','발생site','발생일자','lotno','발생line',
               '담당자','제품타입','폼팩터','발생샘플','개발단계','발생처','불량률','불량수량')
    return not any(k in q for k in blocked)


def _task_from_customer_area(path):
    prs = Presentation(path)

    for sl in prs.slides:
        # Table-based metadata: prefer an explicit 과제명 label first.
        for sh in sl.shapes:
            if not getattr(sh, 'has_table', False):
                continue
            tb = sh.table
            rows = [[_n(tb.cell(r,c).text) for c in range(len(tb.columns))] for r in range(len(tb.rows))]

            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    if '과제명' not in _c(val):
                        continue
                    candidates = row[c+1:]
                    for rr in range(r+1, min(r+4, len(rows))):
                        candidates.append(rows[rr][c])
                    for cand in candidates:
                        if _valid_task(cand):
                            return _n(cand)

            # No explicit label: 고객사 값 다음 칸/다음 위치를 과제명으로 사용.
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    if '고객사' not in _c(val):
                        continue
                    candidates = []
                    # Common layout: [고객사][customer value][과제명 value]
                    if c + 2 < len(row):
                        candidates.append(row[c+2])
                    # If customer and task are in the next row, use the cell after
                    # the customer column first.
                    for rr in range(r+1, min(r+3, len(rows))):
                        if c + 1 < len(rows[rr]):
                            candidates.append(rows[rr][c+1])
                        if c < len(rows[rr]):
                            candidates.append(rows[rr][c])
                    for cand in candidates:
                        if _valid_task(cand):
                            return _n(cand)

        # Shape/text metadata: 고객사 value is followed by task line/value.
        for sh in sl.shapes:
            text = _n(getattr(sh, 'text', ''))
            if not text or '고객사' not in _c(text):
                continue
            lines = [x.strip() for x in text.split('\n') if x.strip()]
            for i, line in enumerate(lines):
                if '고객사' not in _c(line):
                    continue
                # 고객사: xxx / 과제명: yyy
                if i + 1 < len(lines) and '과제명' in _c(lines[i+1]):
                    nxt = re.split(r'[:：]', lines[i+1], maxsplit=1)
                    if len(nxt) == 2 and _valid_task(nxt[1]):
                        return _n(nxt[1])
                # 고객사: xxx / yyy  -> yyy
                if i + 1 < len(lines) and _valid_task(lines[i+1]):
                    return _n(lines[i+1])

    return ''


_previous_extract = base.extract


def extract_v301(path):
    d = _previous_extract(path)
    try:
        task = _task_from_customer_area(path)
        if task:
            d['task_name'] = task
    except Exception:
        pass
    return d


base.extract = extract_v301

if __name__ == '__main__':
    base.App().mainloop()
