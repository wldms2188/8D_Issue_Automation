"""v2.7 semantic 8D extractor for alternate 8D templates."""
import re
import main
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ALIASES={
'2d':['문제 현황','문제현황','불량 현상','불량현상'],
'3d':['임시 조치','임시조치','임시 대책','임시대책','고객 대응','고객대응'],
'4d_cause':['발생 원인','발생원인','원인 분석','원인분석'],
'4d_leak':['유출 원인','유출원인','유출 사유','유출사유'],
'4d_system':['시스템 원인','시스템원인','IT 시스템','IT시스템'],
'5d':['개선 대책','개선대책','영구 개선','영구개선'],
'6d':['효과 검증','효과검증','유효성 검증','유효성검증','유효성 점검','유효성점검'],
'7d':['재발 방지','재발방지','수평 전개','수평전개','관련 표준','관련표준']}
META={'issue_name':['이슈명'],'task_name':['과제명','과제 명'],'customer':['Customer','고객사'],'occurrence_site':['발생 Site','발생Site'],'occurrence_date':['발생 일자','발생일자'],'receipt_date':['접수 일자','접수일자']}

def norm(x): return re.sub(r'\n+','\n',re.sub(r'[ \t]+',' ',str(x or '').replace('\r','\n'))).strip()
def compact(x): return re.sub(r'[^0-9a-z가-힣]','',norm(x)).lower()
def clean(x): return re.sub(r'^[:：\-\s]+','',norm(x)).strip()
def placeholder(x):
 q=compact(x)
 return not q or q in {'000','00','00호기','line00호기','담당자반영일','적용완료일','적용예정일','유첨5why','오창','cnj','cna','cmi','cwa'}
def instruction(x):
 q=compact(x)
 bad=['유첨5why에의거한근본원인','양불사이에4m관점의변경사항','재발불량이라면기존불량의원인분석','표준기준절차서fmeacpsop외의설정여부','it시스템관점에서불량유출된경우','피해유출된사유','4d에서도출된발생원인유출원인시스템원인에대한대책수립과정','각actionitem에대한일정및주관부서가명확하게설정되어있는지','고품처리방안','개선대책으로인한sideeffect는없는지','대책에대한효과가검증되어재현test시불량현상제거및고객spec을만족하는지','관련표준','담당자반영일','적용완료일','적용예정일','손실비용']
 return any(x in q for x in bad) or q.startswith('불량현상') or q.startswith('대책에대한효과가검증되어')
def label(x,aliases):
 q=compact(x); return any(q==compact(a) or q.startswith(compact(a)) for a in aliases)
def flat(sl):
 out=[]
 def rec(sh):
  if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
   for x in sh.shapes: rec(x)
  else: out.append(sh)
 for sh in sl.shapes: rec(sh)
 return out
def table_extract(tb):
 rows=[[norm(tb.cell(r,c).text) for c in range(len(tb.columns))] for r in range(len(tb.rows))]; out={}
 for r,row in enumerate(rows):
  for c,raw in enumerate(row):
   if not raw: continue
   for k,als in ALIASES.items():
    if out.get(k): continue
    for a in sorted(als,key=len,reverse=True):
     m=re.search(re.escape(a)+r'\s*[:：]\s*(.+)$',raw,re.I)
     if m and not placeholder(m.group(1)) and not instruction(m.group(1)): out[k]=clean(m.group(1)); break
    if out.get(k): continue
    if label(raw,als):
     cand=row[c+1:]+[rows[rr][c] for rr in range(r+1,min(r+4,len(rows)))]
     for v in cand:
      v=clean(v)
      if v and not placeholder(v) and not instruction(v) and not any(label(v,a2) for a2 in ALIASES.values()): out[k]=v; break
 return out
def embedded(text):
 s=norm(text); hits=[];out={}
 pats=[('2d',r'(?:^|\n)\s*2\s*D.*?[:：]'),('3d',r'(?:^|\n)\s*3\s*D.*?[:：]'),('4d',r'(?:^|\n)\s*4\s*D.*?[:：]'),('5d',r'(?:^|\n)\s*5\s*D.*?[:：]'),('6d',r'(?:^|\n)\s*6\s*D.*?[:：]'),('7d',r'(?:^|\n)\s*7\s*D.*?[:：]')]
 for k,p in pats:
  for m in re.finditer(p,s,re.I): hits.append((m.start(),m.end(),k))
 hits.sort()
 for i,(a,b,k) in enumerate(hits):
  v=clean(s[b:hits[i+1][0] if i+1<len(hits) else len(s)])
  if v and not instruction(v): out[k]=v
 return out
def extract(path):
 prs=Presentation(path); d={k:'' for k in ['issue_name','task_name','customer','occurrence_site','occurrence_date','receipt_date','problem','temporary_action','customer_response','cause_4d','leak_cause','system_cause','action_5d','verification_6d','spread_7d']}; texts=[]; emb={}
 target={'2d':'problem','3d':'temporary_action','4d_cause':'cause_4d','4d_leak':'leak_cause','4d_system':'system_cause','5d':'action_5d','6d':'verification_6d','7d':'spread_7d'}
 for sl in prs.slides:
  for sh in flat(sl):
   t=norm(getattr(sh,'text',''))
   if t: texts.append(t); emb.update(embedded(t))
   if getattr(sh,'has_table',False):
    f=table_extract(sh.table)
    for k,v in f.items():
     if target.get(k) and not d[target[k]]: d[target[k]]=v
 full='\n'.join(texts)
 for k,als in META.items():
  for a in als:
   if d[k]: break
   m=re.search(re.escape(a)+r'\s*[:：]\s*([^\n]+)',full,re.I)
   if m and not placeholder(m.group(1)) and not instruction(m.group(1)): d[k]=clean(m.group(1))
 for k,s in [('problem','2d'),('temporary_action','3d'),('cause_4d','4d'),('action_5d','5d'),('verification_6d','6d'),('spread_7d','7d')]:
  if not d[k] and emb.get(s): d[k]=emb[s]
 for k in list(d):
  if k not in META and instruction(d[k]): d[k]=''
 if d['cause_4d'] and not d['leak_cause']:
  m=re.search(r'유출\s*원인\s*[:：]?',d['cause_4d'])
  if m: d['leak_cause']=clean(d['cause_4d'][m.end():]); d['cause_4d']=clean(d['cause_4d'][:m.start()])
 return d
main.extract=extract
if callable(getattr(main,'App',None)): main.App().mainloop()
