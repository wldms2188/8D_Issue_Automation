# 8D Issue Automation v3.1.2
# Adds dynamic cascading layout for NEW issues while preserving existing-issue marker positions.
import math, datetime
from pathlib import Path

import main_v311 as v311
import main_v310 as v310
from pptx import Presentation
from pptx.util import Inches, Pt

base = v311.base
N = v310.N
C = v310.C
box = v310.box
walk = v310.walk
EMU = v310.EMU

# Keep the same reference geometry, but allow y/h to change dynamically.
BASE_ZONES = {k: dict(v) for k, v in v310.ZONES.items()}
LEFT = ['2D', '3D', '4D_CAUSE']
RIGHT = ['4D_LEAK', '5D', '6D']
GAP = 0.16
BOTTOM = 7.20
MIN_FONT = 5.5
MAX_FONT = 8.0
MIN_H = {
    '2D': 0.78,
    '3D': 0.86,
    '4D_CAUSE': 1.12,
    '4D_LEAK': 0.76,
    '5D': 1.04,
    '6D': 0.76,
}


def _wrapped_lines(text, width_in, font_pt=8):
    text = N(text)
    if not text:
        return 1
    # Conservative mixed Korean/English estimate for PowerPoint.
    chars = max(8, int(width_in * 10.5 * (8.0 / font_pt)))
    total = 0
    for line in text.split('\n'):
        total += max(1, math.ceil(max(1, len(line)) / chars))
    return total


def _need_height(key, text, has_images, font_pt=8.0):
    z = BASE_ZONES[key]
    text_w = z['w']
    if has_images:
        text_w = max(1.0, z['w'] - min(1.65, z['w'] * .36) - .08)
    lines = _wrapped_lines(text, text_w, font_pt)
    line_h = font_pt / 72.0 * 1.22
    need = lines * line_h + 0.18
    return max(MIN_H[key], need)


def _fit_chain(keys, texts, imgs):
    top = BASE_ZONES[keys[0]]['y']
    heights = [_need_height(k, texts[k], bool(imgs.get(k)), MAX_FONT) for k in keys]
    available = BOTTOM - top - GAP * (len(keys) - 1)
    total = sum(heights)
    font = MAX_FONT

    # Only shrink font if the natural cascading layout exceeds available slide height.
    while total > available and font > MIN_FONT:
        font = round(font - 0.5, 1)
        heights = [_need_height(k, texts[k], bool(imgs.get(k)), font) for k in keys]
        total = sum(heights)

    if total > available:
        # Last-resort proportional compression; text is still preserved, font already at minimum.
        scale = available / total if total else 1.0
        heights = [max(MIN_H[k] * 0.90, h * scale) for k, h in zip(keys, heights)]

    out = {}
    y = top
    for k, h in zip(keys, heights):
        z = dict(BASE_ZONES[k])
        z['y'] = y
        z['h'] = h
        out[k] = z
        y += h + GAP
    return out, font


def _new_layout(d):
    texts = {k: v310.section_text(d, k) for k in BASE_ZONES}
    src = d.get('_section_images', {}) or {}
    imgs = {
        '2D': src.get('2D', []),
        '3D': src.get('3D', []),
        '4D_CAUSE': src.get('4D', []),
        '4D_LEAK': [],
        '5D': src.get('5D', []),
        '6D': src.get('6D', []),
    }
    left, lf = _fit_chain(LEFT, texts, imgs)
    right, rf = _fit_chain(RIGHT, texts, imgs)
    zones = {**left, **right}
    return zones, texts, imgs, min(lf, rf)


def _marker_xy_from_zone(label, zones):
    key = {'2D':'2D','3D':'3D','4D':'4D_LEAK','5D':'5D','6D':'6D'}[label]
    z = zones[key]
    # Marker/title sits just above-left of the corresponding content zone.
    return max(0.10, z['x'] - 0.18), max(0.10, z['y'] - 0.35)


def _add_box(sl, name, z, text, size):
    sh = sl.shapes.add_textbox(Inches(z['x']), Inches(z['y']), Inches(z['w']), Inches(z['h']))
    sh.name = 'AUTO_8D_TEXT_' + name
    tf = sh.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.05); tf.margin_right = Inches(.05)
    tf.margin_top = Inches(.03); tf.margin_bottom = Inches(.03)
    p = tf.paragraphs[0]; p.text = N(text)
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.name = '맑은 고딕'
    return sh


def _render_zone(sl, key, z, text, images, font_size):
    blob = v310.composite(images)
    if blob:
        iw = min(1.65, z['w'] * .36)
        tw = z['w'] - iw - .08
        tz = {'x':z['x'], 'y':z['y'], 'w':tw, 'h':z['h']}
        _add_box(sl, key, tz, text, font_size)
        v310.add_image(sl, blob, z['x'] + tw + .08, z['y'], iw, z['h'], key)
    else:
        _add_box(sl, key, z, text, font_size)


def update_page2(sl, d, g, mode):
    # Existing issues: preserve their current D positions exactly and use v3.1.1 behavior.
    if not v310.is_new(mode):
        return v311.update_page2(sl, d, g, mode)

    # New issues: dynamic cascade, preserving template graphics.
    v310.remove_previous_auto(sl)
    zones, texts, imgs, font_size = _new_layout(d)

    for label in ('2D','3D','4D','5D','6D'):
        x, y = _marker_xy_from_zone(label, zones)
        v310.move_marker_unit(sl, label, x, y)

    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        _render_zone(sl, key, zones[key], texts[key], imgs.get(key, []), font_size)

    # Reuse v3.1.0 metadata/signal logic without re-rendering its fixed zones.
    team = N(g.get('team')); owner = N(g.get('owner'))
    issue = N(d.get('issue_name')); task = N(d.get('task_name'))
    for sh in walk(sl):
        if hasattr(sh, 'text_frame'):
            t = N(getattr(sh, 'text', ''))
            if t.startswith('이슈명'):
                prefix = t.split(':',1)[0] if ':' in t else '이슈명'
                v310.set_text(sh, prefix + ' : ' + issue, 8)
            elif '과제명_이슈 제목' in t:
                v310.set_text(sh, t.replace('과제명_이슈 제목', task), 8)
            elif '00팀 담당자' in t:
                v310.set_text(sh, f'{team} 담당자 : {owner}', 8)
        if getattr(sh, 'has_table', False):
            tb = sh.table
            for r in range(len(tb.rows)):
                for c in range(len(tb.columns)):
                    if C(tb.cell(r,c).text) == 'signal' and c + 1 < len(tb.columns):
                        try: base.signal(tb.cell(r,c+1), base.status(d))
                        except Exception: pass
    v311._fill_occurrence(sl, d, g)


def weekly(src, out, d, g, mode):
    prs = Presentation(src)
    v311.update_page1(prs, d, g)
    if len(prs.slides) > 1:
        update_page2(prs.slides[1], d, g, mode)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(out); saved = out
    except PermissionError:
        p = Path(out)
        saved = str(p.with_name(p.stem + '_' + datetime.datetime.now().strftime('%Y%m%d_%H%M%S') + p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: ' + base.status(d), saved


base.weekly = weekly
base.extract = v310.extract

if __name__ == '__main__':
    base.App().mainloop()
