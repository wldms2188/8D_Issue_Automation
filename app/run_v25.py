import os, re, copy, datetime
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
import app.main as m

SECTION_PATTERNS={'2d':r'2D\s*현상','3d':r'3D\s*임시\s*대책','4d':r'4D\s*원인\s*분석','5d':r'5D\s*개선\s*대책','6d':r'6D\s*(?:유효성\s*점검|효과\s*검증)','7d':r'7D\s*(?:수평\s*전개|재발\s*방지)'}
LABELS={'issue_name':[r'이슈\s*명'],'task_name':[r'과제\s*명'],'customer':[r'Customer',r'고객사'],'occurrence_site':[r'발생\s*Site',r'발생처'],'occurrence_date':[r'발생\s*일자'],'receipt_date':[r'접수\s*일자'],'problem':[r'불량\s*현상',r'문제\s*현황'],'temporary_action':[r'임시\s*조치'],'customer_response':[r'고객\s*대응'],'cause_4d':[r'발생\s*원인'],'leak_cause':[r'유출\s*원인'],'system_cause':[r'IT\s*시스템'],'action_5d':[r'개선\s*대책'],'verification_6d':[r'(?:효과\s*검증|유효성\s*점검)'],'spread_7d':[r'(?:수평\s*전개|재발\s*방지)']}

def norm(x): return re.sub(r'[ \t]+',' ',str(x or '').replace('\r','\n')).strip()
def one(x,n=None):
 s=re.sub(r'\s+',' ',norm(x)); return s if not n or len(s)<=n else s[:n-1]+'…'
def blank(x): return not norm(x) or norm(x) in {'000','00','내용','입력','내용 입력','예시','담당자/ 반영일','00호기'}
def key(x): return re.sub(r'[\s_/\\\-:.,·]+','',one(x).lower())
def iter_shapes(shapes):
 for sh in shapes:
  yield sh
  if getattr(sh,'shape_type',None)==6: yield from iter_shapes(sh.shapes)
def table_cells(sl):
 for sh in iter_shapes(sl.shapes):
  if getattr(sh,'has_table',False):
   tb=sh.table
   for r in range(len(tb.rows)):
    for c in range(len(tb.columns)): yield tb,r,c,tb.cell(r,c)
def extract_labeled(text,pat):
 m=re.search(r'(?:^|\n)\s*(?:'+pat+r')\s*[:：]?\s*',text,re.I|re.S)
 if not m: m=re.search(r'(?:'+pat+r')\s*[:：]?\s*',text,re.I|re.S)
 if not m:return ''
 start=m.end(); poss=[]
 for ps in list(LABELS.values())+[list(SECTION_PATTERNS.values())]:
  for p in ps:
   q=re.search(r'(?:^|\n)\s*(?:'+p+r')\s*[:：]',text[start:],re.I|re.S)
   if q:poss.append(start+q.start())
 end=min(poss) if poss else len(text)
 return norm(text[start:end]).strip(' :：-')
def extract(p):
 prs=Presentation(p); keys=list(LABELS.keys()); d={k:'' for k in keys}; d['standards']={s:'' for s in m.STD}
 blocks=[]
 for sl in prs.slides:
  for sh in iter_shapes(sl.shapes):
   if hasattr(sh,'text') and norm(sh.text):blocks.append(norm(sh.text))
 full='\n'.join(blocks)
 for k,ps in LABELS.items():
  for ptn in ps:
   v=extract_labeled(full,ptn)
   if v:d[k]=v;break
 for sl in prs.slides:
  for tb,r,c,cl in table_cells(sl):
   raw=norm(cl.text)
   for k,ps in LABELS.items():
    if d[k] or not any(re.search(p,raw,re.I) for p in ps):continue
    cand=[]
    if ':' in raw or '：' in raw:cand.append(re.split(r'[:：]',raw,1)[1].strip())
    cand += [norm(tb.cell(r,cc).text) for cc in range(c+1,len(tb.columns))]
    cand += [norm(tb.cell(rr,c).text) for rr in range(r+1,min(r+3,len(tb.rows)))]
    for v in cand:
     if v and not blank(v):d[k]=v;break
 return d

def parse_ym(s):
 s=one(s); q=re.search(r'(20\d{2})\s*[-./]\s*(\d{1,2})',s)
 if q:return q.group(1),q.group(2)
 q=re.search(r'(?<!\d)(\d{2})\s*[-./]\s*(\d{1,2})\s*[-./]\s*\d{1,2}(?!\d)',s)
 return ('20'+q.group(1),q.group(2)) if q else ('','')
def task(d):return (one(d.get('customer'))+' / '+one(d.get('task_name'))).strip(' /')
def header_map(ws):
 aliases={'date':['최종 수정일','수정일'],'plm':['PLM 이슈 번호','PLM','이슈 번호'],'form':['폼팩터'],'product':['제품 타입'],'year':['발생 연도'],'month':['발생 월'],'team':['담당팀'],'owner':['담당자'],'task':['고객사 / 과제명','고객사/과제명','고객사','과제명'],'sample':['발생 샘플'],'stage':['개발 단계'],'site':['발생처','발생 Site'],'problem':['이슈 현상 요약','문제 현상','불량 현상'],'photo':['대표 사진'],'cause':['팩기인 원인','원인'],'action':['개선 대책'],'status':['이슈 상태']}
 best={}; hr=1
 for r in range(1,min(ws.max_row,12)+1):
  for c in range(1,ws.max_column+1):
   v=key(ws.cell(r,c).value)
   for k,names in aliases.items():
    if k not in best and any(key(n) in v or v in key(n) for n in names):best[k]=c;hr=r
 return hr,best
def update_excel(src,out,d,g,new=False):
 wb=m.load_workbook(src);ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active;hr,h=header_map(ws)
 plm=key(g.get('plm_no')); r=None
 if plm and 'plm' in h:
  for rr in range(hr+1,ws.max_row+1):
   if key(ws.cell(rr,h['plm']).value)==plm:r=rr;break
 if not r and not new:
  tt,pp=key(task(d)),key(d.get('problem'))
  for rr in range(hr+1,ws.max_row+1):
   if ('task' in h and tt and key(ws.cell(rr,h['task']).value)==tt) or ('problem' in h and pp and key(ws.cell(rr,h['problem']).value)==pp):r=rr;break
 if new:r=ws.max_row+1
 if not r:raise ValueError('기존 이슈를 Excel에서 찾지 못했습니다. PMS/PLM 이슈번호를 입력하세요.')
 if r>hr+1:
  for c in range(1,ws.max_column+1):ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
 y,mo=parse_ym(d.get('occurrence_date'));cause=' / '.join(x for x in [d.get('cause_4d'),d.get('leak_cause'),d.get('system_cause')] if x)
 vals={'date':datetime.date.today(),'plm':g.get('plm_no',''),'form':g.get('form_factor',''),'product':g.get('product_type',''),'year':y,'month':mo,'team':g.get('team',''),'owner':g.get('owner',''),'task':task(d),'sample':g.get('sample',''),'stage':g.get('stage',''),'site':d.get('occurrence_site',''),'problem':one(d.get('problem'),300),'cause':one(cause,300),'action':one(d.get('action_5d'),300),'status':m.status(d)}
 for k,v in vals.items():
  if k in h:ws.cell(r,h[k]).value=v
 Path(out).parent.mkdir(exist_ok=True)
 try:wb.save(out);saved=out
 except PermissionError:
  p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);wb.save(saved)
 return ('신규 이슈 추가' if new else f'기존 이슈 업데이트 (row {r})'),saved

def fill_sections(sl,d):
 contents={'2d':one(d.get('problem'),700) or '내용 미기재','3d':one(d.get('temporary_action'),600) or '내용 미기재','4d':one(' / '.join(x for x in [d.get('cause_4d'),d.get('leak_cause'),d.get('system_cause')] if x),700) or '내용 미기재','5d':one(d.get('action_5d'),700) or '내용 미기재','6d':one(d.get('verification_6d'),700) or '내용 미기재','7d':one(d.get('spread_7d'),700) or '내용 미기재'}
 shapes=list(iter_shapes(sl.shapes)); labels=[]
 for sh in shapes:
  if not hasattr(sh,'text'):continue
  for k,p in SECTION_PATTERNS.items():
   if re.search(p,norm(sh.text),re.I):labels.append((k,sh));break
 if not labels:return
 # Remove a combined text box containing multiple D labels.
 for sh in list(shapes):
  if sh in [x[1] for x in labels] or not hasattr(sh,'text'):continue
  hits=sum(bool(re.search(p,norm(sh.text),re.I)) for p in SECTION_PATTERNS.values())
  if hits>=2:
   par=sh._element.getparent();par.remove(sh._element)
 labels.sort(key=lambda x:x[1].top)
 for i,(k,lab) in enumerate(labels):
  left=lab.left+lab.width+Inches(.12);top=lab.top
  bottom=labels[i+1][1].top-Inches(.05) if i+1<len(labels) else top+Inches(.75)
  h=max(Inches(.45),bottom-top);w=max(Inches(1),sl.part.presentation.slide_width-left-Inches(.3))
  box=sl.shapes.add_textbox(left,top,w,h);tf=box.text_frame;tf.clear();tf.word_wrap=True
  p=tf.paragraphs[0];p.text=contents[k];p.font.size=Pt(10)
def weekly(src,out,d,g,mode):
 prs=Presentation(src);st=m.status(d)
 for sl in prs.slides:
  for tb,r,c,cl in table_cells(sl):
   if r!=0:continue
   header=' '.join(norm(tb.cell(0,cc).text) for cc in range(len(tb.columns)))
   if '과제명' in header and 'Signal' in header:
    rr=1 if len(tb.rows)>1 else None
    if rr is not None:
     vals=[task(d),one(d.get('issue_name'),80),one(d.get('problem'),120),one(d.get('action_5d') or d.get('verification_6d'),120),'●']
     for cc,v in enumerate(vals[:len(tb.columns)]):tb.cell(rr,cc).text=v
     if len(tb.columns)>=5:
      tb.cell(rr,4).text='●'
      for p in tb.cell(rr,4).text_frame.paragraphs:
       for run in p.runs:run.font.color.rgb=RGBColor(*m.SIGNAL[st]);run.font.bold=True
 target=None
 for sl in prs.slides:
  if any(any(re.search(p,norm(sh.text),re.I) for p in SECTION_PATTERNS.values()) for sh in iter_shapes(sl.shapes) if hasattr(sh,'text')):target=sl;break
 if target is None:target=prs.slides[min(2,len(prs.slides)-1)]
 for sh in iter_shapes(target.shapes):
  if hasattr(sh,'text') and '00팀 담당자 : 000' in sh.text:sh.text=sh.text.replace('00팀 담당자 : 000',f"{g.get('team','')} 담당자 : {g.get('owner','')}")
 fill_sections(target,d)
 Path(out).parent.mkdir(exist_ok=True)
 try:prs.save(out);saved=out
 except PermissionError:
  p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);prs.save(saved)
 return '주간회의 PPT 업데이트 완료: '+st,saved

m.extract=extract;m.update_excel=update_excel;m.weekly=weekly
m.App().mainloop()
