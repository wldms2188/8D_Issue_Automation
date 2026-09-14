# 8D Issue Automation v3.1.7
# Layout policy: preserve 8pt first, move following D blocks downward, shrink only as last resort.
import datetime
from pathlib import Path

import main_v316 as v316
import main_v315 as v315
import main_v314 as v314
import main_v313 as v313
import main_v310 as v310
from pptx import Presentation

base = v316.base
N = v310.N

MAX_FONT = 8.0
MIN_FONT = 6.5
STEP = 0.5
BOTTOM = 7.42
GAP_NORMAL = 0.14
GAP_TIGHT = 0.08

LEFT = ('2D','3D','4D_CAUSE')
RIGHT = ('4D_LEAK','5D','6D')


def _has_section(d, key):
    if key == '4D_CAUSE':
        return bool(N(d.get('cause_4d')))
    if key == '4D_LEAK':
        return bool(N(d.get('leak_cause')) or N(d.get('system_cause')))
    return True


def _text_width(key, has_images, image_ratio=.30):
    z = v310.ZONES[key]
    if not has_images:
        return z['w']
    iw = min(1.45, z['w'] * image_ratio)
    return max(1.0, z['w'] - iw - .07)


def _need_h(key, text, has_images, font, image_ratio=.30):
    # 8pt is the preferred size. Height grows first; font is reduced only if the
    # whole active chain cannot fit after using the available slide height.
    w = _text_width(key, has_images, image_ratio)
    lines = v316.v312._wrapped_lines(text, w, font)
    line_h = font / 72.0 * 1.18
    return max(v316.v312.MIN_H[key], lines * line_h + .15)


def _build_chain(keys, d, texts, imgs):
    active = [k for k in keys if _has_section(d,k)]
    # Keep zone objects for inactive keys too because downstream marker logic expects them.
    inactive = [k for k in keys if k not in active]
    top = v310.ZONES[keys[0]]['y']
    if active and active[0] != keys[0]:
        # If the top 4D leak block is absent, 5D starts at its normal 5D position;
        # do not move it upward into the metadata/header area.
        top = v310.ZONES[active[0]]['y']

    image_ratio = .30
    fonts = {k:MAX_FONT for k in keys}
    gap = GAP_NORMAL

    def calc():
        return {k:_need_h(k,texts[k],bool(imgs.get(k)),fonts[k],image_ratio) for k in active}

    hs = calc()
    available = BOTTOM - top - gap * max(0,len(active)-1)

    # 1) Keep every active D block at 8pt and let its box grow downward.
    # 2) If it barely exceeds the slide, tighten only inter-block gaps.
    if sum(hs.values()) > available:
        gap = GAP_TIGHT
        available = BOTTOM - top - gap * max(0,len(active)-1)

    # 3) Before shrinking text, give text more width by making the image column narrower.
    if sum(hs.values()) > available and any(imgs.get(k) for k in active):
        image_ratio = .24
        hs = calc()

    # 4) Only now shrink the section that benefits the most, one step at a time.
    while sum(hs.values()) > available:
        options=[]
        for k in active:
            if fonts[k] <= MIN_FONT:
                continue
            nf=max(MIN_FONT, fonts[k]-STEP)
            nh=_need_h(k,texts[k],bool(imgs.get(k)),nf,image_ratio)
            gain=hs[k]-nh
            options.append((gain,hs[k],k,nf,nh))
        if not options:
            break
        options.sort(reverse=True)
        _,_,k,nf,nh=options[0]
        fonts[k]=nf
        hs[k]=nh

    # Last resort: preserve all text and fit box heights inside the slide.
    total=sum(hs.values())
    if total > available and total > 0:
        scale=available/total
        for k in active:
            hs[k]=max(v316.v312.MIN_H[k]*.90, hs[k]*scale)
        overflow=sum(hs.values())-available
        if overflow>0 and active:
            k=max(active,key=lambda x:hs[x])
            hs[k]=max(v316.v312.MIN_H[k]*.90,hs[k]-overflow)

    zones={}
    y=top
    for k in active:
        z=dict(v310.ZONES[k])
        z['y']=y
        z['h']=hs[k]
        z['_image_ratio']=image_ratio
        zones[k]=z
        y += hs[k]+gap

    # Inactive sections occupy no flow height. Keep a harmless reference zone for
    # downstream functions; they are not rendered and their 4D marker is deleted.
    for k in inactive:
        z=dict(v310.ZONES[k])
        z['_image_ratio']=image_ratio
        zones[k]=z

    return zones, fonts


def _layout(d, imgs):
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    left_imgs={'2D':imgs.get('2D',[]),'3D':imgs.get('3D',[]),'4D_CAUSE':imgs.get('4D_CAUSE',[])}
    right_imgs={'4D_LEAK':imgs.get('4D_LEAK',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}
    lz,lf=_build_chain(LEFT,d,texts,left_imgs)
    rz,rf=_build_chain(RIGHT,d,texts,right_imgs)
    return {**lz,**rz},texts,{**lf,**rf}


def _render(sl,key,z,text,images,font):
    if not text:
        return
    blob=v313._collage(images)
    if blob:
        ratio=z.get('_image_ratio',.30)
        iw=min(1.45,z['w']*ratio)
        tw=z['w']-iw-.07
        v310.add_box(sl,key,z['x'],z['y'],tw,z['h'],text,font)
        v310.add_image(sl,blob,z['x']+tw+.07,z['y'],iw,z['h'],key)
    else:
        v310.add_box(sl,key,z['x'],z['y'],z['w'],z['h'],text,font)


def update_page2(sl,d,g,mode):
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    zones,texts,fonts=_layout(d,imgs)

    hc=bool(N(d.get('cause_4d')))
    hl=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))

    # Every visible marker/title follows its corresponding dynamically moved block.
    if v310.is_new(mode):
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]
            v310.move_marker_unit(sl,label,max(.10,z['x']-.18),max(.10,z['y']-.35))

    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        if _has_section(d,key):
            _render(sl,key,zones[key],texts[key],imgs.get(key,[]),fonts[key])

    # Keep v3.1.6 exact title/issue-name formatting and occurrence metadata logic.
    v316._meta(sl,d,g)

    # Exact 4D rule: occurrence cause left-bottom, leak/system cause right-top;
    # delete the missing side's marker + title.
    zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']
    left_xy=(max(.10,zl['x']-.18),max(.10,zl['y']-.35))
    right_xy=(max(.10,zr['x']-.18),max(.10,zr['y']-.35))
    v315._ensure_4d_units(sl,hc,hl,left_xy,right_xy)


def weekly(src,out,d,g,mode):
    prs=Presentation(src)
    v314.update_page1(prs,d,g)
    if len(prs.slides)>1:
        update_page2(prs.slides[1],d,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.extract=v316.extract
base.weekly=weekly

if __name__=='__main__':
    base.App().mainloop()
