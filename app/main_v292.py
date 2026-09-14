# 8D Issue Automation v2.9.2
# Fix weekly detail page mapping and grouped 2D/4D representative images.
import io, copy, datetime, re
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from PIL import Image as PILImage

import main_v291 as v291
v29 = v291.v29
base = v29.base

EMU = 914400

def _abs_box(sh):
    return (float(sh.left), float(sh.top), float(sh.width), float(sh.height))

def _picture_children(sh):
    out=[]
    if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.PICTURE:
        out.append(sh)
    elif getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:
        for ch in sh.shapes: out.extend(_picture_children(ch))
    return out

def _composite_group(group):
    pics=_picture_children(group)
    if len(pics) < 2: return None
    boxes=[]
    for p in pics:
        try:
            blob=p.image.blob
            with PILImage.open(io.BytesIO(blob)) as im: boxes.append((p.left,p.top,p.width,p.height,blob))
        except Exception: pass
    if len(boxes)<2:return None
    minx=min(x[0] for x in boxes); miny=min(x[1] for x in boxes)
    maxx=max(x[0]+x[2] for x in boxes); maxy=max(x[1]+x[3] for x in boxes)
    scale=120/EMU
    W=max(1,int((maxx-minx)*scale)); H=max(1,int((maxy-miny)*scale))
    canvas=PILImage.new('RGB',(W,H),'white')
    for x,y,w,h,blob in boxes:
        try:
            with PILImage.open(io.BytesIO(blob)) as im:
                im=im.convert('RGB'); tw=max(1,int(w*scale)); th=max(1,int(h*scale))
                im.thumbnail((tw,th),PILImage.Resampling.LANCZOS)
                canvas.paste(im,(int((x-minx)*scale),int((y-miny)*scale)))
        except Exception: pass
    bio=io.BytesIO();canvas.save(bio,format='PNG');return bio.getvalue()

def _image_candidates(prs):
    out=[]
    for si,sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
                pics=_picture_children(sh)
                if len(pics)>=2:
                    blob=_composite_group(sh)
                    if blob:
                        x,y,w,h=_abs_box(sh);out.append({'slide':si,'x':x,'y':y,'w':w,'h':h,'area':w*h,'blob':blob,'group':True,'count':len(pics)})
            elif sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
                try:
                    with PILImage.open(io.BytesIO(sh.image.blob)) as im:
                        if im.width*im.height>=10000:
                            x,y,w,h=_abs_box(sh);out.append({'slide':si,'x':x,'y':y,'w':w,'h':h,'area':w*h,'blob':sh.image.blob,'group':False,'count':1})
                except Exception:pass
    return out

def _anchors(prs,problem=''):
    anchors=[];pk=base.compact(problem) if problem else ''
    for si,sl in enumerate(prs.slides):
        for sh in sl.shapes:
            txt=' '.join((getattr(ch,'text','') or '') for ch in sh.shapes) if sh.shape_type==MSO_SHAPE_TYPE.GROUP else (getattr(sh,'text','') or '')
            q=base.compact(txt)
            if '2d' in q or '문제현황' in q or '불량현상' in q or (pk and pk[:30] in q):
                anchors.append((si,sh.left+sh.width/2,sh.top+sh.height/2))
    return anchors

def _pick_rep(prs,problem=''):
    cands=_image_candidates(prs)
    if not cands:return None
    anc=_anchors(prs,problem)
    if anc:
        scored=[]
        for c in cands:
            same=[a for a in anc if a[0]==c['slide']]
            if same:
                cx=c['x']+c['w']/2;cy=c['y']+c['h']/2
                dist=min(((cx-a[1])**2+(cy-a[2])**2)**0.5 for a in same)
                # Prefer a multi-picture group when it is on/near the 2D area.
                bonus=-EMU*1.0 if c['group'] and c['count']>=2 else 0
                scored.append((dist+bonus,-c['area'],c))
        if scored:
            scored.sort(key=lambda z:(z[0],z[1]));return scored[0][2]['blob']
    groups=[c for c in cands if c['group']]
    return max(groups,key=lambda c:c['area'])['blob'] if groups else max(cands,key=lambda c:c['area'])['blob']

def extract_v292(path):
    d=v29.v29_extract(path)
    try:
        blob=_pick_rep(Presentation(path),d.get('problem',''))
        if blob:d['_images']=[(1,1,1,blob)]
    except Exception:pass
    return d
base.extract=extract_v292

# Example 2 slide 2 destination mapping, exactly as the supplied example indicates.
DETAIL_BOXES={
 '1':(0.53,2.53,4.76,0.98,'problem'),
 '2':(0.53,3.91,4.76,1.14,'temporary_action'),
 '3-1':(0.53,5.38,4.76,1.79,'cause_4d'),
 '3-2':(5.77,2.47,4.76,0.87,'leak_cause'),
 '4':(5.77,3.74,4.76,1.77,'action_5d'),
 '5':(5.77,5.91,4.78,0.94,'verification_6d'),
}

def _set8(sh,text):
    v29.v29_set_shape_text(sh,text,8)
    try:
        sh.text_frame.word_wrap=True;sh.text_frame.vertical_anchor=MSO_ANCHOR.TOP
        sh.text_frame.margin_left=Inches(.05);sh.text_frame.margin_right=Inches(.05)
        sh.text_frame.margin_top=Inches(.03);sh.text_frame.margin_bottom=Inches(.03)
    except Exception:pass

def _remove_generated(sl):
    for sh in list(sl.shapes):
        try:
            if str(sh.name).startswith('AUTO_8D_DETAIL_'):
                sp=sh._element;sp.getparent().remove(sp)
        except Exception:pass

def _find_number_shape(sl,label):
    for sh in sl.shapes:
        if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.AUTO_SHAPE and base.norm(getattr(sh,'text',''))==label:return sh
    return None

def _ensure_box(sl,label,x,y,w,h):
    old=_find_number_shape(sl,label)
    if old:return old
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name='AUTO_8D_DETAIL_'+label.replace('-','_');return sh

def _detail_text(key,d):
    if key=='leak_cause':return '\n'.join(v29.v29_one(x) for x in [d.get('leak_cause'),d.get('system_cause')] if v29.v29_one(x))
    return v29.v29_one(d.get(key))

def fill_weekly_detail_exact(sl,d,g):
    _remove_generated(sl)
    for label,(x,y,w,h,key) in DETAIL_BOXES.items():
        sh=_ensure_box(sl,label,x,y,w,h);_set8(sh,_detail_text(key,d) or '검토 중')
    for sh in sl.shapes:
        if hasattr(sh,'text'):
            t=base.norm(sh.text)
            if t.startswith('과제명_이슈 제목'): _set8(sh,base.task(d))
            elif t.startswith('이슈명 :'): _set8(sh,'이슈명 : '+v29.v29_one(d.get('issue_name')))
            elif '00팀 담당자' in t:_set8(sh,f"{g.get('team','')} 담당자 : {g.get('owner','')}")
    for sh in sl.shapes:
        if not getattr(sh,'has_table',False):continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                txt=base.compact(tb.cell(r,c).text)
                if txt=='signal' and c+1<len(tb.columns):base.signal(tb.cell(r,c+1),base.status(d))
                elif txt=='발생단계' and c+1<len(tb.columns):v29.v29_set_cell_text(tb.cell(r,c+1),g.get('stage',''),8)
                elif txt=='이슈기인' and c+1<len(tb.columns):v29.v29_set_cell_text(tb.cell(r,c+1),d.get('issue_name',''),8)

def weekly_v292(src,out,d,g,mode):
    prs=Presentation(src);st=base.status(d);title=base.task(d);issue=v29.v29_one(d.get('issue_name'))
    summary=None;summary_idx=None
    for i,sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if getattr(sh,'has_table',False):
                h=' '.join(base.norm(sh.table.cell(0,c).text) for c in range(len(sh.table.columns)))
                if '과제명' in h and 'Signal' in h:summary=sh.table;summary_idx=i;break
        if summary is not None:break
    if summary is not None:
        r=1
        for x in range(1,len(summary.rows)):
            if base.key(title) and base.key(title) in base.key(summary.cell(x,0).text) and base.key(issue) in base.key(summary.cell(x,1).text):r=x;break
            if not any(base.norm(summary.cell(x,c).text) for c in range(min(5,len(summary.columns))) if c!=4):r=x;break
        vals=[title,issue,v29.v29_one(d.get('problem')),v29.v29_progress(d),'●']
        for c,val in enumerate(vals):v29.v29_set_cell_text(summary.cell(r,c),val,8)
        base.signal(summary.cell(r,4),st)
    if summary_idx is None:summary_idx=0
    details=[prs.slides[i] for i in range(summary_idx+1,len(prs.slides))]
    for sl in details:fill_weekly_detail_exact(sl,d,g)
    Path(out).parent.mkdir(exist_ok=True)
    try:prs.save(out);saved=out
    except PermissionError:
        p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);prs.save(saved)
    return '주간회의 PPT 업데이트: '+st,saved

base.weekly=weekly_v292
if __name__=='__main__':base.App().mainloop()
