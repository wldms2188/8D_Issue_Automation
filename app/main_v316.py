# 8D Issue Automation v3.1.6
# - issue_name = source 8D filename stem
# - task_name fallback from issue filename/customer token
# - exact weekly page-2 title/issue formatting based on supplied example
# - adaptive font per D section (8pt preferred, min 6.5pt) instead of chain-wide 5.5pt
import re, datetime
from pathlib import Path

import main_v315 as v315
import main_v314 as v314
import main_v313 as v313
import main_v312 as v312
import main_v310 as v310
from pptx import Presentation
from pptx.util import Pt

base = v315.base
N = v310.N
C = v310.C
EMU = v310.EMU


def _task_from_filename(issue, customer, current=''):
    current = N(current)
    if current and C(current) not in ('상세시험조건','model','packer'):
        return current
    issue = N(issue)
    customer = N(customer)
    if not issue:
        return current
    if customer:
        # 1) customer directly followed by separator: Ford_xxxx -> xxxx
        m = re.search(re.escape(customer) + r'\s*[_\-/／|:]\s*(.+)$', issue, re.I)
        if m:
            val = N(m.group(1)).strip(' _-/／|:：')
            if val and C(val) not in ('상세시험조건','model','packer'):
                return val
        # 2) customer is one token inside a longer filename:
        #    사외_사업부_제품군_Ford_xxxx -> xxxx
        toks = [x.strip() for x in re.split(r'[_\-/／|]+', issue) if x.strip()]
        for i,t in enumerate(toks):
            if C(t) == C(customer) and i+1 < len(toks):
                val = '_'.join(toks[i+1:]).strip()
                if val and C(val) not in ('상세시험조건','model','packer'):
                    return val
    return current


_original_extract = v314.extract

def extract(path):
    d = _original_extract(path)
    # User rule: issue name is the 8D source filename without extension.
    issue = Path(path).stem.strip()
    d['issue_name'] = issue
    try:
        # Re-run the customer-prefix rule after filename issue name has been assigned.
        d = v310.extract_task_exact(path, d)
    except Exception:
        pass
    d['task_name'] = _task_from_filename(issue, d.get('customer'), d.get('task_name'))
    return d


# ---------- adaptive D layout/font ----------
MAX_FONT = 8.0
MIN_FONT = 6.5
STEP = 0.5
GAP = 0.16
BOTTOM = 7.20
CHAINS = (
    ('2D','3D','4D_CAUSE'),
    ('4D_LEAK','5D','6D'),
)


def _text_width(key, has_images):
    z = v310.ZONES[key]
    if not has_images:
        return z['w']
    return max(1.0, z['w'] - min(1.65, z['w'] * .36) - .08)


def _need_h(key, text, has_images, font):
    z = v310.ZONES[key]
    w = _text_width(key, has_images)
    lines = v312._wrapped_lines(text, w, font)
    line_h = font / 72.0 * 1.20
    return max(v312.MIN_H[key], lines * line_h + .17)


def _fit_chain_individual(keys, texts, imgs):
    top = v310.ZONES[keys[0]]['y']
    available = BOTTOM - top - GAP * (len(keys)-1)
    fonts = {k: MAX_FONT for k in keys}

    def heights():
        return {k:_need_h(k,texts[k],bool(imgs.get(k)),fonts[k]) for k in keys}

    hs = heights()
    # Reduce only the section that gains the most height from the next 0.5pt step.
    # This keeps short sections at 8pt instead of shrinking an entire column to 5.5pt.
    while sum(hs.values()) > available:
        options=[]
        for k in keys:
            if fonts[k] <= MIN_FONT:
                continue
            nf=max(MIN_FONT, fonts[k]-STEP)
            nh=_need_h(k,texts[k],bool(imgs.get(k)),nf)
            gain=hs[k]-nh
            options.append((gain, hs[k], k, nf, nh))
        if not options:
            break
        options.sort(reverse=True)
        _,_,k,nf,nh=options[0]
        fonts[k]=nf; hs[k]=nh

    # If extremely long text still exceeds the slide at 6.5pt, compress heights
    # proportionally but never globally force every section to 5.5pt.
    total=sum(hs.values())
    if total > available and total > 0:
        scale=available/total
        for k in keys:
            hs[k]=max(v312.MIN_H[k]*.90, hs[k]*scale)
        # correct rounding overflow on the largest section
        over=sum(hs.values())-available
        if over>0:
            k=max(keys,key=lambda x:hs[x])
            hs[k]=max(v312.MIN_H[k]*.90,hs[k]-over)

    zones={}; y=top
    for k in keys:
        z=dict(v310.ZONES[k]); z['y']=y; z['h']=hs[k]
        zones[k]=z; y += hs[k]+GAP
    return zones, fonts


def _layout(d, imgs):
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    left_imgs={'2D':imgs.get('2D',[]),'3D':imgs.get('3D',[]),'4D_CAUSE':imgs.get('4D_CAUSE',[])}
    right_imgs={'4D_LEAK':imgs.get('4D_LEAK',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}
    lz,lf=_fit_chain_individual(CHAINS[0],texts,left_imgs)
    rz,rf=_fit_chain_individual(CHAINS[1],texts,right_imgs)
    return {**lz,**rz}, texts, {**lf,**rf}


def _render(sl,key,z,text,images,font):
    if not text:
        return
    blob=v313._collage(images)
    if blob:
        iw=min(1.65,z['w']*.36); tw=z['w']-iw-.08
        v310.add_box(sl,key,z['x'],z['y'],tw,z['h'],text,font)
        v310.add_image(sl,blob,z['x']+tw+.08,z['y'],iw,z['h'],key)
    else:
        v310.add_box(sl,key,z['x'],z['y'],z['w'],z['h'],text,font)


# ---------- exact page-2 header formatting ----------
def _set_weekly_title_preserve(sh, task, issue):
    if not hasattr(sh,'text_frame'):
        return False
    t=N(getattr(sh,'text',''))
    if '과제명' not in t or '이슈' not in t or '제목' not in t:
        return False
    task=N(task); issue=N(issue)
    target = (task + '_' + issue).strip('_')
    # Supplied example has three runs: '과제명' / '_' / '이슈 제목'.
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    if len(runs)>=3:
        runs[0].text=task or issue
        # keep the separator run and its original formatting
        sep_idx=None
        for i,r in enumerate(runs):
            if '_' in r.text:
                sep_idx=i; break
        if sep_idx is not None:
            runs[sep_idx].text='_' if task and issue else ''
            after=False
            for i,r in enumerate(runs):
                if i<=sep_idx: continue
                if not after:
                    r.text=issue if task else ''
                    after=True
                else:
                    r.text=''
            return True
    # Fallback: text replacement may alter run splits but keeps shape geometry.
    sh.text_frame.text=target
    return True


def _set_issue_line_preserve(sh, issue):
    if not hasattr(sh,'text_frame'):
        return False
    t=N(getattr(sh,'text',''))
    if not C(t).startswith('이슈명'):
        return False
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    colon_idx=None
    for i,r in enumerate(runs):
        if ':' in r.text or '：' in r.text:
            colon_idx=i; break
    if colon_idx is not None:
        prefix=runs[colon_idx].text.split(':',1)[0].split('：',1)[0]
        runs[colon_idx].text=prefix+': '+N(issue)
        for r in runs[colon_idx+1:]:
            r.text=''
        return True
    # If label and value are separated unusually, preserve first run formatting.
    if runs:
        runs[0].text='이슈명 : '+N(issue)
        for r in runs[1:]: r.text=''
        return True
    return False


def _meta(sl,d,g):
    task=N(d.get('task_name')); issue=N(d.get('issue_name'))
    # Exact textboxes from supplied weekly example.
    for sh in v310.walk(sl):
        if not hasattr(sh,'text_frame'):
            continue
        _set_weekly_title_preserve(sh,task,issue)
        _set_issue_line_preserve(sh,issue)
    # Keep v3.1.4 table metadata logic for 발생단계/Signal/etc.
    v314._meta(sl,d,g)
    # v314._meta may have rewritten the title generically; enforce exact title once more.
    for sh in v310.walk(sl):
        if hasattr(sh,'text_frame'):
            t=N(getattr(sh,'text',''))
            if '과제명' in t and '이슈' in t and '제목' in t:
                _set_weekly_title_preserve(sh,task,issue)
            elif C(t).startswith('이슈명'):
                _set_issue_line_preserve(sh,issue)


def update_page2(sl,d,g,mode):
    # Render using v3.1.6 per-section font/layout, then enforce v3.1.5 4D policy.
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    zones,texts,fonts=_layout(d,imgs)

    hc=bool(N(d.get('cause_4d')))
    hl=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))

    # Move non-4D marker/title groups to their dynamically calculated sections.
    if v310.is_new(mode):
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]
            v310.move_marker_unit(sl,label,max(.10,z['x']-.18),max(.10,z['y']-.35))

    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        _render(sl,key,zones[key],texts[key],imgs.get(key,[]),fonts[key])

    _meta(sl,d,g)

    zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']
    left_xy=(max(.10,zl['x']-.18),max(.10,zl['y']-.35))
    right_xy=(max(.10,zr['x']-.18),max(.10,zr['y']-.35))
    v315._ensure_4d_units(sl,hc,hl,left_xy,right_xy)


def weekly(src,out,d,g,mode):
    prs=Presentation(src)
    v314.update_page1(prs,d,g)
    if len(prs.slides)>1:
        update_page2(prs.slides[1],d,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.extract=extract
base.weekly=weekly

if __name__=='__main__':
    base.App().mainloop()
