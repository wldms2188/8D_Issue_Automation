# 8D Issue Automation v3.2.0
# Adaptive page-2 layout when 4D leak cause is absent:
# - keep 4D occurrence cause on left if 2D/3D/4D fit comfortably at 8pt
# - otherwise move 4D occurrence cause to right-top and stack 5D/6D below it
# - if 4D remains left and right-top is empty, pull 5D upward
import datetime
from pathlib import Path

import main_v319 as v319
import main_v318 as v318
import main_v315 as v315
import main_v310 as v310
from pptx import Presentation

base=v319.base
N=v310.N

MAX_FONT=8.0
BOTTOM=v319.BOTTOM
GAP=v319.GAP


def _fit_at_8(keys,d,texts,imgs):
    active=[k for k in keys if v319._has_section(d,k)]
    if not active:
        return True,0.0,0.0
    top=v310.ZONES[active[0]]['y']
    available=BOTTOM-top-GAP*max(0,len(active)-1)
    need=sum(v319._need_h(k,texts[k],bool(imgs.get(k)),MAX_FONT) for k in active)
    return need<=available,need,available


def _adaptive_layout(d,imgs):
    texts={k:v319.v313._section_text(d,k) for k in v310.ZONES}
    has_cause=bool(N(d.get('cause_4d')))
    has_leak=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))

    li={'2D':imgs.get('2D',[]),'3D':imgs.get('3D',[]),'4D_CAUSE':imgs.get('4D_CAUSE',[])}
    ri={'4D_LEAK':imgs.get('4D_LEAK',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}

    # Default/original two-column case when leak/system cause exists.
    if has_leak:
        lz,lf=v319._fit_chain(('2D','3D','4D_CAUSE'),d,texts,li)
        rz,rf=v319._fit_chain(('4D_LEAK','5D','6D'),d,texts,ri)
        return {**lz,**rz},texts,{**lf,**rf},'normal'

    # No leak/system cause. Decide whether left-side 4D can remain at 8pt without pressure.
    move_cause_right=False
    if has_cause:
        fits,need,available=_fit_at_8(('2D','3D','4D_CAUSE'),d,texts,li)
        # Keep a small safety reserve so real PowerPoint wrapping does not create a 2D/3D overlap.
        reserve=.18
        move_cause_right=(not fits) or (available-need<reserve)

    if has_cause and move_cause_right:
        # Left contains only 2D/3D. Right starts with occurrence-cause 4D, then 5D, then 6D.
        lz,lf=v319._fit_chain(('2D','3D'),d,texts,{'2D':li['2D'],'3D':li['3D']})

        # Reuse right-top 4D geometry for occurrence-cause content.
        d2=dict(d)
        d2['leak_cause']=N(d.get('cause_4d'))
        d2['system_cause']=''
        d2['cause_4d']=''
        texts2=dict(texts)
        texts2['4D_LEAK']=texts['4D_CAUSE']
        ri2={'4D_LEAK':imgs.get('4D_CAUSE',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}
        rz,rf=v319._fit_chain(('4D_LEAK','5D','6D'),d2,texts2,ri2)

        zones={**lz,**rz}
        zones['4D_CAUSE']=dict(rz['4D_LEAK'])
        fonts={**lf,**rf}
        fonts['4D_CAUSE']=rf['4D_LEAK']
        return zones,texts,fonts,'cause_right'

    # Cause remains on left (or does not exist). Right-top 4D is empty, so pull 5D upward.
    lz,lf=v319._fit_chain(('2D','3D','4D_CAUSE'),d,texts,li)

    # Build 5D/6D chain starting at the former right-top 4D y-position.
    right_keys=('5D','6D')
    fonts={k:MAX_FONT for k in right_keys}
    top=v310.ZONES['4D_LEAK']['y']
    active=list(right_keys)
    available=BOTTOM-top-GAP
    hs={k:v319._need_h(k,texts[k],bool(imgs.get(k)),MAX_FONT) for k in active}
    while sum(hs.values())>available+.06:
        options=[]
        for k in active:
            if fonts[k]<=v319.MIN_FONT:
                continue
            nf=max(v319.MIN_FONT,fonts[k]-.5)
            nh=v319._need_h(k,texts[k],bool(imgs.get(k)),nf)
            options.append((hs[k]-nh,k,nf,nh))
        if not options:
            break
        gain,k,nf,nh=max(options,key=lambda x:x[0])
        if gain<=.01:
            break
        fonts[k]=nf; hs[k]=nh
    rz={}; y=top
    for k in active:
        z=dict(v310.ZONES[k]); z['y']=y; z['h']=hs[k]; rz[k]=z
        y+=hs[k]+GAP
    rz['4D_LEAK']=dict(v310.ZONES['4D_LEAK'])
    rf={'4D_LEAK':MAX_FONT,**fonts}
    return {**lz,**rz},texts,{**lf,**rf},'pull_5d_up'


def _update_page2(sl,d,g,mode):
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    zones,texts,fonts,policy=_adaptive_layout(d,imgs)
    hc=bool(N(d.get('cause_4d')))
    hl=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))

    # Move 2D/3D/5D/6D marker+title units to the newly calculated zones.
    if v310.is_new(mode):
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]
            v310.move_marker_unit(sl,label,max(.10,z['x']-.18),max(.10,z['y']-.35))

    # Render standard sections first.
    for key in ('2D','3D','5D','6D'):
        v319._render(sl,key,zones[key],texts[key],imgs.get(key,[]),fonts[key])

    if hl:
        if hc:
            v319._render(sl,'4D_CAUSE',zones['4D_CAUSE'],texts['4D_CAUSE'],imgs.get('4D_CAUSE',[]),fonts['4D_CAUSE'])
        v319._render(sl,'4D_LEAK',zones['4D_LEAK'],texts['4D_LEAK'],imgs.get('4D_LEAK',[]),fonts['4D_LEAK'])
    elif hc:
        if policy=='cause_right':
            # Occurrence cause uses the right-top zone but keeps occurrence-cause text/images.
            v319._render(sl,'4D_CAUSE',zones['4D_CAUSE'],texts['4D_CAUSE'],imgs.get('4D_CAUSE',[]),fonts['4D_CAUSE'])
        else:
            v319._render(sl,'4D_CAUSE',zones['4D_CAUSE'],texts['4D_CAUSE'],imgs.get('4D_CAUSE',[]),fonts['4D_CAUSE'])

    v319._page2_meta(sl,d,g)

    # Strict marker policy, adjusted for adaptive cause-right layout.
    if policy=='cause_right' and hc and not hl:
        zr=zones['4D_CAUSE']
        right_xy=(max(.10,zr['x']-.18),max(.10,zr['y']-.35))
        # Treat the single occurrence-cause 4D as the right-side visible unit.
        v315._ensure_4d_units(sl,False,True,(.1,.1),right_xy)
    else:
        zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']
        left_xy=(max(.10,zl['x']-.18),max(.10,zl['y']-.35))
        right_xy=(max(.10,zr['x']-.18),max(.10,zr['y']-.35))
        v315._ensure_4d_units(sl,hc,hl,left_xy,right_xy)


def weekly(src,out,d,g,mode):
    dd=v319._weekly_display(d)
    prs=Presentation(src)
    v319._update_page1(prs,dd,g)
    if len(prs.slides)>1:
        _update_page2(prs.slides[1],dd,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

# Keep Issue DB/extraction behavior from v3.1.8 unchanged.
base.extract=v318.extract
base.weekly=weekly

if __name__=='__main__':
    base.App().mainloop()
