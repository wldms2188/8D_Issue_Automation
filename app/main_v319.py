# 8D Issue Automation v3.1.9
# Weekly-only presentation refinements on top of v3.1.8.
import re, math, datetime
from pathlib import Path

import main_v318 as v318
import main_v315 as v315
import main_v314 as v314
import main_v313 as v313
import main_v311 as v311
import main_v310 as v310
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

base = v318.base
N = v310.N
C = v310.C
EMU = v310.EMU

MAX_FONT = 8.0
MIN_FONT = 7.0
BOTTOM = 7.47
GAP = 0.10
LEFT = ('2D','3D','4D_CAUSE')
RIGHT = ('4D_LEAK','5D','6D')


def _clean_issue_label(text):
    """Remove document-type words from user-facing issue names."""
    s=N(text)
    if not s:return ''
    s=re.sub(r'(?i)(?<![A-Za-z0-9])(?:8D|REPORT)(?![A-Za-z0-9])',' ',s)
    s=re.sub(r'[ _/|:-]{2,}','_',s)
    s=re.sub(r'\s+',' ',s)
    return s.strip(' _-/|:')

def _trim_before_customer(issue, customer):
    """Keep issue text from the customer token onward; remove document-type words."""
    issue=N(issue); customer=N(customer)
    if not issue or not customer:
        return _clean_issue_label(issue)
    toks=[x.strip() for x in issue.split('_') if x.strip()]
    for i,t in enumerate(toks):
        if C(t)==C(customer):
            return _clean_issue_label('_'.join(toks[i:]))
    # Fallback if customer is embedded in a token or separators are irregular.
    m=re.search(re.escape(customer), issue, re.I)
    return _clean_issue_label(issue[m.start():].strip(' _-/') if m else issue)


def _weekly_task(d):
    customer=N(d.get('customer')); task=N(d.get('task_name'))
    if customer and task:
        if C(task)==C(customer):
            return customer
        pat=re.compile(r'^(?:'+re.escape(customer)+r'\s*[_\-/／|:：]\s*)+',re.I)
        m=pat.match(task)
        if m:
            rest=N(task[m.end():]).strip(' _-/／|:：')
            return customer+('_'+rest if rest else '')
        return customer+'_'+task
    return customer or task


def _title_issue(issue_from_customer):
    """Use a short issue category in the page-2 title; keep details in the issue-name line."""
    q=N(issue_from_customer)
    if '시험' in q:
        return '시험 이슈 발생'
    if '빌드' in q:
        return '빌드 이슈 발생'
    return q


def _clean_detail_line(s):
    s=N(s)
    s=re.sub(r'^[-•·▪◦]\s*','',s)
    s=re.sub(r'^\(?\d+\)?[.)]\s*','',s)
    return s.strip()


def _numbered_detail(text):
    lines=[]
    for raw in N(text).split('\n'):
        v=_clean_detail_line(raw)
        if v:
            lines.append(v)
    return '\n'.join(f'{i}) {v}' for i,v in enumerate(lines,1))


def _progress_text(d):
    cause=[]
    for lab,key in [('발생원인','cause_4d'),('유출원인','leak_cause'),('시스템원인','system_cause')]:
        v=N(d.get(key))
        if v:
            detail=_numbered_detail(v)
            cause.append(f'- {lab}\n{detail}' if detail else f'- {lab}')
    if not cause:
        cause=['- 검토 중']

    progress=[]
    for lab,key in [('임시조치','temporary_action'),('고객대응','customer_response'),('개선대책','action_5d'),('효과검증','verification_6d'),('수평전개','spread_7d')]:
        v=N(d.get(key))
        if v:
            detail=_numbered_detail(v)
            progress.append(f'- {lab}\n{detail}' if detail else f'- {lab}')
    if not progress:
        progress=['- 검토 중']

    req=N(d.get('request') or d.get('request_item') or d.get('requests'))
    if req:
        detail=_numbered_detail(req)
        requests=['- 요청사항\n'+detail] if detail else ['- 요청사항']
    else:
        requests=['- ']

    return '1. 원인\n'+'\n'.join(cause)+'\n\n2. 진행현황\n'+'\n'.join(progress)+'\n\n3. 요청사항\n'+'\n'.join(requests)


def _weekly_display(d):
    dd=dict(d)
    issue=_trim_before_customer(d.get('issue_name'), d.get('customer'))
    dd['issue_name']=issue
    dd['task_name']=_weekly_task(d)
    dd['_title_issue']=_title_issue(issue)
    return dd


def _has_section(d,key):
    if key=='4D_CAUSE':
        return bool(N(d.get('cause_4d')))
    if key=='4D_LEAK':
        return bool(N(d.get('leak_cause')) or N(d.get('system_cause')))
    return True


def _text_width(key, has_images):
    z=v310.ZONES[key]
    iw=min(.95,z['w']*.18) if has_images else 0
    return max(1.0,z['w']-iw-(.06 if has_images else 0))


def _line_count(text, width_in, font):
    # Deliberately conservative for Korean/English mixed text. Previous 14 chars/in
    # underestimated wrapping and caused 2D text to spill into 3D.
    chars=max(10,int(width_in*10.8*(8.0/font)))
    total=0
    for raw in (N(text).split('\n') or ['']):
        line=N(raw)
        # Korean/full-width characters consume more horizontal space than Latin text.
        units=sum(1.0 if ord(ch)>=0x2E80 else (.45 if ch.isspace() else .62) for ch in line)
        total += max(1, math.ceil(units/max(1,chars*.72)))
    return total


def _need_h(key,text,has_images,font=8.0):
    w=_text_width(key,has_images)
    lines=_line_count(text,w,font)
    # 8 pt line height + explicit top/bottom safety margin to prevent text overflow.
    line_h=font/72.0*1.28
    return max(v318.v316.v312.MIN_H[key]*.90, lines*line_h+.20)


def _fit_chain(keys,d,texts,imgs):
    active=[k for k in keys if _has_section(d,k)]
    fonts={k:MAX_FONT for k in keys}
    if not active:
        return {k:dict(v310.ZONES[k]) for k in keys},fonts

    top=v310.ZONES[active[0]]['y']
    available=BOTTOM-top-GAP*max(0,len(active)-1)
    hs={k:_need_h(k,texts[k],bool(imgs.get(k)),MAX_FONT) for k in active}

    # First priority: keep 8 pt and move every following D block downward.
    # Only shrink when the complete active column genuinely cannot fit on the slide.
    while sum(hs.values()) > available+.06:
        candidates=[k for k in active if fonts[k]>MIN_FONT]
        if not candidates:
            break
        # Shrink the block that recovers the most vertical space, not the whole column.
        options=[]
        for k in candidates:
            nf=max(MIN_FONT,fonts[k]-.5)
            nh=_need_h(k,texts[k],bool(imgs.get(k)),nf)
            options.append((hs[k]-nh,k,nf,nh))
        gain,k,nf,nh=max(options,key=lambda x:x[0])
        if gain<=.01:
            break
        fonts[k]=nf; hs[k]=nh

    # If still a tiny amount over, compress only geometry; don't force all text smaller.
    overflow=sum(hs.values())-available
    if 0<overflow<=.18:
        largest=max(active,key=lambda k:hs[k])
        hs[largest]=max(v318.v316.v312.MIN_H[largest]*.88,hs[largest]-overflow)

    zones={}; y=top
    for k in active:
        z=dict(v310.ZONES[k]); z['y']=y; z['h']=hs[k]
        zones[k]=z
        y+=hs[k]+GAP
    for k in keys:
        zones.setdefault(k,dict(v310.ZONES[k]))
    return zones,fonts


def _layout(d,imgs):
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    li={'2D':imgs.get('2D',[]),'3D':imgs.get('3D',[]),'4D_CAUSE':imgs.get('4D_CAUSE',[])}
    ri={'4D_LEAK':imgs.get('4D_LEAK',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}
    lz,lf=_fit_chain(LEFT,d,texts,li)
    rz,rf=_fit_chain(RIGHT,d,texts,ri)
    return {**lz,**rz},texts,{**lf,**rf}


def _render(sl,key,z,text,images,font):
    if not text:
        return
    blob=v313._collage(images)
    if blob:
        iw=min(.95,z['w']*.18); tw=z['w']-iw-.06
        v310.add_box(sl,key,z['x'],z['y'],tw,z['h'],text,font)
        v310.add_image(sl,blob,z['x']+tw+.06,z['y'],iw,z['h'],key)
    else:
        v310.add_box(sl,key,z['x'],z['y'],z['w'],z['h'],text,font)


def _replace_title_runs(sh,task,title_issue):
    if not hasattr(sh,'text_frame'):
        return
    target=(N(task)+'_'+N(title_issue)).strip('_')
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    if not runs:
        sh.text_frame.text=target
        return

    # Preserve the original run formatting from the weekly template.
    runs[0].text=N(task) or N(title_issue)
    sep=None
    for i,r in enumerate(runs):
        if '_' in r.text:
            sep=i; break
    if sep is None and len(runs)>1:
        sep=1
    if sep is not None:
        runs[sep].text='_' if task and title_issue else ''
        if sep+1<len(runs):
            runs[sep+1].text=N(title_issue) if task else ''
            for r in runs[sep+2:]:
                r.text=''
    elif len(runs)==1:
        runs[0].text=target

    tf=sh.text_frame
    tf.word_wrap=False
    try:
        tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass

    # Also pre-fit the font so the saved PPT already opens with a safe one-line title.
    pts=[r.font.size.pt for r in runs if r.font.size]
    base_pt=max(pts) if pts else 20.0
    width=max(.5,float(sh.width)/EMU)
    units=sum(1.0 if ord(ch)>=0x2E80 else (.32 if ch.isspace() else .58) for ch in target)
    # Approximate one-line capacity at the original font size.
    capacity=width*10.2*(20.0/base_pt)
    ratio=min(1.0,capacity/max(units,1.0))
    fitted=max(9.0,math.floor(base_pt*ratio*2)/2.0)
    if fitted<base_pt:
        for r in runs:
            r.font.size=Pt(fitted)


def _set_issue_line(sh,issue):
    if not hasattr(sh,'text_frame'):
        return
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    if runs:
        runs[0].text='이슈명 : '+N(issue)
        for r in runs[1:]: r.text=''
    else:
        sh.text_frame.text='이슈명 : '+N(issue)


def _find_page2_header_shapes(sl):
    title=None; issue=None
    for sh in v310.walk(sl):
        if not hasattr(sh,'text_frame'):
            continue
        t=N(getattr(sh,'text','')); q=C(t)
        if title is None and '과제명' in t and '이슈' in t and '제목' in t:
            title=sh
        if issue is None and q.startswith('이슈명'):
            issue=sh
    return title,issue


def _page2_meta(sl,d,g):
    # Capture template shapes before any generic metadata function can erase placeholders.
    title_sh,issue_sh=_find_page2_header_shapes(sl)
    task=N(d.get('task_name')); issue=N(d.get('issue_name')); title_issue=N(d.get('_title_issue'))

    if title_sh is not None:
        _replace_title_runs(title_sh,task,title_issue)
    if issue_sh is not None:
        _set_issue_line(issue_sh,issue)

    # Update occurrence-stage table without touching the title or issue-line formatting.
    occ=v311._occurrence_value(d,g)
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False):
            continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                q=C(tb.cell(r,c).text)
                if '이슈명' in q and c+1<len(tb.columns):
                    v314._set_cell(tb.cell(r,c+1),issue,8)
                if '발생단계' in q and c+1<len(tb.columns):
                    v314._set_cell(tb.cell(r,c+1),occ,9)
    try:
        v311._fill_occurrence(sl,d,g)
    except Exception:
        pass


def _update_page1(prs,d,g):
    # Reuse the stable page-1 mapping, but use the new structured progress formatter.
    old=v314._progress
    try:
        v314._progress=_progress_text
        v314.update_page1(prs,d,g)
    finally:
        v314._progress=old


def _update_page2(sl,d,g,mode):
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    zones,texts,fonts=_layout(d,imgs)
    hc=bool(N(d.get('cause_4d')))
    hl=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))

    if v310.is_new(mode):
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]
            v310.move_marker_unit(sl,label,max(.10,z['x']-.18),max(.10,z['y']-.35))

    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        if _has_section(d,key):
            _render(sl,key,zones[key],texts[key],imgs.get(key,[]),fonts[key])

    _page2_meta(sl,d,g)

    zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']
    left_xy=(max(.10,zl['x']-.18),max(.10,zl['y']-.35))
    right_xy=(max(.10,zr['x']-.18),max(.10,zr['y']-.35))
    v315._ensure_4d_units(sl,hc,hl,left_xy,right_xy)


def weekly(src,out,d,g,mode):
    dd=_weekly_display(d)
    prs=Presentation(src)
    _update_page1(prs,dd,g)
    if len(prs.slides)>1:
        _update_page2(prs.slides[1],dd,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved


# Keep v3.1.8 extraction/Issue-DB behavior unchanged; only weekly output changes here.
base.extract=v318.extract
base.weekly=weekly

if __name__=='__main__':
    base.App().mainloop()
