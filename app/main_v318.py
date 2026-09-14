# 8D Issue Automation v3.1.8
import re, math, datetime
from pathlib import Path

import main_v317 as v317
import main_v316 as v316
import main_v315 as v315
import main_v314 as v314
import main_v313 as v313
import main_v310 as v310
from pptx import Presentation
from pptx.util import Pt

base=v317.base
N=v310.N
C=v310.C


def clean_issue_name(path):
    s=Path(path).stem.strip()
    s=re.sub(r'\s*\(\d+\)\s*$','',s)
    s=re.sub(r'(?<!\d)(?:19|20)\d{2}[.\-]\d{1,2}[.\-]\d{1,2}(?!\d)','',s)
    s=re.sub(r'(?<!\d)\d{2}[.\-]\d{1,2}[.\-]\d{1,2}(?!\d)','',s)
    parts=[x.strip() for x in s.split('_') if x.strip()]
    out=[]; i=0
    while i<len(parts):
        p=parts[i]; q=re.sub(r'\s+','',p)
        if re.fullmatch(r'\d{1,2}[dD]',q) or re.fullmatch(r'\(\d+\)',q):
            i+=1; continue
        if re.fullmatch(r'(?:19|20)\d{6}',q) or re.fullmatch(r'\d{6}',q):
            i+=1; continue
        if (re.fullmatch(r'(?:19|20)\d{2}',q) or re.fullmatch(r'\d{2}',q)) and i+2<len(parts):
            m=parts[i+1]; d=parts[i+2]
            if re.fullmatch(r'0?[1-9]|1[0-2]',m) and re.fullmatch(r'0?[1-9]|[12]\d|3[01]',d):
                i+=3; continue
        out.append(p); i+=1
    return re.sub(r'_+','_','_'.join(out)).strip(' _-.')


def task_after_customer(issue,customer,current=''):
    toks=[x.strip() for x in N(issue).split('_') if x.strip()]
    customer=N(customer)
    if customer:
        for i,t in enumerate(toks):
            if C(t)==C(customer) and i+1<len(toks):
                return toks[i+1]
    cur=N(current)
    return '' if C(cur) in ('상세시험조건','model','packer') else cur


_old_extract=v314.extract

def extract(path):
    d=_old_extract(path)
    issue=clean_issue_name(path)
    d['issue_name']=issue
    d['task_name']=task_after_customer(issue,d.get('customer'),d.get('task_name'))
    return d


MAX_FONT=8.0
MIN_FONT=7.0
BOTTOM=7.46
GAP=0.07
LEFT=('2D','3D','4D_CAUSE')
RIGHT=('4D_LEAK','5D','6D')


def has_section(d,key):
    if key=='4D_CAUSE': return bool(N(d.get('cause_4d')))
    if key=='4D_LEAK': return bool(N(d.get('leak_cause')) or N(d.get('system_cause')))
    return True


def need_h(key,text,has_images,font=8.0):
    z=v310.ZONES[key]
    iw=min(1.10,z['w']*.20) if has_images else 0
    w=max(1.0,z['w']-iw-(.06 if has_images else 0))
    chars=max(12,int(w*14.0*(8.0/font)))
    lines=sum(max(1,math.ceil(max(1,len(line))/chars)) for line in (N(text).split('\n') or ['']))
    return max(v316.v312.MIN_H[key]*.86,lines*(font/72.0*1.05)+.10)


def fit_chain(keys,d,texts,imgs):
    active=[k for k in keys if has_section(d,k)]
    fonts={k:MAX_FONT for k in keys}
    if not active:
        return {k:dict(v310.ZONES[k]) for k in keys},fonts
    top=v310.ZONES[active[0]]['y']
    available=BOTTOM-top-GAP*max(0,len(active)-1)
    hs={k:need_h(k,texts[k],bool(imgs.get(k)),8.0) for k in active}
    # Protect 8pt unless estimated overflow is substantial.
    while sum(hs.values())-available>.40:
        candidates=[k for k in active if fonts[k]>MIN_FONT]
        if not candidates: break
        k=max(candidates,key=lambda x:hs[x])
        fonts[k]=max(MIN_FONT,fonts[k]-.5)
        hs[k]=need_h(k,texts[k],bool(imgs.get(k)),fonts[k])
    total=sum(hs.values())
    if total>available and total>0:
        scale=available/total
        if scale>=.86:
            hs={k:h*scale for k,h in hs.items()}
    zones={}; y=top
    for k in active:
        z=dict(v310.ZONES[k]); z['y']=y; z['h']=hs[k]; zones[k]=z
        y+=hs[k]+GAP
    for k in keys:
        zones.setdefault(k,dict(v310.ZONES[k]))
    return zones,fonts


def layout(d,imgs):
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    li={'2D':imgs.get('2D',[]),'3D':imgs.get('3D',[]),'4D_CAUSE':imgs.get('4D_CAUSE',[])}
    ri={'4D_LEAK':imgs.get('4D_LEAK',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}
    lz,lf=fit_chain(LEFT,d,texts,li); rz,rf=fit_chain(RIGHT,d,texts,ri)
    return {**lz,**rz},texts,{**lf,**rf}


def render(sl,key,z,text,images,font):
    if not text: return
    blob=v313._collage(images)
    if blob:
        iw=min(1.10,z['w']*.20); tw=z['w']-iw-.06
        v310.add_box(sl,key,z['x'],z['y'],tw,z['h'],text,font)
        v310.add_image(sl,blob,z['x']+tw+.06,z['y'],iw,z['h'],key)
    else:
        v310.add_box(sl,key,z['x'],z['y'],z['w'],z['h'],text,font)


def fit_title(sh,target,min_pt=10.0):
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    if not runs: return
    pts=[r.font.size.pt for r in runs if r.font.size]
    base_pt=max(pts) if pts else 20.0
    width_in=max(.5,float(sh.width)/914400.0)
    weight=sum(1.0 if ord(ch)>=0x2E80 else (.35 if ch in 'ilI1|.,:;! ' else .58) for ch in N(target))
    capacity=width_in*10.8
    pt=base_pt if weight<=capacity else max(min_pt,base_pt*capacity/max(weight,1))
    pt=math.floor(pt*2)/2.0
    for r in runs: r.font.size=Pt(pt)
    sh.text_frame.word_wrap=False


def set_title(sh,task,issue):
    t=N(getattr(sh,'text',''))
    if not hasattr(sh,'text_frame') or '과제명' not in t or '이슈' not in t or '제목' not in t: return False
    task=N(task); issue=N(issue); target=(task+'_'+issue).strip('_')
    runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
    if not runs: return False
    runs[0].text=task or issue
    sep=1 if len(runs)>1 else None
    for i,r in enumerate(runs):
        if '_' in r.text: sep=i; break
    if sep is not None:
        runs[sep].text='_' if task and issue else ''
        if sep+1<len(runs):
            runs[sep+1].text=issue if task else ''
            for r in runs[sep+2:]: r.text=''
    fit_title(sh,target)
    return True


def meta(sl,d,g):
    v316._meta(sl,d,g)
    task=N(d.get('task_name')); issue=N(d.get('issue_name'))
    for sh in v310.walk(sl):
        if hasattr(sh,'text_frame'):
            set_title(sh,task,issue)
            if C(N(getattr(sh,'text',''))).startswith('이슈명'):
                v316._set_issue_line_preserve(sh,issue)


def update_page2(sl,d,g,mode):
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    zones,texts,fonts=layout(d,imgs)
    hc=bool(N(d.get('cause_4d'))); hl=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))
    if v310.is_new(mode):
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]; v310.move_marker_unit(sl,label,max(.10,z['x']-.18),max(.10,z['y']-.35))
    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        if has_section(d,key): render(sl,key,zones[key],texts[key],imgs.get(key,[]),fonts[key])
    meta(sl,d,g)
    zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']
    v315._ensure_4d_units(sl,hc,hl,(max(.10,zl['x']-.18),max(.10,zl['y']-.35)),(max(.10,zr['x']-.18),max(.10,zr['y']-.35)))


def weekly(src,out,d,g,mode):
    prs=Presentation(src)
    v314.update_page1(prs,d,g)
    if len(prs.slides)>1: update_page2(prs.slides[1],d,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try: prs.save(out); saved=out
    except PermissionError:
        p=Path(out); saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix)); prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.extract=extract
base.weekly=weekly

if __name__=='__main__': base.App().mainloop()
