import os, re
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook
import main_recovery_step6 as step6
import main_recovery_step3 as step3
import main_v310 as v310

base=step6.base
N=v310.N
_old_write_row=base.write_row
_EXCEL_ILLEGAL_CONTROL_RE=re.compile(r'[\x00-\x08\x0B\x0C\x0E-\x1F]')

def _excel_safe(text):
    return _EXCEL_ILLEGAL_CONTROL_RE.sub('',N(text))


def _db_summary(d):
    try:
        xs=step3._actual_symptoms(d)
    except Exception:
        xs=[]
    if xs:
        return '\n'.join(xs)
    p=N(d.get('problem'))
    bad=('발생경위','확인사항','시험조건','시험 방법','시험방법','시험절차','시험환경','평가조건','평가방법','온도','습도','soc','dod','soh','전류','전압','cycle','사이클','샘플수')
    out=[]
    for raw in p.split('\n'):
        s=raw.strip()
        q=s.replace(' ','').lower()
        if not s or any(x.replace(' ','').lower() in q for x in bad):
            continue
        if s not in out:
            out.append(s)
        if len(out)>=3:
            break
    return '\n'.join(out) if out else p

def write_row_step7(ws,r,d,g):
    result=_old_write_row(ws,r,d,g)
    if '_db_problem_selected' in g:
        ws.cell(r,14).value=_excel_safe(g['_db_problem_selected'])
    return result
base.write_row=write_row_step7

class ProblemChoiceDialog(tk.Toplevel):
    def __init__(self,parent,summary):
        super().__init__(parent)
        self.title('Issue DB 현상 입력 확인')
        self.geometry('720x470')
        self.result=None
        self.transient(parent)
        self.grab_set()
        ttk.Label(self,text='이렇게 요약됩니다.\n전체 내용을 기재해드릴까요? 아니면 현상 내용만 요약한 것으로 작성해드릴까요?',justify='left').pack(anchor='w',padx=20,pady=(18,10))
        ttk.Label(self,text='현상 요약 미리보기').pack(anchor='w',padx=20,pady=(4,4))
        box=tk.Text(self,height=13,wrap='word')
        box.pack(fill='both',expand=True,padx=20,pady=(0,12))
        box.insert('1.0',summary or '(요약 내용 없음)')
        box.configure(state='disabled')
        row=ttk.Frame(self); row.pack(pady=(0,18))
        ttk.Button(row,text='전체 내용',command=lambda:self.choose('full'),width=16).pack(side='left',padx=6)
        ttk.Button(row,text='현상 요약',command=lambda:self.choose('summary'),width=16).pack(side='left',padx=6)
        ttk.Button(row,text='취소',command=lambda:self.choose(None),width=16).pack(side='left',padx=6)
        self.protocol('WM_DELETE_WINDOW',lambda:self.choose(None))
        self.wait_visibility(); self.focus_set(); self.wait_window(self)
    def choose(self,value):
        self.result=value; self.destroy()

class RecoveryStep7App(step6.RecoveryStep6App):
    def __init__(self):
        super().__init__(); self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP7')
    def run(self):
        g=self.gui()
        if any(not g[k] or not os.path.exists(g[k]) for k in ['ppt8d','pptweekly','xlsx']):
            return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')
        try:
            d=base.extract(g['ppt8d']); mode=self.mode.get()
            summary=_db_summary(d)
            dlg=ProblemChoiceDialog(self,summary)
            if dlg.result is None:
                self.log.delete('1.0','end'); self.log.insert('end','업데이트가 취소되었습니다.'); return
            g['_db_problem_selected']=summary if dlg.result=='summary' else N(d.get('problem'))
            out=Path(g['xlsx']).parent/'자동화_결과'; out.mkdir(exist_ok=True)
            xo=out/(Path(g['xlsx']).stem+'_업데이트.xlsx'); po=out/(Path(g['pptweekly']).stem+'_업데이트.pptx')
            wb=load_workbook(g['xlsx']); ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active
            r,_=base.find(ws,d,g)
            if mode=='existing' and not r:
                return messagebox.showwarning('업데이트 중단','기존 이슈를 Excel에서 정확하게 찾지 못했습니다. PMS/PLM 이슈번호를 입력하거나 이슈 정보를 확인하세요.')
            if mode=='new' and r and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?'):
                return
            a,_=base.update_excel(g['xlsx'],xo,d,g,new=(mode=='new'))
            b,_=base.weekly(g['pptweekly'],po,d,g,mode)
            label='현상 요약' if dlg.result=='summary' else '전체 내용'
            self.log.delete('1.0','end'); self.log.insert('end',f'=== 완료 ===\n이슈 구분: {mode}\nIssue DB 현상: {label}\n상태: {base.status(d)}\n\n{a}\n{b}\n\n결과: {out}')
            messagebox.showinfo('완료',f'자동 업데이트가 완료되었습니다.\n\n{out}')
        except Exception as e:
            messagebox.showerror('실행 오류',repr(e))

if __name__=='__main__':
    RecoveryStep7App().mainloop()
