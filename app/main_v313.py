# 8D Issue Automation v3.1.3
# Exact user-requested weekly mapping + section-specific images + Excel 2D representative collage.
import io, re, copy, math, datetime
from pathlib import Path

import main_v312 as v312
import main_v311 as v311
import main_v310 as v310
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from PIL import Image as PILImage, ImageOps, ImageDraw

base = v312.base
N = v310.N
C = v310.C
box = v310.box
walk = v310.walk
EMU = v310.EMU

# ------------------------- common helpers -------------------------
def _safe_text(sh):
    return N(getattr(sh, 'text', ''))


def _explicit_d(text):
    q = C(text)
    for d in ('2D','3D','4D','5D','6D'):
        if re.search(r'(?<![0-9a-z])' + d.lower() + r'(?![0-9a-z])', N(text).lower()) or d.lower() in q:
            return d
    return None


def _sub4(text):
    q = C(text)
    if any(k in q for k in ('발생원인','발생원인분석')):
        return '4D_CAUSE'
    if any(k in q for k in ('유출원인','유츌원인','시스템원인','유출원인분석')):
        return '4D_LEAK'
    return None


def _picture_blob(sh):
    try:
        if getattr(sh, 'shape_type', None) == MSO_SHAPE_TYPE.PICTURE:
            return sh.image.blob
        if getattr(sh, 'shape_type', None) == MSO_SHAPE_TYPE.GROUP:
            pics = [p for p in walk(sh) if getattr(p, 'shape_type', None) == MSO_SHAPE_TYPE.PICTURE]
            if not pics:
                return None
            minx = min(float(p.left) for p in pics); miny = min(float(p.top) for p in pics)
            maxx = max(float(p.left+p.width) for p in pics); maxy = max(float(p.top+p.height) for p in pics)
            W=max(120,int((maxx-minx)/EMU*140)); H=max(80,int((maxy-miny)/EMU*140))
            canvas=PILImage.new('RGB',(W,H),'white')
            for p in pics:
                with PILImage.open(io.BytesIO(p.image.blob)) as im:
                    im=im.convert('RGB')
                    x=int((float(p.left)-minx)/EMU*140); y=int((float(p.top)-miny)/EMU*140)
                    w=max(1,int(float(p.width)/EMU*140)); h=max(1,int(float(p.height)/EMU*140))
                    im.thumbnail((w,h),PILImage.Resampling.LANCZOS)
                    canvas.paste(im,(max(0,x),max(0,y)))
            b=io.BytesIO(); canvas.save(b,'PNG'); return b.getvalue()
    except Exception:
        return None
    return None


def _image_ok(sh, blob):
    if not blob:
        return False
    try:
        # Reject small logos/icons by both displayed size and pixel size.
        if float(sh.width)/EMU < .45 or float(sh.height)/EMU < .30:
            return False
        with PILImage.open(io.BytesIO(blob)) as im:
            return im.width * im.height >= 12000
    except Exception:
        return False


def _nearest_anchor(cx, cy, anchors):
    best=None; bestd=1e99
    for key, sh in anchors:
        sx,sy,sw,shh=box(sh)
        ax=sx+sw/2; ay=sy+shh/2
        d=((cx-ax)**2+(cy-ay)**2)**.5
        if d<bestd: best=(key,sh); bestd=d
    return best[0] if best else None


def collect_section_images_exact(path):
    """Direct 1:1 section mapping. Explicit D headings first; 4D images split by 발생/유출 subheading."""
    prs=Presentation(path)
    out={k:[] for k in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D')}
    for si,sl in enumerate(prs.slides):
        d_anchors=[]; sub4=[]
        for sh in walk(sl):
            t=_safe_text(sh)
            d=_explicit_d(t)
            if d: d_anchors.append((d,sh))
            s4=_sub4(t)
            if s4: sub4.append((s4,sh))
        if not d_anchors:
            continue
        # If the slide contains exactly one explicit D heading, that D owns all qualifying content images.
        ds=[]
        for d,_ in d_anchors:
            if d not in ds: ds.append(d)
        for sh in sl.shapes:
            if getattr(sh,'shape_type',None) not in (MSO_SHAPE_TYPE.PICTURE,MSO_SHAPE_TYPE.GROUP):
                continue
            blob=_picture_blob(sh)
            if not _image_ok(sh,blob):
                continue
            sx,sy,sw,shh=box(sh); cx=sx+sw/2; cy=sy+shh/2
            if len(ds)==1:
                sec=ds[0]
            else:
                sec=_nearest_anchor(cx,cy,d_anchors)
            if sec=='4D':
                if sub4:
                    target=_nearest_anchor(cx,cy,sub4)
                else:
                    target='4D_CAUSE'
                out[target].append((blob,(sx,sy,sw,shh),si))
            elif sec in out:
                out[sec].append((blob,(sx,sy,sw,shh),si))
    # De-duplicate and preserve top-to-bottom / left-to-right source ordering, max 6 per section.
    for key, vals in out.items():
        vals=sorted(vals,key=lambda z:(z[2],z[1][1],z[1][0]))
        seen=set(); unique=[]
        for item in vals:
            h=hash(item[0])
            if h in seen: continue
            seen.add(h); unique.append(item)
        out[key]=unique[:6]
    return out


def _collage(items, cols=None):
    if not items:
        return None
    ims=[]
    for item in items:
        blob=item[0] if isinstance(item,(tuple,list)) else item
        try:
            with PILImage.open(io.BytesIO(blob)) as im:
                ims.append(im.convert('RGB').copy())
        except Exception:
            pass
    if not ims: return None
    n=len(ims)
    if cols is None:
        cols=2 if n>1 else 1
    rows=math.ceil(n/cols)
    cellw,cellh=560,330
    canvas=PILImage.new('RGB',(cols*cellw,rows*cellh),'white')
    draw=ImageDraw.Draw(canvas)
    for i,im in enumerate(ims):
        x=(i%cols)*cellw; y=(i//cols)*cellh
        fitted=ImageOps.contain(im,(cellw-20,cellh-20))
        canvas.paste(fitted,(x+(cellw-fitted.width)//2,y+(cellh-fitted.height)//2))
        draw.rectangle((x,y,x+cellw-1,y+cellh-1),outline=(200,200,200),width=1)
    b=io.BytesIO(); canvas.save(b,'PNG'); return b.getvalue()


# ------------------------- extraction / Excel image -------------------------
_original_extract = v310.extract

def extract(path):
    d=_original_extract(path)
    try:
        secs=collect_section_images_exact(path)
        d['_section_images']=secs
        # Excel representative image: ALL 2D images combined into one representative collage.
        rep=_collage(secs.get('2D',[]))
        d['_images']=[(1,1,1,rep)] if rep else []
    except Exception:
        d.setdefault('_section_images',{k:[] for k in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D')})
    return d

base.extract=extract


# ------------------------- page 1 -------------------------
def _progress_three_blocks(d):
    causes=[]
    if N(d.get('cause_4d')): causes.append('발생원인: '+N(d.get('cause_4d')))
    if N(d.get('leak_cause')): causes.append('유출원인: '+N(d.get('leak_cause')))
    if N(d.get('system_cause')): causes.append('시스템원인: '+N(d.get('system_cause')))
    current=[]
    for label,key in [('임시조치','temporary_action'),('고객대응','customer_response'),('개선대책','action_5d'),('효과검증','verification_6d'),('수평전개','spread_7d')]:
        if N(d.get(key)): current.append(label+': '+N(d.get(key)))
    request=N(d.get('request') or d.get('request_item') or d.get('requests'))
    return '- 원인\n' + ('\n'.join(causes) if causes else '검토 중') + \
           '\n\n- 진행현황\n' + ('\n'.join(current) if current else '검토 중') + \
           '\n\n- 요청사항\n' + (request if request else '')


def _set_cell(cell,text,size=8):
    cell.text=N(text)
    cell.text_frame.word_wrap=True
    for p in cell.text_frame.paragraphs:
        for run in p.runs:
            run.font.size=Pt(size)


def _replace_team_title(sh,team):
    if not hasattr(sh,'text_frame'): return False
    full=_safe_text(sh)
    if '팀' not in full or '주요' not in full or '논의' not in full: return False
    new=re.sub(r'000\s*팀', N(team)+'팀', full)
    if new==full:
        new=re.sub(r'[^\s]*팀(?=\s*주요\s*논의)', N(team)+'팀', full)
    if new!=full:
        # Preserve formatting when possible.
        done=False
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                before=r.text
                r.text=re.sub(r'000\s*팀',N(team)+'팀',r.text)
                if r.text!=before: done=True
        if not done: sh.text_frame.text=new
        return True
    return False


def update_page1(prs,d,g):
    if not prs.slides: return
    sl=prs.slides[0]; team=N(g.get('team'))
    if team:
        for sh in walk(sl): _replace_team_title(sh,team)
    progress=_progress_three_blocks(d)
    # User template has no task/issue columns. Only fill the existing 진행사항/진행현황 field.
    filled=False
    for sh in walk(sl):
        if not getattr(sh,'has_table',False): continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                q=C(tb.cell(r,c).text)
                if '진행사항' in q or '진행현황' in q:
                    # Prefer the cell below; otherwise the cell to the right.
                    if r+1<len(tb.rows): target=tb.cell(r+1,c)
                    elif c+1<len(tb.columns): target=tb.cell(r,c+1)
                    else: target=tb.cell(r,c)
                    _set_cell(target,progress,8); filled=True; break
            if filled: break
        if filled: break
    if not filled:
        # Textbox label fallback.
        labels=[]
        for sh in walk(sl):
            if hasattr(sh,'text_frame') and ('진행사항' in C(_safe_text(sh)) or '진행현황' in C(_safe_text(sh))): labels.append(sh)
        for lab in labels:
            lx,ly,lw,lh=box(lab); best=None; score=1e99
            for sh in walk(sl):
                if sh is lab or not hasattr(sh,'text_frame'): continue
                sx,sy,sw,shh=box(sh)
                # below or right, closest
                if sy>=ly+lh-.05*EMU:
                    sc=abs(sx-lx)+(sy-(ly+lh))
                elif sx>=lx+lw-.05*EMU:
                    sc=(sx-(lx+lw))+abs(sy-ly)
                else: continue
                if sc<score: best,score=sh,sc
            if best is not None:
                v310.set_text(best,progress,8); break


# ------------------------- page 2 -------------------------
def _section_text(d,key):
    if key=='2D': return N(d.get('problem')) or '검토 중'
    if key=='3D': return '\n'.join(x for x in [N(d.get('temporary_action')),N(d.get('customer_response'))] if x) or '검토 중'
    if key=='4D_CAUSE':
        v=N(d.get('cause_4d')); return ('- 발생원인\n'+v) if v else ''
    if key=='4D_LEAK':
        parts=[]
        if N(d.get('leak_cause')): parts.append('- 유출원인\n'+N(d.get('leak_cause')))
        if N(d.get('system_cause')): parts.append('- 시스템원인\n'+N(d.get('system_cause')))
        return '\n'.join(parts)
    if key=='5D': return N(d.get('action_5d')) or '검토 중'
    if key=='6D': return N(d.get('verification_6d')) or '검토 중'
    return ''


def _find_marker_unit(sl,label):
    return v310.find_marker(sl,label)


def _clone_shape(sl,shape):
    try:
        newel=copy.deepcopy(shape._element)
        sl.shapes._spTree.insert_element_before(newel,'p:extLst')
        # Return the newly appended shape by matching element identity.
        for s in sl.shapes:
            if s._element is newel: return s
    except Exception:
        return None
    return None


def _ensure_second_4d(sl,x,y):
    """Ensure a second 4D marker/title unit exists for the left 발생원인 column without XML regrouping."""
    # Already two visible 4D labels => just use them.
    existing=[]
    for sh in sl.shapes:
        if N(getattr(sh,'text',''))=='4D': existing.append(sh)
        elif getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:
            for ch in walk(sh):
                if ch is not sh and N(getattr(ch,'text',''))=='4D': existing.append(sh); break
    if len(existing)>=2:
        return
    marker,parent=v310.find_marker(sl,'4D')
    source=parent if parent is not None else marker
    if source is None: return
    dup=_clone_shape(sl,source)
    if dup is None: return
    try:
        dup.left=Inches(x); dup.top=Inches(y)
    except Exception:
        pass


def _render_zone(sl,key,z,text,images,font=8):
    if not text:
        return
    blob=_collage(images)
    if blob:
        iw=min(1.72,z['w']*.38); tw=z['w']-iw-.08
        v310.add_box(sl,key,z['x'],z['y'],tw,z['h'],text,font)
        v310.add_image(sl,blob,z['x']+tw+.08,z['y'],iw,z['h'],key)
    else:
        v310.add_box(sl,key,z['x'],z['y'],z['w'],z['h'],text,font)


def _fill_page2_meta(sl,d,g):
    issue=N(d.get('issue_name')); task=N(d.get('task_name')); team=N(g.get('team')); owner=N(g.get('owner'))
    for sh in walk(sl):
        if hasattr(sh,'text_frame'):
            t=_safe_text(sh); q=C(t)
            # title placeholder / generic title
            if '과제명_이슈제목' in q or '과제명이슈제목' in q:
                v310.set_text(sh, task or issue, 8)
            elif t.startswith('이슈명') or q=='이슈명':
                v310.set_text(sh,'이슈명 : '+issue,8)
            elif '00팀담당자' in q:
                v310.set_text(sh,f'{team} 담당자 : {owner}',8)
    v311._fill_occurrence(sl,d,g)


def update_page2(sl,d,g,mode):
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    has_cause=bool(N(d.get('cause_4d')))
    has_leak=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))

    # Geometry follows the requested two-column 4D structure.
    zones={k:dict(v) for k,v in v310.ZONES.items()}
    # If only one 4D side exists, use the matching side only. If both exist, keep both columns.
    if v310.is_new(mode):
        # 2D/3D/5D/6D markers follow their content zones.
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]; v310.move_marker_unit(sl,label,max(.1,z['x']-.18),max(.1,z['y']-.35))
        if has_cause and has_leak:
            # Existing 4D unit -> right (유출), duplicate same unit -> left (발생).
            zr=zones['4D_LEAK']; zl=zones['4D_CAUSE']
            v310.move_marker_unit(sl,'4D',max(.1,zr['x']-.18),max(.1,zr['y']-.35))
            _ensure_second_4d(sl,max(.1,zl['x']-.18),max(.1,zl['y']-.35))
        elif has_cause:
            zl=zones['4D_CAUSE']; v310.move_marker_unit(sl,'4D',max(.1,zl['x']-.18),max(.1,zl['y']-.35))
        elif has_leak:
            zr=zones['4D_LEAK']; v310.move_marker_unit(sl,'4D',max(.1,zr['x']-.18),max(.1,zr['y']-.35))

    _render_zone(sl,'2D',zones['2D'],_section_text(d,'2D'),imgs.get('2D',[]))
    _render_zone(sl,'3D',zones['3D'],_section_text(d,'3D'),imgs.get('3D',[]))
    if has_cause:
        _render_zone(sl,'4D_CAUSE',zones['4D_CAUSE'],_section_text(d,'4D_CAUSE'),imgs.get('4D_CAUSE',[]))
    if has_leak:
        _render_zone(sl,'4D_LEAK',zones['4D_LEAK'],_section_text(d,'4D_LEAK'),imgs.get('4D_LEAK',[]))
    _render_zone(sl,'5D',zones['5D'],_section_text(d,'5D'),imgs.get('5D',[]))
    _render_zone(sl,'6D',zones['6D'],_section_text(d,'6D'),imgs.get('6D',[]))
    _fill_page2_meta(sl,d,g)


# ------------------------- weekly entry -------------------------
def weekly(src,out,d,g,mode):
    prs=Presentation(src)
    update_page1(prs,d,g)
    if len(prs.slides)>1:
        update_page2(prs.slides[1],d,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out); saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix)); prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.weekly=weekly
base.extract=extract

if __name__=='__main__':
    base.App().mainloop()
