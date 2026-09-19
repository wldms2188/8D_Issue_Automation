"""Skip duplicate 8D attachment pages before native PowerPoint COM insertion."""
import hashlib, os, shutil, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import main_recovery_step14_fix2 as core
_original_append=core._append_8d_attachments

def _shape_parts(sh,image_hash=True):
    parts=[str(getattr(sh,'shape_type','')),str(getattr(sh,'left','')),str(getattr(sh,'top','')),str(getattr(sh,'width','')),str(getattr(sh,'height',''))]
    if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:
        for child in sh.shapes:parts.extend(_shape_parts(child,image_hash=image_hash))
    if getattr(sh,'has_table',False):
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):parts.append('T:'+str(tb.cell(r,c).text or '').strip())
    elif hasattr(sh,'text'):parts.append('X:'+str(getattr(sh,'text','') or '').strip())
    if image_hash and getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.PICTURE:
        try:parts.append('I:'+hashlib.sha1(sh.image.blob).hexdigest())
        except Exception:pass
    return parts

def _slide_signature(sl,image_hash=True):
    parts=[]
    for sh in sl.shapes:parts.extend(_shape_parts(sh,image_hash=image_hash))
    return hashlib.sha1('\x1f'.join(parts).encode('utf-8','ignore')).hexdigest()

def _delete_slide(prs,index):
    slide_id=prs.slides._sldIdLst[index]; prs.part.drop_rel(slide_id.rId); prs.slides._sldIdLst.remove(slide_id)

def _remove_current_auto(out_path,prefix):
    prs=Presentation(out_path); changed=False
    for i in range(len(prs.slides)-1,-1,-1):
        if str(getattr(prs.slides[i],'name','') or '').startswith(prefix):_delete_slide(prs,i); changed=True
    if changed:prs.save(out_path)

def _append_without_duplicates(out_path,src8d_path,detail_index,d):
    src=Presentation(src8d_path)
    source_count=len(src.slides)
    if source_count<=1:return 0
    current_prefix='AUTO_8D_ATTACH_'+core._attachment_key(d)+'_'

    # Fast pre-filter: compare text/geometry first. Expensive image-blob hashes are
    # calculated only for destination slides that could actually match a source
    # attachment. This matters a lot for large weekly decks.
    source_cheap={i:_slide_signature(src.slides[i],image_hash=False) for i in range(1,source_count)}
    cheap_needed=set(source_cheap.values())
    dest=Presentation(out_path)
    existing=set()
    for sl in dest.slides:
        if str(getattr(sl,'name','') or '').startswith(current_prefix):
            continue
        cheap=_slide_signature(sl,image_hash=False)
        if cheap not in cheap_needed:
            continue
        existing.add(_slide_signature(sl,image_hash=True))

    keep=[]; seen=set(existing)
    for i in range(1,source_count):
        sig=_slide_signature(src.slides[i],image_hash=True)
        if sig in seen:continue
        seen.add(sig); keep.append(i)

    def call_original(path,count):
        dd=dict(d or {})
        dd['_attachment_source_count']=count
        return _original_append(out_path,path,detail_index,dd)

    if len(keep)==source_count-1:
        return call_original(src8d_path,source_count)
    if not keep:
        _remove_current_auto(out_path,current_prefix)
        return 0

    fd,tmp=tempfile.mkstemp(prefix='8d_attach_',suffix='.pptx'); os.close(fd)
    try:
        shutil.copy2(src8d_path,tmp); reduced=Presentation(tmp); keep_set={0,*keep}
        for idx in range(len(reduced.slides)-1,-1,-1):
            if idx not in keep_set:_delete_slide(reduced,idx)
        reduced.save(tmp)
        return call_original(tmp,len(reduced.slides))
    finally:
        try:Path(tmp).unlink(missing_ok=True)
        except Exception:pass
core._append_8d_attachments=_append_without_duplicates
