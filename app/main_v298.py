# 8D Issue Automation v2.9.8
# Stable weekly-update path: keep the known-good v2.9.3 weekly logic from v2.9.5.
# Representative image: original 8D PPT, page 1 only, restricted to 2D/4D labeled regions.
# No temporary PPT copy is created.
import sys, io
from pathlib import Path

APP_DIR=Path(__file__).resolve().parent
if str(APP_DIR) not in sys.path:
    sys.path.insert(0,str(APP_DIR))

# v2.9.5 contains the known-good weekly/Excel implementation.
import main_v295 as impl
base=impl.base
v29=impl.v29

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image as PILImage

EMU=914400

def _pics(sh):
    out=[]
    typ=getattr(sh,'shape_type',None)
    if typ==MSO_SHAPE_TYPE.PICTURE:
        out.append(sh)
    elif typ==MSO_SHAPE_TYPE.GROUP:
        for c in sh.shapes:
            out.extend(_pics(c))
    return out

def _box(sh):
    return float(sh.left),float(sh.top),float(sh.width),float(sh.height)

def _overlap(a,b):
    ax,ay,aw,ah=a; bx,by,bw,bh=b
    return max(0,min(ax+aw,bx+bw)-max(ax,bx))*max(0,min(ay+ah,by+bh)-max(ay,by))

def _composite_group(group):
    pics=_pics(group)
    if len(pics)<2:return None
    gx,gy,gw,gh=_box(group)
    W=max(1,int(gw/EMU*120)); H=max(1,int(gh/EMU*120))
    canvas=PILImage.new('RGB',(W,H),'white')
    for p in pics:
        try:
            with PILImage.open(io.BytesIO(p.image.blob)) as im:
                im=im.convert('RGB')
                px,py,pw,ph=_box(p)
                x=max(0,int((px-gx)/EMU*120)); y=max(0,int((py-gy)/EMU*120))
                w=max(1,int(pw/EMU*120)); h=max(1,int(ph/EMU*120))
                im.thumbnail((w,h),PILImage.Resampling.LANCZOS)
                canvas.paste(im,(x,y))
        except Exception:
            pass
    b=io.BytesIO();canvas.save(b,'PNG');return b.getvalue()

def _page1_region_image(path):
    prs=Presentation(path)
    if not prs.slides:return None
    sl=prs.slides[0]

    # Explicitly inspect ONLY page 1. 2D has priority over 4D.
    regions=[]
    for sh in sl.shapes:
        q=base.compact(getattr(sh,'text',''))
        if not q:continue
        if any(k in q for k in ('2d','문제현상','문제현황','불량현상')):
            regions.append(('2D',_box(sh)))
        elif any(k in q for k in ('4d','원인분석','발생원인','유출원인')):
            regions.append(('4D',_box(sh)))

    if not regions:return None

    candidates=[]
    for sh in sl.shapes:
        typ=getattr(sh,'shape_type',None)
        if typ==MSO_SHAPE_TYPE.GROUP:
            pics=_pics(sh)
            if len(pics)>=2:
                blob=_composite_group(sh)
                if blob:candidates.append(('group',len(pics),_box(sh),blob))
            elif len(pics)==1:
                p=pics[0]
                try:
                    with PILImage.open(io.BytesIO(p.image.blob)) as im:
                        if im.width*im.height>=10000:candidates.append(('picture',1,_box(sh),p.image.blob))
                except Exception:pass
        elif typ==MSO_SHAPE_TYPE.PICTURE:
            try:
                with PILImage.open(io.BytesIO(sh.image.blob)) as im:
                    if im.width*im.height>=10000:candidates.append(('picture',1,_box(sh),sh.image.blob))
            except Exception:pass

    scored=[]
    for kind,n,vbox,blob in candidates:
        for label,rbox in regions:
            ov=_overlap(vbox,rbox)
            if ov<=0:continue
            priority=2 if label=='2D' else 1
            grouped=1 if kind=='group' else 0
            ratio=ov/max(min(vbox[2]*vbox[3],rbox[2]*rbox[3]),1.0)
            scored.append((priority,grouped,ratio,ov,vbox[2]*vbox[3],blob))
    if not scored:return None
    scored.sort(key=lambda x:(x[0],x[1],x[2],x[3],x[4]),reverse=True)
    return scored[0][5]

def extract_v298(path):
    # Use the existing stable extraction directly. Never create a copied PPT.
    d=impl.base.extract(path)
    try:
        blob=_page1_region_image(path)
        d['_images']=[(1,1,1,blob)] if blob else []
    except Exception:
        d['_images']=[]
    return d

base.extract=extract_v298

if __name__=='__main__':
    base.App().mainloop()
