# 8D Issue Automation v3.1.4
# Fixes: strict D-region images, page1 fields/format, page2 metadata, dynamic layout, dual 4D.
import re, io, copy, math, datetime
from pathlib import Path

import main_v313 as v313
import main_v312 as v312
import main_v311 as v311
import main_v310 as v310
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from PIL import Image as PILImage, ImageOps

base=v313.base
N=v310.N; C=v310.C; box=v310.box; walk=v310.walk; EMU=v310.EMU
SECS=('2D','3D','4D_CAUSE','4D_LEAK','5D','6D')


def _explicit_d(text):
    raw=N(text).lower(); q=C(text)
    for d in ('2D','3D','4D','5D','6D'):
        dl=d.lower()
        if re.search(r'(?<![0-9a-z])'+dl+r'(?![0-9a-z])',raw) or q==dl or q.startswith(dl): return d
    return None


def _sub4(text):
    q=C(text)
    if '발생원인' in q: return '4D_CAUSE'
    if any(k in q for k in ('유출원인','유츌원인','시스템원인')): return '4D_LEAK'
    return None


def _pic_blob(sh):
    return v313._picture_blob(sh)


def _valid_pic(sh,blob):
    if not blob: return False
    try:
        if float(sh.width)/EMU<.30 or float(sh.height)/EMU<.22: return False
        with PILImage.open(io.BytesIO(blob)) as im: return im.width*im.height>=5000
    except Exception: return False


def _bands(sl,slide_h):
    anchors=[]
    for sh in walk(sl):
        d=_explicit_d(N(getattr(sh,'text','')))
        if d:
            x,y,w,h=box(sh); anchors.append((d,x,y,w,h,y+h/2))
    if not anchors: return []
    rows=[]
    for a in sorted(anchors,key=lambda z:z[5]):
        hit=None
        for r in rows:
            if abs(a[5]-r['cy'])<=.28*EMU: hit=r; break
        if hit:
            hit['items'].append(a); hit['cy']=sum(z[5] for z in hit['items'])/len(hit['items'])
        else: rows.append({'cy':a[5],'items':[a]})
    rows.sort(key=lambda r:r['cy'])
    out=[]
    for i,r in enumerate(rows):
        top=min(a[2] for a in r['items'])-.06*EMU
        bottom=(min(a[2] for a in rows[i+1]['items'])-.08*EMU) if i+1<len(rows) else slide_h-.22*EMU
        for a in r['items']: out.append({'d':a[0],'x':a[1],'top':top,'bottom':bottom})
    return out


def collect_images(path):
    prs=Presentation(path); out={k:[] for k in SECS}; slide_h=float(prs.slide_height)
    for si,sl in enumerate(prs.slides):
        bands=_bands(sl,slide_h)
        if not bands: continue
        subs=[]
        for sh in walk(sl):
            s=_sub4(N(getattr(sh,'text','')))
            if s:
                x,y,w,h=box(sh); subs.append((s,x+w/2,y+h/2))
        for sh in sl.shapes:
            if getattr(sh,'shape_type',None) not in (MSO_SHAPE_TYPE.PICTURE,MSO_SHAPE_TYPE.GROUP): continue
            blob=_pic_blob(sh)
            if not _valid_pic(sh,blob): continue
            x,y,w,h=box(sh); cx=x+w/2; cy=y+h/2
            cand=[b for b in bands if b['top']<=cy<b['bottom']]
            if not cand: continue
            b=min(cand,key=lambda q:abs(cx-q['x']))
            sec=b['d']
            if sec=='4D':
                if subs:
                    sec=min(subs,key=lambda z:abs(cx-z[1])+.45*abs(cy-z[2]))[0]
                else:
                    sec='4D_CAUSE' if cx<float(prs.slide_width)/2 else '4D_LEAK'
            if sec in out: out[sec].append((blob,(x,y,w,h),si))
    for k,vals in out.items():
        vals=sorted(vals,key=lambda z:(z[2],z[1][1],z[1][0])); seen=set(); uniq=[]
        for it in vals:
            h=hash(it[0])
            if h in seen: continue
            seen.add(h); uniq.append(it)
        out[k]=uniq[:8]
    return out


def _rep_2d(items):
    if not items: return None
    by={}
    for it in items: by.setdefault(it[2],[]).append(it)
    chosen=max(by.items(),key=lambda kv:(len(kv[1]),sum(x[1][2]*x[1][3] for x in kv[1]),-kv[0]))[1]
    return v313._collage(chosen)

_original_extract=v313.extract

def extract(path):
    d=_original_extract(path)
    try:
        secs=collect_images(path); d['_section_images']=secs
        rep=_rep_2d(secs.get('2D',[])); d['_images']=[(1,1,1,rep)] if rep else []
    except Exception:
        d['_section_images']={k:[] for k in SECS}; d['_images']=[]
    return d
base.extract=extract


def _progress(d):
    a=[]
    for lab,key in [('발생원인','cause_4d'),('유출원인','leak_cause'),('시스템원인','system_cause')]:
        v=N(d.get(key));
        if v: a.append(f'- {lab}: {v}')
    p=[]
    for lab,key in [('임시조치','temporary_action'),('고객대응','customer_response'),('개선대책','action_5d'),('효과검증','verification_6d'),('수평전개','spread_7d')]:
        v=N(d.get(key));
        if v: p.append(f'- {lab}: {v}')
    req=N(d.get('request') or d.get('request_item') or d.get('requests'))
    r=[(x if x.startswith('-') else '- '+x) for x in req.splitlines() if x.strip()] if req else []
    return '1. 원인\n'+('\n'.join(a) if a else '- 검토 중')+'\n\n2. 진행현황\n'+('\n'.join(p) if p else '- 검토 중')+'\n\n3. 요청사항\n'+('\n'.join(r) if r else '- ')


def _set_cell(cell,text,size=8):
    cell.text=N(text); cell.text_frame.word_wrap=True
    for p in cell.text_frame.paragraphs:
        p.space_before=Pt(0); p.space_after=Pt(0)
        for run in p.runs: run.font.size=Pt(size)


def _replace_run_text(sh,old,new):
    if not hasattr(sh,'text_frame'): return False
    for p in sh.text_frame.paragraphs:
        for run in p.runs:
            if old in run.text:
                run.text=run.text.replace(old,new); return True
    return False


def update_page1(prs,d,g):
    if not prs.slides: return
    sl=prs.slides[0]; team=N(g.get('team'))
    if team:
        for sh in walk(sl):
            t=N(getattr(sh,'text',''))
            if '팀' in t and '주요' in t and '논의' in t:
                if not _replace_run_text(sh,'000',team):
                    m=re.search(r'([^\s]+)(?=팀\s*주요\s*논의)',t)
                    if m: _replace_run_text(sh,m.group(1),team)
                break
    for sh in walk(sl):
        if not getattr(sh,'has_table',False): continue
        tb=sh.table; best=None
        for hr in range(min(3,len(tb.rows))):
            hm={}
            for c in range(len(tb.columns)):
                q=C(tb.cell(hr,c).text)
                if '과제명' in q: hm['task']=c
                elif q=='이슈' or '이슈명' in q: hm['issue']=c
                elif '현상' in q or '문제' in q: hm['problem']=c
                elif '진행사항' in q or '진행현황' in q: hm['progress']=c
                elif 'signal' in q: hm['signal']=c
            if best is None or len(hm)>len(best[1]): best=(hr,hm)
        if not best or not best[1]: continue
        hr,hm=best; row=min(hr+1,len(tb.rows)-1)
        vals={'task':N(d.get('task_name')),'issue':N(d.get('issue_name')),'problem':N(d.get('problem')),'progress':_progress(d)}
        for k,v in vals.items():
            if k in hm: _set_cell(tb.cell(row,hm[k]),v,8)
        if 'signal' in hm:
            try: base.signal(tb.cell(row,hm['signal']),base.status(d))
            except Exception: pass
        break


def _section_text(d,key):
    return v313._section_text(d,key)


def _layout(d,imgs):
    texts={k:_section_text(d,k) for k in v310.ZONES}
    left_imgs={'2D':imgs.get('2D',[]),'3D':imgs.get('3D',[]),'4D_CAUSE':imgs.get('4D_CAUSE',[])}
    right_imgs={'4D_LEAK':imgs.get('4D_LEAK',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}
    left,lf=v312._fit_chain(['2D','3D','4D_CAUSE'],texts,left_imgs)
    right,rf=v312._fit_chain(['4D_LEAK','5D','6D'],texts,right_imgs)
    return {**left,**right},texts,min(lf,rf)


def _render(sl,key,z,text,images,font):
    if not text: return
    blob=v313._collage(images)
    if blob:
        iw=min(1.65,z['w']*.36); tw=z['w']-iw-.08
        v310.add_box(sl,key,z['x'],z['y'],tw,z['h'],text,font)
        v310.add_image(sl,blob,z['x']+tw+.08,z['y'],iw,z['h'],key)
    else: v310.add_box(sl,key,z['x'],z['y'],z['w'],z['h'],text,font)


def _clone4d(sl,source,x,y):
    try:
        newel=copy.deepcopy(source._element); sl.shapes._spTree.insert_element_before(newel,'p:extLst')
        dup=list(sl.shapes)[-1]; dup.name='AUTO_8D_4D_LEFT'
        bx,by,_,_=box(dup); dup.left += Inches(x)-int(bx); dup.top += Inches(y)-int(by)
        return dup
    except Exception: return None


def _two4d(sl,left,right):
    marker,parent=v310.find_marker(sl,'4D'); src=parent if parent is not None else marker
    if src is None: return
    dup=_clone4d(sl,src,left[0],left[1])
    bx,by,_,_=box(src); src.left += Inches(right[0])-int(bx); src.top += Inches(right[1])-int(by)
    if dup is None:
        # fallback guarantees a visible left 4D marker even if group clone fails
        sh=sl.shapes.add_shape(1,Inches(left[0]),Inches(left[1]),Inches(.42),Inches(.42)); sh.name='AUTO_8D_4D_LEFT'; v310.set_text(sh,'4D',8)


def _meta(sl,d,g):
    task=N(d.get('task_name')); issue=N(d.get('issue_name')); occ=v311._occurrence_value(d,g)
    for sh in walk(sl):
        if hasattr(sh,'text_frame'):
            t=N(getattr(sh,'text','')); q=C(t)
            if '과제명이슈제목' in q or ('과제명' in q and '제목' in q): v310.set_text(sh,task or issue,8)
            elif q.startswith('이슈명'): v310.set_text(sh,'이슈명 : '+issue,8)
    for sh in walk(sl):
        if not getattr(sh,'has_table',False): continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                q=C(tb.cell(r,c).text)
                if '이슈명' in q and c+1<len(tb.columns): _set_cell(tb.cell(r,c+1),issue,8)
                if '발생단계' in q and c+1<len(tb.columns): _set_cell(tb.cell(r,c+1),occ,9)
    v311._fill_occurrence(sl,d,g)


def update_page2(sl,d,g,mode):
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}; zones,texts,font=_layout(d,imgs)
    hc=bool(N(d.get('cause_4d'))); hl=bool(N(d.get('leak_cause')) or N(d.get('system_cause')))
    if v310.is_new(mode):
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]; v310.move_marker_unit(sl,label,max(.1,z['x']-.18),max(.1,z['y']-.35))
        if hc and hl:
            zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']; _two4d(sl,(max(.1,zl['x']-.18),max(.1,zl['y']-.35)),(max(.1,zr['x']-.18),max(.1,zr['y']-.35)))
        elif hc:
            z=zones['4D_CAUSE']; v310.move_marker_unit(sl,'4D',max(.1,z['x']-.18),max(.1,z['y']-.35))
        elif hl:
            z=zones['4D_LEAK']; v310.move_marker_unit(sl,'4D',max(.1,z['x']-.18),max(.1,z['y']-.35))
    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        _render(sl,key,zones[key],texts[key],imgs.get(key,[]),font)
    _meta(sl,d,g)


def weekly(src,out,d,g,mode):
    prs=Presentation(src); update_page1(prs,d,g)
    if len(prs.slides)>1: update_page2(prs.slides[1],d,g,mode)
    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try: prs.save(out); saved=out
    except PermissionError:
        p=Path(out); saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix)); prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved

base.weekly=weekly
base.extract=extract

if __name__=='__main__': base.App().mainloop()
