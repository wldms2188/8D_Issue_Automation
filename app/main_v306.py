# 8D Issue Automation v3.0.6
import sys, re, copy, datetime
from pathlib import Path
APP_DIR=Path(__file__).resolve().parent
if str(APP_DIR) not in sys.path: sys.path.insert(0,str(APP_DIR))
import main_v305 as v305
base=v305.base
renderer=v305.renderer
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
EMU=914400
DEFAULT=copy.deepcopy(renderer.DEFAULT_LAYOUT)
FLOW={'left':['2D','3D','4D_CAUSE'],'right':['4D_LEAK','5D','6D']}
MIN_H={'2D':.62,'3D':.70,'4D_CAUSE':.90,'4D_LEAK':.58,'5D':.82,'6D':.60}
BOTTOM={'left':7.17,'right':6.85}
IMG={'2D':'2D','3D':'3D','4D_CAUSE':'4D','4D_LEAK':'4D','5D':'5D','6D':'6D'}

def n(x): return str(x or '').replace('\r\n','\n').replace('\r','\n').strip()
def c(x):
    try:return base.compact(x)
    except:return re.sub(r'[^0-9A-Za-z가-힣]','',n(x)).lower()
def valid_task(x):
    q=c(x); return len(n(x))>=2 and not any(k in q for k in ('고객사','이슈명','과제명','model','packer','발생site','발생일자','lotno','발생line','담당자','제품타입','폼팩터','발생샘플','개발단계','발생처','signal'))
def valid_issue(x):
    q=c(x); return len(n(x))>=2 and not any(k in q for k in ('이슈명','과제명','고객사','발생site','발생일자','lotno','signal'))

def extract_meta(path,d):
    prs=Presentation(path); issue=''; task=''
    for sl in prs.slides:
        for sh in sl.shapes:
            if not getattr(sh,'has_table',False): continue
            tb=sh.table; rows=[[n(tb.cell(r,cc).text) for cc in range(len(tb.columns))] for r in range(len(tb.rows))]
            for r,row in enumerate(rows):
                for cc,val in enumerate(row):
                    q=c(val)
                    if '이슈명' in q and not issue:
                        for x in row[cc+1:]:
                            if valid_issue(x): issue=n(x); break
                    if '과제명' in q and not task:
                        for x in row[cc+1:]:
                            if valid_task(x): task=n(x); break
            if not task:
                for r,row in enumerate(rows):
                    for cc,val in enumerate(row):
                        if '고객사' not in c(val): continue
                        cand=[]
                        if cc+2<len(row): cand.append(row[cc+2])
                        for rr in range(r+1,min(r+3,len(rows))):
                            if cc+2<len(rows[rr]): cand.append(rows[rr][cc+2])
                            if cc+1<len(rows[rr]): cand.append(rows[rr][cc+1])
                        for x in cand:
                            if valid_task(x): task=n(x); break
                        if task: break
                    if task: break
        if issue and task: break
    if not issue or not task:
        for sl in prs.slides:
            for sh in sl.shapes:
                lines=[x.strip() for x in n(getattr(sh,'text','')).split('\n') if x.strip()]
                for i,line in enumerate(lines):
                    if not issue and '이슈명' in c(line):
                        if ':' in line or '：' in line:
                            x=re.split(r'[:：]',line,1)[1].strip()
                            if valid_issue(x): issue=x
                        elif i+1<len(lines) and valid_issue(lines[i+1]): issue=lines[i+1]
                    if not task and '과제명' in c(line):
                        if ':' in line or '：' in line:
                            x=re.split(r'[:：]',line,1)[1].strip()
                            if valid_task(x): task=x
                        elif i+1<len(lines) and valid_task(lines[i+1]): task=lines[i+1]
    if issue:d['issue_name']=issue
    if task:d['task_name']=task
    return d

_prev=base.extract
def extract(path):
    d=_prev(path)
    try: d=extract_meta(path,d)
    except Exception: pass
    return d
base.extract=extract

# The previous XML grouping implementation caused list.index(x) errors. Never call it.
_stable=getattr(v305,'_orig_fill',renderer._fill_page2)

def walk(container,parent=None):
    for sh in getattr(container,'shapes',[]):
        yield sh,parent
        if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh,sh)

def marker(sl,label):
    for sh,parent in walk(sl):
        if n(getattr(sh,'text',''))==label:return sh,parent
    return None,None

def existing_layout(sl):
    out=copy.deepcopy(DEFAULT)
    mp={'2D':'2D','3D':'3D','4D':'4D_LEAK','5D':'5D','6D':'6D'}
    for label,key in mp.items():
        sh,_=marker(sl,label)
        if sh is None:continue
        x,y,w,h=float(sh.left)/EMU,float(sh.top)/EMU,float(sh.width)/EMU,float(sh.height)/EMU
        out[key]['x']=x+w+.08; out[key]['y']=y+h/2-out[key]['h']/2
    return out

def text_for(k,d):
    if k=='2D':return d.get('problem') or '검토 중'
    if k=='3D':return '\n'.join(str(x) for x in (d.get('temporary_action'),d.get('customer_response')) if n(x)) or '검토 중'
    if k=='4D_CAUSE':return d.get('cause_4d') or '검토 중'
    if k=='4D_LEAK':return '\n'.join(str(x) for x in (d.get('leak_cause'),d.get('system_cause')) if n(x)) or '검토 중'
    if k=='5D':return d.get('action_5d') or '검토 중'
    return d.get('verification_6d') or '검토 중'

def flow(layout,d):
    out={k:dict(v) for k,v in layout.items()}; imgs=d.get('_section_images',{}) or {}
    for col,keys in FLOW.items():
        cur=min(out[k]['y'] for k in keys)
        for i,k in enumerate(keys):
            text=n(text_for(k,d)); width=out[k]['w']; has=bool(imgs.get(IMG[k],[])); chars=max(7,int(width*(.64 if has else .94)*9.2)); lines=sum(max(1,(len(s)+chars-1)//chars) for s in text.splitlines() or [''])
            need=max(out[k]['h'],.26+lines*.145); room=BOTTOM[col]-cur-sum(MIN_H[z] for z in keys[i+1:])
            out[k]['y']=cur; out[k]['h']=max(MIN_H[k],min(need,max(MIN_H[k],room))); cur=out[k]['y']+out[k]['h']
    return out

def move_unit(sl,label,z):
    sh,parent=marker(sl,label)
    if sh is None:return
    tx=max(.03,z['x']-float(sh.width)/EMU-.08); ty=z['y']+max(0,(z['h']-float(sh.height)/EMU)/2)
    dx=Inches(tx)-sh.left; dy=Inches(ty)-sh.top
    try:
        if parent is not None:
            parent.left+=dx; parent.top+=dy
        else:
            sh.left=Inches(tx); sh.top=Inches(ty)
    except Exception:pass

def clean_labels(sl):
    for sh in list(sl.shapes):
        if str(getattr(sh,'name','')).startswith('AUTO_8D_LABEL_'):
            try:
                el=sh._element; el.getparent().remove(el)
            except Exception:pass

def is_new(mode):
    q=c(mode); return any(x in q for x in ('신규','new','신규이슈'))

def fill(sl,d,g,layout,new):
    use=flow(layout,d) if new else layout
    _stable(sl,d,g,use)
    clean_labels(sl)
    if new:
        for label,key in [('2D','2D'),('3D','3D'),('4D','4D_LEAK'),('5D','5D'),('6D','6D')]:move_unit(sl,label,use[key])
renderer._fill_page2=fill

def headers(tb):
    m={}
    for cc in range(len(tb.columns)):
        q=c(tb.cell(0,cc).text)
        if '과제명' in q:m['task']=cc
        if q=='이슈' or '이슈명' in q:m['issue']=cc
        if '문제' in q or '현상' in q:m['problem']=cc
        if '진행' in q:m['progress']=cc
        if 'signal' in q:m['signal']=cc
    return m

def summary(prs,d):
    vals={'task':n(d.get('task_name')),'issue':n(d.get('issue_name')),'problem':n(d.get('problem'))}
    try:vals['progress']=v305.v303.v301.impl.v29.v29_progress(d)
    except:vals['progress']=''
    for sl in prs.slides:
        for sh in sl.shapes:
            if not getattr(sh,'has_table',False):continue
            tb=sh.table; hm=headers(tb)
            if 'issue' not in hm:continue
            r=1
            for rr in range(1,len(tb.rows)):
                if not any(n(tb.cell(rr,cc).text) for cc in range(len(tb.columns))):r=rr;break
            for k,cc in hm.items():
                if k in vals:tb.cell(r,cc).text=vals[k]
            if 'signal' in hm:base.signal(tb.cell(r,hm['signal']),base.status(d))
            return

def weekly(src,out,d,g,mode):
    prs=Presentation(src); summary(prs,d); new=is_new(mode)
    for i,sl in enumerate(prs.slides):
        if i==0:continue
        layout=copy.deepcopy(DEFAULT) if new else existing_layout(sl)
        fill(sl,d,g,layout,new)
    Path(out).parent.mkdir(exist_ok=True)
    try:prs.save(out);saved=out
    except PermissionError:
        p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);prs.save(saved)
    return '주간회의 PPT 업데이트: '+base.status(d),saved
base.weekly=weekly
if __name__=='__main__':base.App().mainloop()
