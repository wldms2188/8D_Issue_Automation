import datetime
import hashlib
import subprocess
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement

import main_recovery_step14 as s14
import main_recovery_step13 as s13
import main_recovery_step12 as s12
import main_recovery_step10 as s10
import main_v319 as v319
import main_v310 as v310

base=s14.base
N=v310.N


def _project_parts(d):
    return (s13._k(s13._customer_task(d)),s13._k(d.get('task_name')),s13._k(d.get('customer')))

def _project_score(sl,d):
    q=s13._k(s13._slide_text(sl)); full,task,customer=_project_parts(d)
    if full and full in q:return 160
    if task and customer and task in q and customer in q:return 145
    if task and task in q:return 120
    return 0

def _summary_indices(prs):return {si for si,_,_ in s14._summary_pages(prs)}
def _is_detail_like(sl):
    q=s13._k(s13._slide_text(sl)); d_hits=sum(1 for x in ('2d','3d','4d','5d','6d') if x in q); meta_hits=sum(1 for x in ('signal','이슈기인','발생단계') if x in q)
    return d_hits>=3 or (d_hits>=2 and meta_hits>=1)

def _slide_id_pairs(prs):
    out=[]
    for i,sldId in enumerate(prs.slides._sldIdLst):
        try:sid=int(sldId.get('id'))
        except Exception:
            try:sid=int(sldId.id)
            except Exception:continue
        out.append((i,sid))
    return out

def _native_sections(prs):
    root=prs._element; id_to_index={sid:i for i,sid in _slide_id_pairs(prs)}; sections=[]
    try:sec_nodes=root.xpath('.//*[local-name()="sectionLst"]/*[local-name()="section"]')
    except Exception:sec_nodes=[]
    for sec in sec_nodes:
        name=N(sec.get('name')); ids=[]
        try:id_nodes=sec.xpath('.//*[local-name()="sldId"]')
        except Exception:id_nodes=[]
        for n in id_nodes:
            try:ids.append(int(n.get('id')))
            except Exception:continue
        indices=sorted(id_to_index[sid] for sid in ids if sid in id_to_index)
        sections.append({'name':name,'element':sec,'slide_ids':ids,'indices':indices})
    return sections

def _section_match_score(name,d):
    q=s13._k(name); full,task,customer=_project_parts(d)
    if not q:return 0
    if full and q==full:return 300
    if task and q==task:return 280
    if full and full in q:return 260
    if task and customer and task in q and customer in q:return 250
    if task and task in q:return 220
    return 0

def _matching_native_section(prs,d):
    scored=[]
    for sec in _native_sections(prs):
        sc=_section_match_score(sec['name'],d)
        if sc:scored.append((sc,sec))
    if not scored:return None
    scored.sort(key=lambda x:(x[0],len(x[1]['indices'])),reverse=True); return scored[0][1]

def _section_by_name(prs,name):
    target=s13._k(name)
    if not target:return None
    for sec in _native_sections(prs):
        if s13._k(sec.get('name'))==target:
            return sec
    return None

def section_resolution(prs,d):
    """Resolve exact project section first; otherwise return a safe user-confirmable fallback."""
    try:s13._clear_slide_text_cache()
    except Exception:pass
    exact=_matching_native_section(prs,d)
    if exact:
        return {'mode':'exact','name':N(exact.get('name')),'reason':'고객사_과제명 또는 과제명 기준으로 일치하는 구역을 찾았습니다.','section':exact}

    full,task,customer=_project_parts(d)
    candidates=[]
    sections=_native_sections(prs)

    # First fallback: same customer is explicitly present in the native section name.
    for sec in sections:
        q=s13._k(sec.get('name'))
        if customer and customer in q:
            candidates.append((240,sec,'동일 고객사명이 포함된 구역이 확인되었습니다.'))

    # Second fallback: task/customer text is actually present inside detail pages of a section.
    for sec in sections:
        texts=[]
        for i in sec.get('indices',[]):
            if 0<=i<len(prs.slides) and _is_detail_like(prs.slides[i]):
                texts.append(s13._k(s13._slide_text(prs.slides[i])))
        joined=' '.join(texts)
        if not joined:continue
        if task and task in joined:
            candidates.append((220,sec,'동일 과제명 내용이 포함된 상세 구역이 확인되었습니다.'))
        elif customer and customer in joined:
            candidates.append((200,sec,'동일 고객사 내용이 포함된 상세 구역이 확인되었습니다.'))

    if candidates:
        candidates.sort(key=lambda x:(x[0],len(x[1].get('indices',[]))),reverse=True)
        score,sec,reason=candidates[0]
        return {'mode':'suggest','name':N(sec.get('name')),'reason':reason,'section':sec}

    # No native section candidate: detect a nearby detail page by task/customer text only.
    summaries=_summary_indices(prs)
    nearby=[]
    for i,sl in enumerate(prs.slides):
        if i in summaries or not _is_detail_like(sl):continue
        q=s13._k(s13._slide_text(sl))
        if task and task in q:
            nearby.append((220,i,'동일 과제명 내용이 포함된 상세페이지가 확인되었습니다.'))
        elif customer and customer in q:
            nearby.append((190,i,'동일 고객사 내용이 포함된 상세페이지가 확인되었습니다.'))
    if nearby:
        nearby.sort(reverse=True)
        score,i,reason=nearby[0]
        return {'mode':'suggest_nearby','name':f'상세 page {i+1} 주변','reason':reason,'section':None,'insert_after':i+1}

    return {'mode':'new','name':N(s13._customer_task(d)) or N(d.get('task_name')) or N(d.get('customer')) or '신규 과제','reason':'일치하거나 확인 가능한 기존 구역을 찾지 못했습니다. 신규 구역을 생성합니다.','section':None}

def _selected_section(prs,d,g):
    if str(g.get('_weekly_create_new_section') or '').lower() in ('1','true','yes'):
        return None
    name=N(g.get('_weekly_section_override_name'))
    if name:
        sec=_section_by_name(prs,name)
        if sec:return sec
    return _matching_native_section(prs,d)

def _new_section_name(d):
    return N(s13._customer_task(d)) or N(d.get('task_name')) or N(d.get('customer')) or '신규 과제'

def _create_native_section(prs,name,slide_index):
    """Create a real PowerPoint native section containing the specified detail slide."""
    sid=_slide_id_at(prs,slide_index)
    if sid is None:return None
    root=prs._element
    try:
        lists=root.xpath('./*[local-name()="sectionLst"]')
    except Exception:
        lists=[]
    if lists:
        section_lst=lists[0]
    else:
        section_lst=OxmlElement('p:sectionLst')
        try:root.insert_element_before(section_lst,'p:sldSz','p:notesSz','p:defaultTextStyle','p:extLst')
        except Exception:root.append(section_lst)

    base_name=N(name) or '신규 과제'
    existing={s13._k(x.get('name')) for x in _native_sections(prs)}
    final_name=base_name
    n=2
    while s13._k(final_name) in existing:
        final_name=f'{base_name} ({n})'; n+=1

    sec=OxmlElement('p:section')
    sec.set('name',final_name)
    sec.set('id','{'+str(uuid.uuid4()).upper()+'}')
    sld_lst=OxmlElement('p:sldIdLst')
    sld=OxmlElement('p:sldId'); sld.set('id',str(sid))
    sld_lst.append(sld); sec.append(sld_lst); section_lst.append(sec)
    return {'name':final_name,'element':sec,'slide_ids':[sid],'indices':[slide_index]}

def _slide_id_at(prs,index):
    sldId=prs.slides._sldIdLst[index]
    try:return int(sldId.get('id'))
    except Exception:
        try:return int(sldId.id)
        except Exception:return None

def _add_slide_to_native_section(prs,section,slide_index):
    if not section:return False
    sid=_slide_id_at(prs,slide_index)
    if sid is None:return False
    sec=section['element']
    try:
        for n in sec.xpath('.//*[local-name()="sldId"]'):
            if str(n.get('id'))==str(sid):return True
        lists=sec.xpath('./*[local-name()="sldIdLst"]')
        if not lists:return False
        lst=lists[0]; existing=lst.xpath('./*[local-name()="sldId"]')
        if not existing:return False
        import copy
        node=copy.deepcopy(existing[-1]); node.set('id',str(sid)); lst.append(node); return True
    except Exception:return False

def _project_detail_indices(prs,d):
    sec=_matching_native_section(prs,d); summaries=_summary_indices(prs)
    if sec and sec['indices']:return [i for i in sec['indices'] if i not in summaries and _is_detail_like(prs.slides[i])]
    return [i for i,sl in enumerate(prs.slides) if i not in summaries and _is_detail_like(sl) and _project_score(sl,d)>=120]

def _find_existing_detail(prs,d):
    pages=_project_detail_indices(prs,d)
    if not pages:return None
    issue=s13._k(s13._issue_display(d)); best=None
    for i in pages:
        text=s13._slide_text(prs.slides[i]); q=s13._k(text); score=200+s13._detail_structure_score(prs.slides[i])*4
        if issue and issue in q:score+=150
        else:
            from difflib import SequenceMatcher
            ratios=[SequenceMatcher(None,issue,s13._k(x)).ratio() for x in text.splitlines() if s13._k(x)] if issue else []
            if ratios:score+=int(70*max(ratios))
        if best is None or score>best[0]:best=(score,i)
    return best[1] if best and best[0]>=285 else None

def _new_detail_position(prs,d):
    sec=_matching_native_section(prs,d)
    if sec and sec['indices']:return max(sec['indices'])+1
    pages=_project_detail_indices(prs,d); return max(pages)+1 if pages else len(prs.slides)

def _template_detail_index(prs,d):
    summaries=_summary_indices(prs); same=_project_detail_indices(prs,d)
    if same:return same[-1]
    candidates=[]
    for i,sl in enumerate(prs.slides):
        if i in summaries or not _is_detail_like(sl):continue
        structure=s13._detail_structure_score(sl); q=s13._k(s13._slide_text(sl)); meta=sum(1 for x in ('signal','이슈기인','발생단계') if x in q); candidates.append((structure*10+meta*8,i))
    if not candidates:raise ValueError('주간회의 PPT에서 복제할 수 있는 2D~6D 상세 페이지 양식을 찾지 못했습니다.')
    candidates.sort(reverse=True); return candidates[0][1]

def _overlap_ratio(sh,zone):
    try:x=float(sh.left)/v310.EMU; y=float(sh.top)/v310.EMU; w=float(sh.width)/v310.EMU; h=float(sh.height)/v310.EMU
    except Exception:return 0.0
    zx,zy,zw,zh=zone; ix=max(0,min(x+w,zx+zw)-max(x,zx)); iy=max(0,min(y+h,zy+zh)-max(y,zy)); return ix*iy/max(w*h,1e-6)

def _content_zones():return [(.38,2.35,5.05,1.30),(.38,3.72,5.05,1.50),(.38,5.18,5.05,2.05),(5.58,2.30,5.10,1.25),(5.58,3.56,5.10,2.10),(5.58,5.72,5.10,1.25)]
def _static_detail_label(text):return s13._k(text) in {'2d','3d','4d','5d','6d','7d','현상','임시대책필요시','원인분석','개선대책','유효성점검','수평전개','signal','이슈기인','발생단계'}
def _clear_cloned_issue_content(sl):
    v310.remove_previous_auto(sl); zones=_content_zones()
    for sh in list(sl.shapes):
        if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.GROUP:continue
        if not any(_overlap_ratio(sh,z)>=0.35 for z in zones):continue
        if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.PICTURE:
            try:sh._element.getparent().remove(sh._element)
            except Exception:pass
            continue
        if hasattr(sh,'text_frame'):
            old=N(getattr(sh,'text',''))
            if old and not _static_detail_label(old):
                try:sh.text=''
                except Exception:
                    try:sh.text_frame.clear()
                    except Exception:pass

def _clone_detail_shell(prs,d,insert_at,matched_section=None):
    template_index=_template_detail_index(prs,d)
    s13._clone_slide_with_rels(prs,template_index); s13._move_last_slide_to(prs,insert_at)
    sl=prs.slides[insert_at]; _clear_cloned_issue_content(sl)
    if matched_section:_add_slide_to_native_section(prs,matched_section,insert_at)
    return sl,template_index,matched_section

def _attachment_key(d):
    raw=(s13._customer_task(d)+'|'+s13._issue_display(d)).encode('utf-8','ignore'); return hashlib.sha1(raw).hexdigest()[:12]
def _ps_quote(path):return str(path).replace("'","''")

def _append_8d_attachments(out_path,src8d_path,detail_index,d):
    source_count=(d or {}).get('_attachment_source_count')
    if source_count in (None,''):
        try:source_count=len(Presentation(src8d_path).slides)
        except Exception as e:raise RuntimeError('8D 원본의 유첨 페이지 수를 확인하지 못했습니다: '+repr(e))
    try:source_count=int(source_count)
    except Exception:source_count=0
    if source_count<=1:return 0
    insert_after=detail_index+1; key=_attachment_key(d); prefix='AUTO_8D_ATTACH_'+key+'_'; out_q=_ps_quote(Path(out_path).resolve()); src_q=_ps_quote(Path(src8d_path).resolve()); prefix_q=prefix.replace("'","''")
    script=f"""
$ErrorActionPreference = 'Stop'
$ppt = $null
$pres = $null
try {{
    $ppt = New-Object -ComObject PowerPoint.Application
    try { $ppt.DisplayAlerts = 1 } catch {}
    try { $ppt.AutomationSecurity = 3 } catch {}
    $pres = $ppt.Presentations.Open('{out_q}', 0, 0, 0)
    for ($i = $pres.Slides.Count; $i -ge 1; $i--) {{
        $nm = [string]$pres.Slides.Item($i).Name
        if ($nm.StartsWith('{prefix_q}')) {{ $pres.Slides.Item($i).Delete() }}
    }}
    $added = $pres.Slides.InsertFromFile('{src_q}', {insert_after}, 2, {source_count})
    for ($j = 1; $j -le $added; $j++) {{ $pres.Slides.Item({insert_after} + $j).Name = '{prefix_q}' + $j }}
    $pres.Save()
    Write-Output $added
}}
finally {{
    if ($pres -ne $null) {{ try {{ $pres.Close() }} catch {{}} }}
    if ($ppt -ne $null) {{ try {{ $ppt.Quit() }} catch {{}} }}
}}
"""
    last_detail=''
    for attempt in range(2):
        try:
            p=subprocess.run(
                ['powershell.exe','-NoProfile','-Sta','-ExecutionPolicy','Bypass','-Command',script],
                capture_output=True,text=True,timeout=60
            )
        except FileNotFoundError:
            raise RuntimeError('Windows PowerShell을 찾지 못해 8D 유첨 페이지를 붙이지 못했습니다.')
        except subprocess.TimeoutExpired:
            # Do not repeat a hung Office COM call; one timeout is enough.
            raise RuntimeError('PowerPoint 유첨 페이지 복사 작업이 60초를 초과해 중단되었습니다.')
        if p.returncode==0:
            nums=[int(x.strip()) for x in (p.stdout or '').splitlines() if x.strip().isdigit()]
            return nums[-1] if nums else source_count-1
        last_detail=(p.stderr or p.stdout or '').strip()
        rpc=('RPC' in last_detail.upper() or '0X800706BA' in last_detail.upper() or '0X80010108' in last_detail.upper())
        if attempt==0 and rpc:
            continue
        break
    raise RuntimeError('8D 유첨 페이지 자동 복사 실패: '+last_detail[-800:])

def _clean_updated_summary_placeholder(prs,summary_index):
    """Remove empty content placeholders only on the summary page we just touched."""
    if summary_index is None or summary_index<0 or summary_index>=len(prs.slides):
        return False
    changed=False
    sl=prs.slides[summary_index]
    for sh in list(sl.shapes):
        try:
            if sh.shape_type!=MSO_SHAPE_TYPE.PLACEHOLDER:
                continue
            if getattr(sh,'has_table',False):
                continue
            if str(getattr(sh,'text','') or '').strip():
                continue
            sh._element.getparent().remove(sh._element)
            changed=True
        except Exception:
            pass
    return changed

def _weekly_progress(g,text):
    cb=(g or {}).get('_weekly_progress_callback')
    if callable(cb):
        try:cb(text)
        except Exception:pass

def _append_8d_attachments_safe(out_path,src8d_path,detail_index,d):
    """Attachment failure must not discard the already-saved weekly update."""
    try:
        return _append_8d_attachments(out_path,src8d_path,detail_index,d),''
    except Exception as e:
        return 0,str(e)

def weekly_fix5(src,out,d,g,mode):
    try:s13._clear_slide_text_cache()
    except Exception:pass
    _weekly_progress(g,'주간회의 파일을 불러오는 중...')
    prs=Presentation(src)

    _weekly_progress(g,'주간회의 요약 페이지를 확인하는 중...')
    summary_index,_,summary_action=s14._update_summary_by_task(prs,d,g,mode)

    _weekly_progress(g,'상세 페이지 위치를 확인하는 중...')
    matched_section=_selected_section(prs,d,g)
    force_new=str(g.get('_weekly_create_new_section') or '').lower() in ('1','true','yes')
    nearby_after=g.get('_weekly_insert_after_override')

    # Existing issue lookup is constrained by an explicitly chosen native section when possible.
    target=None
    if mode=='existing' and not force_new:
        if matched_section and matched_section.get('indices'):
            pages=[i for i in matched_section['indices'] if 0<=i<len(prs.slides) and _is_detail_like(prs.slides[i])]
            issue=s13._k(s13._issue_display(d)); best=None
            for i in pages:
                text=s13._slide_text(prs.slides[i]); q=s13._k(text)
                score=200+s13._detail_structure_score(prs.slides[i])*4
                if issue and issue in q:score+=150
                if best is None or score>best[0]:best=(score,i)
            if best and best[0]>=285:target=best[1]
        else:
            target=_find_existing_detail(prs,d)

    detail_action='업데이트'
    if target is None:
        if matched_section and matched_section.get('indices'):
            insert_at=max(matched_section['indices'])+1
        elif nearby_after not in (None,''):
            try:insert_at=max(0,min(len(prs.slides),int(nearby_after)))
            except Exception:insert_at=len(prs.slides)
        else:
            insert_at=len(prs.slides) if force_new else _new_detail_position(prs,d)

        _,template_index,_=_clone_detail_shell(prs,d,insert_at,matched_section)
        target=insert_at

        if force_new:
            matched_section=_create_native_section(prs,_new_section_name(d),target)
            sec_name=N(matched_section.get('name')) if matched_section else '(신규 텍스트 구역)'
            detail_action=f'신규 과제 구역 [{sec_name}] 생성 후 상세 페이지 추가 (양식 원본 slide {template_index+1})'
        else:
            sec_name=N(matched_section.get('name')) if matched_section else '(확인된 상세페이지 주변)'
            detail_action=f'선택 구역 [{sec_name}]에 상세 페이지 추가 (양식 원본 slide {template_index+1})'

    _weekly_progress(g,'상세 페이지 내용을 반영하는 중...')
    s13._update_detail_slide(prs.slides[target],d,g,mode)
    try:s13._clear_slide_text_cache(prs.slides[target])
    except Exception:pass
    # Remove placeholder UI artifacts before the first save. This avoids the
    # expensive post-save reopen/resave pass on large weekly decks.
    if _clean_updated_summary_placeholder(prs,summary_index):
        g['_weekly_placeholder_cleaned']='1'
    else:
        # The page was inspected even if it had nothing to remove.
        g['_weekly_placeholder_cleaned']='1'

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    _weekly_progress(g,'주간회의 파일을 저장하는 중...')
    try:prs.save(out); saved=out
    except PermissionError:
        p=Path(out); saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix)); prs.save(saved)

    _weekly_progress(g,'8D 유첨 페이지를 확인하는 중...')
    attached,attach_error=_append_8d_attachments_safe(saved,g.get('ppt8d',''),target,d)
    sec_name=N(matched_section.get('name')) if matched_section else ('신규 구역' if force_new else '확인된 상세페이지 주변')
    if attach_error:
        attach_msg=' / 8D 유첨 추가 실패(주간회의 본문은 저장됨): '+attach_error[:220]
    else:
        attach_msg=f' / 8D 유첨 {attached}페이지 추가' if attached else ' / 8D 유첨 없음'
    _weekly_progress(g,'주간회의 업데이트 마무리 중...')
    return ('주간회의 PPT: 요약 '+summary_action+' / 상세 '+detail_action+f' / 배치 구역={sec_name} / 과제명={N(d.get("task_name"))}'+attach_msg,saved)

base.weekly=weekly_fix5
class RecoveryStep14Fix5App(s10.RecoveryStep10App):
    def __init__(self):super().__init__(); self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP14 FIX5')
if __name__=='__main__':RecoveryStep14Fix5App().mainloop()
