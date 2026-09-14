import os
from pathlib import Path
from tkinter import messagebox
from openpyxl import load_workbook

import main_recovery_step9 as step9
import main_recovery_step8 as step8
import main_recovery_step7 as step7

base=step9.base
N=step9.N
OCCURRENCE_SITES=step9.OCCURRENCE_SITES


class RecoveryStep10App(step9.RecoveryStep9App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP10')

    def run(self):
        g=self.gui()
        if any(not g[k] or not os.path.exists(g[k]) for k in ['ppt8d','pptweekly','xlsx']):
            return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')
        if g.get('occurrence_site') not in OCCURRENCE_SITES:
            return messagebox.showwarning('확인','발생처를 선택하세요.')

        try:
            d=base.extract(g['ppt8d'])
            mode=self.mode.get()

            # Check Excel matching BEFORE DB-only dialogs.
            wb=load_workbook(g['xlsx'])
            ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active
            r,_=base.find(ws,d,g)

            weekly_only=False
            if mode=='existing' and not r:
                prompt=(
                    '기존 이슈를 Issue DB Excel에서 정확하게 찾지 못했습니다.\n\n'
                    'Excel 업데이트는 건너뛰고 주간회의 자료만 업데이트할까요?\n\n'
                    '예 : Excel은 변경하지 않고 주간회의 PPT만 업데이트\n'
                    '아니오 : 전체 업데이트 중단'
                )
                if not messagebox.askyesno('기존 이슈 미확인',prompt,parent=self):
                    self.log.delete('1.0','end')
                    self.log.insert('end','업데이트가 취소되었습니다.')
                    return
                weekly_only=True

            if mode=='new' and r and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?',parent=self):
                return

            # Weekly page2 issue-origin is needed whether Excel is updated or not.
            od=step8.OriginDialog(self,d)
            if od.result is None:
                self.log.delete('1.0','end')
                self.log.insert('end','업데이트가 취소되었습니다.')
                return
            g['_issue_origin_selected']=od.result

            pd_result=None
            final_status=None

            # These two confirmations are Issue-DB-only, so skip them for weekly-only mode.
            if not weekly_only:
                summary=step7._db_summary(d)
                pd=step7.ProblemChoiceDialog(self,summary)
                if pd.result is None:
                    self.log.delete('1.0','end')
                    self.log.insert('end','업데이트가 취소되었습니다.')
                    return
                pd_result=pd.result
                g['_db_problem_selected']=summary if pd.result=='summary' else N(d.get('problem'))

                judged,reason=step9._judge_issue_status(d)
                final_status=judged
                if judged=='close':
                    sixd=reason or '(6D 내용 없음)'
                    prompt=(
                        '6D까지 작성되어 있으며, 아래 6D 내용에 진행 중/예정 표현이 없어 close로 판단됩니다.\n\n'
                        '판단 이유(6D 내용)\n'
                        '────────────────────\n'
                        f'{sixd}\n'
                        '────────────────────\n\n'
                        '이슈 상태를 close로 처리하시겠습니까?\n'
                        '아니오를 선택하면 open으로 처리합니다.'
                    )
                    if not messagebox.askyesno('이슈 상태 확인',prompt,parent=self):
                        final_status='open'
                g['_issue_status_selected']=final_status

            out=Path(g['xlsx']).parent/'자동화_결과'
            out.mkdir(exist_ok=True)
            xo=out/(Path(g['xlsx']).stem+'_업데이트.xlsx')
            po=out/(Path(g['pptweekly']).stem+'_업데이트.pptx')

            if weekly_only:
                a='Issue DB Excel 업데이트: 건너뜀 (기존 이슈 미확인)'
            else:
                a,_=base.update_excel(g['xlsx'],xo,d,g,new=(mode=='new'))

            b,_=base.weekly(g['pptweekly'],po,d,g,mode)

            self.log.delete('1.0','end')
            lines=[
                '=== 완료 ===',
                f'이슈 구분: {mode}',
                f'업데이트 방식: {"주간회의 PPT만" if weekly_only else "Excel + 주간회의 PPT"}',
                f'발생처: {g.get("occurrence_site","")}',
                f'이슈기인: {od.result}',
            ]
            if not weekly_only:
                phenomenon_label='현상 요약' if pd_result=='summary' else '전체 내용'
                lines.extend([
                    f'Issue DB 현상: {phenomenon_label}',
                    f'Issue DB 상태: {final_status}',
                ])
            lines.extend(['',a,b,'',f'결과: {out}'])
            self.log.insert('end','\n'.join(lines))

            done=(
                '자동 업데이트가 완료되었습니다.\n\n'
                + ('Issue DB Excel: 업데이트 안 함\n주간회의 PPT: 업데이트 완료\n\n' if weekly_only else 'Issue DB Excel: 업데이트 완료\n주간회의 PPT: 업데이트 완료\n\n')
                + f'이슈기인: {od.result}\n\n{out}'
            )
            messagebox.showinfo('완료',done,parent=self)

        except Exception as e:
            messagebox.showerror('실행 오류',repr(e),parent=self)


if __name__=='__main__':
    RecoveryStep10App().mainloop()
