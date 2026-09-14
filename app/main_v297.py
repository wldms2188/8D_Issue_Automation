# 8D Issue Automation v2.9.7 - diagnostic runner
# Keeps v2.9.6 processing logic unchanged, but exposes the exact GUI execution stage
# and writes a full traceback so a newly introduced regression can be isolated without
# requiring terminal output.
import os
import traceback
from pathlib import Path
from openpyxl import load_workbook
from tkinter import messagebox

import main_v296 as impl
base = impl.base


class App(base.App):
    def run(self):
        stage = 'GUI 입력값 확인'
        g = None
        log_path = None
        try:
            g = self.gui()
            for k in ('ppt8d', 'pptweekly', 'xlsx'):
                if not g.get(k):
                    messagebox.showwarning('확인', '파일을 모두 선택해주세요.')
                    return

            stage = '1/4 8D PPT 읽기'
            if hasattr(self, 'log'):
                self.log.delete('1.0', 'end')
                self.log.insert('end', '[1/4] 8D PPT 읽는 중...\n')
                self.log.update_idletasks()
            d = base.extract(g['ppt8d'])

            stage = '2/4 출력 경로 및 Excel 준비'
            mode = self.mode.get()
            outdir = Path(g['xlsx']).parent / '자동화_결과'
            outdir.mkdir(exist_ok=True)
            xo = outdir / (Path(g['xlsx']).stem + '_자동화.xlsx')
            po = outdir / (Path(g['pptweekly']).stem + '_자동화.pptx')
            wb = load_workbook(g['xlsx'])
            ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active

            if hasattr(self, 'log'):
                self.log.insert('end', '[2/4] Excel 기존 이슈 확인 중...\n')
                self.log.update_idletasks()
            r, _ = base.find(ws, d, g)

            stage = '3/4 Excel 이슈 DB 업데이트'
            if hasattr(self, 'log'):
                self.log.insert('end', '[3/4] Excel 업데이트 중...\n')
                self.log.update_idletasks()
            a, _ = base.update_excel(g['xlsx'], xo, d, g, new=(mode == 'new'))

            stage = '4/4 주간회의 PPT 업데이트'
            if hasattr(self, 'log'):
                self.log.insert('end', '[4/4] 주간회의 PPT 업데이트 중...\n')
                self.log.update_idletasks()
            b, _ = base.weekly(g['pptweekly'], po, d, g, mode)

            if hasattr(self, 'log'):
                self.log.insert('end', '\n완료\n' + str(a) + '\n' + str(b) + '\n')
            messagebox.showinfo('완료', f'자동 업데이트가 완료되었습니다.\n\nExcel: {xo}\nPPT: {po}')

        except Exception as e:
            # Never hide the actual call location behind repr(e). This is intentionally
            # kept in the GUI build so company-PC testing does not require terminal use.
            tb = traceback.format_exc()
            if g and g.get('xlsx'):
                outdir = Path(g['xlsx']).parent / '자동화_결과'
            else:
                outdir = Path.cwd() / '자동화_결과'
            outdir.mkdir(exist_ok=True)
            log_path = outdir / '자동업데이트_오류.log'
            try:
                log_path.write_text(
                    '[8D Issue Automation v2.9.7]\n'
                    f'발생 단계: {stage}\n'
                    f'오류: {type(e).__name__}: {e}\n\n'
                    '--- Traceback ---\n' + tb,
                    encoding='utf-8'
                )
            except Exception:
                pass
            if hasattr(self, 'log'):
                self.log.insert('end', f'\n오류 단계: {stage}\n{type(e).__name__}: {e}\n')
            messagebox.showerror(
                '자동 업데이트 오류',
                f'발생 단계:\n{stage}\n\n'
                f'{type(e).__name__}: {e}\n\n'
                f'상세 오류 로그:\n{log_path}'
            )


if __name__ == '__main__':
    App().mainloop()
