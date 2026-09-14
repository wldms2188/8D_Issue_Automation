# 8D Issue Automation v3.0.4
# Fixes task extraction, preserves template 2D~6D circular markers,
# and prevents unnecessary text shrinking.
import sys, re
from pathlib import Path
APP_DIR=Path(__file__).resolve().parent
if str(APP_DIR) not in sys.path: sys.path.insert(0,str(APP_DIR))
import main_v303 as v303
base=v303.base
renderer=v303.renderer
from pptx import Presentation
from pptx.util import Inches

def _n(x): return str(x or '').replace('\r\n','\n').replace('\r','\n').strip()
def _c(x):
    try: return base.compact(x)
    except Exception: return re.sub(r'[^0-9A-Za-z가-힣]','',_n(x)).lower()
def _valid_task(x):
    q=_c(x)
    if not x or len(_n(x))<2: return False
    blocked=('고객사','이슈명','과제명','model','packer','발생site','발생일자','lotno','발생line','담당자','제품타입','폼팩터','발생샘플','개발단계','발생처','불량률','불량수량','발생현상','signal')
    return not any(k in q for k in blocked)

def _task_from_customer_area(path):
    prs=Presentation(path)
    # Authoritative order: 고객사 -> 고객사 값 -> 과제명 -> 과제명 값.
    for sl in prs.slides:
        for sh in sl.shapes:
            if not getattr(sh,'has_table',False): continue
            tb=sh.table
            rows=[[_n(tb.cell(r,c).text) for c in range(len(tb.columns))] for r in range(len(tb.rows))]
            for r,row in enumerate(rows):
                for c,val in enumerate(row):
                    if '고객사' not in _c(val): continue
                    right=row[c+1:]
                    # First look for an explicit 과제명 label after 고객사.
                    for j,cand in enumerate(right):
                        if '과제명' in _c(cand) and j+1<len(right) and _valid_task(right[j+1]):
                            return _n(right[j+1])
                    # Then support [고객사][고객][과제].
                    vals=[x for x in right if _n(x)]
                    if len(vals)>=2 and _valid_task(vals[1]): return _n(vals[1])
                    # If the label/value pair is split to the next row.
                    for rr in range(r+1,min(r+4,len(rows))):
                        below=rows[rr]
                        for cc,cand in enumerate(below):
                            if '과제명' in _c(cand) and cc+1<len(below) and _valid_task(below[cc+1]):
                                return _n(below[cc+1])
    # Text-box form.
    for sl in prs.slides:
        for sh in sl.shapes:
            text=_n(getattr(sh,'text',''))
            if not text or '고객사' not in _c(text): continue
            lines=[x.strip() for x in text.split('\n') if x.strip()]
            for i,line in enumerate(lines):
                if '고객사' not in _c(line): continue
                for j in range(i+1,min(i+5,len(lines))):
                    if '과제명' in _c(lines[j]):
                        p=re.split(r'[:：]',lines[j],maxsplit=1)
                        if len(p)==2 and _valid_task(p[1]): return _n(p[1])
                        if j+1<len(lines) and _valid_task(lines[j+1]): return _n(lines[j+1])
                    elif _valid_task(lines[j]): return _n(lines[j])
    return ''

_prev_extract=base.extract
def extract_v304(path):
    d=_prev_extract(path)
    try:
        task=_task_from_customer_area(path)
        if task: d['task_name']=task
    except Exception: pass
    return d
base.extract=extract_v304

# Keep 8pt whenever it fits; only shrink if the actual box cannot contain the text.
_orig_fit=renderer._fit_text
def _fit_text_v304(sh,text,max_size=8,min_size=5,prefix=None):
    return _orig_fit(sh,text,max_size=max_size,min_size=max(7.0,min_size),prefix=prefix)
renderer._fit_text=_fit_text_v304

# Find existing template markers recursively; never create replacement 2D~6D labels.
def _walk(container):
    for sh in getattr(container,'shapes',[]):
        yield sh
        if hasattr(sh,'shapes'):
            yield from _walk(sh)

def _move_markers(sl,layout):
    used=set()
    for label,zkey in [('2D','2D'),('3D','3D'),('4D','4D_CAUSE'),('5D','5D'),('6D','6D')]:
        z=layout.get(zkey)
        if not z: continue
        candidates=[]
        for sh in _walk(sl):
            if _n(getattr(sh,'text',''))==label and not str(getattr(sh,'name','')).startswith('AUTO_8D_'):
                candidates.append(sh)
        target=next((x for x in candidates if id(x) not in used),None)
        if target is None: continue
        used.add(id(target))
        try:
            tw=target.width/914400.0; th=target.height/914400.0
            target.left=Inches(max(0.05,z['x']-tw-0.08))
            target.top=Inches(z['y']+max(0,(z['h']-th)/2))
        except Exception: pass

_orig_fill=renderer._fill_page2
def _fill_page2_v304(sl,d,g,layout):
    _orig_fill(sl,d,g,layout)
    # v300 creates textual labels internally; remove those only.
    for sh in list(sl.shapes):
        if str(getattr(sh,'name','')).startswith('AUTO_8D_LABEL_'):
            el=sh._element; el.getparent().remove(el)
    _move_markers(sl,layout)
renderer._fill_page2=_fill_page2_v304

if __name__=='__main__': base.App().mainloop()
