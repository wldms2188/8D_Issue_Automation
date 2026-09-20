import copy
import datetime
from pathlib import Path

from pptx import Presentation

import main_recovery_step10 as step10
import main_recovery_step4 as step4
import main_recovery_step8 as step8
import main_v320 as v320
import main_v319 as v319
import main_v315 as v315
import main_v313 as v313
import main_v310 as v310

base=step10.base
N=v310.N


def _layout_with_4d_placeholders(d, imgs):
    """Keep BOTH 4D slots in the normal two-column chain even when 4D is not written yet."""
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    # Explicit requested placeholder. Do not let missing 4D collapse the layout.
    if not N(d.get('cause_4d')):
        texts['4D_CAUSE']='검토 중'
    if not (N(d.get('leak_cause')) or N(d.get('system_cause'))):
        texts['4D_LEAK']='검토 중'

    # v319._fit_chain normally removes an empty 4D. Feed placeholders only to the
    # LAYOUT copy so 2D-3D-4D and 4D-5D-6D always keep their normal vertical order.
    dl=dict(d)
    if not N(dl.get('cause_4d')):
        dl['cause_4d']='검토 중'
    if not (N(dl.get('leak_cause')) or N(dl.get('system_cause'))):
        dl['leak_cause']='검토 중'

    li={'2D':imgs.get('2D',[]),'3D':imgs.get('3D',[]),'4D_CAUSE':imgs.get('4D_CAUSE',[])}
    ri={'4D_LEAK':imgs.get('4D_LEAK',[]),'5D':imgs.get('5D',[]),'6D':imgs.get('6D',[])}
    lz,lf=v319._fit_chain(('2D','3D','4D_CAUSE'),dl,texts,li)
    rz,rf=v319._fit_chain(('4D_LEAK','5D','6D'),dl,texts,ri)
    return {**lz,**rz},texts,{**lf,**rf}


def _adaptive_cascade_layout(d,imgs):
    """Cascade D blocks using actual content height while keeping them non-overlapping.

    Short 2D/3D content releases vertical space to the following section, so a long
    4D can move upward.  If 4D still needs more room, its box grows downward up to
    the safe slide bottom.  Text is auto-fit only after available geometry is used.
    """
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    if not N(d.get('cause_4d')):
        texts['4D_CAUSE']='검토 중'
    if not (N(d.get('leak_cause')) or N(d.get('system_cause'))):
        texts['4D_LEAK']='검토 중'

    chains=(('2D','3D','4D_CAUSE'),('4D_LEAK','5D','6D'))
    zones={}; fonts={}
    bottom=v319.BOTTOM
    gap=v319.GAP

    for keys in chains:
        top=v310.ZONES[keys[0]]['y']
        available=bottom-top-gap*(len(keys)-1)
        fs={k:v319.MAX_FONT for k in keys}
        hs={k:v319._need_h(k,texts[k],bool(imgs.get(k)),fs[k]) for k in keys}

        # Use geometry first; reduce only the section that gains the most room.
        while sum(hs.values())>available+.01:
            options=[]
            for k in keys:
                if fs[k]<=6.0:
                    continue
                nf=max(6.0,fs[k]-.5)
                nh=v319._need_h(k,texts[k],bool(imgs.get(k)),nf)
                options.append((hs[k]-nh,k,nf,nh))
            if not options:
                break
            gain,k,nf,nh=max(options,key=lambda x:x[0])
            if gain<=.005:
                break
            fs[k]=nf; hs[k]=nh

        # If content is still larger than the slide, fit the geometry to the
        # remaining column height. Auto-fit in step12 preserves all text inside.
        overflow=sum(hs.values())-available
        if overflow>0:
            mins={k:v318.v316.v312.MIN_H[k]*.82 for k in keys}
            room=overflow
            for k in sorted(keys,key=lambda x:hs[x]-mins[x],reverse=True):
                if room<=.001:
                    break
                reducible=max(0.0,hs[k]-mins[k])
                cut=min(reducible,room)
                hs[k]-=cut
                room-=cut
            if room>0:
                scale=available/max(sum(hs.values()),.01)
                hs={k:max(.42,hs[k]*scale) for k in keys}

        y=top
        for k in keys:
            z=dict(v310.ZONES[k])
            z['y']=y
            z['h']=max(.42,hs[k])
            zones[k]=z
            fonts[k]=fs[k]
            y+=z['h']+gap

        # Hard safety: last section may touch, but never cross, the slide bottom.
        last=keys[-1]
        end=zones[last]['y']+zones[last]['h']
        if end>bottom:
            zones[last]['h']=max(.42,bottom-zones[last]['y'])

    return zones,texts,fonts


def _strict_template_layout(d,imgs):
    """Keep every D inside its original weekly-template content rectangle.

    The previous dynamic layout could enlarge 2D downward when its text/image was
    long, while the 3D marker/title stayed at the template position.  Visually this
    made a genuine 2D image look as if it belonged to 3D.  This layout never lets a
    section borrow vertical space from the next D section.
    """
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    if not N(d.get('cause_4d')):
        texts['4D_CAUSE']='검토 중'
    if not (N(d.get('leak_cause')) or N(d.get('system_cause'))):
        texts['4D_LEAK']='검토 중'

    zones={k:dict(v310.ZONES[k]) for k in v310.ZONES}
    fonts={}
    for key,z in zones.items():
        font=v319.MAX_FONT
        need=v319._need_h(key,texts[key],bool(imgs.get(key)),font)
        while need>z['h'] and font>6.0:
            font=max(6.0,font-.5)
            need=v319._need_h(key,texts[key],bool(imgs.get(key)),font)
        fonts[key]=font
    return zones,texts,fonts


def _english_template_anchored_layout(d,imgs):
    """Keep the weekly template's native D regions as hard boundaries for English text.

    English sentences can be much longer than Korean.  Use each native zone's
    original top position, allow the text box to grow only into the existing gap
    before the next D zone, and never cross that next zone.  PowerPoint auto-fit
    handles any remaining excess text inside the assigned region.
    """
    texts={k:v313._section_text(d,k) for k in v310.ZONES}
    if not N(d.get('cause_4d')):
        texts['4D_CAUSE']='검토 중'
    if not (N(d.get('leak_cause')) or N(d.get('system_cause'))):
        texts['4D_LEAK']='검토 중'

    chains=(('2D','3D','4D_CAUSE'),('4D_LEAK','5D','6D'))
    zones={}
    fonts={}
    bottom=v319.BOTTOM
    safe_gap=.12

    for keys in chains:
        for i,key in enumerate(keys):
            z0=dict(v310.ZONES[key])
            next_y=(v310.ZONES[keys[i+1]]['y'] if i+1<len(keys) else bottom)
            max_h=max(.35,next_y-z0['y']-safe_gap)

            font=v319.MAX_FONT
            need=v319._need_h(key,texts[key],bool(imgs.get(key)),font)
            # Try reducing only this section's font before constraining height.
            while need>max_h and font>6.0:
                font=max(6.0,font-.5)
                need=v319._need_h(key,texts[key],bool(imgs.get(key)),font)

            z0['h']=min(max(z0['h'],need),max_h)
            zones[key]=z0
            fonts[key]=font

    return zones,texts,fonts


def _restore_missing_4d_unit(sl, zones):
    """Never delete a 4D marker/title. If one native unit remains, clone it only when needed.
    No fallback rectangle is created.
    """
    try:
        units=v315._top_4d_units(sl)
    except Exception:
        return
    if len(units)>=2 or not units:
        return

    unit=units[0]
    try:
        ux=float(unit.left)/v310.EMU
    except Exception:
        ux=0

    zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']
    left_xy=(max(.10,zl['x']-.18),max(.10,zl['y']-.35))
    right_xy=(max(.10,zr['x']-.18),max(.10,zr['y']-.35))
    target=left_xy if ux>4.0 else right_xy

    # Clone the existing native marker/title unit. This preserves its appearance and
    # does not construct a new PowerPoint XML group.
    try:
        newel=copy.deepcopy(unit._element)
        sl.shapes._spTree.insert_element_before(newel,'p:extLst')
        dup=None
        for sh in sl.shapes:
            if sh._element is newel:
                dup=sh
                break
        if dup is not None:
            dup.name='AUTO_8D_4D_RESTORED'
            bx,by,_,_=v310.box(dup)
            from pptx.util import Inches
            dup.left += Inches(target[0])-int(bx)
            dup.top += Inches(target[1])-int(by)
    except Exception:
        pass


def _update_page2_step11(sl,d,g,mode):
    # Remove only prior AUTO text/images; keep the template/native 2D~6D graphics.
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    zones,texts,fonts=_layout_with_4d_placeholders(d,imgs)

    # Existing issue: preserve the current marker/title positions.
    # New issue: retain the same established movement behavior for 2D/3D/5D/6D.
    if v310.is_new(mode):
        for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
            z=zones[key]
            v310.move_marker_unit(sl,label,max(.10,z['x']-.18),max(.10,z['y']-.35))

    # 4D is ALWAYS rendered. Missing source content becomes '검토 중'.
    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        v319._render(sl,key,zones[key],texts[key],imgs.get(key,[]),fonts[key])

    v319._page2_meta(sl,d,g)

    # Previous versions deleted 4D units when the cause was empty. Never do that now.
    # If the input is already a prior generated file with only one 4D unit, restore
    # the missing unit by cloning the surviving native unit rather than drawing a box.
    _restore_missing_4d_unit(sl,zones)


def weekly_step11(src,out,d,g,mode):
    dd=v319._weekly_display(d)
    p1=dict(dd)
    p1['issue_name']=step4.step1._page1_issue(d)
    p1['problem']=step4._full_problem(d)

    prs=Presentation(src)
    v319._update_page1(prs,p1,g)

    if len(prs.slides)>1:
        _update_page2_step11(prs.slides[1],dd,g,mode)
        step4._force_page2_header(prs.slides[1],d,g)

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)

    # Keep STEP8 issue-origin formatting/update after the PPT has been saved.
    origin=N(g.get('_issue_origin_selected'))
    if origin:
        step8._set_origin_in_ppt(saved,origin)

    return '주간회의 PPT 업데이트: '+base.status(d),saved


base.weekly=weekly_step11


class RecoveryStep11App(step10.RecoveryStep10App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP11')


if __name__=='__main__':
    RecoveryStep11App().mainloop()
