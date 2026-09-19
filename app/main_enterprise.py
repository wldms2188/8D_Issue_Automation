"""Enterprise presentation layer for 8D Issue Automation.
Keeps the validated processing core and replaces presentation/UI only.
"""
import os
import tkinter as tk
from tkinter import ttk, filedialog
from pathlib import Path
from openpyxl import load_workbook

import main_final as legacy
import ui_enterprise as ui

try:
    from tkinterdnd2 import TkinterDnD
    _RootBase=TkinterDnD.Tk
except Exception:
    _RootBase=tk.Tk

base=legacy.base; N=legacy.N

def target_ready_status(ppt8d,current):
    current=str(current or '').strip()
    if ppt8d and ('8D 원본을 선택' in current or not current):
        return 'READY  ·  8D 원본 선택 완료 · 미리보기 가능'
    if not ppt8d and current.startswith('READY'):
        return 'READY  ·  8D 원본을 선택해 주세요.'
    return current
def weekly_status_from_choice(d,choice):
    if str(choice or '').lower()=='close':
        return '개선 완료'
    if N(d.get('action_5d')) or N(d.get('verification_6d')):
        return '개선 검증중'
    return '원인/개선 미확인'

def issue_db_status_reason(d,judged):
    """Human-readable reason for the Issue DB open/close proposal."""
    judged='close' if str(judged or '').lower()=='close' else 'open'
    sixd=N(d.get('verification_6d'))
    if not sixd:
        return '6D 효과검증 내용이 없어 완료 여부를 확인할 수 없으므로 open으로 판단했습니다.'
    q=sixd.lower().replace(' ','')
    if judged=='close':
        if any(x in q for x in ('완료','정상','이상없음','검증완료','verificationcomplete','completed','passed','validated','verified','noabnormality')):
            return '6D에서 완료·정상·검증완료 의미의 표현을 확인하여 close로 판단했습니다.'
        return '6D 효과검증 결과가 완료 상태로 판정되어 close로 판단했습니다.'
    if any(x in q for x in ('진행중','검증중','확인중','예정','계획','inprogress','ongoing','pending','planned','scheduled','tbd')):
        return '6D에서 진행 중·검증 중·예정 의미의 표현을 확인하여 open으로 판단했습니다.'
    if any(x in q for x in ('불량','이상','실패','재발','abnormal','failed','failure','nok','oos')):
        return '6D에서 이상·실패·재발 의미의 표현을 확인하여 open으로 판단했습니다.'
    return '6D에 완료를 확정할 표현이 명확하지 않아 open으로 판단했습니다.'

def weekly_status_reason(d,weekly_status):
    """Human-readable reason for the weekly-meeting Signal proposal."""
    st=N(weekly_status)
    if st=='개선 완료':
        return '6D 효과검증이 완료 상태로 판단되어 주간회의 Signal을 개선 완료로 판단했습니다.'
    if st=='개선 검증중':
        if N(d.get('verification_6d')):
            return '6D 효과검증 내용은 있으나 완료가 확정되지 않아 주간회의 Signal을 개선 검증중으로 판단했습니다.'
        return '5D 개선대책은 있으나 6D 완료 검증이 확인되지 않아 주간회의 Signal을 개선 검증중으로 판단했습니다.'
    return '5D 개선대책과 6D 효과검증 내용이 확인되지 않아 주간회의 Signal을 원인/개선 미확인으로 판단했습니다.'




class EnterpriseOriginDialog(tk.Toplevel):
    def __init__(self,parent,d):
        super().__init__(parent); self.withdraw(); self.result=None; self.title('이슈기인 확인'); self.configure(bg=ui.WHITE); self.resizable(False,False); self.transient(parent)
        cause=legacy.step8._cause_text(d); rec,reason=legacy.step8._recommend_origin(d)
        head=tk.Frame(self,bg=ui.NAVY,height=58); head.pack(fill='x'); head.pack_propagate(False)
        tk.Label(head,text='이슈기인 확인',bg=ui.NAVY,fg='white',font=('Malgun Gothic',12,'bold')).pack(side='left',padx=22)
        body=tk.Frame(self,bg='white'); body.pack(fill='both',expand=True,padx=24,pady=18)
        tk.Label(body,text='8D 원인 내용을 기준으로 이슈기인을 확인합니다.',bg='white',fg=ui.TEXT,font=('Malgun Gothic',10,'bold')).pack(anchor='w')
        tk.Label(body,text='8D 원인 내용',bg='white',fg=ui.MUTED,font=('Malgun Gothic',8,'bold')).pack(anchor='w',pady=(14,5))
        box=tk.Text(body,height=10,wrap='word',bg='#F7F9FB',fg=ui.TEXT,relief='flat',highlightthickness=1,highlightbackground=ui.BORDER,font=('Malgun Gothic',9),padx=10,pady=8)
        box.pack(fill='x'); box.insert('1.0',cause or '(4D 원인 내용 없음)'); box.configure(state='disabled')
        card=tk.Frame(body,bg='#EEF5FA',highlightbackground='#D5E3EE',highlightthickness=1); card.pack(fill='x',pady=12)
        tk.Label(card,text=f'추천  {rec}',bg='#EEF5FA',fg=ui.NAVY,font=('Malgun Gothic',10,'bold')).pack(anchor='w',padx=12,pady=(9,2))
        tk.Label(card,text=reason,bg='#EEF5FA',fg='#4E6375',font=('Malgun Gothic',8),wraplength=690,justify='left').pack(anchor='w',padx=12,pady=(0,9))
        self.var=tk.StringVar(value=rec)
        if rec=='TBD':
            tk.Label(body,text='원인 미확정 상태이므로 TBD로 반영됩니다.',bg='white',fg=ui.MUTED,font=('Malgun Gothic',9)).pack(anchor='w')
        else:
            row=tk.Frame(body,bg='white'); row.pack(fill='x',pady=(2,5))
            for value in legacy.step8.ORIGINS:
                ttk.Radiobutton(row,text=value+('  · 추천' if value==rec else ''),variable=self.var,value=value,style='Mode.TRadiobutton').pack(side='left',padx=(0,14))
        foot=tk.Frame(self,bg='#F6F8FA',height=62); foot.pack(fill='x'); foot.pack_propagate(False)
        b=tk.Frame(foot,bg='#F6F8FA'); b.pack(side='right',padx=20,pady=12)
        self._button(b,'취소',self.cancel,False).pack(side='left',padx=4); self._button(b,'확인',self.ok,True).pack(side='left',padx=4)
        self.protocol('WM_DELETE_WINDOW',self.cancel); ui.center_window(self,parent,760,610); self.deiconify(); self.grab_set(); self.focus_force(); parent.wait_window(self)
    def _button(self,p,text,cmd,primary): return tk.Button(p,text=text,command=cmd,width=12,bd=0,font=('Malgun Gothic',9,'bold'),bg=ui.BLUE if primary else '#E5EBF0',fg='white' if primary else ui.TEXT,pady=7,cursor='hand2')
    def ok(self): self.result=self.var.get(); self.destroy()
    def cancel(self): self.result=None; self.destroy()


class EnterpriseProblemDialog(tk.Toplevel):
    def __init__(self,parent,summary):
        super().__init__(parent); self.withdraw(); self.result=None; self.title('Issue DB 현상 입력'); self.configure(bg='white'); self.resizable(False,False); self.transient(parent)
        head=tk.Frame(self,bg=ui.NAVY,height=58); head.pack(fill='x'); head.pack_propagate(False); tk.Label(head,text='Issue DB 현상 입력',bg=ui.NAVY,fg='white',font=('Malgun Gothic',12,'bold')).pack(side='left',padx=22)
        body=tk.Frame(self,bg='white'); body.pack(fill='both',expand=True,padx=24,pady=18)
        tk.Label(body,text='Issue DB의 현상 항목에 입력할 범위를 선택해 주세요.',bg='white',fg=ui.TEXT,font=('Malgun Gothic',10,'bold')).pack(anchor='w')
        tk.Label(body,text='주간회의 현상은 선택과 관계없이 원문 전체를 유지합니다.',bg='white',fg=ui.MUTED,font=('Malgun Gothic',8)).pack(anchor='w',pady=(3,13))
        tk.Label(body,text='현상 요약 미리보기',bg='white',fg=ui.MUTED,font=('Malgun Gothic',8,'bold')).pack(anchor='w',pady=(0,5))
        box=tk.Text(body,height=12,wrap='word',bg='#F7F9FB',fg=ui.TEXT,relief='flat',highlightthickness=1,highlightbackground=ui.BORDER,font=('Malgun Gothic',9),padx=10,pady=8); box.pack(fill='both',expand=True); box.insert('1.0',summary or '(요약 내용 없음)'); box.configure(state='disabled')
        foot=tk.Frame(self,bg='#F6F8FA',height=66); foot.pack(fill='x'); foot.pack_propagate(False); b=tk.Frame(foot,bg='#F6F8FA'); b.pack(side='right',padx=20,pady=13)
        self._button(b,'취소',lambda:self.choose(None),False).pack(side='left',padx=4); self._button(b,'전체 내용',lambda:self.choose('full'),False).pack(side='left',padx=4); self._button(b,'현상 요약',lambda:self.choose('summary'),True).pack(side='left',padx=4)
        self.protocol('WM_DELETE_WINDOW',lambda:self.choose(None)); ui.center_window(self,parent,720,500); self.deiconify(); self.grab_set(); self.focus_force(); parent.wait_window(self)
    def _button(self,p,text,cmd,primary): return tk.Button(p,text=text,command=cmd,width=12,bd=0,font=('Malgun Gothic',9,'bold'),bg=ui.BLUE if primary else '#E5EBF0',fg='white' if primary else ui.TEXT,pady=7,cursor='hand2')
    def choose(self,v): self.result=v; self.destroy()


class EnterpriseStatusDialog(tk.Toplevel):
    def __init__(self,parent,d,judged,for_excel=True,for_weekly=True):
        super().__init__(parent); self.withdraw(); self.result=None; self.title('상태 판단 최종 확인'); self.configure(bg='white'); self.resizable(False,False); self.transient(parent)
        judged='close' if str(judged).lower()=='close' else 'open'
        weekly=weekly_status_from_choice(d,judged)
        self.db_var=tk.StringVar(value=judged)
        self.weekly_var=tk.StringVar(value=weekly)
        self.db_confirm=tk.BooleanVar(value=False)
        self.weekly_confirm=tk.BooleanVar(value=False)
        self.for_excel=bool(for_excel); self.for_weekly=bool(for_weekly)

        head=tk.Frame(self,bg=ui.NAVY,height=58); head.pack(fill='x'); head.pack_propagate(False)
        tk.Label(head,text='상태 판단 최종 확인',bg=ui.NAVY,fg='white',font=('Malgun Gothic',12,'bold')).pack(side='left',padx=22)
        body=tk.Frame(self,bg='white'); body.pack(fill='both',expand=True,padx=24,pady=16)
        tk.Label(body,text='주간회의와 Issue DB의 상태를 각각 확인해 주세요. 상태를 변경한 뒤에도 각 항목의 확인 체크가 필요합니다.',bg='white',fg=ui.TEXT,font=('Malgun Gothic',10,'bold'),wraplength=790,justify='left').pack(anchor='w',pady=(0,10))

        if self.for_weekly:
            self._weekly_card(body,d,weekly)
        if self.for_excel:
            self._db_card(body,d,judged)

        foot=tk.Frame(self,bg='#F6F8FA',height=70); foot.pack(fill='x'); foot.pack_propagate(False)
        b=tk.Frame(foot,bg='#F6F8FA'); b.pack(side='right',padx=20,pady=13)
        tk.Button(b,text='취소',command=self.cancel,width=12,bd=0,font=('Malgun Gothic',9,'bold'),bg='#E5EBF0',fg=ui.TEXT,pady=8,cursor='hand2').pack(side='left',padx=4)
        self.ok_btn=tk.Button(b,text='상태 확인 완료',command=self.ok,width=16,bd=0,font=('Malgun Gothic',9,'bold'),bg='#AAB7C2',fg='white',pady=8,cursor='arrow',state='disabled')
        self.ok_btn.pack(side='left',padx=4)

        if self.for_excel: self.db_confirm.trace_add('write',lambda *_:self._refresh_ok())
        if self.for_weekly: self.weekly_confirm.trace_add('write',lambda *_:self._refresh_ok())
        self.protocol('WM_DELETE_WINDOW',self.cancel)
        height=690 if self.for_excel and self.for_weekly else 480
        ui.center_window(self,parent,840,height); self.deiconify(); self.grab_set(); self.focus_force(); parent.wait_window(self)

    def _card_shell(self,parent,title,auto_status,reason):
        card=tk.Frame(parent,bg='#F8FAFC',highlightbackground='#D5E3EE',highlightthickness=1)
        card.pack(fill='x',pady=(0,10))
        tk.Label(card,text=title,bg='#F8FAFC',fg=ui.NAVY,font=('Malgun Gothic',10,'bold')).pack(anchor='w',padx=14,pady=(10,2))
        tk.Label(card,text=f'자동 판단  :  {auto_status}',bg='#F8FAFC',fg=ui.TEXT,font=('Malgun Gothic',9,'bold')).pack(anchor='w',padx=14,pady=(2,3))
        tk.Label(card,text='판단 근거',bg='#F8FAFC',fg=ui.MUTED,font=('Malgun Gothic',8,'bold')).pack(anchor='w',padx=14,pady=(5,1))
        tk.Label(card,text=reason,bg='#F8FAFC',fg='#4E6375',font=('Malgun Gothic',8),wraplength=760,justify='left').pack(anchor='w',padx=14,pady=(0,8))
        return card

    def _weekly_card(self,parent,d,auto_status):
        card=self._card_shell(parent,'주간회의 Signal',auto_status,weekly_status_reason(d,auto_status))
        row=tk.Frame(card,bg='#F8FAFC'); row.pack(fill='x',padx=14,pady=(0,7))
        for val in ('원인/개선 미확인','개선 검증중','개선 완료'):
            ttk.Radiobutton(row,text=val,variable=self.weekly_var,value=val,style='Mode.TRadiobutton',command=lambda:self.weekly_confirm.set(False)).pack(side='left',padx=(0,18))
        tk.Checkbutton(card,text='주간회의 상태를 이 선택으로 확인',variable=self.weekly_confirm,command=self._refresh_ok,bg='#F8FAFC',fg=ui.TEXT,activebackground='#F8FAFC',font=('Malgun Gothic',9,'bold')).pack(anchor='w',padx=12,pady=(0,10))

    def _db_card(self,parent,d,auto_status):
        card=self._card_shell(parent,'Issue DB 상태',auto_status,issue_db_status_reason(d,auto_status))
        row=tk.Frame(card,bg='#F8FAFC'); row.pack(fill='x',padx=14,pady=(0,7))
        ttk.Radiobutton(row,text='open  ·  진행/검증 필요',variable=self.db_var,value='open',style='Mode.TRadiobutton',command=lambda:self.db_confirm.set(False)).pack(side='left',padx=(0,24))
        ttk.Radiobutton(row,text='close  ·  개선 완료',variable=self.db_var,value='close',style='Mode.TRadiobutton',command=lambda:self.db_confirm.set(False)).pack(side='left')
        tk.Checkbutton(card,text='Issue DB 상태를 이 선택으로 확인',variable=self.db_confirm,command=self._refresh_ok,bg='#F8FAFC',fg=ui.TEXT,activebackground='#F8FAFC',font=('Malgun Gothic',9,'bold')).pack(anchor='w',padx=12,pady=(0,10))

    def _refresh_ok(self):
        ready=(not self.for_excel or self.db_confirm.get()) and (not self.for_weekly or self.weekly_confirm.get())
        if ready:
            self.ok_btn.configure(state='normal',bg=ui.BLUE,cursor='hand2')
        else:
            self.ok_btn.configure(state='disabled',bg='#AAB7C2',cursor='arrow')

    def ok(self):
        if not ((not self.for_excel or self.db_confirm.get()) and (not self.for_weekly or self.weekly_confirm.get())):
            return
        self.result={
            'issue_db': self.db_var.get() if self.for_excel else None,
            'weekly': self.weekly_var.get() if self.for_weekly else None,
        }
        self.destroy()

    def cancel(self): self.result=None; self.destroy()


class EnterpriseApp(legacy.FinalApp, _RootBase):
    def __init__(self):
        _RootBase.__init__(self); self.title('8D Issue Automation 1.0 | Pack 개발품질'); self.geometry('1180x850'); self.minsize(1080,780); self.configure(bg=ui.BG)
        self.vars={}; self.mode=tk.StringVar(value='existing'); self.status_var=tk.StringVar(value='READY  ·  8D 원본을 선택해 주세요.')
        self._configure_styles(); self._build_enterprise_ui(); self._center_main()
    def _center_main(self):
        self.update_idletasks(); w=1180; h=850; self.geometry(f'{w}x{h}+{max(0,(self.winfo_screenwidth()-w)//2)}+{max(0,(self.winfo_screenheight()-h)//2)}')
    def _build_enterprise_ui(self):
        head=tk.Frame(self,bg=ui.NAVY,height=96); head.pack(fill='x'); head.pack_propagate(False)
        tk.Label(head,text='8D Issue Automation',bg=ui.NAVY,fg='white',font=('Malgun Gothic',20,'bold')).place(x=32,y=18); tk.Label(head,text='Issue Management Productivity Tool  ·  v1.0',bg=ui.NAVY,fg='#C9D8E6',font=('Malgun Gothic',9)).place(x=34,y=58); tk.Label(head,text='PACK DEVELOPMENT QUALITY',bg=ui.NAVY,fg='#D7E4EF',font=('Segoe UI',8,'bold')).place(relx=1,x=-32,y=38,anchor='e')
        body=tk.Frame(self,bg=ui.BG); body.pack(fill='both',expand=True,padx=26,pady=18)
        mode=tk.Frame(body,bg='white',highlightbackground=ui.BORDER,highlightthickness=1); mode.pack(fill='x',pady=(0,12)); tk.Label(mode,text='이슈 처리 방식',bg='white',fg=ui.NAVY,font=('Malgun Gothic',10,'bold')).pack(side='left',padx=(18,24),pady=13)
        for text,value in [('기존 이슈 업데이트','existing'),('신규 이슈 등록','new')]: ttk.Radiobutton(mode,text=text,variable=self.mode,value=value,style='Mode.TRadiobutton').pack(side='left',padx=(0,20))
        tk.Label(mode,text='기존 이슈는 동일 이슈 행/페이지를 찾아 업데이트합니다.',bg='white',fg=ui.MUTED,font=('Malgun Gothic',8)).pack(side='right',padx=18)
        content=tk.Frame(body,bg=ui.BG); content.pack(fill='x'); left=tk.Frame(content,bg='white',highlightbackground=ui.BORDER,highlightthickness=1); left.pack(side='left',fill='both',expand=True,padx=(0,7)); right=tk.Frame(content,bg='white',highlightbackground=ui.BORDER,highlightthickness=1); right.pack(side='left',fill='both',expand=True,padx=(7,0))
        self._section_title(left,'01','입력 파일','Drag & Drop 지원'); filebody=tk.Frame(left,bg='white'); filebody.pack(fill='x',padx=16,pady=(0,14))
        for key in ('ppt8d','pptweekly','xlsx'): self.vars[key]=tk.StringVar()
        self.dropzones={}; specs=[('8D 원본 PPT','ppt8d',('.pptx',),True),('주간회의 PPT','pptweekly',('.pptx',),False),('Issue DB Excel','xlsx',('.xlsx',),False)]
        for title,key,exts,req in specs:
            z=ui.DropZone(filebody,title,key,self.vars[key],exts,self.pick,self._refresh_target,req); z.pack(fill='x',pady=5); self.dropzones[key]=z
        tk.Label(filebody,text=f"Drag & Drop: {'사용 가능' if ui.DND_AVAILABLE else '미사용 · [찾기] 버튼 사용'}  |  8D는 필수, 나머지는 선택 입력",bg='white',fg=ui.MUTED,font=('Malgun Gothic',8)).pack(anchor='w',pady=(4,0))
        self._section_title(right,'02','담당 및 분류 정보','표준 입력'); fields=tk.Frame(right,bg='white'); fields.pack(fill='x',padx=18,pady=(0,10))
        for key in ('team','task_name','owner','sample','plm_no'): self.vars[key]=tk.StringVar()
        self._enterprise_entry(fields,'담당팀','team','예: Pack개발품질1팀'); self._enterprise_entry(fields,'고객사 과제명','task_name','담당팀 연계 과제 선택'); self._enterprise_entry(fields,'담당자','owner','예: 홍길동'); self._enterprise_entry(fields,'발생 샘플','sample','예: DUT3 / Sample No.'); self._enterprise_entry(fields,'PMS/PLM 이슈번호','plm_no','기존값이 있을 때 입력'); tk.Frame(fields,bg=ui.BORDER,height=1).pack(fill='x',pady=10)
        self.vars['form_factor']=tk.StringVar(value='파우치형'); self.vars['product_type']=tk.StringVar(value='EV Pack'); self.vars['occurrence_site']=tk.StringVar(value=legacy.OCCURRENCE_SITES[0]); self.vars['stage']=tk.StringVar(value='DV')
        self._enterprise_combo(fields,'폼팩터','form_factor',legacy.FORM_FACTORS); self._enterprise_combo(fields,'제품 타입','product_type',legacy.PRODUCT_TYPES); self._enterprise_combo(fields,'발생처','occurrence_site',legacy.OCCURRENCE_SITES); self._enterprise_combo(fields,'개발 단계','stage',legacy.STAGES)
        target=tk.Frame(body,bg='white',highlightbackground=ui.BORDER,highlightthickness=1); target.pack(fill='x',pady=12); tk.Label(target,text='UPDATE TARGET',bg='white',fg=ui.MUTED,font=('Segoe UI',8,'bold')).pack(side='left',padx=(16,12),pady=12); self.target_label=tk.Label(target,text='',bg='white',fg=ui.NAVY,font=('Malgun Gothic',9,'bold')); self.target_label.pack(side='left'); self.target_weekly=tk.Label(target,text='',bg='#EDF1F4',font=('Segoe UI',8,'bold'),padx=9,pady=4); self.target_weekly.pack(side='right',padx=(5,16)); self.target_excel=tk.Label(target,text='',bg='#EDF1F4',font=('Segoe UI',8,'bold'),padx=9,pady=4); self.target_excel.pack(side='right')
        actions=tk.Frame(body,bg=ui.BG); actions.pack(fill='x',pady=(0,10)); tk.Button(actions,text='8D 내용 미리보기',command=self.preview,bg='#E5EBF0',fg=ui.NAVY,bd=0,font=('Malgun Gothic',9,'bold'),padx=18,pady=9,cursor='hand2').pack(side='left'); tk.Button(actions,text='자동 업데이트 실행  →',command=self.run,bg=ui.BLUE,fg='white',bd=0,font=('Malgun Gothic',10,'bold'),padx=24,pady=10,cursor='hand2').pack(side='right')
        result=tk.Frame(body,bg='white',highlightbackground=ui.BORDER,highlightthickness=1); result.pack(fill='both',expand=True); rt=tk.Frame(result,bg='white'); rt.pack(fill='x',padx=16,pady=(12,7)); tk.Label(rt,text='실행 결과',bg='white',fg=ui.NAVY,font=('Malgun Gothic',10,'bold')).pack(side='left'); tk.Label(rt,textvariable=self.status_var,bg='#E9F1F8',fg='#315A7D',font=('Malgun Gothic',8),padx=9,pady=4).pack(side='right'); self.log=tk.Text(result,height=7,wrap='word',font=('Consolas',9),bg='#FBFCFD',fg='#283A49',relief='flat',highlightthickness=1,highlightbackground='#E1E7EC',padx=10,pady=8); self.log.pack(fill='both',expand=True,padx=16,pady=(0,14)); self.log.insert('end','8D 원본을 Drag & Drop하거나 [찾기]로 선택해 주세요.\n')
        foot=tk.Frame(self,bg='#E8EEF3',height=30); foot.pack(fill='x'); foot.pack_propagate(False); tk.Label(foot,text='8D Issue Automation 1.0   ·   원본 파일은 직접 덮어쓰지 않습니다.',bg='#E8EEF3',fg='#607180',font=('Malgun Gothic',8)).pack(side='left',padx=26); self._refresh_target()
    def _section_title(self,parent,no,title,sub):
        f=tk.Frame(parent,bg='white'); f.pack(fill='x',padx=18,pady=(15,10)); tk.Label(f,text=no,bg=ui.NAVY,fg='white',font=('Segoe UI',8,'bold'),width=3,pady=3).pack(side='left'); tk.Label(f,text=title,bg='white',fg=ui.NAVY,font=('Malgun Gothic',11,'bold')).pack(side='left',padx=9); tk.Label(f,text=sub,bg='white',fg=ui.MUTED,font=('Malgun Gothic',8)).pack(side='right')
    def _enterprise_entry(self,parent,label,key,hint):
        row=tk.Frame(parent,bg='white'); row.pack(fill='x',pady=4); tk.Label(row,text=label,bg='white',fg=ui.TEXT,font=('Malgun Gothic',9),width=17,anchor='w').pack(side='left'); ttk.Entry(row,textvariable=self.vars[key]).pack(side='left',fill='x',expand=True); tk.Label(parent,text=hint,bg='white',fg='#98A3AD',font=('Malgun Gothic',7)).pack(anchor='e')
    def _enterprise_combo(self,parent,label,key,values):
        row=tk.Frame(parent,bg='white'); row.pack(fill='x',pady=5); tk.Label(row,text=label,bg='white',fg=ui.TEXT,font=('Malgun Gothic',9),width=17,anchor='w').pack(side='left'); ttk.Combobox(row,textvariable=self.vars[key],values=values,state='readonly').pack(side='left',fill='x',expand=True)
    def pick(self,k,desc=None,pattern=None):
        types=[('Excel (*.xlsx)','*.xlsx')] if k=='xlsx' else [('PowerPoint (*.pptx)','*.pptx')]; f=filedialog.askopenfilename(title='파일 선택',filetypes=types+[('모든 파일','*.*')],parent=self)
        if f: self.vars[k].set(f); self._refresh_target()
    def _refresh_target(self):
        if not hasattr(self,'target_label'): return
        ppt8d=(self.vars.get('ppt8d').get().strip() if self.vars.get('ppt8d') else '')
        weekly=bool(self.vars.get('pptweekly') and self.vars['pptweekly'].get().strip()); excel=bool(self.vars.get('xlsx') and self.vars['xlsx'].get().strip()); text='주간회의 PPT + Issue DB Excel' if weekly and excel else '주간회의 PPT만 업데이트' if weekly else 'Issue DB Excel만 업데이트' if excel else '주간회의 또는 Issue DB를 선택해 주세요.'; self.target_label.configure(text=text); self.target_weekly.configure(text='WEEKLY  ON' if weekly else 'WEEKLY  OFF',bg='#E7F4ED' if weekly else '#EDF1F4',fg=ui.GREEN if weekly else ui.MUTED); self.target_excel.configure(text='ISSUE DB  ON' if excel else 'ISSUE DB  OFF',bg='#E7F4ED' if excel else '#EDF1F4',fg=ui.GREEN if excel else ui.MUTED)
        current=self.status_var.get().strip()
        refreshed=target_ready_status(ppt8d,current)
        if refreshed!=current:self.status_var.set(refreshed)

    def preview(self):
        g=self.gui()
        if not g.get('ppt8d'): return ui.warning(self,'입력 확인','8D 원본 PPT를 선택해 주세요.')
        try:
            self.status_var.set('ANALYZING  ·  8D 내용을 추출하고 있습니다...'); self.update_idletasks(); d=base.extract(g['ppt8d']); labels=[('이슈명','issue_name'),('과제명','task_name'),('고객사','customer'),('발생일자','occurrence_date'),('2D 현상','problem'),('3D 임시조치','temporary_action'),('4D 발생원인','cause_4d'),('4D 유출원인','leak_cause'),('5D 개선대책','action_5d'),('6D 효과검증','verification_6d')]; self.log.delete('1.0','end'); self.log.insert('end','[ 8D 추출 결과 ]\n'+'─'*72+'\n')
            for label,key in labels: self.log.insert('end',f'{label:<12} : {N(d.get(key)) or "-"}\n')
            self.status_var.set('READY  ·  8D 추출 완료')
        except Exception as e: self.status_var.set('ERROR  ·  추출 실패'); ui.error(self,'미리보기 오류',repr(e))

    def run(self):
        g=self.gui(); ppt8d=g.get('ppt8d',''); weekly=g.get('pptweekly',''); xlsx=g.get('xlsx','')
        if not ppt8d or not os.path.exists(ppt8d): return ui.warning(self,'입력 확인','8D 원본 PPT는 반드시 선택해 주세요.')
        if not weekly and not xlsx: return ui.warning(self,'출력 확인','주간회의 PPT 또는 Issue DB Excel 중 하나 이상을 선택해 주세요.')
        if weekly and not os.path.exists(weekly): return ui.warning(self,'입력 확인','선택한 주간회의 PPT 파일을 찾을 수 없습니다.')
        if xlsx and not os.path.exists(xlsx): return ui.warning(self,'입력 확인','선택한 Issue DB Excel 파일을 찾을 수 없습니다.')
        if g.get('occurrence_site') not in legacy.OCCURRENCE_SITES: return ui.warning(self,'입력 확인','발생처를 선택해 주세요.')
        try:
            self.status_var.set('RUNNING  ·  8D 원본 분석 중...'); self.update_idletasks(); d=base.extract(ppt8d); selected_task=g.get('task_name','').strip();
            if selected_task: d['task_name']=selected_task
            mode=self.mode.get(); do_weekly=bool(weekly); do_excel=bool(xlsx); excel_row=None
            if do_excel:
                self.status_var.set('RUNNING  ·  Issue DB 대상 확인 중...'); self.update_idletasks(); wb=load_workbook(xlsx); ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active; excel_row,_=base.find(ws,d,g)
                if mode=='existing' and not excel_row:
                    if do_weekly:
                        if ui.ask_yes_no(self,'기존 이슈 미확인','Issue DB에서 동일 이슈를 찾지 못했습니다.\n\nExcel은 변경하지 않고 주간회의 PPT만 업데이트할까요?'): do_excel=False
                        else: self.status_var.set('READY  ·  사용자가 실행을 취소했습니다.'); return
                    else: self.status_var.set('READY  ·  기존 이슈 미확인'); return ui.warning(self,'기존 이슈 미확인','Issue DB에서 동일 이슈를 찾지 못했습니다.\n기존 이슈는 임의로 신규 행을 만들지 않습니다.')
                if mode=='new' and excel_row and not ui.ask_yes_no(self,'중복 가능성',f'유사 이슈(row {excel_row})가 있습니다.\n그래도 신규 이슈로 추가할까요?'): self.status_var.set('READY  ·  사용자가 실행을 취소했습니다.'); return
            od=None
            if do_weekly:
                self.status_var.set('WAITING  ·  이슈기인 확인 필요'); self.update_idletasks(); od=EnterpriseOriginDialog(self,d)
                if od.result is None: self.status_var.set('READY  ·  사용자가 실행을 취소했습니다.'); return
                g['_issue_origin_selected']=od.result
            pd_result=None; db_status=None; weekly_status=None
            if do_excel:
                summary=legacy.step7._db_summary(d); self.status_var.set('WAITING  ·  Issue DB 현상 입력 확인 필요'); self.update_idletasks(); pd=EnterpriseProblemDialog(self,summary)
                if pd.result is None: self.status_var.set('READY  ·  사용자가 실행을 취소했습니다.'); return
                pd_result=pd.result; g['_db_problem_selected']=summary if pd.result=='summary' else N(d.get('problem'))
            if do_excel or do_weekly:
                judged,_reason=legacy.step9._judge_issue_status(d)
                self.status_var.set('WAITING  ·  상태 판단 최종 확인 필요'); self.update_idletasks()
                sd=EnterpriseStatusDialog(self,d,judged,do_excel,do_weekly)
                if sd.result is None: self.status_var.set('READY  ·  사용자가 실행을 취소했습니다.'); return
                db_status=sd.result.get('issue_db')
                weekly_status=sd.result.get('weekly')
                if do_excel: g['_issue_status_selected']=db_status
                if do_weekly: g['_weekly_status_selected']=weekly_status
            anchor=xlsx if do_excel else weekly; out=Path(anchor).parent/'자동화_결과'; out.mkdir(exist_ok=True); results=[]
            if do_excel: self.status_var.set('RUNNING  ·  Issue DB 업데이트 중...'); self.update_idletasks(); xo=out/(Path(xlsx).stem+'_업데이트.xlsx'); a,_=base.update_excel(xlsx,xo,d,g,new=(mode=='new')); results.append(a)
            if do_weekly: self.status_var.set('RUNNING  ·  주간회의 PPT 업데이트 중...'); self.update_idletasks(); po=out/(Path(weekly).stem+'_업데이트.pptx'); b,_=base.weekly(weekly,po,d,g,mode); results.append(b)
            self.log.delete('1.0','end'); targets=[]
            if do_excel: targets.append('Issue DB Excel')
            if do_weekly: targets.append('주간회의 PPT')
            lines=['[ 업데이트 완료 ]','─'*72,f'이슈 구분       : {"신규 이슈" if mode=="new" else "기존 이슈"}',f'업데이트 대상   : {" + ".join(targets)}'];
            if do_weekly and od: lines.append(f'이슈기인        : {od.result}')
            if do_weekly: lines.append(f'주간회의 상태   : {weekly_status}')
            if do_excel: lines.extend([f'Issue DB 현상   : {"현상 요약" if pd_result=="summary" else "전체 내용"}',f'Issue DB 상태   : {db_status}'])
            lines.extend(['']+results+['',f'결과 폴더       : {out}']); self.log.insert('end','\n'.join(lines)); self.status_var.set('COMPLETE  ·  선택한 자료의 업데이트가 완료되었습니다.'); ui.info(self,'업데이트 완료','선택한 자료만 업데이트했습니다.\n\n'+'\n'.join('✓ '+x for x in targets)+f'\n\n결과 폴더\n{out}')
        except Exception as e: self.status_var.set('ERROR  ·  실행 중 오류가 발생했습니다.'); ui.error(self,'실행 오류',repr(e))


if __name__=='__main__': EnterpriseApp().mainloop()
