import os
import tkinter as tk
from pathlib import Path
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook

import main_recovery_step14_fix2 as core
import main_recovery_step9 as step9
import main_recovery_step8 as step8
import main_recovery_step7 as step7

base=core.base
N=core.N
OCCURRENCE_SITES=step9.OCCURRENCE_SITES

FORM_FACTORS=('원통형','파우치형')
PRODUCT_TYPES=('ESS BPU','ESS Link','EV Cell-Unit','EV Module','EV Pack','IT Pack','LEV Pack','PHEV Pack')
STAGES=('CV','DV','PD','MP')


class FinalApp(core.RecoveryStep14Fix5App):
    """Final internal-deployment UI. Processing logic remains on the validated core."""
    def __init__(self):
        # Do not build the legacy packed UI; only initialise Tk and our state.
        tk.Tk.__init__(self)
        self.title('8D Issue Automation 1.0 | Pack 개발품질')
        self.geometry('1040x820')
        self.minsize(960,760)
        self.configure(bg='#F3F6F9')
        self.vars={}
        self.mode=tk.StringVar(value='existing')
        self.status_var=tk.StringVar(value='Ready  |  8D 원본을 선택해 주세요.')
        self._configure_styles()
        self._build_ui()

    def _configure_styles(self):
        s=ttk.Style(self)
        try: s.theme_use('clam')
        except Exception: pass
        s.configure('App.TFrame',background='#F3F6F9')
        s.configure('Card.TFrame',background='#FFFFFF')
        s.configure('Header.TLabel',background='#173A5E',foreground='#FFFFFF',font=('Malgun Gothic',19,'bold'))
        s.configure('SubHeader.TLabel',background='#173A5E',foreground='#C9D8E6',font=('Malgun Gothic',9))
        s.configure('Section.TLabel',background='#FFFFFF',foreground='#173A5E',font=('Malgun Gothic',11,'bold'))
        s.configure('Field.TLabel',background='#FFFFFF',foreground='#394B5A',font=('Malgun Gothic',9))
        s.configure('Hint.TLabel',background='#FFFFFF',foreground='#7A8793',font=('Malgun Gothic',8))
        s.configure('Status.TLabel',background='#E9F1F8',foreground='#315A7D',font=('Malgun Gothic',9))
        s.configure('Primary.TButton',font=('Malgun Gothic',10,'bold'),padding=(18,10),foreground='#FFFFFF',background='#1769AA')
        s.map('Primary.TButton',background=[('active','#12598F')])
        s.configure('Secondary.TButton',font=('Malgun Gothic',9),padding=(12,8),foreground='#173A5E',background='#E8EFF5')
        s.map('Secondary.TButton',background=[('active','#DCE8F1')])
        s.configure('Browse.TButton',font=('Malgun Gothic',8,'bold'),padding=(10,5))
        s.configure('Mode.TRadiobutton',background='#FFFFFF',foreground='#263746',font=('Malgun Gothic',9))
        s.configure('TEntry',padding=5)
        s.configure('TCombobox',padding=4)

    def _build_ui(self):
        # Header
        head=tk.Frame(self,bg='#173A5E',height=88)
        head.pack(fill='x')
        head.pack_propagate(False)
        ttk.Label(head,text='8D Issue Automation',style='Header.TLabel').place(x=30,y=18)
        ttk.Label(head,text='Issue DB · 주간회의 자료 자동 업데이트  |  v1.0',style='SubHeader.TLabel').place(x=32,y=55)
        tk.Label(head,text='Pack 개발품질',bg='#173A5E',fg='#D7E4EF',font=('Malgun Gothic',9,'bold')).place(relx=1.0,x=-30,y=33,anchor='e')

        body=ttk.Frame(self,style='App.TFrame',padding=(24,18,24,12))
        body.pack(fill='both',expand=True)

        # Mode card
        mode_card=ttk.Frame(body,style='Card.TFrame',padding=(18,12))
        mode_card.pack(fill='x',pady=(0,10))
        ttk.Label(mode_card,text='이슈 처리 방식',style='Section.TLabel').pack(side='left',padx=(0,25))
        ttk.Radiobutton(mode_card,text='기존 이슈 업데이트',variable=self.mode,value='existing',style='Mode.TRadiobutton').pack(side='left',padx=(0,22))
        ttk.Radiobutton(mode_card,text='신규 이슈 등록',variable=self.mode,value='new',style='Mode.TRadiobutton').pack(side='left')
        ttk.Label(mode_card,text='기존 이슈는 동일 이슈를 찾아 업데이트합니다.',style='Hint.TLabel').pack(side='right')

        # Main two-column area
        cols=ttk.Frame(body,style='App.TFrame')
        cols.pack(fill='x')
        left=ttk.Frame(cols,style='Card.TFrame',padding=(18,14))
        right=ttk.Frame(cols,style='Card.TFrame',padding=(18,14))
        left.pack(side='left',fill='both',expand=True,padx=(0,6))
        right.pack(side='left',fill='both',expand=True,padx=(6,0))

        ttk.Label(left,text='01  입력 파일',style='Section.TLabel').pack(anchor='w',pady=(0,10))
        self._file_row(left,'8D 원본 PPT  *','ppt8d','PowerPoint (*.pptx)','*.pptx','필수')
        self._file_row(left,'주간회의 PPT','pptweekly','PowerPoint (*.pptx)','*.pptx','선택')
        self._file_row(left,'Issue DB Excel','xlsx','Excel (*.xlsx)','*.xlsx','선택')
        ttk.Label(left,text='* 8D 원본은 필수입니다. 주간회의/Issue DB는 필요한 자료만 선택하면 됩니다.',style='Hint.TLabel').pack(anchor='w',pady=(4,0))

        ttk.Separator(left).pack(fill='x',pady=14)
        ttk.Label(left,text='02  담당 정보',style='Section.TLabel').pack(anchor='w',pady=(0,9))
        self._entry_row(left,'담당팀','team','예: Pack개발품질1팀')
        self._entry_row(left,'담당자','owner','예: 홍길동')
        self._entry_row(left,'발생 샘플','sample','예: DUT3 / Sample No.')
        self._entry_row(left,'PMS/PLM 이슈번호','plm_no','기존값이 있을 때 입력')

        ttk.Label(right,text='03  이슈 분류',style='Section.TLabel').pack(anchor='w',pady=(0,10))
        self._combo_row(right,'폼팩터','form_factor',FORM_FACTORS,'파우치형')
        self._combo_row(right,'제품 타입','product_type',PRODUCT_TYPES,'EV Pack')
        self._combo_row(right,'발생처','occurrence_site',OCCURRENCE_SITES,OCCURRENCE_SITES[0])
        self._combo_row(right,'개발 단계','stage',STAGES,'DV')

        ttk.Separator(right).pack(fill='x',pady=14)
        ttk.Label(right,text='업데이트 대상',style='Section.TLabel').pack(anchor='w',pady=(0,8))
        self.target_hint=ttk.Label(right,text='선택한 파일에 따라 자동 결정됩니다.',style='Hint.TLabel')
        self.target_hint.pack(anchor='w')
        target_box=tk.Frame(right,bg='#F5F8FB',highlightbackground='#D9E3EC',highlightthickness=1)
        target_box.pack(fill='x',pady=(9,0))
        self.target_label=tk.Label(target_box,text='8D 원본 + [주간회의 / Issue DB 선택]',bg='#F5F8FB',fg='#315A7D',font=('Malgun Gothic',9,'bold'),anchor='w',padx=12,pady=11)
        self.target_label.pack(fill='x')

        # Actions
        actions=ttk.Frame(body,style='App.TFrame')
        actions.pack(fill='x',pady=(12,8))
        ttk.Button(actions,text='8D 내용 미리보기',style='Secondary.TButton',command=self.preview).pack(side='left')
        ttk.Button(actions,text='자동 업데이트 실행',style='Primary.TButton',command=self.run).pack(side='right')

        # Log/result card
        log_card=ttk.Frame(body,style='Card.TFrame',padding=(16,12))
        log_card.pack(fill='both',expand=True)
        top=ttk.Frame(log_card,style='Card.TFrame')
        top.pack(fill='x',pady=(0,7))
        ttk.Label(top,text='실행 결과',style='Section.TLabel').pack(side='left')
        ttk.Label(top,textvariable=self.status_var,style='Status.TLabel',padding=(8,3)).pack(side='right')
        self.log=tk.Text(log_card,height=8,wrap='word',font=('Consolas',9),bg='#FBFCFD',fg='#283A49',relief='flat',highlightthickness=1,highlightbackground='#DDE5EC',padx=10,pady=8)
        self.log.pack(fill='both',expand=True)
        self.log.insert('end','8D 원본을 선택한 뒤 [8D 내용 미리보기]로 추출 결과를 확인할 수 있습니다.\n')

        foot=tk.Frame(self,bg='#E8EEF3',height=28)
        foot.pack(fill='x')
        foot.pack_propagate(False)
        tk.Label(foot,text='8D Issue Automation 1.0   |   원본 파일은 직접 덮어쓰지 않습니다.',bg='#E8EEF3',fg='#607180',font=('Malgun Gothic',8)).pack(side='left',padx=24)

    def _file_row(self,parent,label,key,desc,pattern,badge):
        self.vars[key]=tk.StringVar()
        row=ttk.Frame(parent,style='Card.TFrame')
        row.pack(fill='x',pady=4)
        top=ttk.Frame(row,style='Card.TFrame'); top.pack(fill='x')
        ttk.Label(top,text=label,style='Field.TLabel').pack(side='left')
        tk.Label(top,text=badge,bg='#E8F1F8' if badge=='선택' else '#E7F4ED',fg='#3E6C8E' if badge=='선택' else '#2F7A55',font=('Malgun Gothic',7,'bold'),padx=6,pady=1).pack(side='right')
        line=ttk.Frame(row,style='Card.TFrame'); line.pack(fill='x',pady=(3,0))
        ttk.Entry(line,textvariable=self.vars[key]).pack(side='left',fill='x',expand=True)
        ttk.Button(line,text='찾기',style='Browse.TButton',command=lambda k=key,d=desc,p=pattern:self.pick(k,d,p)).pack(side='left',padx=(6,0))
        self.vars[key].trace_add('write',lambda *_:self._refresh_target())

    def _entry_row(self,parent,label,key,hint):
        self.vars[key]=tk.StringVar()
        row=ttk.Frame(parent,style='Card.TFrame'); row.pack(fill='x',pady=4)
        ttk.Label(row,text=label,style='Field.TLabel',width=18).pack(side='left')
        ttk.Entry(row,textvariable=self.vars[key]).pack(side='left',fill='x',expand=True)
        ttk.Label(parent,text=hint,style='Hint.TLabel').pack(anchor='e',pady=(0,1))

    def _combo_row(self,parent,label,key,values,default):
        self.vars[key]=tk.StringVar(value=default)
        row=ttk.Frame(parent,style='Card.TFrame'); row.pack(fill='x',pady=6)
        ttk.Label(row,text=label,style='Field.TLabel',width=18).pack(side='left')
        ttk.Combobox(row,textvariable=self.vars[key],values=values,state='readonly').pack(side='left',fill='x',expand=True)

    def pick(self,k,desc=None,pattern=None):
        if desc is None:
            desc='Excel (*.xlsx)' if k=='xlsx' else 'PowerPoint (*.pptx)'
            pattern='*.xlsx' if k=='xlsx' else '*.pptx'
        f=filedialog.askopenfilename(title='파일 선택',filetypes=[(desc,pattern),('모든 파일','*.*')])
        if f:self.vars[k].set(f)

    def _refresh_target(self):
        if not hasattr(self,'target_label'): return
        weekly=bool(self.vars.get('pptweekly') and self.vars['pptweekly'].get().strip())
        excel=bool(self.vars.get('xlsx') and self.vars['xlsx'].get().strip())
        if weekly and excel: text='업데이트 대상  |  주간회의 PPT + Issue DB Excel'
        elif weekly: text='업데이트 대상  |  주간회의 PPT만'
        elif excel: text='업데이트 대상  |  Issue DB Excel만'
        else: text='업데이트 대상  |  주간회의 또는 Issue DB를 선택해 주세요.'
        self.target_label.configure(text=text)

    def gui(self):
        return {k:v.get().strip() for k,v in self.vars.items()}

    def preview(self):
        g=self.gui()
        if not g.get('ppt8d'):
            return messagebox.showwarning('입력 확인','8D 원본 PPT를 선택해 주세요.',parent=self)
        try:
            self.status_var.set('Analyzing  |  8D 내용을 추출하고 있습니다...')
            self.update_idletasks()
            d=base.extract(g['ppt8d'])
            labels=[
                ('이슈명','issue_name'),('과제명','task_name'),('고객사','customer'),('발생일자','occurrence_date'),
                ('현상','problem'),('3D 임시조치','temporary_action'),('4D 발생원인','cause_4d'),('4D 유출원인','leak_cause'),
                ('5D 개선대책','action_5d'),('6D 효과검증','verification_6d')]
            self.log.delete('1.0','end')
            self.log.insert('end','[ 8D 추출 결과 ]\n'+'─'*72+'\n')
            for label,key in labels:
                val=N(d.get(key)) or '-'
                self.log.insert('end',f'{label:<12} : {val}\n')
            self.status_var.set('Ready  |  8D 추출 완료')
        except Exception as e:
            self.status_var.set('Error  |  추출 실패')
            messagebox.showerror('미리보기 오류',repr(e),parent=self)

    def run(self):
        g=self.gui()
        ppt8d=g.get('ppt8d',''); weekly=g.get('pptweekly',''); xlsx=g.get('xlsx','')
        if not ppt8d or not os.path.exists(ppt8d):
            return messagebox.showwarning('입력 확인','8D 원본 PPT는 반드시 선택해 주세요.',parent=self)
        if not weekly and not xlsx:
            return messagebox.showwarning('출력 확인','주간회의 PPT 또는 이슈 DB Excel 중 하나 이상을 선택해 주세요.',parent=self)
        if weekly and not os.path.exists(weekly):
            return messagebox.showwarning('입력 확인','선택한 주간회의 PPT 파일을 찾을 수 없습니다.',parent=self)
        if xlsx and not os.path.exists(xlsx):
            return messagebox.showwarning('입력 확인','선택한 이슈 DB Excel 파일을 찾을 수 없습니다.',parent=self)
        if g.get('occurrence_site') not in OCCURRENCE_SITES:
            return messagebox.showwarning('입력 확인','발생처를 선택해 주세요.',parent=self)
        try:
            self.status_var.set('Running  |  자동 업데이트를 진행하고 있습니다...'); self.update_idletasks()
            d=base.extract(ppt8d); mode=self.mode.get(); do_weekly=bool(weekly); do_excel=bool(xlsx)
            excel_row=None
            if do_excel:
                wb=load_workbook(xlsx); ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active
                excel_row,_=base.find(ws,d,g)
                if mode=='existing' and not excel_row:
                    if do_weekly:
                        if messagebox.askyesno('기존 이슈 미확인','Issue DB에서 동일 이슈를 찾지 못했습니다.\n\nExcel은 변경하지 않고 주간회의 PPT만 업데이트할까요?',parent=self): do_excel=False
                        else: self.status_var.set('Ready  |  사용자가 실행을 취소했습니다.'); return
                    else:
                        self.status_var.set('Ready  |  기존 이슈 미확인')
                        return messagebox.showwarning('기존 이슈 미확인','Issue DB에서 동일 이슈를 찾지 못했습니다.\n기존 이슈는 임의로 신규 행을 만들지 않습니다.',parent=self)
                if mode=='new' and excel_row and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {excel_row})가 있습니다. 그래도 신규 추가할까요?',parent=self):
                    self.status_var.set('Ready  |  사용자가 실행을 취소했습니다.'); return
            if do_weekly:
                od=step8.OriginDialog(self,d)
                if od.result is None: self.status_var.set('Ready  |  사용자가 실행을 취소했습니다.'); return
                g['_issue_origin_selected']=od.result
            else: od=None
            pd_result=None; final_status=None
            if do_excel:
                summary=step7._db_summary(d); pd=step7.ProblemChoiceDialog(self,summary)
                if pd.result is None: self.status_var.set('Ready  |  사용자가 실행을 취소했습니다.'); return
                pd_result=pd.result; g['_db_problem_selected']=summary if pd.result=='summary' else N(d.get('problem'))
                judged,reason=step9._judge_issue_status(d); final_status=judged
                if judged=='close':
                    sixd=reason or '(6D 내용 없음)'
                    prompt='6D 내용 기준으로 close로 판단되었습니다.\n\n판단 이유(6D 원문)\n────────────────────\n'+sixd+'\n────────────────────\n\n이슈 상태를 close로 처리하시겠습니까?\n아니오를 선택하면 open으로 처리합니다.'
                    if not messagebox.askyesno('이슈 상태 확인',prompt,parent=self): final_status='open'
                g['_issue_status_selected']=final_status
            anchor=xlsx if do_excel else weekly; out=Path(anchor).parent/'자동화_결과'; out.mkdir(exist_ok=True)
            results=[]
            if do_excel:
                xo=out/(Path(xlsx).stem+'_업데이트.xlsx'); a,_=base.update_excel(xlsx,xo,d,g,new=(mode=='new')); results.append(a)
            if do_weekly:
                po=out/(Path(weekly).stem+'_업데이트.pptx'); b,_=base.weekly(weekly,po,d,g,mode); results.append(b)
            self.log.delete('1.0','end'); targets=[]
            if do_excel: targets.append('Issue DB Excel')
            if do_weekly: targets.append('주간회의 PPT')
            lines=['[ 업데이트 완료 ]','─'*72,f'이슈 구분       : {"신규 이슈" if mode=="new" else "기존 이슈"}',f'업데이트 대상   : {" + ".join(targets)}']
            if do_weekly and od: lines.append(f'이슈기인        : {od.result}')
            if do_excel: lines.extend([f'Issue DB 현상   : {"현상 요약" if pd_result=="summary" else "전체 내용"}',f'Issue DB 상태   : {final_status}'])
            lines.extend(['']+results+['',f'결과 폴더       : {out}']); self.log.insert('end','\n'.join(lines))
            self.status_var.set('Complete  |  선택한 자료의 업데이트가 완료되었습니다.')
            messagebox.showinfo('업데이트 완료','선택한 자료만 업데이트했습니다.\n\n'+'\n'.join('✓ '+x for x in targets)+f'\n\n결과 폴더\n{out}',parent=self)
        except Exception as e:
            self.status_var.set('Error  |  실행 중 오류가 발생했습니다.')
            messagebox.showerror('실행 오류',repr(e),parent=self)


if __name__=='__main__':
    FinalApp().mainloop()
