# 8D Issue Automation v2.8
# This file is the corrected v2.8 implementation.
# It is intentionally self-contained so run.bat can execute it directly.

import os, re, copy, datetime, io
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Pt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

STATUS={'complete':'개선 완료','verify':'개선 검증중','unknown':'원인/개선 미확인'}
SIGNAL={'개선 완료':(0,176,80),'개선 검증중':(255,192,0),'원인/개선 미확인':(255,0,0)}
STD=['D-FMEA','P-FMEA','Control plan','작업 표준','제품규격']
ALIASES={'2d':['2D 현상','2D현상','문제 현황','문제현황','불량 현상','불량현상'],'3d':['3D 임시대책','3D 임시 대책','3D임시대책','임시 조치','임시조치','임시 대책','임시대책','고객 대응','고객대응'],'4d_cause':['4D 원인분석','4D 원인 분석','4D원인분석','발생 원인','발생원인','원인 분석','원인분석'],'4d_leak':['유출 원인','유출원인','유출 사유','유출사유'],'4d_system':['시스템 원인','시스템원인','IT 시스템','IT시스템'],'5d':['5D 개선대책','5D 개선 대책','5D개선대책','개선 대책','개선대책','영구 개선','영구개선','개선 사항','개선사항'],'6d':['6D 유효성점검','6D 유효성 점검','6D유효성점검','효과 검증','효과검증','유효성 검증','유효성검증','유효성 점검','유효성점검','효과성 검증','효과성검증'],'7d':['7D 수평전개','7D 수평 전개','7D수평전개','수평 전개','수평전개','재발 방지','재발방지','재발/방지','재발 방지/수평전개','재발/방지/수평전개']}
EXACT={'task_name':['과제명','과제 명'],'model':['Model','모델'],'customer':['Customer','고객사'],'occurrence_site':['발생 Site','발생Site','발생처'],'occurrence_date':['발생 일자','발생일자'],'receipt_date':['접수 일자','접수일자'],'lot_no':['Lot no.','Lot No.','Lot . No.'],'packer':['Packer'],'line':['발생 Line','발생 호기','호기 정보'],'defect_rate':['불량률'],'defect_qty':['불량 수량'],'history':['발생 이력'],'issue_grade':['이슈 등급'],'problem':['불량 현상','문제 현황','문제/현황'],'temporary_action':['임시 조치','임시조치','임시 대책','임시대책'],'customer_response':['고객 대응','고객대응'],'cause_4d':['발생 원인','발생원인'],'leak_cause':['유출 원인','유출원인'],'system_cause':['시스템 원인','시스템원인','IT 시스템','IT시스템'],'action_5d':['개선 대책','개선대책','개선/대책','개선 사항','개선사항'],'verification_6d':['효과 검증','효과검증','유효성 검증','유효성검증','유효성 점검','유효성점검','효과성 검증','효과성검증'],'spread_7d':['재발 방지','재발방지','재발/방지','수평 전개','수평전개','재발 방지/수평전개','재발/방지/수평전개']}
STRUCT={compact if False else '' for _ in []}

def norm(x):
    if x is None:return ''
    return re.sub(r'\n+','\n',re.sub(r'[ \t]+',' ',str(x).replace('\r','\n'))).strip()
def compact(x):return re.sub(r'[^0-9A-Za-z가-힣]','',norm(x)).lower()
def one(x,n=None):
    s=re.sub(r'\s+',' ',norm(x));return s if not n or len(s)<=n else s[:n-1]+'…'
def clean(x):return re.sub(r'^[:：\-\s]+','',norm(x)).strip()
def blank(x):return not compact(x) or compact(x) in {'000','00','00호기','line00호기','담당자반영일','적용완료일','적용예정일','유첨5why','오창','cnj','cna','cmi','cwa'}
def instruction(x):
    q=compact(x);bad=['유첨5why에의거한근본원인','양불사이에4m관점의변경사항','재발불량이라면기존불량의원인분석','표준기준절차서fmeacpsop외의설정여부','it시스템관점에서불량유출된경우','피해유출된사유','4d에서도출된발생원인유출원인시스템원인에대한대책수립과정','각actionitem에대한일정및주관부서가명확하게설정되어있는지','고품처리방안','개선대책으로인한sideeffect는없는지','대책에대한효과가검증되어재현test시불량현상제거및고객spec을만족하는지','관련표준','담당자반영일','적용완료일','적용예정일','손실비용','10개lot모니터링등','본이슈로인한비용및산출근거','fact와data에근거한불량현상정의','작업자interview및itsystem내기록','불량과추정원인및임시조치에대해고객에게통보협의합의','발생원인개선','검출력개선및고품처리','개선대책','개선대책에대한효과성검증','대책에대한효과성검증','개선전후효과가있는지없는지결과로비교','mece관점에서잠재요인을발굴하고근본요인을도출','설비검사에서검출이안되고유출된이유','공정설계guidedesign리뷰iqclqcoqc','불량현상','임시조치','고객으로의불량유출과추가확산을막기위한조치','관련표준개정','d-fmea','p-fmea','controlplan','제품규격','작업표준','pfmeaform','현재예방관리','현재검출관리','추가예방조치','추가검출조치','발생요인재현결과에대한대책수립']
    return any(compact(x) in q for x in bad) or q.startswith('불량현상') or q.startswith('mece관점에서')
STRUCT=set(compact(x) for x in ['문제/현황','문제 현황','임시/조치','임시 조치','원인/분석','원인 분석','원인/파악','원인 파악','개선/대책','개선 대책','효과검증','효과 검증','유효성 검증','재발/방지','재발 방지','재발/방지/수평전개','재발 방지/수평 전개','수평전개','수평 전개','관련 표준','관련표준','개선 사항','진행 사항','검출력 개선 및 고품 처리','발생 Plant','발생 Line','발생 호기','신규/재발차수','사내/사외','Member (*원인 공정)','팀 구성','Lot no.','Lot . No.','불량 수량','부품 공급사','이슈 등급','Model','Customer','Packer'])
def structural(x):return compact(x) in STRUCT

def flat(sl):
    out=[]
    def rec(sh):
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
            for x in sh.shapes:rec(x)
        else:out.append(sh)
    for sh in sl.shapes:rec(sh)
    return out

def embedded(text):
    s=norm(text);hits=[];out={}
    for k,p in [('2d',r'(?:^|\n)\s*2\s*D.*?[:：]'),('3d',r'(?:^|\n)\s*3\s*D.*?[:：]'),('4d',r'(?:^|\n)\s*4\s*D.*?[:：]'),('5d',r'(?:^|\n)\s*5\s*D.*?[:：]'),('6d',r'(?:^|\n)\s*6\s*D.*?[:：]'),('7d',r'(?:^|\n)\s*7\s*D.*?[:：]')]:
        for m in re.finditer(p,s,re.I):hits.append((m.start(),m.end(),k))
    hits.sort()
    for i,(_,b,k) in enumerate(hits):out[k]=clean(s[b:hits[i+1][0] if i+1<len(hits) else len(s)])
    return out

def exact_label(x,als):return compact(x) in {compact(a) for a in als}
def candidates(rows,r,c):
    raw=rows[r][c];vals=[]
    if ':' in raw:vals.append(raw.split(':',1)[1].strip())
    vals += [rows[rr][c] for rr in range(r+1,min(r+5,len(rows)))]
    vals += rows[r][c+1:]
    return [clean(v) for v in vals]
def table_extract(tb):
    rows=[[norm(tb.cell(r,c).text) for c in range(len(tb.columns))] for r in range(len(tb.rows))];out={};section={'problem','temporary_action','customer_response','cause_4d','leak_cause','system_cause','action_5d','verification_6d','spread_7d'}
    for r,row in enumerate(rows):
        for c,raw in enumerate(row):
            if not raw:continue
            for k,als in EXACT.items():
                if out.get(k) or not exact_label(raw,als):continue
                vals=([clean(raw.split(':',1)[1])] if ':' in raw else [clean(v) for v in row[c+1:]]) if k in section else candidates(rows,r,c)
                for v in vals:
                    if not v or blank(v) or instruction(v) or structural(v):continue
                    if any(exact_label(v,a2) for a2 in EXACT.values()):continue
                    out[k]=v;break
    return out

def extract(path):
    prs=Presentation(path);d={k:'' for k in ['issue_name']+list(EXACT)};d['standards']={x:'' for x in STD};d['_images']=[];texts=[];emb={}
    for sl in prs.slides:
        for sh in flat(sl):
            t=norm(getattr(sh,'text',''))
            if t:
                texts.append(t);emb.update(embedded(t))
                if '이슈명' in t and ':' in t and not d['issue_name']:
                    m=re.search(r'이슈명\s*[:：]\s*([^\n]+)',t,re.I)
                    if m:d['issue_name']=clean(m.group(1))
            if getattr(sh,'has_table',False):
                f=table_extract(sh.table)
                for k,v in f.items():
                    if not d.get(k):d[k]=v
            if getattr(sh,'shape_type',None)==MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob=sh.image.blob
                    with PILImage.open(io.BytesIO(blob)) as im:
                        w,h=im.size
                        if w*h>=10000:d['_images'].append((w*h,w,h,blob))
                except Exception:pass
    full='\n'.join(texts)
    for k,als in {'issue_name':['이슈명','이슈 제목'],**EXACT}.items():
        if d.get(k):continue
        for a in als:
            m=re.search(re.escape(a)+r'\s*[:：]\s*([^\n]+)',full,re.I)
            if m:
                v=clean(m.group(1))
                if v and not blank(v) and not instruction(v):d[k]=v;break
    for k,ek in [('problem','2d'),('temporary_action','3d'),('cause_4d','4d'),('action_5d','5d'),('verification_6d','6d'),('spread_7d','7d')]:
        if not d[k] and emb.get(ek):d[k]=emb[ek]
    if d['temporary_action']:
        m=re.search(r'고객\s*대응\s*[:：]',d['temporary_action'])
        if m:d['customer_response']=clean(d['temporary_action'][m.end():]);d['temporary_action']=clean(d['temporary_action'][:m.start()])
    if d['cause_4d']:
        m=re.search(r'유출\s*원인\s*[:：]',d['cause_4d'])
        if m:
            if not d['leak_cause']:d['leak_cause']=clean(d['cause_4d'][m.end():])
            d['cause_4d']=clean(d['cause_4d'][:m.start()])
    if not d.get('task_name') and d.get('issue_name') and not blank(d['issue_name']):d['task_name']=d['issue_name']
    return d

def parse_year_month(text):
    s=one(text)
    for p in [r'(20\d{2})\s*[-./]\s*(\d{1,2})',r'(?<!\d)(\d{2})\s*[-./]\s*(\d{1,2})\s*[-./]\s*(\d{1,2})(?!\d)',r'(?<!\d)(\d{4})(\d{2})(\d{2})(?!\d)']:
        m=re.search(p,s)
        if m:return (m.group(1) if len(m.group(1))==4 else '20'+m.group(1),m.group(2))
    return '',''
def status(d):
    v=one(d.get('verification_6d'))
    if re.search(r'완료\s*(예정|추정)',v):return STATUS['verify']
    if re.search(r'검증\s*완료|(?<![가-힣])완료(?![가-힣])',v):return STATUS['complete']
    if re.search(r'진행\s*중|진행중|검증\s*중|검증중|예정|추정',v):return STATUS['verify']
    if blank(d.get('cause_4d')) and blank(d.get('action_5d')):return STATUS['unknown']
    return STATUS['verify']
def progress(d):
    cause=' / '.join(x for x in [d.get('cause_4d'),d.get('leak_cause'),d.get('system_cause')] if x)
    prog=' / '.join(x for x in [d.get('customer_response'),d.get('temporary_action'),d.get('action_5d'),d.get('verification_6d')] if x)
    return '\n'.join(['원인: '+one(cause,180) if cause else '원인: 검토 중','진행 현황: '+one(prog,320) if prog else '진행 현황: 검토 중','요청 사항:'])
def task(d):return ('%s / %s'%(one(d.get('customer')),one(d.get('task_name')))).strip(' /')
def key(x):return re.sub(r'[\s_/\\\-:.,·]+','',one(x).lower())
def sig(ws,r):return {'plm':key(ws.cell(r,3).value),'task':key(ws.cell(r,10).value),'problem':key(ws.cell(r,14).value),'cause':key(ws.cell(r,16).value),'action':key(ws.cell(r,17).value)}
def find(ws,d,g):
    p=key(g.get('plm_no'))
    if p:
        for r in range(8,ws.max_row+1):
            if sig(ws,r)['plm']==p:return r,100
    tt,pp,cc,aa=key(task(d)),key(d.get('problem')),key(d.get('cause_4d')),key(d.get('action_5d'));best=(None,0)
    for r in range(8,ws.max_row+1):
        s=sig(ws,r);score=(5 if tt and s['task']==tt else 0)+(5 if pp and s['problem']==pp else 0)+(2 if cc and s['cause']==cc else 0)+(2 if aa and s['action']==aa else 0)
        if score>best[1]:best=(r,score)
    return best if best[1]>=10 else (None,0)
def copyfmt(ws,a,b):
    ws.row_dimensions[b].height=ws.row_dimensions[a].height
    for c in range(1,ws.max_column+1):ws.cell(b,c)._style=copy.copy(ws.cell(a,c)._style)
def representative_image(d):
    imgs=d.get('_images') or [];return max(imgs,key=lambda x:x[0])[3] if imgs else None
def add_excel_photo(ws,row,blob):
    if not blob:return
    try:
        with PILImage.open(io.BytesIO(blob)) as im:
            w,h=im.size;max_w=max(80,min(210,int((ws.column_dimensions['O'].width or 20)*7)));max_h=95;scale=min(max_w/w,max_h/h,1.0)
            img=XLImage(io.BytesIO(blob));img.width=max(1,int(w*scale));img.height=max(1,int(h*scale));ws.add_image(img,f'O{row}');ws.row_dimensions[row].height=max(ws.row_dimensions[row].height or 15,min(100,img.height*0.75+8))
    except Exception:pass
def write_row(ws,r,d,g):
    y,m=parse_year_month(d.get('occurrence_date'));st=status(d);vals={2:datetime.date.today(),3:g.get('plm_no',''),4:g.get('form_factor',''),5:g.get('product_type',''),6:y,7:m,8:g.get('team',''),9:g.get('owner',''),10:task(d),11:g.get('sample',''),12:g.get('stage',''),13:d.get('occurrence_site',''),14:one(d.get('problem'),120),15:'',16:one(' / '.join(x for x in [d.get('cause_4d'),d.get('leak_cause'),d.get('system_cause')] if x),160),17:one(d.get('action_5d'),160),18:st}
    for c,v in vals.items():ws.cell(r,c).value=v
    add_excel_photo(ws,r,representative_image(d));return st
def update_excel(src,out,d,g,new=False):
    wb=load_workbook(src);ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active;r,score=find(ws,d,g)
    if r and not new:
        old=ws.cell(r,3).value;write_row(ws,r,d,g)
        if not g.get('plm_no') and old:ws.cell(r,3).value=old
        msg=f'기존 이슈 업데이트 (row {r}, match score {score})'
    else:
        r=ws.max_row+1
        if r>7:copyfmt(ws,r-1,r)
        write_row(ws,r,d,g);msg=f'신규 이슈 추가 (row {r})'
    Path(out).parent.mkdir(exist_ok=True)
    try:wb.save(out);saved=out
    except PermissionError:
        p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);wb.save(saved);msg+=f'\n※ 기존 결과 파일이 열려 있어 새 파일로 저장: {saved.name}'
    return msg,saved

def font_run(run,size=8):
    run.font.name='맑은 고딕';run.font.size=Pt(size)
    try:r=run._r.get_or_add_rPr();r.set('a:latin','맑은 고딕');r.set('a:ea','맑은 고딕')
    except Exception:pass
def set_shape_text(sh,text,size=8):
    sh.text=text
    if hasattr(sh,'text_frame'):
        sh.text_frame.word_wrap=True
        for p in sh.text_frame.paragraphs:
            for run in p.runs:font_run(run,size)
def set_cell_text(cl,text,size=8):
    cl.text=text;cl.text_frame.word_wrap=True
    for p in cl.text_frame.paragraphs:
        for run in p.runs:font_run(run,size)
def signal(cl,st):
    set_cell_text(cl,'●',8);rgb=SIGNAL[st]
    for p in cl.text_frame.paragraphs:
        for run in p.runs:run.font.color.rgb=RGBColor(*rgb);run.font.bold=True
def table_cells(sl):
    for sh in flat(sl):
        if getattr(sh,'has_table',False):
            tb=sh.table
            for r in range(len(tb.rows)):
                for c in range(len(tb.columns)):yield tb,r,c,tb.cell(r,c)
def section_key(text):
    s=compact(text)
    for k,als in [('2d',ALIASES['2d']),('3d',ALIASES['3d']),('4d_cause',ALIASES['4d_cause']),('4d_leak',ALIASES['4d_leak']),('4d_system',ALIASES['4d_system']),('5d',ALIASES['5d']),('6d',ALIASES['6d']),('7d',ALIASES['7d'])]:
        if any(compact(a) in s for a in als):return k
    return None
def section_contents(d):return {'2d':one(d.get('problem'),700),'3d':one(' / '.join(x for x in [d.get('temporary_action'),d.get('customer_response')] if x),700),'4d_cause':one(d.get('cause_4d'),700),'4d_leak':one(d.get('leak_cause'),700),'4d_system':one(d.get('system_cause'),700),'5d':one(d.get('action_5d'),700),'6d':one(d.get('verification_6d'),700),'7d':one(d.get('spread_7d'),700)}
def fill_semantic_table(tb,d,used=None):
    used=used or set();contents=section_contents(d)
    for r in range(len(tb.rows)):
        for c in range(len(tb.columns)):
            k=section_key(tb.cell(r,c).text)
            if not k or not contents.get(k):continue
            for cc in range(c+1,len(tb.columns)):
                cl=tb.cell(r,cc);txt=norm(cl.text)
                if id(cl) in used:continue
                if not txt or txt in ['내용','입력','내용 입력','예시'] or instruction(txt):set_cell_text(cl,contents[k],8);used.add(id(cl));break
def fill_metadata_table(tb,d,g):
    for r in range(len(tb.rows)):
        for c in range(len(tb.columns)):
            txt=norm(tb.cell(r,c).text)
            for k,als in EXACT.items():
                if not txt or not exact_label(txt,als):continue
                val=d.get(k,'') or g.get(k,'')
                if not val:continue
                dest=None
                for rr in range(r+1,min(r+5,len(tb.rows))):
                    cl=tb.cell(rr,c)
                    if not norm(cl.text) or blank(cl.text) or instruction(cl.text):dest=cl;break
                if dest is None:
                    for cc in range(c+1,len(tb.columns)):
                        cl=tb.cell(r,cc)
                        if not norm(cl.text) or blank(cl.text) or instruction(cl.text):dest=cl;break
                if dest is not None:set_cell_text(dest,one(val,300),8);break
def fill_numbered_boxes(sl,d):
    mapping={'1':'2d','2':'3d','3-1':'4d_cause','3-2':'customer_response','4':'5d','5':'6d','6':'6d','7':'7d'};contents=section_contents(d)
    for sh in flat(sl):
        txt=norm(getattr(sh,'text',''))
        if txt in mapping:set_shape_text(sh,d.get(mapping[txt],'') if mapping[txt]=='customer_response' else contents.get(mapping[txt],''),8)
def replace_metadata_shapes(sl,d,g):
    title=task(d);issue=one(d.get('issue_name'))
    for sh in flat(sl):
        if not hasattr(sh,'text') or not sh.text:continue
        old=sh.text;new=old
        if '과제명_이슈 제목' in new:new=new.replace('과제명_이슈 제목',f'{title}_{issue}'.strip('_'))
        if '이슈명' in new and ':' in new:new=new.split(':',1)[0]+' : '+issue
        if '00팀 담당자 : 000' in new:new=new.replace('00팀 담당자 : 000',f"{g.get('team','')} 담당자 : {g.get('owner','')}")
        if new!=old:set_shape_text(sh,new,8)
def fill_detail_slide(sl,d,g):
    replace_metadata_shapes(sl,d,g);fill_numbered_boxes(sl,d);used=set()
    for tb,_,_,_ in table_cells(sl):fill_metadata_table(tb,d,g);fill_semantic_table(tb,d,used)
    for sh in flat(sl):
        if hasattr(sh,'text_frame') and norm(sh.text):
            for p in sh.text_frame.paragraphs:
                for run in p.runs:font_run(run,8)
def clone_slide(prs,sl):
    ns=prs.slides.add_slide(sl.slide_layout)
    for sh in sl.shapes:ns.shapes._spTree.insert_element_before(copy.deepcopy(sh.element),'p:extLst')
    return ns
def weekly(src,out,d,g,mode):
    prs=Presentation(src);st=status(d);title=task(d);issue=one(d.get('issue_name'));summary=None
    for sl in prs.slides:
        for sh in flat(sl):
            if getattr(sh,'has_table',False):
                header=' '.join(norm(sh.table.cell(0,c).text) for c in range(len(sh.table.columns)))
                if '과제명' in header and 'Signal' in header:summary=sh.table;break
        if summary:break
    if summary:
        r=None
        for x in range(1,len(summary.rows)):
            if key(title) and key(title) in key(summary.cell(x,0).text) and key(issue) in key(summary.cell(x,1).text):r=x;break
        if r is None:
            for x in range(1,len(summary.rows)):
                if not any(norm(summary.cell(x,c).text) for c in range(min(5,len(summary.columns))) if c!=4):r=x;break
        if r is None:summary._tbl.append(copy.deepcopy(summary.rows[-1]._tr));r=len(summary.rows)-1
        for c,v in enumerate([title,issue,one(d.get('problem'),140),progress(d),'●']):set_cell_text(summary.cell(r,c),v,8)
        signal(summary.cell(r,4),st)
    summary_idx=None
    for i,sl in enumerate(prs.slides):
        headers=[' '.join(norm(sh.table.cell(0,c).text) for c in range(len(sh.table.columns))) for sh in flat(sl) if getattr(sh,'has_table',False)]
        if any('과제명' in h and 'Signal' in h for h in headers):summary_idx=i;break
    if summary_idx is None:summary_idx=0
    details=[prs.slides[i] for i in range(summary_idx+1,len(prs.slides))]
    if not details:details=[clone_slide(prs,prs.slides[1] if len(prs.slides)>1 else prs.slides[0])]
    for sl in details:fill_detail_slide(sl,d,g)
    for sl in details:
        for tb,r,c,cl in table_cells(sl):
            if 'signal' in compact(cl.text) and c+1<len(tb.columns):signal(tb.cell(r,c+1),st)
    Path(out).parent.mkdir(exist_ok=True)
    try:prs.save(out);saved=out
    except PermissionError:
        p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);prs.save(saved)
    return '주간회의 PPT 업데이트: '+st,saved

class App(tk.Tk):
    def __init__(self):
        super().__init__();self.title('8D 이슈 자동화 v2.8');self.geometry('820x700');self.vars={};self.mode=tk.StringVar(value='existing')
        fields=[('8D PPT','ppt8d',1),('주간회의 PPT','pptweekly',1),('이슈 DB Excel','xlsx',1),('담당팀','team',0),('담당자','owner',0),('폼팩터','form_factor',0),('제품 타입','product_type',0),('발생 샘플','sample',0),('개발 단계','stage',0),('PMS/PLM 이슈번호(기존값이 있으면 입력)','plm_no',0)]
        q=ttk.Frame(self);q.pack(fill='x',padx=18,pady=8);ttk.Label(q,text='이슈 구분',width=31).pack(side='left');ttk.Radiobutton(q,text='기존 이슈',variable=self.mode,value='existing').pack(side='left');ttk.Radiobutton(q,text='신규 이슈',variable=self.mode,value='new').pack(side='left')
        for lab,k,isfile in fields:
            self.vars[k]=tk.StringVar();r=ttk.Frame(self);r.pack(fill='x',padx=18,pady=4);ttk.Label(r,text=lab,width=31).pack(side='left');ttk.Entry(r,textvariable=self.vars[k],width=69).pack(side='left')
            if isfile:ttk.Button(r,text='찾기',command=lambda x=k:self.pick(x)).pack(side='left',padx=5)
        ttk.Button(self,text='8D 내용 미리보기',command=self.preview).pack(pady=4);ttk.Button(self,text='자동 업데이트 실행',command=self.run).pack(pady=4);self.log=tk.Text(self,height=15,width=100);self.log.pack(padx=18,pady=10)
    def pick(self,k):
        ft=[('Excel','*.xlsx')] if k=='xlsx' else [('PowerPoint','*.pptx')];f=filedialog.askopenfilename(filetypes=ft)
        if f:self.vars[k].set(f)
    def gui(self):return {k:v.get().strip() for k,v in self.vars.items()}
    def preview(self):
        g=self.gui()
        if not g['ppt8d']:return messagebox.showwarning('확인','8D PPT를 선택하세요.')
        try:
            d=extract(g['ppt8d']);self.log.delete('1.0','end');self.log.insert('end','=== 8D 추출 결과 ===\n')
            for k in ['issue_name','task_name','model','customer','occurrence_site','occurrence_date','problem','temporary_action','customer_response','cause_4d','leak_cause','system_cause','action_5d','verification_6d','spread_7d']:self.log.insert('end',f'{k}: {d.get(k)}\n')
            self.log.insert('end',f'\n=== 상태 판정 ===\n{status(d)}\n\n=== 주간회의 진행사항 ===\n{progress(d)}')
        except Exception as e:messagebox.showerror('오류',repr(e))
    def run(self):
        g=self.gui()
        if any(not g[k] or not os.path.exists(g[k]) for k in ['ppt8d','pptweekly','xlsx']):return messagebox.showwarning('확인','8D PPT, 주간회의 PPT, 이슈 DB Excel을 모두 선택하세요.')
        try:
            d=extract(g['ppt8d']);mode=self.mode.get();out=Path(g['xlsx']).parent/'자동화_결과';out.mkdir(exist_ok=True);xo=out/(Path(g['xlsx']).stem+'_업데이트.xlsx');po=out/(Path(g['pptweekly']).stem+'_업데이트.pptx');wb=load_workbook(g['xlsx']);ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active;r,_=find(ws,d,g)
            if mode=='existing' and not r:return messagebox.showwarning('업데이트 중단','기존 이슈를 Excel에서 정확하게 찾지 못했습니다. PMS/PLM 이슈번호를 입력하거나 이슈 정보를 확인하세요.')
            if mode=='new' and r and not messagebox.askyesno('중복 가능성',f'유사 이슈(row {r})가 있습니다. 그래도 신규 추가할까요?'):return
            a,_=update_excel(g['xlsx'],xo,d,g,new=(mode=='new'));b,_=weekly(g['pptweekly'],po,d,g,mode);self.log.delete('1.0','end');self.log.insert('end',f'=== 완료 ===\n이슈 구분: {mode}\n상태: {status(d)}\n\n{a}\n{b}\n\n결과: {out}');messagebox.showinfo('완료',f'자동 업데이트가 완료되었습니다.\n\n{out}')
        except Exception as e:messagebox.showerror('실행 오류',repr(e))

if __name__=='__main__':App().mainloop()
