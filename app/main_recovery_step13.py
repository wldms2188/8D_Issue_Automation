import copy
import datetime
import os
import re
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

import main_recovery_step12 as step12
import main_recovery_step10 as step10
import main_recovery_step8 as step8
import main_recovery_step4 as step4
import main_v319 as v319
import main_v310 as v310

base=step12.base
N=v310.N
C=v310.C


# -----------------------------------------------------------------------------
# Shared normalization
# -----------------------------------------------------------------------------
def _k(x):
    return re.sub(r'[^0-9A-Za-z가-힣]+','',N(x)).lower()


def _split_top_level_customer_task(value):
    s=N(value)
    depth=0
    for i,ch in enumerate(s):
        if ch=='(':
            depth+=1
        elif ch==')' and depth:
            depth-=1
        elif ch=='_' and depth==0:
            return s[:i].strip(),s[i+1:].strip()
    return '',s

def _customer_task(d):
    customer=N(d.get('customer'))
    task=N(d.get('task_name'))

    # User contract: top-level A_B already means customer_project.
    # Therefore never prepend d.customer again, even if an earlier extractor
    # populated customer incorrectly.
    prefix,rest=_split_top_level_customer_task(task)
    if prefix and rest:
        # Collapse accidental repeated prefix: A_A_B -> A_B.
        pat=re.compile(r'^(?:'+re.escape(prefix)+r'\s*[_\-/／|:：]\s*)+',re.I)
        m=pat.match(rest)
        while m:
            rest=N(rest[m.end():]).strip(' _-/／|:：')
            m=pat.match(rest)
        return prefix+'_'+rest if rest else prefix

    if customer and task:
        if C(task)==C(customer):
            return customer
        return customer+'_'+task
    return customer or task

# Make the active v2.8/v2.9 base use the same canonical customer_task everywhere.
# This fixes Issue DB matching/writing cases such as A / A_B -> A_B.
base.task=_customer_task


def _issue_display(d):
    return step4.step1._page1_issue(d)


# -----------------------------------------------------------------------------
# Issue DB Excel
# NEW: insert immediately ABOVE the C-column "↑ ... 행 삽입 ..." guide row.
# EXISTING: keep existing exact/similarity matching and update that row only.
# -----------------------------------------------------------------------------
def _is_insert_guide(value):
    raw=N(value)
    q=_k(raw)
    if '행삽입' not in q:
        return False
    # Accept common upward-arrow glyphs used in Excel templates.
    return any(ch in raw for ch in ('↑','⬆','⇧','⇑','↟','⭡')) or '위쪽' in raw or '위방향' in raw


def _find_insert_guide_row(ws):
    for r in range(1,ws.max_row+1):
        if _is_insert_guide(ws.cell(r,3).value):
            return r
    # Safe fallback for templates where the arrow glyph was lost but the guide text remains.
    for r in range(1,ws.max_row+1):
        if '행삽입' in _k(ws.cell(r,3).value):
            return r
    return None


def _copy_row_format(ws,source_row,target_row):
    if source_row<1:
        return
    try:
        ws.row_dimensions[target_row].height=ws.row_dimensions[source_row].height
        ws.row_dimensions[target_row].hidden=ws.row_dimensions[source_row].hidden
    except Exception:
        pass
    for c in range(1,ws.max_column+1):
        src=ws.cell(source_row,c)
        dst=ws.cell(target_row,c)
        try: dst._style=copy.copy(src._style)
        except Exception: pass
        try: dst.number_format=src.number_format
        except Exception: pass
        try: dst.alignment=copy.copy(src.alignment)
        except Exception: pass
        try: dst.protection=copy.copy(src.protection)
        except Exception: pass
        # Values are intentionally NOT copied.
        dst.value=None


def update_excel_step13(src,out,d,g,new=False):
    wb=load_workbook(src)
    ws=wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active

    if new:
        marker=_find_insert_guide_row(ws)
        if marker is None:
            raise ValueError("Issue DB C열에서 '↑ ... 행 삽입 ...' 안내 행을 찾지 못했습니다. 원본 양식을 확인하세요.")

        # openpyxl insert_rows(marker) creates the new blank row at marker and pushes
        # the original guide row down, i.e. exactly ABOVE the guide row.
        ws.insert_rows(marker,1)
        source=max(1,marker-1)
        _copy_row_format(ws,source,marker)
        base.write_row(ws,marker,d,g)
        msg=f'신규 이슈 추가 (C열 행 삽입 안내 바로 위, row {marker})'
    else:
        r,score=base.find(ws,d,g)
        if not r:
            raise ValueError('기존 이슈를 Issue DB에서 찾지 못했습니다. 신규 행은 자동 추가하지 않았습니다.')
        old_plm=ws.cell(r,3).value
        base.write_row(ws,r,d,g)
        if not N(g.get('plm_no')) and old_plm:
            ws.cell(r,3).value=old_plm
        msg=f'기존 이슈 업데이트 (row {r}, match score {score})'

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        wb.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        wb.save(saved)
        msg+='\n※ 기존 결과 파일이 열려 있어 새 파일로 저장: '+Path(saved).name
    return msg,saved


base.update_excel=update_excel_step13


# -----------------------------------------------------------------------------
# Weekly PPT helpers
# -----------------------------------------------------------------------------
def _slide_text(sl):
    parts=[]
    for sh in v310.walk(sl):
        t=N(getattr(sh,'text',''))
        if t:
            parts.append(t)
        if getattr(sh,'has_table',False):
            tb=sh.table
            for r in range(len(tb.rows)):
                for c in range(len(tb.columns)):
                    v=N(tb.cell(r,c).text)
                    if v:
                        parts.append(v)
    return '\n'.join(parts)


def _detail_structure_score(sl):
    text=_slide_text(sl)
    q=_k(text)
    score=0
    for token in ('2d','3d','4d','5d','6d'):
        if token in q:
            score+=1
    for token in ('signal','이슈기인','발생단계'):
        if token in q:
            score+=1
    return score


def _find_summary(prs):
    for si,sl in enumerate(prs.slides):
        for sh in v310.walk(sl):
            if not getattr(sh,'has_table',False):
                continue
            tb=sh.table
            for hr in range(min(3,len(tb.rows))):
                qs=[C(tb.cell(hr,c).text) for c in range(len(tb.columns))]
                if any('과제명' in q for q in qs) and any('signal'==q for q in qs):
                    return si,tb,hr
    return None,None,None


def _summary_map(tb,hr):
    hm={}
    for c in range(len(tb.columns)):
        q=C(tb.cell(hr,c).text)
        if '과제명' in q: hm['task']=c
        elif q=='이슈' or '이슈명' in q: hm['issue']=c
        elif '현상' in q or '문제' in q: hm['problem']=c
        elif '진행사항' in q or '진행현황' in q: hm['progress']=c
        elif q=='signal': hm['signal']=c
    return hm


def _row_text(tb,r,c):
    try: return N(tb.cell(r,c).text)
    except Exception: return ''


def _row_match_score(tb,r,hm,task,issue):
    tk=_k(task); ik=_k(issue)
    rt=_k(_row_text(tb,r,hm.get('task',0))) if 'task' in hm else ''
    ri=_k(_row_text(tb,r,hm.get('issue',0))) if 'issue' in hm else ''
    score=0
    if tk and rt==tk: score+=100
    elif tk and (tk in rt or rt in tk): score+=60
    if ik and ri==ik: score+=100
    elif ik and (ik in ri or ri in ik): score+=70
    elif ik and ri:
        score+=int(50*SequenceMatcher(None,ik,ri).ratio())
    return score


def _append_summary_row(tb,hr):
    # User requirement: NEW issue is added at the VERY END of the summary table,
    # not at the first blank row. Clone the final data row to preserve exact styling.
    source=len(tb.rows)-1
    if source<=hr:
        source=hr
    new_tr=copy.deepcopy(tb.rows[source]._tr)
    tb._tbl.append(new_tr)
    r=len(tb.rows)-1
    for c in range(len(tb.columns)):
        try: tb.cell(r,c).text=''
        except Exception: pass
    return r


def _update_summary(prs,d,g,mode):
    si,tb,hr=_find_summary(prs)
    if tb is None:
        raise ValueError('주간회의 PPT에서 과제명/Signal이 있는 1페이지 요약표를 찾지 못했습니다.')

    hm=_summary_map(tb,hr)
    task=_customer_task(d)
    issue=_issue_display(d)
    target=None

    if mode=='existing':
        best=(None,-1)
        for r in range(hr+1,len(tb.rows)):
            sc=_row_match_score(tb,r,hm,task,issue)
            if sc>best[1]: best=(r,sc)
        if best[0] is not None and best[1]>=140:
            target=best[0]

    if target is None:
        # NEW always appends. EXISTING also appends only when that issue is not yet
        # present in the weekly summary; it does not overwrite a different issue.
        target=_append_summary_row(tb,hr)

    vals={
        'task':task,
        'issue':issue,
        'problem':step4._full_problem(d),
        'progress':v319._progress_text(d),
    }
    for k,v in vals.items():
        if k in hm:
            try:
                # preserve existing row style; only replace the text
                tb.cell(target,hm[k]).text=N(v)
                for p in tb.cell(target,hm[k]).text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size=v319.Pt(8) if False else run.font.size
            except Exception:
                pass

    if 'signal' in hm:
        step12.base.signal(tb.cell(target,hm['signal']),step12._signal_status_from_g(d,g))
    return si,target


def _detail_match_score(sl,d):
    structure=_detail_structure_score(sl)
    if structure<5:
        return -1

    text=_slide_text(sl)
    q=_k(text)
    task=_k(_customer_task(d))
    issue=_k(_issue_display(d))
    title=_k(v319._weekly_display(d).get('_title_issue'))

    score=structure*5
    if task and task in q: score+=80
    if issue and issue in q: score+=100
    elif title and title in q: score+=50
    elif issue:
        # Compare against each visible line so unrelated long slide text does not win.
        ratios=[SequenceMatcher(None,issue,_k(line)).ratio() for line in text.splitlines() if _k(line)]
        if ratios: score+=int(60*max(ratios))
    return score


def _find_existing_detail(prs,d,summary_index):
    best=(None,-1)
    for i,sl in enumerate(prs.slides):
        if i==summary_index:
            continue
        sc=_detail_match_score(sl,d)
        if sc>best[1]: best=(i,sc)
    # Require both a detail-like structure and meaningful title/task similarity.
    return best[0] if best[0] is not None and best[1]>=115 else None


def _find_detail_template(prs,summary_index):
    # Prefer the originally shared blank/example-style detail page when present.
    preferred=[]; fallback=[]
    for i,sl in enumerate(prs.slides):
        if i==summary_index:
            continue
        score=_detail_structure_score(sl)
        if score<5:
            continue
        text=_slide_text(sl)
        q=_k(text)
        if '과제명이슈제목' in q or ('이슈명' in q and ('8불량명' in q or '불량명' in q)):
            preferred.append((score,i))
        fallback.append((score,i))
    pool=preferred or fallback
    if not pool:
        raise ValueError('주간회의 PPT에서 2D~6D 상세 양식 페이지를 찾지 못했습니다.')
    pool.sort(reverse=True)
    return pool[0][1]


def _clone_slide_with_rels(prs,src_index):
    src=prs.slides[src_index]
    dst=prs.slides.add_slide(src.slide_layout)

    # Re-create non-layout relationships and rewrite rIds in copied XML. This keeps
    # logos/pictures in the supplied weekly template intact when the page is cloned.
    rid_map={}
    for rid,rel in src.part.rels.items():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        try:
            nr=dst.part.rels._add_relationship(rel.reltype,rel._target,rel.is_external)
            rid_map[rid]=nr
        except Exception:
            pass

    for sh in src.shapes:
        el=copy.deepcopy(sh._element)
        for node in el.iter():
            for attr,val in list(node.attrib.items()):
                if val in rid_map:
                    node.attrib[attr]=rid_map[val]
        dst.shapes._spTree.insert_element_before(el,'p:extLst')
    return dst


def _move_last_slide_to(prs,index):
    sldIdLst=prs.slides._sldIdLst
    new_id=sldIdLst[-1]
    sldIdLst.remove(new_id)
    sldIdLst.insert(index,new_id)


def _task_slide_indices(prs,d,summary_index):
    tk=_k(_customer_task(d))
    out=[]
    if not tk:
        return out
    for i,sl in enumerate(prs.slides):
        if i==summary_index:
            continue
        if _detail_structure_score(sl)>=5 and tk in _k(_slide_text(sl)):
            out.append(i)
    return out


def _new_detail_position(prs,d,summary_index):
    same=_task_slide_indices(prs,d,summary_index)
    if same:
        # Same customer_task section exists -> add after its final detail page.
        return max(same)+1
    # No section yet -> create a new customer_task section at the end of detail pages.
    return len(prs.slides)


def _set_origin_on_slide(sl,origin):
    if not origin:
        return False
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False):
            continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                if C(tb.cell(r,c).text)=='이슈기인' and c+1<len(tb.columns):
                    target=tb.cell(r,c+1)
                    reference=step8._reference_value_cell(tb,r,c+1)
                    step8._set_cell_like_reference(target,origin,reference,9)
                    return True
    return False


def _set_detail_signal(sl,d,g):
    st=step12._signal_status_from_g(d,g)
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False):
            continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                if C(tb.cell(r,c).text)=='signal' and c+1<len(tb.columns):
                    base.signal(tb.cell(r,c+1),st)


def _update_detail_slide(sl,d,g,mode):
    dd=v319._weekly_display(d)
    step12._update_page2_step12(sl,dd,g,mode)
    step4._force_page2_header(sl,d,g)
    _set_detail_signal(sl,d,g)
    _set_origin_on_slide(sl,N(g.get('_issue_origin_selected')))


def weekly_step13(src,out,d,g,mode):
    prs=Presentation(src)
    summary_index,_row=_update_summary(prs,d,g,mode)

    target_index=None
    if mode=='existing':
        target_index=_find_existing_detail(prs,d,summary_index)

    action='업데이트'
    if target_index is None:
        template_index=_find_detail_template(prs,summary_index)
        insert_at=_new_detail_position(prs,d,summary_index)
        _clone_slide_with_rels(prs,template_index)
        _move_last_slide_to(prs,insert_at)
        target_index=insert_at
        action='추가'

    _update_detail_slide(prs.slides[target_index],d,g,mode)

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)

    return f'주간회의 PPT {action}: {_customer_task(d)} / 상세 page {target_index+1}',saved


base.weekly=weekly_step13


class RecoveryStep13App(step10.RecoveryStep10App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP13')


if __name__=='__main__':
    RecoveryStep13App().mainloop()
