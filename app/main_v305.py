# 8D Issue Automation v3.0.5
# Latest correction: task/issue mapping, template marker grouping, clean weekly images.
import sys,re,copy,datetime
from pathlib import Path
APP_DIR=Path(__file__).resolve().parent
if str(APP_DIR) not in sys.path: sys.path.insert(0,str(APP_DIR))
import main_v303 as v303
base=v303.base
renderer=v303.renderer
from pptx import Presentation
from pptx.util import Inches

def _n(x): return str(x or '').replace('\r\n','\n').replace('\r','\n').strip()
def _c(x):
    try:return base.compact(x)
    except Exception:return re.sub(r'[^0-9A-Za-z가-힣]','',_n(x)).lower()
def _clean_task(x):
    s=_n(x); s=re.sub(r'^과제명\s*[:：]\s*','',s); return s.strip(' /／|:：')
def _valid_task(x):
    q=_c(x)
    if len(_n(x))<2:return False
    return not any(k in q for k in ('고객사','이슈명','과제명','model','packer','발생site','발생일자','lotno','발생line','담당자','제품타입','폼팩터','발생샘플','개발단계','발생처','불량률','불량수량','signal'))

def _task_from_issue_name(d):
    issue=_n(d.get('issue_name'))
    if not issue:return ''
    # 8D issue-name area commonly stores 고객사/과제명. The task is strictly the
    # part after 고객사, not the whole issue string.
    m=re.search(r'고객사\s*[/／|>→:-]\s*(.+)$',issue)
    if m:return _clean_task(m.group(1))
    parts=re.split(r'[/／|]',issue,maxsplit=1)
    if len(parts)==2 and _valid_task(parts[1]):return _clean_task(parts[1])
    return ''

def _task_from_ppt(path):
    prs=Presentation(path)
    for sl in prs.slides:
        # Text box form: explicitly search the customer/task line.
        for sh in sl.shapes:
            text=_n(getattr(sh,'text',''))
            if not text or '고객사' not in _c(text):continue
            m=re.search(r'고객사\s*[/／|>→:-]\s*(.+)$',text.replace('\n',' '))
            if m and _valid_task(m.group(1)):return _clean_task(m.group(1))
            lines=[x.strip() for x in text.split('\n') if x.strip()]
            for i,line in enumerate(lines):
                if '고객사' not in _c(line):continue
                for j in range(i+1,min(i+5,len(lines))):
                    if '과제명' in _c(lines[j]):
                        p=re.split(r'[:：]',lines[j],maxsplit=1)
                        if len(p)==2 and _valid_task(p[1]):return _clean_task(p[1])
        # Table form.
        for sh in sl.shapes:
            if not getattr(sh,'has_table',False):continue
            tb=sh.table; rows=[[_n(tb.cell(r,c).text) for c in range(len(tb.columns))] for r in range(len(tb.rows))]
            for r,row in enumerate(rows):
                for c,val in enumerate(row):
                    if '고객사' not in _c(val):continue
                    right=[_n(x) for x in row[c+1:] if _n(x)]
                    for j,x in enumerate(right):
                        if '과제명' in _c(x) and j+1<len(right) and _valid_task(right[j+1]):return _clean_task(right[j+1])
                    # [고객사][고객사 값][과제명 값]
                    if len(right)>=2 and _valid_task(right[1]):return _clean_task(right[1])
                    for rr in range(r+1,min(r+4,len(rows))):
                        for cc,x in enumerate(rows[rr]):
                            if '과제명' in _c(x) and cc+1<len(rows[rr]) and _valid_task(rows[rr][cc+1]):return _clean_task(rows[rr][cc+1])
    return ''

_prev_extract=base.extract
def extract_v305(path):
    d=_prev_extract(path)
    try:
        t=_task_from_issue_name(d)
        if not t:t=_task_from_ppt(path)
        if t:d['task_name']=t
    except Exception:pass
    # Weekly should not scatter every source image; one largest meaningful image per D.
    try:
        imgs=d.get('_section_images',{}) or {}; out={k:[] for k in ('2D','3D','4D','5D','6D')}
        for sec,vals in imgs.items():
            if vals:
                vals=list(vals); vals.sort(key=lambda v:(v[1][2]*v[1][3],v[2]),reverse=True); out[sec]=[vals[0]]
        d['_section_images']=out
    except Exception:pass
    return d
base.extract=extract_v305

# ---------------------------------------------------------------------------
# Existing D circles: preserve the template shape, move it with the flowed zone.
# ---------------------------------------------------------------------------
def _walk(container):
    for sh in getattr(container,'shapes',[]):
        yield sh
        if hasattr(sh,'shapes'):yield from _walk(sh)

def _move_markers(sl,layout):
    mapping=[('2D','2D'),('3D','3D'),('4D','4D_LEAK'),('5D','5D'),('6D','6D')]; used=set()
    for label,zkey in mapping:
        z=layout.get(zkey)
        if not z:continue
        cand=[s for s in _walk(sl) if _n(getattr(s,'text',''))==label and not str(getattr(s,'name','')).startswith('AUTO_8D_')]
        target=next((s for s in cand if id(s) not in used),None)
        if target is None:continue
        used.add(id(target))
        try:
            tw=target.width/914400.0; th=target.height/914400.0
            target.left=Inches(max(.03,z['x']-tw-.08)); target.top=Inches(z['y']+max(0,(z['h']-th)/2))
        except Exception:pass

def _remove_generated_labels(sl):
    for sh in list(sl.shapes):
        if str(getattr(sh,'name','')).startswith('AUTO_8D_LABEL_'):
            el=sh._element; el.getparent().remove(el)

def _ungroup_old(sl):
    # Remove groups created by this version before the next run, returning their children
    # to the slide so the renderer can replace the old content cleanly.
    tree=sl.shapes._spTree
    for grp in list(tree):
        try:name=grp.cNvPr.get('name')
        except Exception:name=''
        if not str(name).startswith('8D_GROUP_'):continue
        children=list(grp)[2:]
        # grpSpPr is usually child index 1; move all shape children only.
        for child in children:
            grp.remove(child); tree.insert(list(tree).index(grp),child)
        tree.remove(grp)


def _make_group(sl,shapes,name):
    shapes=[s for s in shapes if s is not None and getattr(s,'_element',None) is not None]
    if len(shapes)<2:return
    from pptx.oxml.xmlchemy import OxmlElement
    tree=sl.shapes._spTree
    xs=[int(s.left) for s in shapes]; ys=[int(s.top) for s in shapes]
    xe=[int(s.left+s.width) for s in shapes]; ye=[int(s.top+s.height) for s in shapes]
    minx,miny,maxx,maxy=min(xs),min(ys),max(xe),max(ye); cx=maxx-minx; cy=maxy-miny
    grp=OxmlElement('p:grpSp'); nv=OxmlElement('p:nvGrpSpPr'); cNvPr=OxmlElement('p:cNvPr')
    ids=[int(x.get('id')) for x in tree if x.get('id') and str(x.get('id')).isdigit()]
    cNvPr.set('id',str(max(ids+[1])+1)); cNvPr.set('name',name)
    nv.append(cNvPr); nv.append(OxmlElement('p:cNvGrpSpPr')); nv.append(OxmlElement('p:nvPr')); grp.append(nv)
    gp=OxmlElement('p:grpSpPr'); xf=OxmlElement('a:xfrm')
    off=OxmlElement('a:off');off.set('x',str(minx));off.set('y',str(miny));ext=OxmlElement('a:ext');ext.set('cx',str(cx));ext.set('cy',str(cy))
    choff=OxmlElement('a:chOff');choff.set('x',str(minx));choff.set('y',str(miny));chext=OxmlElement('a:chExt');chext.set('cx',str(cx));chext.set('cy',str(cy))
    xf.extend([off,ext,choff,chext]);gp.append(xf);grp.append(gp)
    first=min(list(tree).index(s._element) for s in shapes);tree.insert(first,grp)
    for s in shapes:
        if s._element.getparent() is tree:tree.remove(s._element)
        grp.append(s._element)

def _group_sections(sl):
    # Use direct top-level generated shapes and the existing template circle.
    direct=list(sl.shapes)
    markers={}
    for sh in _walk(sl):
        t=_n(getattr(sh,'text',''))
        if t in ('2D','3D','4D','5D','6D') and not str(getattr(sh,'name','')).startswith('AUTO_8D_'):markers.setdefault(t,sh)
    mapping={
        '2D':['TEXT_2D','IMG_2D'],
        '3D':['TEXT_3D','IMG_3D'],
        '4D':['TEXT_4D_CAUSE','IMG_4D_CAUSE','TEXT_4D_LEAK','IMG_4D_LEAK'],
        '5D':['TEXT_5D','IMG_5D'],
        '6D':['TEXT_6D','IMG_6D']}
    for label,prefixes in mapping.items():
        arr=[]
        if markers.get(label):arr.append(markers[label])
        arr += [s for s in direct if any(p in str(getattr(s,'name','')) for p in prefixes)]
        _make_group(sl,arr,'8D_GROUP_'+label)

# ---------------------------------------------------------------------------
# Renderer wrapper: collision flow remains active, but template circles are the labels.
# ---------------------------------------------------------------------------
_orig_fill=renderer._fill_page2
def _fill_page2_v305(sl,d,g,layout):
    _ungroup_old(sl)
    _orig_fill(sl,d,g,layout)
    _remove_generated_labels(sl)
    _move_markers(sl,layout)
    _group_sections(sl)
renderer._fill_page2=_fill_page2_v305

# ---------------------------------------------------------------------------
# Summary table: read the existing headers. Never assume Issue/Task column order.
# ---------------------------------------------------------------------------
def _header_map(tb):
    m={}
    for c in range(len(tb.columns)):
        q=_c(tb.cell(0,c).text)
        if '과제명' in q:m['task']=c
        if '이슈' in q:m['issue']=c
        if '문제' in q or '현상' in q:m['problem']=c
        if '진행' in q:m['progress']=c
        if 'signal' in q:m['signal']=c
    return m

def _progress(d):
    try:return v303.v301.impl.v29.v29_progress(d)
    except Exception:return ''

def _set_summary(prs,d):
    vals={'task':_n(d.get('task_name')),'issue':_n(d.get('issue_name')),'problem':_n(d.get('problem')),'progress':_progress(d)}
    for sl in prs.slides:
        for sh in sl.shapes:
            if not getattr(sh,'has_table',False):continue
            tb=sh.table;hm=_header_map(tb)
            if 'issue' not in hm or ('task' not in hm and 'problem' not in hm):continue
            row=None
            for r in range(1,len(tb.rows)):
                if not any(_n(tb.cell(r,c).text) for c in range(len(tb.columns))):row=r;break
            if row is None:row=1 if len(tb.rows)>1 else 0
            for k,c in hm.items():
                if k in vals:tb.cell(row,c).text=vals[k]
            if 'signal' in hm:base.signal(tb.cell(row,hm['signal']),base.status(d))
            return

def weekly_v305(src,out,d,g,mode):
    prs=Presentation(src);_set_summary(prs,d);layout=renderer._load_layout();st=base.status(d)
    summary_idx=0
    for i,sl in enumerate(prs.slides):
        found=False
        for sh in sl.shapes:
            if getattr(sh,'has_table',False) and 'issue' in _header_map(sh.table):found=True;break
        if found:summary_idx=i;break
    for sl in [prs.slides[i] for i in range(summary_idx+1,len(prs.slides))]:renderer._fill_page2(sl,d,g,layout)
    Path(out).parent.mkdir(exist_ok=True)
    try:prs.save(out);saved=out
    except PermissionError:
        p=Path(out);saved=p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix);prs.save(saved)
    return '주간회의 PPT 업데이트: '+st,saved
base.weekly=weekly_v305
if __name__=='__main__':base.App().mainloop()
