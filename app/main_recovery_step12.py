import datetime
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

import main_recovery_step11 as step11
import main_recovery_step10 as step10
import main_recovery_step8 as step8
import main_recovery_step4 as step4
import main_v319 as v319
import main_v310 as v310
import main_v315 as v315

base=step11.base
N=v310.N
C=v310.C


def _signal_status_from_g(d,g):
    """Use the independently confirmed weekly-meeting Signal when available."""
    weekly=N(g.get('_weekly_status_selected'))
    if weekly in ('원인/개선 미확인','개선 검증중','개선 완료'):
        return weekly

    # Backward-compatible fallback for older callers that only supply Issue DB status.
    selected=N(g.get('_issue_status_selected')).lower()
    if selected=='close':
        return '개선 완료'
    if selected=='open':
        return '개선 검증중' if N(d.get('action_5d')) or N(d.get('verification_6d')) else '원인/개선 미확인'
    return base.status(d)


def _sync_page2_signal_only(prs,d,g):
    """Update ONLY page-2 Signal.

    Page 1 is already updated by v319._update_page1(), which knows the correct
    issue row. Rewriting page 1 here is unsafe for existing issues because the
    active issue may not be the row immediately below the header; doing so can
    overwrite the 유첨/비고 cell. Page 2 must simply mirror the same status.
    """
    st=_signal_status_from_g(d,g)
    if len(prs.slides)<2:
        return st

    sl=prs.slides[1]
    for sh in v310.walk(sl):
        if not getattr(sh,'has_table',False):
            continue
        tb=sh.table
        for r in range(len(tb.rows)):
            for c in range(len(tb.columns)):
                if C(tb.cell(r,c).text)=='signal' and c+1<len(tb.columns):
                    base.signal(tb.cell(r,c+1),st)
    return st


def _labeled_4d_text(d,key):
    if key=='4D_CAUSE':
        body=step11._compact_content_text(d.get('cause_4d')) or '검토 중'
        return step11._compact_content_text('- 발생원인\n'+body)
    parts=[step11._compact_content_text(x) for x in (d.get('leak_cause'),d.get('system_cause'))]
    body='\n'.join(x for x in parts if x) or '검토 중'
    return step11._compact_content_text('- 유출원인\n'+body)


def _update_page2_step12(sl,d,g,mode):
    v310.remove_previous_auto(sl)
    imgs=d.get('_section_images',{}) or {}
    english_mode=bool(d.get('_english_mode'))
    # Cascade by actual content height. Short 3D allows 4D to move upward;
    # long 4D can grow downward to the safe slide bottom. Images remain inside
    # the same calculated D box and all marker/title units are moved with it.
    zones,texts,fonts=step11._adaptive_cascade_layout(d,imgs)

    texts['4D_CAUSE']=_labeled_4d_text(d,'4D_CAUSE')
    texts['4D_LEAK']=_labeled_4d_text(d,'4D_LEAK')

    # Align marker/title units for BOTH new and existing issues. A previously
    # generated page may contain markers moved by an older dynamic layout.
    for label,key in [('2D','2D'),('3D','3D'),('5D','5D'),('6D','6D')]:
        z=zones[key]
        v310.move_marker_unit(sl,label,max(.10,z['x']-.18),max(.10,z['y']-.35))

    for key in ('2D','3D','4D_CAUSE','4D_LEAK','5D','6D'):
        v319._render(sl,key,zones[key],texts[key],imgs.get(key,[]),fonts[key])

    # Prevent visual text overflow beyond each calculated D geometry for both
    # Korean and English. Images are fit inside the same calculated rectangle.
    for sh in sl.shapes:
        name=str(getattr(sh,'name',''))
        if name.startswith('AUTO_8D_TEXT_') and hasattr(sh,'text_frame'):
            try:
                sh.text_frame.word_wrap=True
                sh.text_frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            except Exception:
                pass

    v319._page2_meta(sl,d,g)

    # 7D is a native template marker/title unit.  It must not constrain 5D/6D
    # geometry; simply keep the whole existing 7D circle + horizontal-deployment
    # unit below 6D.  Move the native group (or marker+nearby title) together.
    try:
        marker7,parent7=v310.find_marker(sl,'7D')
        if marker7 is not None:
            _,y7,_,_=v310.box(parent7 if parent7 is not None else marker7)
            current_y7=float(y7)/v310.EMU
            z6=zones['6D']
            # The template's 7D unit itself was too high. Move the whole
            # native marker+title visibly downward; do not squeeze or cap 6D.
            unit=parent7 if parent7 is not None else marker7
            x7,_,_,h7=v310.box(unit)
            try:
                slide_h=float(sl.part.package.presentation_part.presentation.slide_height)/v310.EMU
            except Exception:
                slide_h=7.5
            desired_y7=max(current_y7+.24, z6['y']+z6['h']+.28)
            desired_y7=min(desired_y7, max(.10,slide_h-float(h7)/v310.EMU-.06))
            if abs(desired_y7-current_y7)>.01:
                v310.move_marker_unit(sl,'7D',float(x7)/v310.EMU,desired_y7)
    except Exception:
        pass

    # Both 4D blocks are always present in the detail layout.  Keep the native
    # 4D circle + item-name as a group and position each group just above the
    # upper-left corner of its corresponding content region.
    zl=zones['4D_CAUSE']; zr=zones['4D_LEAK']
    left_xy=(max(.10,zl['x']-.18),max(.10,zl['y']-.35))
    right_xy=(max(.10,zr['x']-.18),max(.10,zr['y']-.35))
    v315._ensure_4d_units(sl,True,True,left_xy,right_xy)


def weekly_step12(src,out,d,g,mode):
    dd=v319._weekly_display(d)
    p1=dict(dd)
    p1['issue_name']=step4.step1._page1_issue(d)
    p1['problem']=step4._full_problem(d)

    prs=Presentation(src)

    # Page 1 owns its own correct row-selection logic. Do not touch it again later.
    v319._update_page1(prs,p1,g)

    if len(prs.slides)>1:
        _update_page2_step12(prs.slides[1],dd,g,mode)
        step4._force_page2_header(prs.slides[1],d,g)

    # Synchronize page 2 only. This keeps page 1 유첨/비고 completely untouched.
    _sync_page2_signal_only(prs,d,g)

    Path(out).parent.mkdir(parents=True,exist_ok=True)
    try:
        prs.save(out); saved=out
    except PermissionError:
        p=Path(out)
        saved=str(p.with_name(p.stem+'_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+p.suffix))
        prs.save(saved)

    origin=N(g.get('_issue_origin_selected'))
    if origin:
        step8._set_origin_in_ppt(saved,origin)

    return '주간회의 PPT 업데이트: '+_signal_status_from_g(d,g),saved


base.weekly=weekly_step12


class RecoveryStep12App(step10.RecoveryStep10App):
    def __init__(self):
        super().__init__()
        self.title('8D 이슈 자동화 v3.2.0 RECOVERY STEP12 FIX2')


if __name__=='__main__':
    RecoveryStep12App().mainloop()
